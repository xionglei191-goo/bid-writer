from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    args = parser.parse_args()
    document = Document(args.path)
    for index, paragraph in enumerate(document.paragraphs):
        drawings = paragraph._p.xpath(".//w:drawing")
        if not drawings:
            continue
        rel_ids = paragraph._p.xpath(".//a:blip/@r:embed")
        targets = []
        for rel_id in rel_ids:
            relation = document.part.rels.get(rel_id)
            targets.append(str(relation.target_ref) if relation else rel_id)
        before = [document.paragraphs[pos].text for pos in range(max(0, index - 2), index)]
        after = [document.paragraphs[pos].text for pos in range(index + 1, min(len(document.paragraphs), index + 3))]
        print(index, len(drawings), targets, {"before": before, "text": paragraph.text, "after": after})


if __name__ == "__main__":
    main()
