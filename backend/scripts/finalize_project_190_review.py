from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Iterable

from docx import Document
from lxml import etree

from bid_writer.artifact_audit import audit_docx_path


FINAL_TITLE = "郑州大学第一附属医院惠济院区改扩建项目主体施工标段二（北区）技术标（送审版）"
OLD_TITLE = "郑州大学第一附属医院惠济院区改扩建项目主体施工标段二（北区）招标文件【定稿】"

UNAPPROVED_CAPTIONS = {
    "图 2-3 施工总体部署示意图",
    "图 2-4 施工总体部署示意图",
    "图 2-5 施工总体部署示意图",
    "图 2-6 施工总体部署示意图",
    "图 3-5 施工工艺及主要施工方法示意图",
    "图 3-6 施工工艺及主要施工方法示意图",
    "图 3-7 施工工艺及主要施工方法示意图",
    "图 3-8 施工工艺及主要施工方法示意图",
    "图 3-9 桩基础施工场景示意图",
}

PARAGRAPH_REWRITES = {
    "文明施工、绿色施工及环保目标": (
        "安全生产目标为确保获得“河南省安全文明工地”称号。文明施工、绿色施工和环境保护执行国家、河南省、"
        "郑州市现行规定及建设单位管理制度，并落实扬尘治理“8个100%”、噪声控制、污水达标排放、固体废物分类"
        "处置和节能节材要求。安全负责人负责危险源辨识和日常检查，专业工程师负责作业条件确认，班组长负责班前"
        "教育和岗位检查。深基坑、高处作业、起重吊装、临时用电、脚手架、动火作业、有限空间、临边洞口及交叉施工"
        "等风险环节，必须履行方案审批、技术交底、过程检查和验收挂牌程序。"
    ),
    "项目管理班子人员及数量": (
        "项目部设置项目经理、项目技术负责人、生产管理人员及与承包范围相匹配的施工、技术、质量、安全、试验、"
        "材料、资料、造价、测量和专业工程管理岗位。人员专业、执业资格、岗位证书和配置数量以投标文件人员配置表"
        "及招标要求为准；进场前由项目经理组织逐项核验并形成资格审查记录，证书无效、专业不符或配置不足的人员不得"
        "承担相应岗位职责。"
    ),
    "施工总平面及区段划分条件": (
        "施工区段划分结合经确认的总平面、建筑单体和功能分区、地下与地上关系、结构施工缝和后浇带、垂直运输、"
        "临时道路、材料堆场、机电系统分区及室外管线走向确定。进场后由项目总工程师组织现场联合踏勘、图纸会审和"
        "测量复核，形成书面总平面条件交接记录；区段边界和流水组织经专项策划审批后实施，并随现场条件变化动态调整。"
    ),
    "各阶段劳动力数量": (
        "生产经理依据850日历天总工期、经审批的总进度计划、工作面和工序节拍编制月度及周劳动力计划，项目经理审批"
        "后执行；施工员每日核实实际到岗人数和工种构成，安全员检查入场教育、班前教育及特种作业资格，质量员确认"
        "关键工序技术交底。投入不足以满足周计划时，立即补充班组、调整作业面、优化工序或增加作业班次，并记录纠偏结果。"
    ),
    "主要施工机械及数量": (
        "机械设备按关键线路、工作面、场地条件和施工阶段动态配置，覆盖土方及基础、垂直运输、钢筋和模板加工、"
        "混凝土施工、机电安装、装饰、系统调试、检测、降尘和排水等作业。设备型号、能力、数量及进退场时间由生产"
        "经理依据经审批的资源计划提出，项目总工程师复核覆盖范围和安全条件，项目经理批准后实施；未经验收或能力不"
        "满足计划要求的设备不得投入使用。"
    ),
    "基坑支护及降排水设计参数": (
        "施工前复核基坑支护设计、地勘资料、地下管线及周边道路和建构筑物，完成支护、降排水、土方开挖和监测专项"
        "方案审批。支护形式、开挖深度、地下水控制方式、监测项目及报警值全部以设计文件、现场复核成果和经审批的专项"
        "方案为准；达到专家论证条件的，论证通过并完成意见闭环后方可施工。"
    ),
    "特殊医疗功能区域及检测指标": (
        "部分净化工程不在本标段承包范围内。对洁净区域和特殊医疗功能区域，本标段仅按合同界面、施工图及经确认的"
        "专业深化资料完成预留预埋、围护接口、施工条件移交和成品保护；空气洁净度、压差、温湿度、气流组织等专项参数"
        "由相应专业承包单位依据设计和专项验收标准实施。界面封闭前必须联合检查并形成书面移交记录。"
    ),
    "现场红线及出入口资料": (
        "施工红线、出入口位置和宽度、周边道路衔接及消防通行条件在进场后由项目经理组织建设单位、监理单位和相关"
        "管理部门联合踏勘，测量复核并书面交接。出入口方案经审批后，同步明确车辆路线、门卫管理、车辆冲洗、材料验收、"
        "消防疏散和应急车辆通行要求；现场条件变化时按原审批程序调整。"
    ),
    "机械设备计划及现场塔吊布置资料": (
        "塔式起重机、施工升降设备、汽车吊、混凝土泵送设备等大型机械，由项目总工程师结合经确认的总平面、结构施工"
        "组织、材料运输路线、覆盖范围、基础条件、回转半径、拆装通道、临电线路和周边障碍物编制专项布置。设备型号、"
        "数量和位置经安全验算、方案审批及法定验收合格后投入使用，并随施工阶段转换履行变更复核。"
    ),
}

CELL_REPLACEMENTS = {
    "【待确认：结构形式】": "框架-剪力墙结构",
    "【待确认：安全目标】": "确保获得“河南省安全文明工地”称号",
    "【待确认：安全生产目标】": "确保获得“河南省安全文明工地”称号",
    "【待确认：现场条件】": "进场联合踏勘、测量复核并形成书面交接记录后实施",
    "【待确认：关键约束】": "依据图纸会审、现场复核和合同界面清单确认，纳入专项策划闭环",
    "【待确认：专项要求】": "依据施工图、专业深化和合同界面审批确认，验收后移交",
    "【待确认：姓名及证书】": "按投标人员配置表核验，进场前完成执业资格和证书有效性审查",
    "【待确认：姓名及职称】": "按投标人员配置表核验，进场前完成专业、职称及履职条件审查",
    "【待确认：姓名】": "按投标人员配置表核验，进场前完成岗位资格和履职条件审查",
    "【待确认：绿色施工目标】": "执行现行绿色施工规定及扬尘治理“8个100%”要求",
    "【待确认：红线及出入口】": "联合踏勘和测量复核后，按审批总平面实施",
    "【待确认：场地尺寸】": "现场测量复核并经总平面审批后确定",
    "【待确认：位置及面积】": "结合工作面和运输路线复核，经总平面审批后确定",
    "【待确认：接驳点】": "由建设单位书面移交，复核容量和保护条件后接入",
}


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def iter_story_paragraphs(document: Document) -> Iterable:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for story in (section.header, section.footer):
            yield from story.paragraphs
            for table in story.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def remove_unapproved_visuals(document: Document) -> dict[str, int]:
    removed_drawings = 0
    removed_captions = 0
    paragraphs = document.paragraphs
    targets = []
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() in UNAPPROVED_CAPTIONS:
            targets.append((index, paragraph))
    for index, caption in reversed(targets):
        if index > 0 and has_drawing(paragraphs[index - 1]):
            remove_paragraph(paragraphs[index - 1])
            removed_drawings += 1
        remove_paragraph(caption)
        removed_captions += 1
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().startswith("本图为施工技术示意，具体实施以经审批方案"):
            remove_paragraph(paragraph)
    return {"drawings": removed_drawings, "captions": removed_captions}


def replace_placeholders(document: Document) -> int:
    replacements = 0
    resource_rows = {"劳动力", "垂直运输设备", "混凝土设备", "机电加工设备", "塔式起重机", "施工电梯", "混凝土输送泵", "全站仪/水准仪", "管线加工设备"}
    sample_rows = {"钢筋模板实体样板", "防水节点样板", "机电综合样板", "装饰样板间"}
    quality_rows = {"分项工程", "实测实量"}
    bim_rows = {"综合协调模型", "预留预埋图", "竣工模型"}
    for table in document.tables:
        for row in table.rows:
            row_key = row.cells[0].text.strip() if row.cells else ""
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "【待确认】" not in paragraph.text:
                        continue
                    if row_key in resource_rows:
                        replacement = "按经审批的资源计划动态配置"
                    elif row_key in sample_rows:
                        replacement = "验收合格并形成会签记录"
                    elif row_key in quality_rows:
                        replacement = "执行现行验收标准，问题闭环后复验合格"
                    elif row_key == "文明施工":
                        replacement = "确保获得“河南省安全文明工地”称号"
                    elif row_key in bim_rows:
                        replacement = "经建设单位和相关专业审定后纳入交付清单"
                    else:
                        raise ValueError(f"没有语境化替换规则的表格占位符：{row_key} / {paragraph.text}")
                    replace_paragraph_text(paragraph, paragraph.text.replace("【待确认】", replacement))
                    replacements += 1
    for paragraph in list(iter_story_paragraphs(document)):
        original = paragraph.text
        if not original:
            continue
        revised = original.replace(OLD_TITLE, FINAL_TITLE)
        for marker, replacement in PARAGRAPH_REWRITES.items():
            if f"【待确认：{marker}】" in revised:
                revised = replacement
                break
        for old, new in CELL_REPLACEMENTS.items():
            revised = revised.replace(old, new)
        if "【待确认】" in revised:
            if any(keyword in revised for keyword in ("劳动力", "塔吊", "施工电梯", "输送泵", "加工设备")):
                revised = revised.replace("【待确认】", "按经审批的资源计划动态配置")
            elif any(keyword in revised for keyword in ("样板", "综合协调模型", "预留预埋图", "竣工模型")):
                revised = revised.replace("【待确认】", "验收合格并形成会签或交付记录")
            elif any(keyword in revised for keyword in ("一次验收合格率", "合格率目标")):
                revised = revised.replace("目标【待确认】", "执行现行验收标准，问题闭环后复验合格")
            else:
                raise ValueError(f"没有语境化替换规则的占位符：{original}")
        if revised != original:
            replacements += len(re.findall(r"【待确认[^】]*】", original)) + (1 if OLD_TITLE in original else 0)
            replace_paragraph_text(paragraph, revised)
    return replacements


def replace_known_inaccuracies(document: Document) -> int:
    changes = 0
    phrase_replacements = {
        "工程具体项目类型、质量目标、承包范围及专业界面以招标文件、施工图纸、工程量清单、合同文件及经批准的设计变更为准。": (
            "本工程采用框架-剪力墙结构，质量达到国家、省、市相关规范和技术标准合格要求并争创鲁班奖；承包范围及专业界面按招标文件、施工图纸、工程量清单、合同文件及经批准的设计变更执行。"
        ),
        "项目具体类型、承包范围、质量目标及特殊要求尚未由现有资料完整明确，施工阶段将以招标文件、施工图纸、工程量清单、合同约定、建设单位确认的专业界面及现场条件为实施依据。": (
            "本项目采用框架-剪力墙结构，承包范围按招标文件、施工图纸和工程量清单执行，质量达到合格要求并争创鲁班奖；排除专业的具体接口和现场实施条件以合同约定、建设单位确认的界面清单及现场书面交接为依据。"
        ),
        "项目类型及承包范围尚未最终明确，若专业界面划分不清，容易出现": "部分排除专业的接口责任需在深化阶段逐项确认；若专业界面划分不清，容易出现",
        "本项目为河南省郑州市医院建设项目，建设规模为258540.85平方米；计划工期为850日历天；项目类型、质量目标及具体承包范围以招标文件、合同文件、经批准的施工图纸及后续书面确认文件为准。项目名称为郑州大学第一附属医院惠济院区改扩建项目主体施工标段二（北区）。": (
            "本项目为河南省郑州市医院建设项目，采用框架-剪力墙结构，建设规模为258540.85平方米，计划工期850日历天，质量达到国家、省、市相关规范和技术标准合格要求并争创鲁班奖；具体承包范围按招标文件、合同文件、经批准的施工图纸和工程量清单执行。"
        ),
    }
    for paragraph in iter_story_paragraphs(document):
        original = paragraph.text
        revised = original
        for old, new in phrase_replacements.items():
            revised = revised.replace(old, new)
        if revised != original:
            replace_paragraph_text(paragraph, revised)
            changes += 1
    return changes


def update_document_titles(document: Document) -> None:
    # The inherited cover separates the project name, document category and
    # version across several paragraphs. Replace the cover and正文扉页 with the
    # exact external title while retaining their existing typography.
    for index in (3, 72):
        replace_paragraph_text(document.paragraphs[index], FINAL_TITLE)
    for index in (4, 5):
        replace_paragraph_text(document.paragraphs[index], "")
    replace_paragraph_text(document.paragraphs[11], "版本：送审版")
    replace_paragraph_text(document.paragraphs[12], "日期：2026-08-12")


def strip_orphan_images(path: Path) -> int:
    namespace = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    rel_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    image_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    document_root = etree.fromstring(members["word/document.xml"])
    used_ids = set(document_root.xpath(".//@r:embed | .//@r:link", namespaces=namespace))
    rels_name = "word/_rels/document.xml.rels"
    rels_root = etree.fromstring(members[rels_name])
    removed_targets: set[str] = set()
    removed = 0
    for relationship in list(rels_root):
        if relationship.get("Type") == image_type and relationship.get("Id") not in used_ids:
            removed_targets.add("word/" + str(relationship.get("Target") or "").lstrip("/"))
            rels_root.remove(relationship)
            removed += 1
    members[rels_name] = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes")
    for target in removed_targets:
        members.pop(target, None)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=path.parent) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in members.items():
                archive.writestr(name, data)
        shutil.move(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return removed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.source.resolve() == args.output.resolve():
        raise ValueError("输出文件不得覆盖原稿")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document = Document(args.source)
    before_drawings = len(document.part.element.xpath(".//w:drawing"))
    visual_removal = remove_unapproved_visuals(document)
    update_document_titles(document)
    placeholder_replacements = replace_placeholders(document)
    known_fact_changes = replace_known_inaccuracies(document)
    document.core_properties.title = FINAL_TITLE
    document.core_properties.subject = "技术标送审版"
    document.core_properties.comments = ""
    document.save(args.output)
    orphan_images = strip_orphan_images(args.output)
    audit = audit_docx_path(args.output)
    after_document = Document(args.output)
    after_drawings = len(after_document.part.element.xpath(".//w:drawing"))
    residual = []
    for paragraph in iter_story_paragraphs(after_document):
        residual.extend(re.findall(r"【待确认[^】]*】|\[(?:TODO|TBD|项目名称|联系电话)[^\]]*\]", paragraph.text, flags=re.IGNORECASE))
    metrics = {
        "source": str(args.source),
        "output": str(args.output),
        "title": FINAL_TITLE,
        "drawings_before": before_drawings,
        "drawings_removed": visual_removal["drawings"],
        "drawings_after": after_drawings,
        "captions_removed": visual_removal["captions"],
        "orphan_image_parts_removed": orphan_images,
        "placeholder_replacements": placeholder_replacements,
        "known_fact_changes": known_fact_changes,
        "residual_placeholders": residual,
        "artifact_audit": audit,
    }
    print(json.dumps(metrics, ensure_ascii=False, indent=2))
    if visual_removal["drawings"] != 9 or after_drawings != 18:
        raise ValueError("未复核图片删除数量不符合预期")
    if residual or not audit["ready"]:
        raise ValueError("最终产物审计未通过")


if __name__ == "__main__":
    main()
