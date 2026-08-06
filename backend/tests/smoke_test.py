from __future__ import annotations

import os
import sys
import zipfile
import json
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

test_db = BACKEND_DIR.parent / "data" / "test_smoke.sqlite"
test_db.parent.mkdir(parents=True, exist_ok=True)
for suffix in ("", "-wal", "-shm"):
    target = Path(str(test_db) + suffix)
    if target.exists():
        target.unlink()
os.environ["BID_WRITER_DB"] = str(test_db)
os.environ["BID_WRITER_DISABLE_USER_ENV"] = "1"
os.environ.pop("OPENAI_API_KEY", None)

from bid_writer.db import connect, init_db, stats
from bid_writer.bid_strategy import (
    generate_bid_strategy,
    render_bid_strategy_markdown,
    update_bid_strategy,
)
from bid_writer.case_assets import (
    create_case_asset_from_draft,
    create_case_assets_from_tender,
    list_case_assets,
    render_case_assets_markdown,
    search_case_assets,
    update_case_asset,
)
from bid_writer.channel_ops import build_channel_ops, render_channel_ops_markdown
from bid_writer.client_delivery_confirmation import build_client_delivery_confirmation, confirm_client_delivery
from bid_writer.client_delivery_preparation import prepare_client_delivery
from bid_writer.client_package_validation import validate_latest_client_package
from bid_writer.closure_confirmation import build_closure_confirmation, record_closure_confirmation
from bid_writer.communications import (
    create_communication,
    list_communications,
    render_communications_markdown,
    suggest_communication_reply,
    update_communication,
)
from bid_writer.command_center import build_command_center
from bid_writer.coverage_report import build_coverage_report
from bid_writer.construction_methods import method_outline_for
from bid_writer.delivery_assistant import build_delivery_assistant, render_delivery_instruction
from bid_writer.delivery_records import list_delivery_records, update_delivery_record
from bid_writer.delivery_release import build_delivery_release
from bid_writer.delivery_review import build_delivery_review
from bid_writer.document_processing import list_document_processing_records
from bid_writer.document_blocks import (
    delete_document_block,
    generate_document_blocks,
    list_document_blocks,
    update_document_block,
)
from bid_writer.document_settings import get_document_settings, render_document_settings_markdown, update_document_settings
from bid_writer.draft_polish import apply_replacement, list_replacement_records, preview_replacement, scan_project_polish
from bid_writer.draft_versions import list_draft_versions, restore_draft_version
from bid_writer.enterprise_profiles import get_enterprise_profile, render_enterprise_profile_markdown, update_enterprise_profile
from bid_writer.exporter import _append_markdown_content, _append_markdown_line, export_client_package, export_docx, export_markdown, export_package
from bid_writer.final_document import build_final_document, render_final_document_markdown, update_final_document
from bid_writer.final_checklist import build_final_checklist, update_final_check_item
from bid_writer.feedback import create_feedback_item, feedback_summary, update_feedback_item
from bid_writer.feedback_rework import apply_feedback_rework_plan, build_feedback_rework, execute_feedback_rework
from bid_writer.generator import generate_draft
from bid_writer.intake_assistant import build_intake_assistant, create_intake_tender, render_intake_markdown
from bid_writer.kb_audit import build_kb_audit, render_kb_audit_markdown
from bid_writer.kb_importer import import_knowledge_base
from bid_writer.kb_ocr_queue import build_kb_ocr_queue
from bid_writer.llm_config import llm_settings, probe_llm
from bid_writer.materials import create_material_item, list_material_items, material_summary, update_material_item
from bid_writer.order_dashboard import build_order_dashboard
from bid_writer.order_confirmation import build_order_confirmation
from bid_writer.order_wizard import build_order_wizard, render_order_wizard_markdown
from bid_writer.package_validation import validate_latest_package
from bid_writer.payments import create_payment_record, payment_summary, render_payments_markdown, update_payment_record
from bid_writer.parsing import parse_tender_text
from bid_writer.pricing import build_price_quote, ensure_default_pricing_rules, list_quotation_records, update_pricing_rule
from bid_writer.task_status import sync_all_production_status, sync_production_status
from bid_writer.production_readiness import build_production_readiness, render_production_readiness_markdown
from bid_writer.production_starter import build_production_starter, render_production_starter_markdown
from bid_writer.planner import (
    add_section_plan,
    audit_section_plan,
    build_section_plan,
    delete_section_plan,
    generate_all_from_plan,
    generate_from_plan,
    list_section_plans,
    repair_section_plan,
    render_plan_audit_markdown,
    update_section_plan,
)
from bid_writer.production_pipeline import run_production_pipeline
from bid_writer.production_tasks import get_production_task, update_production_task
from bid_writer.project_overview import build_project_overview
from bid_writer.project_timeline import build_project_timeline, render_project_timeline_markdown
from bid_writer.project_profiles import get_project_profile, update_project_profile
from bid_writer.requirement_scope import requirement_scope
from bid_writer.quality_gate import build_quality_gate
from bid_writer.retrieval import search_chunks
from bid_writer.review import review_draft, review_tender
from bid_writer.retrospective import build_project_retrospective, build_retrospective_dashboard, update_project_retrospective
from bid_writer.response_matrix import auto_link_response_matrix, build_response_matrix, render_response_matrix_markdown
from bid_writer.revision_tasks import (
    list_revision_tasks,
    render_revision_tasks_markdown,
    sync_revision_tasks,
    update_revision_task,
)
from bid_writer.section_templates import (
    all_templates,
    delete_section_template,
    save_section_template,
    template_coverage_report,
    template_for,
    template_terms,
    update_section_template,
)
from bid_writer.source_audit import build_source_audit, render_source_audit_markdown
from bid_writer.source_preview import preview_source
from bid_writer.tenders import (
    create_requirement,
    create_tender,
    delete_requirement,
    delete_tender,
    get_draft,
    get_requirements,
    get_tender,
    list_drafts,
    parse_tender,
    update_draft,
    update_requirement,
    update_tender_source,
)
from bid_writer.workflow import build_workflow_status
from bid_writer.visual_assets import list_visual_assets, resolve_asset_path, update_visual_asset_review
from bid_writer.visual_pipeline import build_visual_plan, generate_hybrid_visual, image_generation_settings

conn = connect()
init_db(conn)
conn.close()

print("import", flush=True)
counters = import_knowledge_base(reset=True, limit_documents=20)
assert counters["documents"] > 0, counters
assert counters["chunks"] > 0, counters

kb_audit = build_kb_audit(limit_directories=12)
assert kb_audit["summary"]["title"] == "知识库体检", kb_audit["summary"]
assert kb_audit["summary"]["manifest_rows"] >= kb_audit["summary"]["manifest_ok"] >= counters["documents"], kb_audit["summary"]
assert kb_audit["summary"]["indexed_documents"] == counters["documents"], kb_audit["summary"]
assert kb_audit["directories"], kb_audit
assert any(item["reason"] for item in kb_audit["top_reasons"]), kb_audit["top_reasons"]
assert "知识库体检报告" in render_kb_audit_markdown(kb_audit), kb_audit["markdown"][:200]
ocr_queue = build_kb_ocr_queue(limit=5)
assert ocr_queue["summary"]["title"] == "OCR 补录队列", ocr_queue["summary"]
assert ocr_queue["summary"]["pending_count"] >= 0, ocr_queue["summary"]
assert "execution" in ocr_queue["policy"], ocr_queue["policy"]

print("search", flush=True)
results = search_chunks("工程概况 施工部署", limit=5)
assert results, "search returned no chunks"
kb_preview = preview_source(results[0])
assert kb_preview["source_type"] == "kb_chunk", kb_preview
assert kb_preview["content"] and kb_preview["source_path"], kb_preview
assert kb_preview["file_exists"] and kb_preview["markdown_excerpt"], kb_preview

print("llm", flush=True)
llm = llm_settings()
assert llm["provider"] == "openai-compatible", llm
assert llm["generation_mode"] == "local_fallback", llm
probe = probe_llm()
assert not probe["ok"] and "OPENAI_API_KEY" in probe["error"], probe

print("enterprise", flush=True)
enterprise = update_enterprise_profile(
    {
        "profile_name": "冒烟测试投标单位",
        "bidder_name": "测试建设集团有限公司",
        "qualification_summary": "建筑工程施工总承包壹级，具备机电安装、装饰装修等专业履约能力。",
        "capability_summary": "具备医院综合楼项目组织、资源调配、深化设计和总承包协调能力。",
        "quality_system": "执行公司质量管理体系，落实样板引路、三检制和过程验收。",
        "safety_system": "执行安全文明施工标准化管理，落实风险分级管控和隐患闭环整改。",
        "key_personnel": "拟投入项目经理、技术负责人、质量负责人、安全负责人及机电专业工程师。",
        "equipment_resources": "配置塔吊、施工电梯、测量仪器、检测设备和周转材料资源。",
        "similar_projects": "具有医院、公共建筑、综合楼等类似项目技术标编制和履约经验。",
        "service_commitment": "配合客户完成格式调整、资料补充和交付前人工复核。",
    }
)
assert get_enterprise_profile()["bidder_name"] == "测试建设集团有限公司", enterprise
assert "测试建设集团有限公司" in render_enterprise_profile_markdown(), enterprise

print("processing", flush=True)
processing_source = test_db.parent / "processing_source.txt"
processing_source.write_text(
    "项目名称：处理记录测试项目\n技术标评分要求：施工总体部署、质量保证措施、安全文明施工。\n",
    encoding="utf-8",
)
processing_tender = create_tender("", file_path=str(processing_source), industry="医院类")
processing_records = list_document_processing_records(int(processing_tender["id"]))
assert processing_records and processing_records[0]["status"] == "success", processing_records
assert processing_records[0]["action"] == "plain_text", processing_records
assert processing_records[0]["text_chars"] > 20, processing_records
deleted_processing = delete_tender(int(processing_tender["id"]))
assert deleted_processing["deleted"], deleted_processing
processing_source.unlink(missing_ok=True)

print("templates", flush=True)
template = save_section_template(
    {
        "name": "吊装专项施工方案",
        "keywords": "吊装\n起重吊装\n大型设备",
        "intent": "说明大型设备和构件吊装的施工组织、技术控制和安全保障。",
        "outline": "吊装条件分析\n吊装流程\n安全控制\n验收复核",
        "quality_points": "起重机械参数需人工复核\n吊装半径和道路承载需核验",
    }
)
assert template["id"] and template["source"] == "custom", template
assert template_for("大型设备起重吊装").name == "吊装专项施工方案"
assert "吊装" in template_terms("吊装专项施工方案")
updated_template = update_section_template(
    int(template["id"]),
    {
        "name": "吊装专项施工方案",
        "keywords": ["吊装专项", "大型设备"],
        "intent": "更新后的吊装专项模板。",
        "outline": ["吊装部署", "过程监测"],
        "quality_points": ["专项方案需人工复核"],
    },
)
assert updated_template["intent"] == "更新后的吊装专项模板。", updated_template
assert any(item["name"] == "吊装专项施工方案" and item["source"] == "custom" for item in all_templates())
deleted_template = delete_section_template(int(template["id"]))
assert deleted_template["deleted"], deleted_template
coverage_report = template_coverage_report()
assert coverage_report["summary"]["construction_method_template_covered"], coverage_report
assert coverage_report["summary"]["missing_total"] == 0, coverage_report
assert any(item["name"] == "施工工艺及主要施工方法" and item["covered"] for item in coverage_report["items"]), coverage_report
assert any("医院" in item["keywords"] for item in coverage_report["construction_methods"]["groups"]), coverage_report
assert template_for("施工方案与技术措施（10分）").name == "施工工艺及主要施工方法"
assert template_for("资源配备计划（5分）").name == "施工总体部署"
assert template_for("风险管理措施（10分）").name == "工程重点难点分析及对策"
assert template_for("扬尘污染防治方案及建筑垃圾处置方案").name == "安全文明施工及环境保护"

print("tender", flush=True)
multiline_scoring = parse_tender_text(
    """
项目名称：评分表解析测试工程
计划工期：850 日历天
施工组织设计评分标准（总分100分）
评分因素 参考评分标准
施工方案与技
术措施
10
各项主要内容的措施是否科学先进、计划是否合理可行。
质量管理体系
与措施
10 确保工程质量的技术组织措施合理、可行。
风险管理措施 10
风险预控符合规范要求，各阶段应急措施得力。
25
水性漆的应用方案
3 水性漆的应用方案合理、可行
非道路移动机械排放污染的管控措施
2
禁止使用2013 年 10 月以前生产的不合格机械。
1、各档次的标准设定如下：
综合标（总分100分）
""",
    "评分表解析测试工程",
)
scoring_items = [item for item in multiline_scoring["requirements"] if item["kind"] == "scoring"]
assert len(scoring_items) == 5, scoring_items
assert scoring_items[0]["content"].startswith("施工方案与技术措施（10分）"), scoring_items
assert "确保工程质量" in scoring_items[1]["content"], scoring_items
assert scoring_items[-1]["content"].startswith("非道路移动机械排放污染的管控措施（2分）"), scoring_items
assert "水性漆的应用方案合理、可行" in scoring_items[-2]["content"], scoring_items
assert not any("（25分）" in item["content"] for item in scoring_items), scoring_items
assert multiline_scoring["overview"]["duration_days"] == 850, multiline_scoring["overview"]
assert requirement_scope({"kind": "risk", "content": "未提交投标保证金将否决投标"}) == "compliance"
assert requirement_scope({"kind": "scoring", "content": "风险管理措施（10分）"}) == "chapter"
tender_text = """
项目名称：测试医院综合楼工程
建设规模：总建筑面积约 50000 平方米，地下2层，地上12层，框架剪力墙结构。
计划工期：540 日历天。
质量目标：确保合格，争创省优工程。
安全目标：杜绝重伤及以上事故，创建安全文明工地。
招标范围：土建、装饰装修、机电安装、室外配套工程。
现场条件：医院不停诊，施工场地狭小，周边交通繁忙。
特殊要求：医疗专项机电、洁污分流、净化区域施工需重点控制。
技术标评分要求：施工总体部署、工程重难点分析、质量保证措施、安全文明施工措施。
主要施工工艺要求：测量放线、基坑土方、主体结构、防水屋面、装饰装修、机电安装等主要施工方法必须完整响应。
投标文件必须响应工期、质量、安全目标，不得出现与本项目无关的历史项目名称。
"""
tender = create_tender("测试医院综合楼工程", tender_text, industry="医院类")
parsed = parse_tender(int(tender["id"]))
assert parsed["requirements"], "requirements not parsed"
assert parsed["profile"]["duration_days"] == 540, parsed["profile"]
assert "50000" in parsed["profile"]["building_area"], parsed["profile"]
assert "地下2层" in parsed["profile"]["floor_info"] and "地上12层" in parsed["profile"]["floor_info"], parsed["profile"]
assert "框架剪力墙" in parsed["profile"]["structure_type"], parsed["profile"]
assert "省优" in parsed["profile"]["quality_target"], parsed["profile"]
assert "安全文明" in parsed["profile"]["safety_target"], parsed["profile"]
assert "土建" in parsed["profile"]["contract_scope"], parsed["profile"]
assert "不停诊" in parsed["profile"]["site_conditions"], parsed["profile"]
assert "医疗专项机电" in parsed["profile"]["special_requirements"], parsed["profile"]
manual_requirement = create_requirement(
    int(tender["id"]),
    {
        "kind": "评分点",
        "content": "人工补充：医疗专项机电调试计划需单独响应。",
        "source_hint": "人工校正",
        "priority": "high",
    },
)
assert manual_requirement["id"] and manual_requirement["priority"] == "high", manual_requirement
manual_requirement = update_requirement(
    int(manual_requirement["id"]),
    {"content": "人工补充：医疗专项机电调试计划和联动测试需单独响应。", "status": "confirmed"},
)
assert manual_requirement["status"] == "confirmed" and "联动测试" in manual_requirement["content"], manual_requirement
requirements = get_requirements(int(tender["id"]))
assert any(item["id"] == manual_requirement["id"] for item in requirements), requirements
status = sync_production_status(int(tender["id"]))
assert status["current_status"] == "待生成目录", status
task = get_production_task(int(tender["id"]))
assert task["deliverable_format"] == "DOCX + ZIP 交付包", task
task = update_production_task(
    int(tender["id"]),
    {
        "customer_name": "测试客户",
        "source_platform": "内部验证",
        "order_no": "TEST-001",
        "deadline": "2026-08-10 18:00",
        "budget": "1200 元",
        "deliverable_format": "DOCX + ZIP 交付包",
        "delivery_status": "生产中",
        "delivery_notes": "用于冒烟测试的临时生产任务。",
    },
)
assert task["customer_name"] == "测试客户", task
dashboard = build_order_dashboard()
assert dashboard["summary"]["total"] == 1, dashboard
assert dashboard["items"][0]["customer_name"] == "测试客户", dashboard
intake = build_intake_assistant(
    int(tender["id"]),
    {
        "customer_message": "客户从闲鱼咨询：医院综合楼技术标，明天前要 Word 初稿，预算 1200 元，需要施工工艺和质量安全章节。",
        "source_platform": "闲鱼",
        "deadline": "2026-08-10 18:00",
        "budget_expectation": "1200 元",
    },
)
assert intake["acceptance"]["decision"] in {"可承接", "谨慎承接"}, intake
assert "闲鱼" == intake["suggested_task"]["source_platform"], intake
assert intake["estimate"]["line_items"], intake["estimate"]
assert "客户回复" in render_intake_markdown(intake), intake
pricing_rules = ensure_default_pricing_rules()
assert any(rule["rule_key"] == "base_price" for rule in pricing_rules), pricing_rules
base_rule = next(rule for rule in pricing_rules if rule["rule_key"] == "base_price")
updated_rule = update_pricing_rule(int(base_rule["id"]), {"value": "720"})
assert updated_rule["value"] == "720", updated_rule
quote = build_price_quote(
    int(tender["id"]),
    {
        "customer_message": "客户从闲鱼咨询：医院综合楼技术标，明天前要 Word 初稿，预算 1200 元，需要施工工艺和质量安全章节。",
        "source_platform": "闲鱼",
    },
    save_record=True,
)
assert quote["suggested_price"] and quote["line_items"], quote
assert "报价测算" in quote["markdown"], quote["markdown"][:200]
quote_records = list_quotation_records(int(tender["id"]))
assert len(quote_records) == 1 and quote_records[0]["suggested_price"] == quote["suggested_price"], quote_records
confirmation = build_order_confirmation(int(tender["id"]), {"agreed_price": "1200 元", "revision_rounds": "2"})
assert "订单确认单" in confirmation["markdown"], confirmation["markdown"][:200]
assert "不承诺中标结果" in confirmation["customer_message"], confirmation["customer_message"]
payment = create_payment_record(
    int(tender["id"]),
    {
        "amount": "600",
        "payment_stage": "定金",
        "payment_method": "闲鱼",
        "status": "已收款",
        "proof": "冒烟测试收款凭证",
    },
)
payment_status = payment_summary(int(tender["id"]))
assert payment_status["status"] == "partial" and payment_status["outstanding_amount"] == 600, payment_status
payment = update_payment_record(int(payment["id"]), {"amount": "1200", "payment_stage": "全款", "status": "已确认"})
payment_status = payment_summary(int(tender["id"]))
assert payment_status["status"] == "paid" and payment_status["delivery_authorized"], payment_status
assert "已收齐" in render_payments_markdown(int(tender["id"]))
communication_suggestion = suggest_communication_reply(
    int(tender["id"]),
    {
        "stage": "quote",
        "channel": "闲鱼",
        "customer_message": "客户问：这份医院技术标多少钱，明天能不能交？",
    },
)
assert communication_suggestion["reply"] and communication_suggestion["stage"] == "报价", communication_suggestion
communication = create_communication(
    int(tender["id"]),
    {
        "stage": "quote",
        "channel": "闲鱼",
        "customer_message": "客户问：这份医院技术标多少钱，明天能不能交？",
        "system_reply": communication_suggestion["reply"],
        "status": "待发送",
        "notes": "冒烟测试报价沟通。",
    },
)
assert communication["id"] and communication["status"] == "待发送", communication
communication = update_communication(int(communication["id"]), {"status": "已发送"})
assert communication["status"] == "已发送", communication
communications = list_communications(int(tender["id"]))
assert communications and communications[0]["stage"] == "报价", communications
communications_markdown = render_communications_markdown(int(tender["id"]))
assert "客户沟通记录" in communications_markdown and "系统建议回复" in communications_markdown, communications_markdown
overview = build_project_overview(int(tender["id"]))
assert overview["readiness"]["label"] and overview["metrics"]["communications"] >= 1, overview
assert "项目总览" in overview["markdown"], overview["markdown"][:200]
starter = build_production_starter(int(tender["id"]))
assert starter["summary"]["strategy"] == "标书生产工具优先", starter["summary"]
assert starter["customer_reply"] and starter["start_checklist"], starter
assert "不自动登录平台" in starter["channel_policy"]["automation_boundary"], starter["channel_policy"]
assert "生产启动包" in render_production_starter_markdown(starter), starter["markdown"][:200]
wizard = build_order_wizard(int(tender["id"]))
assert wizard["summary"]["title"] == "新单生产向导", wizard["summary"]
assert wizard["steps"] and wizard["current_step"]["action"]["action_code"], wizard
assert wizard["copyables"]["customer_reply"], wizard["copyables"]
assert "新单生产向导" in render_order_wizard_markdown(wizard), wizard["markdown"][:200]
channel_ops = build_channel_ops(int(tender["id"]))
assert channel_ops["summary"]["title"] == "渠道运营助手", channel_ops["summary"]
assert channel_ops["summary"]["copy_assets"] >= 6, channel_ops["summary"]
assert any(item["key"] == "listing_description" and "不承诺中标结果" in item["text"] for item in channel_ops["scripts"]), channel_ops["scripts"]
assert "不自动登录平台" in channel_ops["policy"]["automation_boundary"], channel_ops["policy"]
assert "渠道运营助手" in render_channel_ops_markdown(channel_ops), channel_ops["markdown"][:200]
profile = get_project_profile(int(tender["id"]))
assert profile["project_name"] == "测试医院综合楼工程", profile
doc_settings = update_document_settings(
    int(tender["id"]),
    {
        "document_title": "测试医院综合楼工程技术标",
        "document_subtitle": "投标文件技术部分",
        "bidder_name": "测试投标单位",
        "prepared_by": "测试编制人",
        "reviewed_by": "测试复核人",
        "header_text": "测试医院综合楼工程 / 技术标",
        "footer_text": "测试页脚：正式投标前人工复核",
        "include_cover": True,
        "include_toc": True,
        "section_page_break": True,
    },
)
assert doc_settings["document_title"] == "测试医院综合楼工程技术标", doc_settings
assert get_document_settings(int(tender["id"]))["include_cover"], doc_settings
assert "成稿格式设置" in render_document_settings_markdown(int(tender["id"]))
materials = list_material_items(int(tender["id"]))
assert any(item["name"] == "完整招标文件" for item in materials), materials
custom_material = create_material_item(
    int(tender["id"]),
    {
        "category": "格式资料",
        "name": "企业技术标格式模板",
        "status": "待补充",
        "source": "客户稍后提供",
        "notes": "用于统一封面和页眉页脚。",
    },
)
custom_material = update_material_item(int(custom_material["id"]), {"status": "已具备", "source": "客户微信"})
assert custom_material["status"] == "已具备", custom_material
materials_summary = material_summary(int(tender["id"]))
assert materials_summary["total"] >= 6, materials_summary
profile = update_project_profile(
    int(tender["id"]),
    {
        "project_type": "医院类综合楼",
        "structure_type": "框架剪力墙结构",
        "building_area": "约 50000 平方米",
        "floor_info": "地下2层，地上12层",
        "duration_days": 540,
        "quality_target": "确保合格，争创省优工程",
        "safety_target": "杜绝重伤及以上事故，创建安全文明工地",
        "contract_scope": "土建、装饰装修、机电安装、室外配套工程",
        "special_requirements": "医疗专项机电、洁污分流、净化区域施工需重点控制",
    },
)
assert profile["structure_type"] == "框架剪力墙结构", profile
assert any("医疗专项机电" in item for item in method_outline_for(profile)), method_outline_for(profile)
workflow = build_workflow_status(int(tender["id"]))
assert workflow["summary"]["requirements"] > 0, workflow
assert workflow["summary"]["task_missing"] == [], workflow
assert workflow["summary"]["profile_missing"] == [], workflow
command_center = build_command_center(int(tender["id"]))
assert command_center["tender"]["id"] == tender["id"], command_center
assert command_center["cards"] and any(item["key"] == "requirements" for item in command_center["cards"]), command_center
assert command_center["progress"] and command_center["primary_actions"], command_center
plan = build_section_plan(int(tender["id"]))
assert plan, "section plan not built"
status = sync_production_status(int(tender["id"]))
assert status["current_status"] == "生产中", status
assert any(item["section_title"] == "施工工艺及主要施工方法" for item in plan), "construction method template missing"
plan_audit = audit_section_plan(int(tender["id"]))
assert plan_audit["summary"]["missing_standard"] == 0, plan_audit
assert plan_audit["summary"]["construction_method_present"], plan_audit["summary"]
assert "目录完整性审计报告" in render_plan_audit_markdown(plan_audit)
construction_plan = next(item for item in list_section_plans(int(tender["id"])) if item["section_title"] == "施工工艺及主要施工方法")
deleted_construction_plan = delete_section_plan(int(construction_plan["id"]))
assert deleted_construction_plan["deleted"], deleted_construction_plan
missing_plan_audit = audit_section_plan(int(tender["id"]))
assert missing_plan_audit["summary"]["missing_standard"] >= 1, missing_plan_audit
assert not missing_plan_audit["summary"]["construction_method_present"], missing_plan_audit["summary"]
repaired_plan = repair_section_plan(int(tender["id"]))
assert repaired_plan["added_count"] >= 1, repaired_plan
assert repaired_plan["after"]["summary"]["missing_standard"] == 0, repaired_plan["after"]["summary"]
assert repaired_plan["after"]["summary"]["construction_method_present"], repaired_plan["after"]["summary"]
response_matrix = build_response_matrix(int(tender["id"]))
assert response_matrix["summary"]["requirements"] >= 1, response_matrix
assert response_matrix["summary"]["planned_requirements"] >= 1, response_matrix["summary"]
assert "响应矩阵挂接报告" in render_response_matrix_markdown(response_matrix)
late_requirement = create_requirement(
    int(tender["id"]),
    {
        "kind": "评分点",
        "content": "后续补充：BIM深化设计、碰撞检查和智慧建造应用需单独响应。",
        "source_hint": "后续补充",
        "priority": "high",
    },
)
response_matrix_before_link = build_response_matrix(int(tender["id"]))
late_item_before = next(item for item in response_matrix_before_link["requirements"] if item["id"] == late_requirement["id"])
assert not late_item_before["planned_sections"] and late_item_before["auto_linkable"], late_item_before
linked_matrix = auto_link_response_matrix(int(tender["id"]))
assert linked_matrix["linked_count"] >= 1, linked_matrix
late_item_after = next(item for item in linked_matrix["after"]["requirements"] if item["id"] == late_requirement["id"])
assert late_item_after["planned_sections"], late_item_after
deleted_late_requirement = delete_requirement(int(late_requirement["id"]))
assert deleted_late_requirement["deleted"], deleted_late_requirement
plans = list_section_plans(int(tender["id"]))
assert plans and plans[0]["section_title"], "section plans not listed"
linked_plan = add_section_plan(
    int(tender["id"]),
    "医疗专项机电调试计划",
    requirement_ids=[int(manual_requirement["id"])],
)
assert int(manual_requirement["id"]) in linked_plan["requirement_ids"], linked_plan
deleted_requirement = delete_requirement(int(manual_requirement["id"]))
assert deleted_requirement["deleted"], deleted_requirement
linked_plan_after = next(item for item in list_section_plans(int(tender["id"])) if item["id"] == linked_plan["id"])
assert int(manual_requirement["id"]) not in linked_plan_after["requirement_ids"], linked_plan_after
deleted_plan = delete_section_plan(int(linked_plan["id"]))
assert deleted_plan["deleted"], deleted_plan
custom_plan = add_section_plan(int(tender["id"]), "成品保护措施")
updated_plan = update_section_plan(int(custom_plan["id"]), section_title="成品保护及交付保障", order_no=99)
assert updated_plan["section_title"] == "成品保护及交付保障"
deleted_plan = delete_section_plan(int(custom_plan["id"]))
assert deleted_plan["deleted"], deleted_plan
plans = list_section_plans(int(tender["id"]))

strategy = generate_bid_strategy(int(tender["id"]))
assert strategy["positioning"] and strategy["section_focus"], strategy
assert any("施工工艺" in item["section_title"] for item in strategy["section_focus"]), strategy["section_focus"]
strategy = update_bid_strategy(
    int(tender["id"]),
    {
        "writing_tone": "测试策略口径：强调项目化、可执行措施、来源可追溯。",
        "status": "approved",
    },
)
assert strategy["status"] == "approved" and "测试策略口径" in strategy["writing_tone"], strategy
assert "投标响应策略" in render_bid_strategy_markdown(int(tender["id"]))

print("generate", flush=True)
draft = generate_draft(int(tender["id"]), "施工总体部署", category="医院类")
assert "施工总体部署" in draft["content"]
assert "投标响应策略" in draft["content"], draft["content"][:800]
assert "投标单位资料" in draft["content"] and "测试建设集团有限公司" in draft["content"], draft["content"][:1200]
assert "项目化控制要点" in draft["content"], draft["content"][:800]
assert "540 日历天" in draft["content"], draft["content"][:1200]
assert "省优" in draft["content"] and "安全文明" in draft["content"], draft["content"][:1200]
assert draft["generation"]["mode"] == "local_fallback", draft["generation"]
assert "OPENAI_API_KEY" in draft["generation"]["error"], draft["generation"]
method_draft = generate_draft(int(tender["id"]), "施工工艺及主要施工方法", category="医院类")
assert "框架剪力墙结构" in method_draft["content"], method_draft["content"][:500]
assert "医疗专项机电" in method_draft["content"], method_draft["content"][:800]
assert "不停诊" in method_draft["content"], method_draft["content"][:1200]
assert "专项要求" in method_draft["content"], method_draft["content"][:1200]
assert method_draft["generation"]["mode"] == "local_fallback", method_draft["generation"]
draft_citation_preview = preview_source(method_draft["citations"][0])
assert draft_citation_preview["content"] and draft_citation_preview["source_path"], draft_citation_preview
good_findings = review_draft(int(draft["id"]))
assert not any(item.get("type") == "project_consistency" for item in good_findings), good_findings
bad_consistency_draft = update_draft(int(draft["id"]), "# 施工总体部署\n\n本章根据实际情况组织施工，后续按相关要求执行。")
bad_findings = review_draft(int(bad_consistency_draft["id"]))
assert any(item.get("field") == "duration_days" for item in bad_findings), bad_findings
assert any(item.get("field") == "quality_target" for item in bad_findings), bad_findings
consistency_gate = build_quality_gate(int(tender["id"]))
assert any(item.get("scope") == "项目一致性" and "工期" in item.get("title", "") for item in consistency_gate["revision_tasks"]), consistency_gate["revision_tasks"]
draft = update_draft(int(draft["id"]), draft["content"])
risky_draft = update_draft(int(draft["id"]), draft["content"] + "\n\n风险表述：确保中标。")
assert "确保中标" in risky_draft["content"], risky_draft["content"][-100:]
polish_report = scan_project_polish(int(tender["id"]))
assert any(item["search_text"] == "确保中标" for item in polish_report["suggestions"]), polish_report
replacement_preview = preview_replacement(int(tender["id"]), "确保中标", "确保技术标响应完整")
assert replacement_preview["changed_count"] >= 1, replacement_preview
replacement_result = apply_replacement(int(tender["id"]), "确保中标", "确保技术标响应完整", notes="冒烟测试风险承诺清理")
assert replacement_result["record"] and replacement_result["changed_count"] >= 1, replacement_result
assert "确保中标" not in get_draft(int(draft["id"]))["content"]
replacement_records = list_replacement_records(int(tender["id"]))
assert replacement_records and replacement_records[0]["changed_count"] >= 1, replacement_records
versions = list_draft_versions(int(draft["id"]))
assert len(versions) >= 3, versions
planned_draft = generate_from_plan(int(plans[0]["id"]))
assert planned_draft["id"], "planned draft not generated"
batch = generate_all_from_plan(int(tender["id"]))
assert batch["generated_count"] >= 1, batch
rich_document = generate_document_blocks(int(tender["id"]), regenerate=True)
assert rich_document["summary"]["tables"] >= 20, rich_document["summary"]
assert rich_document["summary"]["organization_charts"] >= 1, rich_document["summary"]
assert rich_document["summary"]["flow_charts"] >= 10, rich_document["summary"]
assert rich_document["summary"]["gantt_charts"] >= 1, rich_document["summary"]
image_settings = image_generation_settings()
assert not image_settings["configured"] and "IMAGE_MODEL" in image_settings["note"], image_settings
visual_plan = build_visual_plan(int(tender["id"]))
assert visual_plan["summary"]["technical_diagrams"] >= 1, visual_plan["summary"]
assert visual_plan["summary"]["ai_hybrid_plans"] >= 4, visual_plan["summary"]
assert visual_plan["summary"]["real_only_requirements"] == 2, visual_plan["summary"]
assert visual_plan["policy"]["real_material"], visual_plan["policy"]
technical_asset = next(item for item in list_visual_assets(int(tender["id"])) if item["visual_class"] == "technical_diagram")
scene_plan = visual_plan["ai_hybrid_plans"][0]
hybrid_result = generate_hybrid_visual(
    int(tender["id"]),
    scene_key=scene_plan["scene_key"],
    source_path=str(resolve_asset_path(technical_asset["file_path"])),
    section_title=scene_plan["section_title"],
)
hybrid_asset = hybrid_result["asset"]
assert hybrid_asset["visual_class"] == "ai_scene", hybrid_asset
assert not hybrid_asset["can_be_evidence"], hybrid_asset
assert hybrid_asset["review_status"] == "待复核" and hybrid_asset["disclaimer"], hybrid_asset
assert hybrid_asset["overlay"]["steps"], hybrid_asset
reviewed_hybrid = update_visual_asset_review(
    int(hybrid_asset["id"]), review_status="需修改", review_notes="冒烟测试：技术负责人复核后再使用。"
)
assert reviewed_hybrid["review_status"] == "需修改" and not reviewed_hybrid["can_be_evidence"], reviewed_hybrid
rich_blocks = list_document_blocks(int(tender["id"]))
assert any(item.get("asset", {}).get("visual_class") == "ai_scene" for item in rich_blocks), rich_blocks
editable_block = next(item for item in rich_blocks if item["block_type"] == "table")
updated_block = update_document_block(int(editable_block["id"]), {"caption": "冒烟测试专业表格"})
assert updated_block["caption"] == "冒烟测试专业表格", updated_block
manual_block = update_document_block(int(editable_block["id"]), {"caption": editable_block["caption"]})
assert manual_block["caption"] == editable_block["caption"], manual_block
status = sync_production_status(int(tender["id"]))
assert status["current_status"] == "待导出", status
workflow = build_workflow_status(int(tender["id"]))
assert workflow["summary"]["generated_plans"] >= workflow["summary"]["plans"] >= 1, workflow
saved = update_draft(int(draft["id"]), draft["content"] + "\n\n补充：本段为保存测试。")
assert "保存测试" in saved["content"]
versions = list_draft_versions(int(draft["id"]))
assert len(versions) >= 4, versions
assert any(str(item["origin"]).startswith("replacement:") for item in versions), versions
restored = restore_draft_version(int(draft["id"]), int(versions[-1]["id"]))
assert "保存测试" not in restored["content"], restored["content"][-100:]
loaded = get_draft(int(draft["id"]))
assert loaded["id"] == draft["id"]
assert loaded["generation"]["mode"] == "local_fallback", loaded["generation"]
drafts = list_drafts(int(tender["id"]))
assert any(item["id"] == draft["id"] for item in drafts)
assert any(item["id"] == draft["id"] and item["generation_mode"] == "local_fallback" for item in drafts), drafts
full_tender = get_tender(int(tender["id"]))
assert full_tender["drafts"], "draft list missing from tender detail"
case_asset = create_case_asset_from_draft(
    int(method_draft["id"]),
    {"tags": "医院类 / 施工工艺 / 高复用", "reusable_score": 5},
)
assert case_asset["section_title"] == "施工工艺及主要施工方法", case_asset
case_asset = update_case_asset(int(case_asset["id"]), {"summary": "医院类施工工艺章节，可复用医疗专项机电调试段落。"})
assert "医疗专项机电" in case_asset["summary"], case_asset
case_assets = list_case_assets(tender_id=int(tender["id"]))
assert any(item["id"] == case_asset["id"] for item in case_assets), case_assets
case_hits = search_case_assets("医疗专项机电", category="医院类", heading="施工工艺", limit=5)
assert any(item["source_type"] == "case_asset" for item in case_hits), case_hits
case_preview = preview_source(next(item for item in case_hits if item["source_type"] == "case_asset"))
assert case_preview["source_type"] == "case_asset" and "医疗专项机电" in case_preview["content"], case_preview
retrieved_with_assets = search_chunks("医疗专项机电", category="医院类", heading="施工工艺", limit=8)
assert any(item.get("source_type") == "case_asset" for item in retrieved_with_assets), retrieved_with_assets
case_assets_markdown = render_case_assets_markdown(int(tender["id"]))
assert "案例资产" in case_assets_markdown and "施工工艺及主要施工方法" in case_assets_markdown, case_assets_markdown
bulk_assets = create_case_assets_from_tender(
    int(tender["id"]),
    {"tags": "项目成稿 / 可复用", "reusable_score": 4, "source_status": "批量沉淀"},
)
assert bulk_assets["processed_count"] >= len(drafts), bulk_assets
assert bulk_assets["updated_count"] >= 1, bulk_assets
findings = review_tender(int(tender["id"]))
assert isinstance(findings, list)
coverage = build_coverage_report(int(tender["id"]))
assert coverage["summary"]["generated_sections"] >= 1, coverage
assert coverage["requirements"], "coverage requirements missing"
delivery_report = build_delivery_review(int(tender["id"]))
assert delivery_report["summary"]["generated_sections"] >= 1, delivery_report
assert delivery_report["summary"]["task_missing"] == [], delivery_report
assert "交付审查报告" in delivery_report["markdown"], delivery_report["markdown"][:200]
active_plan = next(item for item in list_section_plans(int(tender["id"])) if item.get("draft_id"))
active_draft_id = int(active_plan["draft_id"])
draft_with_revision_issue = update_draft(
    active_draft_id,
    get_draft(active_draft_id)["content"] + "\n\n修订台账测试风险：确保中标。",
)
assert "确保中标" in draft_with_revision_issue["content"], draft_with_revision_issue["content"][-120:]
revision_report = sync_revision_tasks(int(tender["id"]))
assert revision_report["summary"]["active"] >= 1 and revision_report["summary"]["high_open"] >= 1, revision_report
assert "修订任务台账" in render_revision_tasks_markdown(revision_report)
revision_item = next(item for item in revision_report["items"] if item["severity"] == "high")
revision_item = update_revision_task(
    int(revision_item["id"]),
    {"status": "处理中", "owner": "测试员", "notes": "冒烟测试正在处理风险承诺。"},
)
assert revision_item["status"] == "处理中" and revision_item["owner"] == "测试员", revision_item
cleanup_revision_issue = apply_replacement(int(tender["id"]), "确保中标", "确保技术标响应完整", notes="修订任务台账测试清理")
assert cleanup_revision_issue["changed_count"] >= 1, cleanup_revision_issue
revision_report = sync_revision_tasks(int(tender["id"]))
assert revision_report["summary"]["resolved_by_sync"] >= 1, revision_report["summary"]
revision_report = list_revision_tasks(int(tender["id"]))
assert any(not item["active"] and item["status"] == "已消除" for item in revision_report["items"]), revision_report
quality_gate = build_quality_gate(int(tender["id"]))
assert quality_gate["summary"]["total_sections"] >= 1, quality_gate
assert "materials_pending_required" in quality_gate["summary"], quality_gate
assert "质量门禁报告" in quality_gate["markdown"], quality_gate["markdown"][:200]
final_checklist = build_final_checklist(int(tender["id"]))
assert final_checklist["items"] and final_checklist["summary"]["required_total"] >= 1, final_checklist
updated_final_check = update_final_check_item(
    int(final_checklist["items"][0]["id"]),
    {"status": "已确认", "owner": "测试员", "notes": "冒烟测试确认。"},
)
assert updated_final_check["status"] == "已确认" and updated_final_check["owner"] == "测试员", updated_final_check
final_document = build_final_document(int(tender["id"]))
assert final_document["content"] and final_document["summary"]["generated_sections"] >= 1, final_document
assert "施工总体部署" in final_document["content"], final_document["content"][:500]
final_document = update_final_document(
    int(tender["id"]),
    {"status": "approved", "approved_by": "测试员", "notes": "冒烟测试定稿确认。"},
)
assert final_document["summary"]["approval_status"] == "approved", final_document
assert final_document["approval"]["approved_by"] == "测试员", final_document
assert "成稿确认报告" in render_final_document_markdown(int(tender["id"]))
readiness_before_package = build_production_readiness(int(tender["id"]))
assert readiness_before_package["summary"]["blockers"] >= 1, readiness_before_package["summary"]
assert any(item["key"] == "package" for item in readiness_before_package["blockers"]), readiness_before_package["blockers"]
assert "项目可交付性评估" in render_production_readiness_markdown(readiness_before_package)
print("export", flush=True)
exported = export_markdown(int(tender["id"]))
assert Path(exported["path"]).exists()
exported_text = Path(exported["path"]).read_text(encoding="utf-8")
assert "交付审查报告" in exported_text
assert "投标单位资料" in exported_text and "测试建设集团有限公司" in exported_text
assert "收款记录" in exported_text and "已收齐" in exported_text
exported_docx = export_docx(int(tender["id"]))
assert Path(exported_docx["path"]).exists()
import docx  # type: ignore

inline_doc = docx.Document()
_append_markdown_line(inline_doc, "**重点**与*说明*")
inline_paragraph = inline_doc.paragraphs[-1]
assert inline_paragraph.text == "重点与说明", inline_paragraph.text
assert inline_paragraph.runs[0].bold and inline_paragraph.runs[-1].italic, [run.text for run in inline_paragraph.runs]
table_doc = docx.Document()
_append_markdown_content(
    table_doc,
    "| 施工阶段 | 主要工作内容 | 进度控制重点 |\n|---|---|---|\n| 施工准备阶段 | 图纸会审、现场复测 | 完成开工条件核查 |",
)
assert len(table_doc.tables) == 1, len(table_doc.tables)
assert table_doc.tables[0].cell(1, 0).text == "施工准备阶段", table_doc.tables[0].cell(1, 0).text
assert not any("|---|" in paragraph.text for paragraph in table_doc.paragraphs), [paragraph.text for paragraph in table_doc.paragraphs]

docx_text = "\n".join(paragraph.text for paragraph in docx.Document(exported_docx["path"]).paragraphs)
rich_docx = docx.Document(exported_docx["path"])
assert len(rich_docx.tables) >= 20, len(rich_docx.tables)
assert len(rich_docx.inline_shapes) >= 10, len(rich_docx.inline_shapes)
assert "桩基础施工场景示意图" in docx_text, docx_text[-2000:]
assert "AI生成" not in docx_text and "历史资料示例图" not in docx_text, docx_text[-2000:]
assert "测试医院综合楼工程技术标" in docx_text and "目录" in docx_text and "测试投标单位" in docx_text, docx_text[:500]
assert "投标单位资料" not in docx_text and "收款记录" not in docx_text and "响应矩阵" not in docx_text, docx_text[:1200]
assert "参考来源摘要" not in docx_text and "需人工确认" not in docx_text, docx_text[-2000:]
package = export_package(int(tender["id"]))
assert Path(package["path"]).exists()
assert package["delivery_record"]["status"] == "已导出", package
readiness_after_package = build_production_readiness(int(tender["id"]))
assert readiness_after_package["summary"]["requirements"] >= 1, readiness_after_package["summary"]
assert readiness_after_package["summary"]["generated_sections"] >= 1, readiness_after_package["summary"]
assert readiness_after_package["summary"]["package_readiness"] in {"ready", "warning"}, readiness_after_package["summary"]
package_validation = validate_latest_package(int(tender["id"]))
assert package_validation["summary"]["package_exists"], package_validation["summary"]
assert package_validation["summary"]["zip_readable"], package_validation["summary"]
assert package_validation["summary"]["core_missing"] == 0, package_validation["summary"]
assert package_validation["summary"]["invalid_json"] == 0, package_validation["summary"]
assert "交付包清单核验报告" in package_validation["markdown"], package_validation["markdown"][:200]
delivery_release = build_delivery_release(int(tender["id"]))
assert delivery_release["summary"]["release_status"] in {"ready", "conditional", "blocked"}, delivery_release["summary"]
assert delivery_release["checks"] and "交付放行单" in delivery_release["markdown"], delivery_release["markdown"][:200]
command_center_after_package = build_command_center(int(tender["id"]))
assert command_center_after_package["package"]["summary"]["package_exists"], command_center_after_package["package"]
assert command_center_after_package["headline"]["next_step_title"], command_center_after_package["headline"]
retrospective = build_project_retrospective(int(tender["id"]))
assert retrospective["summary"]["status"] == "待复盘", retrospective
assert retrospective["retrospective"]["actual_price"] == "1200 元", retrospective
closure = build_closure_confirmation(int(tender["id"]))
assert closure["summary"]["closure_status"] == "待客户确认", closure
status = sync_production_status(int(tender["id"]))
assert status["current_status"] == "待客户确认", status
assistant = build_delivery_assistant(int(tender["id"]))
assert assistant["summary"]["has_package"], assistant
assert "测试客户" in assistant["customer_message"], assistant["customer_message"]
instruction = render_delivery_instruction(int(tender["id"]))
assert "交付说明" in instruction and "客户发货说明" in instruction, instruction[:300]
delivery_records = list_delivery_records(int(tender["id"]))
assert len(delivery_records) == 1, delivery_records
delivered = update_delivery_record(
    int(delivery_records[0]["id"]),
    {
        "status": "已交付",
        "delivery_channel": "内部验证",
        "recipient": "测试客户",
        "notes": "冒烟测试标记已交付。",
    },
)
assert delivered["status"] == "已交付" and delivered["delivered_at"], delivered
status = sync_production_status(int(tender["id"]))
assert status["current_status"] == "已交付", status
workflow = build_workflow_status(int(tender["id"]))
assert workflow["summary"]["delivered_records"] == 1, workflow
dashboard = build_order_dashboard()
assert dashboard["summary"]["delivered"] == 1, dashboard
closure = build_closure_confirmation(int(tender["id"]), {"reply_deadline": "2026-08-11 12:00"})
assert closure["summary"]["can_record_closure"], closure
assert "最终交付版本" in closure["customer_message"], closure["customer_message"]
feedback = create_feedback_item(
    int(tender["id"]),
    {
        "feedback_text": "客户要求补充医疗专项机电调试计划。",
        "related_section": "施工工艺及主要施工方法",
        "priority": "high",
        "action_plan": "补充医疗专项机电调试、联动测试和成品保护措施。",
    },
)
assert feedback["status"] == "待处理", feedback
status = sync_production_status(int(tender["id"]))
assert status["current_status"] == "返工处理", status
summary = feedback_summary(int(tender["id"]))
assert summary["open"] == 1, summary
closure = build_closure_confirmation(int(tender["id"]))
assert not closure["summary"]["can_close"] and closure["summary"]["closure_status"] == "需返工处理", closure
workflow = build_workflow_status(int(tender["id"]))
assert workflow["summary"]["feedback"]["open"] == 1, workflow
dashboard = build_order_dashboard()
assert dashboard["summary"]["feedback_open"] == 1, dashboard
assistant = build_delivery_assistant(int(tender["id"]))
assert "客户反馈未解决" in "\n".join(assistant["actions"]), assistant
feedback_rework = build_feedback_rework(int(tender["id"]))
assert feedback_rework["summary"]["open_feedback"] == 1, feedback_rework["summary"]
assert feedback_rework["summary"]["affected_sections"] >= 1, feedback_rework
assert "医疗专项机电" in feedback_rework["customer_message"], feedback_rework["customer_message"]
applied_rework = apply_feedback_rework_plan(int(tender["id"]), {"owner": "测试员"})
assert applied_rework["summary"]["applied_updates"] == 1, applied_rework["summary"]
assert applied_rework["items"][0]["status"] == "处理中", applied_rework["items"]
executed_rework = execute_feedback_rework(int(tender["id"]), {"owner": "测试员"})
assert executed_rework["execution"]["updated_drafts"] >= 1, executed_rework["execution"]
assert executed_rework["items"][0]["status"] == "需复核", executed_rework["items"]
rework_draft_id = int(executed_rework["execution"]["touched"][0]["draft_id"])
rework_draft = get_draft(rework_draft_id)
assert "客户反馈响应补充" in rework_draft["content"], rework_draft["content"][-800:]
rework_versions = list_draft_versions(rework_draft_id)
assert any("feedback_rework" in item["origin"] for item in rework_versions), rework_versions
resolved_feedback = update_feedback_item(int(feedback["id"]), {"status": "已解决"})
assert resolved_feedback["status"] == "已解决" and resolved_feedback["resolved_at"], resolved_feedback
status = sync_production_status(int(tender["id"]))
assert status["current_status"] == "已交付", status
closure = build_closure_confirmation(int(tender["id"]))
assert closure["summary"]["can_record_closure"], closure
closure = record_closure_confirmation(
    int(tender["id"]),
    {
        "confirmed_by": "测试客户",
        "reply_deadline": "2026-08-11 12:00",
        "confirmation_note": "客户确认最终交付版本无误。",
    },
)
assert closure["summary"]["closure_status"] == "已结案", closure
assert closure["latest_closure"]["status"] == "客户已确认", closure
status = sync_production_status(int(tender["id"]))
assert status["current_status"] == "已结案", status
retrospective = update_project_retrospective(
    int(tender["id"]),
    {
        "status": "已复盘",
        "actual_price": "1200 元",
        "actual_cost": "200 元",
        "work_hours": 4,
        "revision_count": 1,
        "satisfaction": "高",
        "risk_level": "低",
        "reusable_score": 5,
        "industry_tags": "医院类 / 施工工艺 / 高复用",
        "reusable_assets": "施工工艺章节、医疗专项机电调试段落、交付话术。",
        "lessons": "急单需先确认评分办法和图纸清单。",
        "next_action": "沉淀医院类施工工艺模板。",
    },
)
assert retrospective["summary"]["status"] == "已复盘", retrospective
assert retrospective["summary"]["gross_margin_value"] == 1000.0, retrospective
assert retrospective["summary"]["can_archive_case"], retrospective
retrospective_dashboard = build_retrospective_dashboard()
assert retrospective_dashboard["summary"]["reviewed"] == 1, retrospective_dashboard
assert retrospective_dashboard["summary"]["total_revenue"] >= 1200, retrospective_dashboard
workflow = build_workflow_status(int(tender["id"]))
assert workflow["summary"]["retrospective"]["status"] == "已复盘", workflow
all_status = sync_all_production_status()
assert all_status["total"] == 1, all_status
summary = feedback_summary(int(tender["id"]))
assert summary["open"] == 0 and summary["resolved"] == 1, summary
dashboard = build_order_dashboard()
assert dashboard["summary"]["feedback_open"] == 0, dashboard
assert dashboard["summary"]["closed"] == 1, dashboard
timeline = build_project_timeline(int(tender["id"]))
timeline_kinds = {item["kind"] for item in timeline["events"]}
assert {
    "tender",
    "task",
    "quote",
    "payment",
    "communication",
    "plan",
    "draft",
    "polish",
    "final_document",
    "delivery",
    "feedback",
    "closure",
    "retrospective",
    "case_asset",
} <= timeline_kinds, timeline
assert timeline["summary"]["payment_events"] >= 1 and timeline["summary"]["draft_events"] >= 1, timeline["summary"]
assert "项目生产履历" in render_project_timeline_markdown(timeline)
source_audit = build_source_audit(int(tender["id"]))
assert source_audit["summary"]["drafts"] >= 1 and source_audit["summary"]["total_citations"] >= 1, source_audit
assert source_audit["summary"]["invalid_citations"] == 0, source_audit["summary"]
assert source_audit["summary"]["kb_citations"] >= 1, source_audit["summary"]
assert "引用来源审计报告" in render_source_audit_markdown(source_audit)
with zipfile.ZipFile(package["path"]) as archive:
    names = set(archive.namelist())
    package_starter = json.loads(archive.read("生产启动包.json").decode("utf-8"))
    package_wizard = json.loads(archive.read("新单生产向导.json").decode("utf-8"))
    package_channel_ops = json.loads(archive.read("渠道运营助手.json").decode("utf-8"))
    package_timeline = json.loads(archive.read("项目生产履历.json").decode("utf-8"))
    package_plan_audit = json.loads(archive.read("目录完整性审计.json").decode("utf-8"))
    package_response_matrix = json.loads(archive.read("响应矩阵挂接报告.json").decode("utf-8"))
    package_source_audit = json.loads(archive.read("引用来源审计.json").decode("utf-8"))
    package_revision_tasks = json.loads(archive.read("修订任务台账.json").decode("utf-8"))
    package_feedback_rework = json.loads(archive.read("客户反馈返工处理单.json").decode("utf-8"))
    package_readiness = json.loads(archive.read("项目可交付性评估.json").decode("utf-8"))
    package_validation_json = json.loads(archive.read("交付包清单核验.json").decode("utf-8"))
    package_delivery_release = json.loads(archive.read("交付放行单.json").decode("utf-8"))
assert {"技术标初稿.docx", "技术标初稿.md", "生产任务.json", "项目资料.json", "响应覆盖报告.json", "生产流程状态.json", "README.txt"} <= names, names
assert {"项目总览.md", "项目总览.json"} <= names, names
assert {"项目生产履历.md", "项目生产履历.json"} <= names, names
assert {"生产启动包.md", "生产启动包.json"} <= names, names
assert {"新单生产向导.md", "新单生产向导.json"} <= names, names
assert {"渠道运营助手.md", "渠道运营助手.json"} <= names, names
assert {"目录完整性审计.md", "目录完整性审计.json"} <= names, names
assert {"响应矩阵挂接报告.md", "响应矩阵挂接报告.json"} <= names, names
assert {"引用来源审计.md", "引用来源审计.json"} <= names, names
assert {"客户反馈返工处理单.md", "客户反馈返工处理单.json"} <= names, names
assert "total_feedback" in package_feedback_rework["summary"], package_feedback_rework["summary"]
assert {"成稿格式设置.md", "成稿格式设置.json"} <= names, names
assert {"投标单位资料.md", "投标单位资料.json"} <= names, names
assert {"投标响应策略.md", "投标响应策略.json"} <= names, names
assert {"交付审查报告.md", "交付审查报告.json", "交付说明.md"} <= names, names
assert {"接单评估.md", "接单评估.json"} <= names, names
assert {"客户沟通记录.md", "客户沟通记录.json"} <= names, names
assert {"报价测算.md", "报价测算.json"} <= names, names
assert {"收款记录.md", "收款记录.json"} <= names, names
assert {"订单确认单.md", "订单确认单.json"} <= names, names
assert {"质量门禁报告.md", "质量门禁报告.json"} <= names, names
assert {"修订任务台账.md", "修订任务台账.json"} <= names, names
assert {"最终核对清单.md", "最终核对清单.json"} <= names, names
assert {"成稿确认报告.md", "成稿确认报告.json"} <= names, names
assert {"项目化校正记录.md", "项目化校正记录.json"} <= names, names
assert {"资料清单.md", "资料清单.json"} <= names, names
assert {"结案确认单.md", "结案确认单.json"} <= names, names
assert {"项目复盘.md", "项目复盘.json"} <= names, names
assert {"案例资产.md", "案例资产.json"} <= names, names
assert {"项目可交付性评估.md", "项目可交付性评估.json"} <= names, names
assert {"交付放行单.md", "交付放行单.json"} <= names, names
assert {"交付包清单核验.md", "交付包清单核验.json"} <= names, names
assert package_starter["summary"]["strategy"] == "标书生产工具优先", package_starter["summary"]
assert "不自动登录平台" in package_starter["channel_policy"]["automation_boundary"], package_starter["channel_policy"]
assert package_wizard["summary"]["title"] == "新单生产向导", package_wizard["summary"]
assert package_wizard["steps"] and package_wizard["copyables"]["customer_reply"], package_wizard
assert package_channel_ops["summary"]["title"] == "渠道运营助手", package_channel_ops["summary"]
assert package_channel_ops["scripts"] and "不自动登录平台" in package_channel_ops["policy"]["automation_boundary"], package_channel_ops
assert package_timeline["summary"]["payment_events"] >= 1 and package_timeline["summary"]["draft_events"] >= 1, package_timeline["summary"]
assert package_plan_audit["summary"]["construction_method_present"], package_plan_audit["summary"]
assert package_plan_audit["summary"]["missing_standard"] == 0, package_plan_audit["summary"]
assert package_response_matrix["summary"]["requirements"] >= 1, package_response_matrix["summary"]
assert package_response_matrix["summary"]["high_unplanned_requirements"] == 0, package_response_matrix["summary"]
assert package_source_audit["summary"]["total_citations"] >= 1 and package_source_audit["summary"]["invalid_citations"] == 0, package_source_audit["summary"]
assert package_revision_tasks["summary"]["total"] >= 1, package_revision_tasks["summary"]
assert any(not item["active"] and item["status"] == "已消除" for item in package_revision_tasks["items"]), package_revision_tasks
assert package_readiness["summary"]["requirements"] >= 1, package_readiness["summary"]
assert package_readiness["summary"]["generated_sections"] >= 1, package_readiness["summary"]
assert package_validation_json["summary"]["core_missing"] == 0, package_validation_json["summary"]
assert package_validation_json["summary"]["invalid_json"] == 0, package_validation_json["summary"]
assert package_delivery_release["summary"]["release_status"] in {"ready", "conditional", "blocked"}, package_delivery_release["summary"]
assert package_delivery_release["checks"], package_delivery_release
client_package = export_client_package(int(tender["id"]))
assert Path(client_package["path"]).exists() and client_package["format"] == "client_zip", client_package
assert client_package["delivery_record"]["package_format"] == "client_zip", client_package["delivery_record"]
with zipfile.ZipFile(client_package["path"]) as archive:
    client_names = set(archive.namelist())
    client_md = archive.read("技术标初稿_客户版.md").decode("utf-8")
assert {"技术标初稿_客户版.docx", "技术标初稿_客户版.md", "客户发货说明.md", "客户文件清单.txt", "README.txt"} <= client_names, client_names
assert not {"报价测算.json", "收款记录.json", "引用来源审计.json", "质量门禁报告.json", "交付放行单.json"} & client_names, client_names
assert "收款记录" not in client_md and "引用来源" not in client_md and "交付审查报告" not in client_md, client_md[:1000]
client_package_validation = validate_latest_client_package(int(tender["id"]))
assert client_package_validation["summary"]["readiness"] == "ready", client_package_validation["summary"]
assert client_package_validation["summary"]["forbidden_files"] == 0, client_package_validation["summary"]
assert client_package_validation["summary"]["content_findings"] == 0, client_package_validation
client_delivery_confirmation = build_client_delivery_confirmation(int(tender["id"]))
assert client_delivery_confirmation["summary"]["has_client_package"], client_delivery_confirmation["summary"]
assert client_delivery_confirmation["summary"]["client_package_ready"], client_delivery_confirmation["summary"]
assert client_delivery_confirmation["summary"]["can_send"], client_delivery_confirmation
assert "客户版 Word 初稿" in client_delivery_confirmation["customer_message"], client_delivery_confirmation["customer_message"]
client_delivery_confirmed = confirm_client_delivery(
    int(tender["id"]),
    {
        "delivery_channel": "内部验证",
        "recipient": "测试客户",
        "confirmation_note": "测试确认客户包已人工发出。",
        "customer_message": client_delivery_confirmation["customer_message"],
    },
)
assert client_delivery_confirmed["summary"]["delivery_confirmed"], client_delivery_confirmed["summary"]
assert client_delivery_confirmed["summary"]["latest_client_delivery_status"] == "已交付", client_delivery_confirmed["summary"]
latest_client_record = client_delivery_confirmed["latest_client_package"]
assert latest_client_record["package_format"] == "client_zip" and latest_client_record["delivered_at"], latest_client_record
release_after_client_package = build_delivery_release(int(tender["id"]))
assert any(item["key"] == "client_package_validation" for item in release_after_client_package["checks"]), release_after_client_package["checks"]
internal_validation_after_client_package = validate_latest_package(int(tender["id"]))
assert internal_validation_after_client_package["summary"]["core_missing"] == 0, internal_validation_after_client_package["summary"]
for output in (
    client_package["path"],
    client_package["components"]["client_markdown"],
    client_package["components"]["client_docx"],
):
    Path(output).unlink(missing_ok=True)
for output in (exported["path"], exported_docx["path"], package["path"]):
    Path(output).unlink(missing_ok=True)
deleted = delete_tender(int(tender["id"]))
assert deleted["deleted"], deleted

intake_created = create_intake_tender(
    {
        "customer_message": "闲鱼客户咨询：项目名称：接单草稿医院技术标。需要技术标 Word 初稿，明天前交，预算 900 元，招标文件稍后发。",
        "customer_name": "接单草稿客户",
        "source_platform": "闲鱼",
        "budget_expectation": "900 元",
        "deliverable_format": "DOCX + ZIP 交付包",
        "industry": "医院类",
    }
)
intake_tender = intake_created["tender"]
assert intake_tender["id"], intake_created
assert intake_created["task"]["customer_name"] == "接单草稿客户", intake_created["task"]
assert intake_created["intake"]["acceptance"]["decision"] in {"暂缓接单", "谨慎承接", "可承接"}, intake_created["intake"]
intake_communications = list_communications(int(intake_tender["id"]))
assert intake_communications and intake_communications[0]["stage"] == "询盘", intake_communications
source_updated = update_tender_source(
    int(intake_tender["id"]),
    text="""
项目名称：接单草稿医院技术标
技术标评分要求：施工总体部署、施工工艺及主要施工方法、质量保证措施、安全文明施工、施工进度计划。
主要施工工艺要求：测量放线、主体结构、防水屋面、装饰装修、医疗专项机电安装必须完整响应。
废标风险：投标文件不得出现与本项目无关的历史项目名称。
""",
    name="接单草稿医院技术标",
    industry="医院类",
)
assert source_updated["tender"]["id"] == intake_tender["id"], source_updated
assert len(source_updated["requirements"]) >= 3, source_updated["requirements"]
assert source_updated["source_update"]["reset_plan"], source_updated["source_update"]
deleted = delete_tender(int(intake_tender["id"]))
assert deleted["deleted"], deleted

pipeline_tender = create_tender(
    "流水线测试医院项目",
    """
项目名称：流水线测试医院项目
技术标评分要求：施工总体部署、施工工艺及主要施工方法、质量保证措施、安全文明施工、施工进度计划。
主要施工工艺要求：测量放线、主体结构、装饰装修、机电安装、医疗专项机电必须响应。
""",
    industry="医院类",
)
update_production_task(
    int(pipeline_tender["id"]),
    {
        "customer_name": "流水线客户",
        "source_platform": "闲鱼",
        "deadline": "2026-08-12 18:00",
        "deliverable_format": "DOCX + ZIP 交付包",
    },
)
update_project_profile(
    int(pipeline_tender["id"]),
    {
        "project_type": "医院综合楼",
        "structure_type": "框架结构",
        "building_area": "约 18000 平方米",
        "quality_target": "合格",
        "safety_target": "无重伤事故",
        "contract_scope": "土建、装饰、机电安装",
    },
)
pipeline = run_production_pipeline(int(pipeline_tender["id"]), {"export_package": True})
assert pipeline["status"]["current_status"] == "待客户确认", pipeline
assert pipeline["quality_gate"]["summary"]["total_sections"] >= 1, pipeline
assert pipeline["package"]["path"] and Path(pipeline["package"]["path"]).exists(), pipeline
pipeline_validation = validate_latest_package(int(pipeline_tender["id"]))
assert pipeline_validation["summary"]["zip_readable"] and pipeline_validation["summary"]["core_missing"] == 0, pipeline_validation["summary"]
pipeline_readiness = build_production_readiness(int(pipeline_tender["id"]))
assert pipeline_readiness["summary"]["requirements"] >= 1 and pipeline_readiness["summary"]["package_readiness"] in {"ready", "warning"}, pipeline_readiness["summary"]
client_delivery_preparation = prepare_client_delivery(int(pipeline_tender["id"]), {"export_internal_package": False})
assert client_delivery_preparation["summary"]["ready_to_send"], client_delivery_preparation["summary"]
assert client_delivery_preparation["summary"]["client_package_ready"], client_delivery_preparation["summary"]
assert Path(client_delivery_preparation["client_package"]["path"]).exists(), client_delivery_preparation["client_package"]
assert client_delivery_preparation["client_package_validation"]["summary"]["readiness"] == "ready", client_delivery_preparation["client_package_validation"]["summary"]
assert client_delivery_preparation["client_delivery_confirmation"]["summary"]["can_send"], client_delivery_preparation["client_delivery_confirmation"]["summary"]
with zipfile.ZipFile(pipeline["package"]["path"]) as archive:
    pipeline_names = set(archive.namelist())
assert {"技术标初稿.docx", "技术标初稿.md", "项目总览.md", "项目生产履历.md", "项目生产履历.json", "生产启动包.md", "生产启动包.json", "新单生产向导.md", "新单生产向导.json", "渠道运营助手.md", "渠道运营助手.json", "引用来源审计.md", "引用来源审计.json", "成稿格式设置.md", "投标响应策略.md", "接单评估.md", "客户沟通记录.md", "报价测算.md", "资料清单.md", "质量门禁报告.md", "最终核对清单.md", "成稿确认报告.md", "项目化校正记录.md", "交付说明.md", "项目可交付性评估.md", "项目可交付性评估.json", "交付放行单.md", "交付放行单.json", "结案确认单.md", "项目复盘.md", "案例资产.md", "案例资产.json", "交付包清单核验.md", "交付包清单核验.json"} <= pipeline_names, pipeline_names
for output in (
    pipeline["package"]["path"],
    pipeline["package"]["components"]["markdown"],
    pipeline["package"]["components"]["docx"],
    client_delivery_preparation["client_package"]["path"],
    client_delivery_preparation["client_package"]["components"]["client_markdown"],
    client_delivery_preparation["client_package"]["components"]["client_docx"],
):
    Path(output).unlink(missing_ok=True)
deleted = delete_tender(int(pipeline_tender["id"]))
assert deleted["deleted"], deleted

conn = connect()
final_stats = stats(conn)
assert final_stats["tenders"] == 0, final_stats
print(final_stats)
conn.close()
