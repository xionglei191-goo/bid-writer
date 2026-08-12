from __future__ import annotations

import base64
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from bid_writer.artifact_audit import audit_docx_path, audit_text, audit_zip_path


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


class ArtifactAuditTest(unittest.TestCase):
    def test_scans_body_table_header_and_footer_placeholders(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "audit.docx"
            document = Document()
            document.add_paragraph("正文【待确认：项目经理】")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "设备[TBD]"
            document.sections[0].header.paragraphs[0].text = "[投标单位]"
            document.sections[0].footer.paragraphs[0].text = "<联系电话>"
            document.save(path)

            report = audit_docx_path(path)

        self.assertFalse(report["ready"])
        self.assertEqual(report["counts"]["placeholder"], 4)
        sources = {item["source"] for item in report["findings"]}
        self.assertTrue(any(source.endswith(":body") for source in sources))
        self.assertTrue(any(":header:" in source for source in sources))
        self.assertTrue(any(":footer:" in source for source in sources))

    def test_legal_url_is_not_a_placeholder(self) -> None:
        report = audit_text("参考[官方规范](https://example.com/spec?q=TODO)及 http://127.0.0.1:8876/api/health/ready")
        self.assertTrue(report["ready"], report)

    def test_blocks_test_data_internal_paths_and_unapproved_visuals(self) -> None:
        report = audit_text(
            "默认投标单位；来源路径 F:\\internal\\bid_writer_app\\data\\draft.json；"
            "历史资料示例图（使用前人工核对适用性）；AI生成施工示意图"
        )
        self.assertGreaterEqual(report["counts"]["test_data"], 1)
        self.assertGreaterEqual(report["counts"]["internal_content"], 1)
        self.assertGreaterEqual(report["counts"]["unapproved_visual"], 2)

    def test_blocks_unapproved_drawing_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image_path = root / "pixel.png"
            image_path.write_bytes(PNG_1X1)
            docx_path = root / "drawing.docx"
            document = Document()
            document.add_picture(str(image_path))
            document.inline_shapes[0]._inline.docPr.set("descr", "AI生成施工示意图")
            document.save(docx_path)
            report = audit_docx_path(docx_path)
        self.assertFalse(report["ready"])
        self.assertGreaterEqual(report["counts"]["unapproved_visual"], 1)

    def test_review_package_whitelist_and_embedded_content(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            package = Path(temp_dir) / "review.zip"
            with zipfile.ZipFile(package, "w") as archive:
                archive.writestr("送审说明.md", "请复核后签章。")
                archive.writestr("文件清单.txt", "送审说明.md")
                archive.writestr("内部审计资料.json", "{}")
            report = audit_zip_path(
                package,
                allowed_names={"送审说明.md", "文件清单.txt"},
                require_review_package_types=True,
            )
        self.assertFalse(report["ready"])
        self.assertGreaterEqual(report["counts"]["package_whitelist"], 1)


if __name__ == "__main__":
    unittest.main()
