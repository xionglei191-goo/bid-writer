from __future__ import annotations

import tempfile
import unittest
import os
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from docx import Document

from bid_writer.artifact_audit import audit_docx_path
from bid_writer_v2.production.service import ProductionService


class DeliveryCoverTest(unittest.TestCase):
    """Exercise the renderer without constructing an app or opening any database."""

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="delivery-cover-")
        self.path = Path(self.temp.name) / "cover.docx"
        self.service = object.__new__(ProductionService)
        self.project = {
            "name": "医院门诊综合楼施工组织设计",
            "profile": {
                "bidder_name": "华建建设工程有限公司",
                "authorized_signatory": "陈工",
                "professional_reviewer": "李工",
                "delivery_confirmation": {"confirmed_at": "2026-09-07T23:50:24+00:00"},
                "seal_requirements": "来源路径仅用于内部办理说明",
            },
            "sections": [{"title": "施工总体部署", "draft": {"content": "# 施工总体部署\n\n施工前组织现场复核与图纸会审，安排技术交底并形成过程记录。"}}],
        }

    def tearDown(self) -> None:
        self.temp.cleanup()

    def render(self, mode: str) -> tuple[Document, str]:
        self.service._write_docx(self.project, self.path, mode=mode)
        document = Document(self.path)
        return document, "\n".join(paragraph.text for paragraph in document.paragraphs)

    def test_formal_cover_uses_supplied_identity_and_actual_confirmation_date(self) -> None:
        document, text = self.render("formal")
        self.assertIn("正式版", text)
        self.assertNotIn("专业送审版", text)
        self.assertIn("投标单位：华建建设工程有限公司", text)
        self.assertIn("授权签字人：陈工", text)
        self.assertIn("专业复核人：李工", text)
        self.assertIn("定稿日期：2026年09月08日", text)
        self.assertNotIn("来源路径", text)
        self.assertNotIn("盖章完成", text)
        self.assertEqual(len(document.inline_shapes), 0)
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertEqual(document.styles["Title"]._element.xpath("./w:pPr/w:pBdr"), [])
        self.assertTrue(audit_docx_path(self.path)["ready"])

    def test_review_cover_is_explicitly_review_without_finalization_date(self) -> None:
        _, text = self.render("review")
        self.assertIn("专业送审版", text)
        self.assertNotIn("正式版", text)
        self.assertNotIn("定稿日期", text)
        self.assertTrue(audit_docx_path(self.path)["ready"])

    def test_missing_identity_does_not_invent_people_signatures_or_dates(self) -> None:
        self.project["profile"] = {}
        document, text = self.render("review")
        for label in ("投标单位", "授权签字人", "专业复核人", "定稿日期", "编制日期", "技术负责人"):
            self.assertNotIn(label, text)
        self.assertEqual(len(document.inline_shapes), 0)
        self.assertTrue(audit_docx_path(self.path)["ready"])

    def test_existing_identity_aliases_are_rendered_and_heading_colors_are_black(self) -> None:
        self.project["profile"] = {"tenderer_name": "华建建设工程有限公司", "signatory": "陈工", "reviewed_by": "李工", "document_date": "2026年09月08日"}
        document, text = self.render("review")
        self.assertIn("投标单位：华建建设工程有限公司", text)
        self.assertIn("授权签字人：陈工", text)
        self.assertIn("专业复核人：李工", text)
        self.assertIn("编制日期：2026年09月08日", text)
        for style in ("Title", "Subtitle", "Heading 1", "Heading 3", "Heading 9"):
            self.assertEqual(str(document.styles[style].font.color.rgb), "000000")

    def test_unsupported_mode_cannot_create_an_ambiguous_cover(self) -> None:
        with self.assertRaisesRegex(ValueError, "封面模式"):
            self.service._write_docx(self.project, self.path, mode="unknown")
        self.assertFalse(self.path.exists())

    def test_docx_contains_actual_cached_titles_in_a_native_toc_field(self) -> None:
        self.project["sections"].append({"title": "施工准备", "draft": {"content": "# 施工准备\n\n开展现场交底并形成检查记录。"}})
        document, _ = self.render("review")
        paragraphs = document.paragraphs
        contents_start = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text == "目录")
        body_start = next(index for index, paragraph in enumerate(paragraphs) if paragraph.style.name == "Heading 1")
        cached = "\n".join(paragraph.text for paragraph in paragraphs[contents_start:body_start])
        self.assertIn("施工总体部署", cached)
        self.assertIn("施工准备", cached)
        self.assertIn("w:instrText", document._element.xml)
        self.assertIn('w:fldCharType="separate"', document._element.xml)
        self.assertNotIn('w:instr="TOC', document._element.xml)

    @unittest.skipUnless(os.environ.get("BID_WRITER_TEST_LIBREOFFICE") == "1", "Requires the application container LibreOffice and python3-uno")
    def test_pdf_toc_uses_actual_pages_after_a_long_chapter(self) -> None:
        from pypdf import PdfReader

        self.project["sections"][0]["draft"]["content"] += "\n\n" + "\n\n".join(
            "施工期间按专业划分工作任务，做好作业衔接、现场管理和资料归档。" for _ in range(90)
        )
        self.project["sections"].append({"title": "施工准备", "draft": {"content": "# 施工准备\n\n开展现场交底并形成检查记录。"}})
        self.render("formal")
        pdf = self.service._convert_pdf(self.path)
        pages = [page.extract_text() for page in PdfReader(pdf).pages]
        actual_page = next(index for index, text in enumerate(pages, 1) if text.startswith("施工准备\n"))
        self.assertGreater(actual_page, 4)
        toc = next(text for text in pages if text.startswith("目录\n"))
        number = re.search(r"施工准备\.+\s*(\d+)", toc)
        self.assertIsNotNone(number)
        self.assertEqual(int(number.group(1)), actual_page)

    @unittest.skipUnless(os.environ.get("BID_WRITER_TEST_LIBREOFFICE") == "1", "Requires the application container LibreOffice and python3-uno")
    def test_parallel_pdf_exports_do_not_share_an_office_profile_or_pipe(self) -> None:
        from pypdf import PdfReader

        review_path = self.path.with_name("review.docx")
        formal_path = self.path.with_name("formal.docx")
        self.service._write_docx(self.project, review_path, mode="review")
        self.service._write_docx(self.project, formal_path, mode="formal")
        with ThreadPoolExecutor(max_workers=2) as pool:
            review_pdf, formal_pdf = list(pool.map(self.service._convert_pdf, [review_path, formal_path]))
        self.assertIn("专业送审版", PdfReader(review_pdf).pages[0].extract_text())
        self.assertIn("正式版", PdfReader(formal_pdf).pages[0].extract_text())


if __name__ == "__main__":
    unittest.main()
