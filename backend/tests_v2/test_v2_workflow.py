from __future__ import annotations

import ast
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from bid_writer_v2.app import create_app
from bid_writer_v2.database import Database
from bid_writer_v2.knowledge.service import KnowledgeService
from bid_writer_v2.production.service import ProductionService
from bid_writer_v2.settings import Settings


def build_settings(root: Path) -> Settings:
    app_root = root / "03_生产系统" / "bid_writer_app"
    data_root = app_root / "data"
    delivery_root = root / "04_交付与报告"
    return Settings(
        workspace_root=root,
        app_root=app_root,
        raw_root=root / "01_原始标书库",
        knowledge_root=root / "02_知识库",
        delivery_root=delivery_root,
        data_root=data_root,
        db_path=data_root / "test.sqlite3",
        upload_root=data_root / "uploads",
        cache_root=data_root / "cache",
        export_root=delivery_root / "项目交付",
        qa_root=delivery_root / "质量验收",
        operations_enabled=False,
    )


def create_source_docx(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("主要施工方案与技术措施", level=1)
    document.add_paragraph("施工前完成技术交底、图纸会审、现场复核和作业条件确认。施工过程中执行样板引路、旁站检查、实测实量和安全巡检，施工进度、劳动力、材料、机械和作业面统一协调。各工序完成后按检验批组织验收，发现偏差立即整改并形成闭环记录，确保质量、安全、进度和环境管理目标同步实现。")
    document.add_heading("混凝土施工工艺", level=2)
    document.add_paragraph("依次完成施工准备、测量放线、模板检查、钢筋隐蔽验收、混凝土运输、分层浇筑、机械振捣、表面收整、保温保湿养护和成品保护。浇筑前核对配合比、坍落度和预留预埋条件，浇筑过程中控制分层厚度和间歇时间，浇筑完成后按规范留置试件并记录养护条件。")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "施工工序"
    table.cell(0, 1).text = "控制点"
    table.cell(0, 2).text = "检查方式"
    table.cell(1, 0).text = "浇筑"
    table.cell(1, 1).text = "连续施工"
    table.cell(1, 2).text = "旁站检查"
    document.save(path)


class V2WorkflowTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.settings = build_settings(self.root)
        self.settings.ensure_directories()
        self.db = Database(self.settings.db_path)
        self.db.migrate()
        self.knowledge = KnowledgeService(self.db, self.settings)
        self.production = ProductionService(self.db, self.settings, self.knowledge)
        create_source_docx(self.settings.raw_root / "005、学校类" / "学校教学楼技术标.docx")

    def tearDown(self) -> None:
        self.temp.cleanup()

    def publish_first_unit(self) -> dict:
        self.knowledge.scan_sources()
        jobs = self.knowledge.create_jobs(limit=10)
        self.assertEqual(jobs["created"], 1)
        result = self.knowledge.run_job(jobs["job_ids"][0])
        self.assertEqual(result["status"], "completed")
        units = self.knowledge.list_units(limit=50)
        self.assertGreater(len(units), 0)
        publication = None
        selected = next((item for item in units if item["unit_type"] == "construction_method"), units[0])
        for unit in units:
            detail = self.knowledge.get_unit(unit["id"])
            version = detail["versions"][0]
            self.knowledge.review_unit(unit["id"], version["id"], "approve", "测试审核人")
            current = self.knowledge.publish_unit(unit["id"], version["id"], "测试发布人")
            if unit["id"] == selected["id"]:
                publication = current
        return {"unit": self.knowledge.get_unit(selected["id"]), "publication": publication}

    def test_knowledge_requires_review_before_search(self) -> None:
        self.knowledge.scan_sources()
        jobs = self.knowledge.create_jobs(limit=10)
        self.knowledge.run_job(jobs["job_ids"][0])
        self.assertEqual(self.knowledge.search("施工工艺", "学校"), [])
        unit = self.knowledge.list_units(limit=20)[0]
        detail = self.knowledge.get_unit(unit["id"])
        version = detail["versions"][0]
        self.knowledge.review_unit(unit["id"], version["id"], "approve", "测试审核人")
        publication = self.knowledge.publish_unit(unit["id"], version["id"], "测试发布人")
        results = self.knowledge.search("施工", "学校")
        self.assertTrue(any(item["unit_id"] == unit["id"] for item in results))
        self.assertTrue(results[0]["sources"])
        self.knowledge.retire_publication(publication["publication_id"], "测试管理员")
        self.assertEqual(self.knowledge.search("施工", "学校"), [])

    def test_project_generation_quality_and_native_docx(self) -> None:
        published = self.publish_first_unit()
        old_disable = os.environ.get("BID_WRITER_DISABLE_LLM")
        os.environ["BID_WRITER_DISABLE_LLM"] = "1"
        try:
            project = self.production.create_project(
                {
                    "name": "学校教学楼技术标测试项目",
                    "industry": "学校",
                    "project_type": "房屋建筑",
                    "region": "河南",
                    "source_text": "投标人必须编制主要施工方案与技术措施。",
                    "profile": {"structure_type": "框架结构", "quality_target": "合格"},
                }
            )
            self.production.parse_requirements(project["id"])
            outline = self.production.build_outline(project["id"])
            target = outline["sections"][3]
            generated = self.production.generate_section(project["id"], target["id"])
            self.assertTrue(generated["citations"])
            self.assertNotIn("[项目名称]", generated["content"])
            self.assertIn("| ---", generated["content"])
            self.production.confirm_draft(
                generated["draft_id"],
                "测试技术负责人",
                [{"index": 0, "resolution": "已核对知识来源和项目资料，同意采用当前表述"}],
            )

            with self.db.connect() as conn:
                base = conn.execute("SELECT * FROM project_drafts WHERE id=?", (generated["draft_id"],)).fetchone()
                copied_draft_ids = []
                for section in outline["sections"]:
                    if section["id"] == target["id"]:
                        continue
                    cursor = conn.execute(
                        """
                        INSERT INTO project_drafts(
                            project_id,section_id,content,citations_json,confirmations_json,version_no,status,content_hash,evidence_status
                        ) VALUES (?,?,?,?, '[]',1,'draft',?,?)
                        """,
                        (project["id"], section["id"], base["content"], base["citations_json"], base["content_hash"], base["evidence_status"]),
                    )
                    copied_draft_ids.append(int(cursor.lastrowid))
                    conn.execute("UPDATE project_sections SET status='drafted' WHERE id=?", (section["id"],))
            for draft_id in copied_draft_ids:
                self.production.confirm_draft(draft_id, "测试技术负责人")

            quality = self.production.quality_gate(project["id"])
            self.assertTrue(quality["review_ready"], quality)
            self.assertFalse(quality["formal_ready"], quality)
            self.assertFalse(quality["ready"], quality)
            self.assertEqual(quality["readiness_level"], "review")
            with self.assertRaisesRegex(ValueError, "正式"):
                self.production.export_project(project["id"], "docx")
            review_export = self.production.export_project(project["id"], "docx", "review")
            self.assertEqual(review_export["mode"], "review")
            review_package = self.production.export_project(project["id"], "package", "review")
            with zipfile.ZipFile(review_package["file_path"]) as archive:
                names = set(archive.namelist())
            self.assertEqual(len(names), 4)
            self.assertTrue(any(name.endswith(".docx") for name in names))
            self.assertTrue(any(name.endswith(".pdf") for name in names))
            self.assertIn("送审说明.md", names)
            self.assertIn("文件清单.txt", names)
            self.production.update_project(
                project["id"],
                {
                    "profile": {
                        "bidder_name": "河南建设工程有限公司",
                        "professional_reviewer": "注册建造师复核人",
                        "compliance_confirmed": True,
                        "manual_finalized": True,
                    }
                },
            )
            quality = self.production.quality_gate(project["id"])
            self.assertTrue(quality["formal_ready"], quality)
            self.assertTrue(quality["ready"], quality)
            exported = self.production.export_project(project["id"], "docx")
            path = Path(exported["file_path"])
            self.assertTrue(path.exists())
            rendered = Document(path)
            self.assertGreater(len(rendered.tables), 0)
            self.assertTrue(any("学校教学楼技术标测试项目" in paragraph.text for paragraph in rendered.paragraphs))
            self.assertEqual(published["unit"]["status"], "published")
        finally:
            if old_disable is None:
                os.environ.pop("BID_WRITER_DISABLE_LLM", None)
            else:
                os.environ["BID_WRITER_DISABLE_LLM"] = old_disable

    def test_confirmation_is_bound_to_content_and_resolved_individually(self) -> None:
        self.publish_first_unit()
        old_disable = os.environ.get("BID_WRITER_DISABLE_LLM")
        os.environ["BID_WRITER_DISABLE_LLM"] = "1"
        try:
            project = self.production.create_project(
                {"name": "签审测试", "industry": "学校", "source_text": "投标人必须编制主要施工方案与技术措施。"}
            )
            self.production.parse_requirements(project["id"])
            outline = self.production.build_outline(project["id"])
            generated = self.production.generate_section(project["id"], outline["sections"][3]["id"])

            with self.assertRaisesRegex(ValueError, "必须逐项处理"):
                self.production.confirm_draft(generated["draft_id"], "审核人")

            confirmed = self.production.confirm_draft(
                generated["draft_id"],
                "审核人",
                [{"index": 0, "resolution": "已复核，采用当前内容"}],
            )
            self.assertEqual(confirmed["resolved_confirmations"], 1)
            detail = self.production.get_project(project["id"])
            draft = next(item["draft"] for item in detail["sections"] if item["id"] == outline["sections"][3]["id"])
            self.assertEqual(len(draft["confirmations"]), 1)
            self.assertEqual(draft["confirmation_resolutions"][0]["resolver"], "审核人")
            reviewed_section = next(item for item in detail["sections"] if item["id"] == outline["sections"][3]["id"])
            self.assertEqual(reviewed_section["status"], "reviewed")

            changed = self.production.update_draft(generated["draft_id"], generated["content"] + "\n新增内容。")
            self.assertTrue(changed["quality_confirmations_invalidated"])
            detail = self.production.get_project(project["id"])
            changed_section = next(item for item in detail["sections"] if item["id"] == outline["sections"][3]["id"])
            self.assertEqual(changed_section["status"], "drafted")
            quality = self.production.quality_gate(project["id"])
            self.assertEqual(quality["metrics"]["covered_requirements"], 0)
            self.assertGreaterEqual(quality["metrics"]["unreviewed_drafts"], 1)
        finally:
            if old_disable is None:
                os.environ.pop("BID_WRITER_DISABLE_LLM", None)
            else:
                os.environ["BID_WRITER_DISABLE_LLM"] = old_disable

    def test_production_service_has_no_duplicate_methods(self) -> None:
        source_path = Path(__file__).parents[1] / "bid_writer_v2" / "production" / "service.py"
        module = ast.parse(source_path.read_text(encoding="utf-8"))
        service = next(node for node in module.body if isinstance(node, ast.ClassDef) and node.name == "ProductionService")
        methods = [node.name for node in service.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))]
        duplicates = sorted({name for name in methods if methods.count(name) > 1})
        self.assertEqual(duplicates, [])

    def test_corpus_slot_reconciler_refills_missing_ai_ticks(self) -> None:
        app = create_app(self.settings)
        with app.state.db.connect() as conn:
            run_id = int(
                conn.execute(
                    "INSERT INTO corpus_runs(snapshot_hash,status,stage,policy_json) VALUES (?,'running','ai',?)",
                    ("r" * 64, json.dumps({"ai_concurrency": 3})),
                ).lastrowid
            )
            for index in range(3):
                source_id = int(
                    conn.execute(
                        "INSERT INTO source_files(absolute_path,relative_path,file_name,extension,sha256,family_key) VALUES (?,?,?,?,?,?)",
                        (str(self.root / f"r{index}.md"), f"r{index}.md", f"r{index}.md", ".md", str(index) * 64, f"r{index}"),
                    ).lastrowid
                )
                conn.execute(
                    "INSERT INTO corpus_run_items(run_id,source_id,source_hash,item_kind,stage,status) VALUES (?,?,?,'document','ai','pending')",
                    (run_id, source_id, str(index) * 64),
                )
        self.assertEqual(app.state.reconcile_corpus_slots(), 1)
        active = app.state.db.row(
            "SELECT COUNT(*) AS count FROM app_jobs WHERE job_type='knowledge.corpus.tick' AND target_id=? AND status='pending'",
            (f"{run_id}:ai",),
        )
        self.assertEqual(int(active["count"]), 3)
        app.state.reconcile_corpus_slots()
        active = app.state.db.row(
            "SELECT COUNT(*) AS count FROM app_jobs WHERE job_type='knowledge.corpus.tick' AND target_id=? AND status='pending'",
            (f"{run_id}:ai",),
        )
        self.assertEqual(int(active["count"]), 3)

    def test_http_status_uses_v2_database(self) -> None:
        app = create_app(self.settings)
        with TestClient(app) as client:
            response = client.get("/api/status")
            self.assertEqual(response.status_code, 200)
            payload = response.json()
            self.assertEqual(payload["version"], "2.0.0")
            self.assertEqual(payload["architecture"], "knowledge-engineering-first")
            self.assertFalse(payload["operations_enabled"])
            self.assertIn("ai", payload)
            self.assertEqual(client.get("/api/ai/metrics").status_code, 200)
            self.assertEqual(client.get("/api/ai/runs").status_code, 200)
            self.assertEqual(client.get("/api/ai/prompts").status_code, 200)
            self.assertEqual(client.get("/api/evaluation/cases").status_code, 200)
            self.assertEqual(client.get("/api/evaluation/runs").status_code, 200)
            self.assertEqual(client.get("/api/knowledge-ai/metrics").status_code, 200)
            self.assertEqual(client.get("/api/knowledge-ai/candidates").status_code, 200)
            self.assertEqual(client.get("/api/knowledge-ai/tasks").status_code, 200)
            metrics = client.get("/metrics")
            self.assertEqual(metrics.status_code, 200)
            self.assertIn("bid_writer_ai_runs", metrics.text)
            self.assertIn("bid_writer_retrieval_latency_ms", metrics.text)
            self.assertIn("bid_writer_audit_chain_valid", metrics.text)

    def test_unhandled_errors_do_not_leak_internal_details(self) -> None:
        app = create_app(self.settings)
        app.state.knowledge.metrics = lambda: (_ for _ in ()).throw(RuntimeError("secret-path-and-sql"))
        with TestClient(app, raise_server_exceptions=False) as client:
            response = client.get("/api/status", headers={"X-Request-ID": "test-request-id"})
        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.json()["code"], "internal_error")
        self.assertEqual(response.json()["request_id"], "test-request-id")
        self.assertNotIn("secret-path-and-sql", response.text)

    def test_unknown_api_uses_stable_json_error(self) -> None:
        app = create_app(self.settings)
        response = TestClient(app).get("/api/definitely-missing", headers={"X-Request-ID": "missing-test"})
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["code"], "http_404")
        self.assertEqual(response.json()["request_id"], "missing-test")


if __name__ == "__main__":
    unittest.main()
