from __future__ import annotations

import json
import re
import shutil
import sqlite3
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from .bid_strategy import get_bid_strategy, render_bid_strategy_markdown
from .artifact_audit import audit_docx_path, audit_zip_path
from .case_assets import list_case_assets, render_case_assets_markdown
from .closure_confirmation import build_closure_confirmation
from .communications import list_communications, render_communications_markdown
from .channel_ops import build_channel_ops, render_channel_ops_markdown
from .coverage_report import build_coverage_report
from .db import connect, row_to_dict, rows_to_dicts
from .delivery_assistant import render_delivery_instruction
from .delivery_records import record_delivery_export
from .delivery_release import build_delivery_release, render_delivery_release_markdown
from .delivery_review import build_delivery_review
from .document_blocks import blocks_by_section, render_section_blocks_markdown
from .document_settings import get_document_settings, render_document_settings_markdown
from .document_templates import template_for_tender
from .docx_audit import audit_docx, build_preflight_signature, save_docx_audit
from .docx_blocks import append_markdown_table, append_section_blocks, collect_caption_index
from .draft_polish import list_replacement_records, render_replacement_records_markdown
from .enterprise_profiles import enterprise_summary, get_enterprise_profile, render_enterprise_profile_markdown
from .final_document import build_final_document, render_final_document_markdown
from .final_checklist import build_final_checklist
from .intake_assistant import build_intake_assistant, render_intake_markdown
from .feedback_rework import build_feedback_rework, render_feedback_rework_markdown
from .materials import list_material_items, render_materials_markdown
from .order_confirmation import build_order_confirmation
from .order_wizard import build_order_wizard, render_order_wizard_markdown
from .payments import payment_summary, render_payments_markdown
from .package_validation import render_package_validation_markdown, validate_package_path
from .planner import audit_section_plan, render_plan_audit_markdown
from .production_tasks import get_production_task
from .project_overview import build_project_overview
from .project_timeline import build_project_timeline, render_project_timeline_markdown
from .project_profiles import get_project_profile
from .quality_gate import build_quality_gate
from .pricing import build_price_quote
from .production_readiness import build_production_readiness, render_production_readiness_markdown
from .production_starter import build_production_starter, render_production_starter_markdown
from .response_matrix import build_response_matrix, render_response_matrix_markdown
from .retrospective import build_project_retrospective
from .revision_tasks import render_revision_tasks_markdown, sync_revision_tasks
from .settings import EXPORT_DIR
from .source_audit import build_source_audit, render_source_audit_markdown
from .workflow import build_workflow_status


def _safe_name(text: str) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", text)
    text = re.sub(r"\s+", "_", text).strip("._")
    return text[:80] or "tender"


def _save_docx(document: Any, path: Path) -> Path:
    try:
        document.save(str(path))
        return path
    except PermissionError:
        revised = path.with_name(f"{path.stem}_修正版_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        document.save(str(revised))
        return revised


def _new_docx_document(docx_module: Any, tender_id: int, conn: sqlite3.Connection) -> Any:
    template = template_for_tender(tender_id, conn=conn)
    source_value = str(template.get("source_docx_path") or "").strip()
    if not source_value:
        return docx_module.Document()
    source = Path(source_value)
    if not source.is_file() or source.suffix.lower() != ".docx":
        return docx_module.Document()
    document = docx_module.Document(str(source))
    body = document._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)
    return document


def _content_with_heading(content: str, section_title: str) -> str:
    lines = content.strip().splitlines()
    if not lines:
        return f"# {section_title}\n\n> 本章尚未生成。"
    if lines[0].startswith("# "):
        lines[0] = f"# {section_title}"
        return "\n".join(lines).strip()
    return f"# {section_title}\n\n{content.strip()}"


def _client_safe_content(content: str) -> str:
    skip_titles = {"参考来源摘要", "引用来源", "需人工确认", "响应矩阵", "收款记录", "内部审查", "投标单位资料"}
    result: list[str] = []
    skip_level = 0
    for line in content.splitlines():
        match = re.match(r"^(#{1,6})\s*(.*?)\s*$", line.strip())
        if match:
            level = len(match.group(1))
            title = match.group(2).strip("：:")
            if title in skip_titles:
                skip_level = level
                continue
            if skip_level and level <= skip_level:
                skip_level = 0
        if skip_level:
            continue
        result.append(line)
    return "\n".join(result).strip()


def _ordered_draft_entries(conn: sqlite3.Connection, tender_id: int) -> list[dict[str, Any]]:
    planned_rows = rows_to_dicts(
        conn.execute(
            """
            SELECT
                sp.id AS plan_id,
                sp.order_no,
                sp.section_title AS plan_section_title,
                sp.status AS plan_status,
                sp.draft_id AS plan_draft_id,
                d.id AS draft_id,
                d.section_title AS draft_section_title,
                d.content,
                d.citations_json
            FROM section_plans sp
            LEFT JOIN drafts d ON d.id = sp.draft_id
            WHERE sp.tender_id = ?
            ORDER BY sp.order_no, sp.id
            """,
            (tender_id,),
        ).fetchall()
    )
    if not planned_rows:
        drafts = rows_to_dicts(conn.execute("SELECT * FROM drafts WHERE tender_id = ? ORDER BY id", (tender_id,)).fetchall())
        return [
            {
                "section_title": draft["section_title"],
                "content": draft["content"],
                "citations": json.loads(draft.get("citations_json") or "[]"),
                "draft_id": draft["id"],
                "missing": False,
            }
            for draft in drafts
        ]

    entries: list[dict[str, Any]] = []
    used_draft_ids: set[int] = set()
    for row in planned_rows:
        draft_id = row.get("draft_id")
        if draft_id:
            used_draft_ids.add(int(draft_id))
            entries.append(
                {
                    "section_title": row["plan_section_title"] or row["draft_section_title"],
                    "content": row.get("content") or "",
                    "citations": json.loads(row.get("citations_json") or "[]"),
                    "draft_id": int(draft_id),
                    "plan_id": row["plan_id"],
                    "order_no": row["order_no"],
                    "missing": False,
                }
            )
        else:
            entries.append(
                {
                    "section_title": row["plan_section_title"],
                    "content": "",
                    "citations": [],
                    "draft_id": None,
                    "plan_id": row["plan_id"],
                    "order_no": row["order_no"],
                    "missing": True,
                }
            )

    if used_draft_ids:
        placeholders = ",".join("?" for _ in used_draft_ids)
        orphan_rows = rows_to_dicts(
            conn.execute(
                f"""
                SELECT *
                FROM drafts
                WHERE tender_id = ? AND id NOT IN ({placeholders})
                ORDER BY id
                """,
                [tender_id, *sorted(used_draft_ids)],
            ).fetchall()
        )
    else:
        orphan_rows = rows_to_dicts(conn.execute("SELECT * FROM drafts WHERE tender_id = ? ORDER BY id", (tender_id,)).fetchall())
    for draft in orphan_rows:
        entries.append(
            {
                "section_title": draft["section_title"],
                "content": draft["content"],
                "citations": json.loads(draft.get("citations_json") or "[]"),
                "draft_id": draft["id"],
                "missing": False,
                "orphan": True,
            }
        )
    return entries


def export_markdown(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")
    entries = _ordered_draft_entries(conn, tender_id)
    rich_blocks = blocks_by_section(tender_id, conn=conn)
    enterprise_markdown = render_enterprise_profile_markdown(conn=conn)
    lines = [f"# {tender['name']}", "", "## 投标单位资料", ""]
    for line in enterprise_markdown.splitlines()[1:]:
        lines.append(line)
    payments_markdown = render_payments_markdown(tender_id, conn=conn)
    lines.extend(["", "## 收款记录", ""])
    for line in payments_markdown.splitlines()[1:]:
        lines.append(line)
    lines.extend(["", "## 响应矩阵", ""])
    requirements = rows_to_dicts(conn.execute("SELECT * FROM requirements WHERE tender_id = ? ORDER BY id", (tender_id,)).fetchall())
    for req in requirements:
        lines.append(f"- [{req['priority']}] {req['kind']}：{req['content']}")
    lines.append("")
    lines.append("## 技术标章节")
    for entry in entries:
        lines.append("")
        if entry.get("missing"):
            lines.append(f"# {entry['section_title']}")
            lines.append("")
            lines.append("> 本章尚未生成。请在目录规划中生成本章后再导出正式稿。")
            continue
        content = str(entry.get("content") or "").strip()
        lines.append(_content_with_heading(content, str(entry["section_title"])))
        rich_markdown = render_section_blocks_markdown(rich_blocks.get(str(entry["section_title"]), []))
        if rich_markdown:
            lines.extend(["", rich_markdown])
        citations = entry.get("citations") or []
        if citations:
            lines.append("")
            lines.append("### 引用来源")
            for citation in citations:
                lines.append(f"- {citation.get('source_path')} / {citation.get('heading_text')}")
    delivery_review = build_delivery_review(tender_id, conn=conn)
    lines.append("")
    lines.append(delivery_review["markdown"])
    path = EXPORT_DIR / f"{_safe_name(tender['name'])}_{tender_id}.md"
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    if own_conn:
        conn.close()
    return {"path": str(path), "format": "markdown"}


def export_docx(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")
    entries = _ordered_draft_entries(conn, tender_id)
    rich_blocks = blocks_by_section(tender_id, conn=conn)
    try:
        import docx  # type: ignore
    except Exception as exc:
        raise RuntimeError("DOCX export requires python-docx") from exc
    document = _new_docx_document(docx, tender_id, conn)
    settings = get_document_settings(tender_id, conn=conn)
    apply_docx_styles(document, settings=settings)
    _apply_header_footer(document, settings)
    if settings.get("include_cover"):
        _append_cover_page(document, tender, settings)
    if settings.get("include_toc"):
        _append_toc_page(document, entries, rich_blocks)
    title_paragraph = document.add_heading(settings.get("document_title") or tender["name"], level=0)
    if settings.get("include_toc"):
        title_paragraph.paragraph_format.page_break_before = True
    document.add_heading("技术标章节", level=1)
    for index, entry in enumerate(entries):
        if settings.get("section_page_break") and index > 0:
            document.add_page_break()
        if entry.get("missing"):
            document.add_heading(str(entry["section_title"]), level=1)
            document.add_paragraph("本章尚未生成。请在目录规划中生成本章后再导出正式稿。")
            continue
        safe_content = _client_safe_content(str(entry.get("content") or ""))
        content = _content_with_heading(safe_content, str(entry["section_title"]))
        _append_markdown_content(document, content)
        append_section_blocks(
            document,
            rich_blocks.get(str(entry["section_title"]), []),
            index + 1,
            formal=True,
            section_title=str(entry["section_title"]),
        )
    path = EXPORT_DIR / f"{_safe_name(tender['name'])}_{tender_id}.docx"
    path = _save_docx(document, path)
    audit = audit_docx(path, formal=True)
    final_hash = str(build_final_document(tender_id, include_content=False, conn=conn).get("summary", {}).get("hash") or "")
    audit = save_docx_audit(tender_id, audit, build_preflight_signature(tender_id, final_hash, conn))
    if not audit["ready"]:
        raise ValueError("正式DOCX结构审计未通过：" + "；".join(audit["blockers"]))
    if own_conn:
        conn.close()
    return {"path": str(path), "format": "docx", "layout_audit": audit}


def export_client_markdown(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")
    entries = _ordered_draft_entries(conn, tender_id)
    rich_blocks = blocks_by_section(tender_id, conn=conn)
    lines = [f"# {tender['name']}", "", "## 技术标章节"]
    for entry in entries:
        lines.append("")
        if entry.get("missing"):
            lines.append(f"# {entry['section_title']}")
            lines.append("")
            lines.append("> 本章尚未生成。")
            continue
        content = _client_safe_content(str(entry.get("content") or ""))
        lines.append(_content_with_heading(content, str(entry["section_title"])))
        rich_markdown = render_section_blocks_markdown(rich_blocks.get(str(entry["section_title"]), []))
        if rich_markdown:
            lines.extend(["", rich_markdown])
    path = EXPORT_DIR / f"{_safe_name(tender['name'])}_{tender_id}_客户版.md"
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    if own_conn:
        conn.close()
    return {"path": str(path), "format": "client_markdown"}


def export_client_docx(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")
    entries = _ordered_draft_entries(conn, tender_id)
    rich_blocks = blocks_by_section(tender_id, conn=conn)
    try:
        import docx  # type: ignore
    except Exception as exc:
        raise RuntimeError("DOCX export requires python-docx") from exc

    document = _new_docx_document(docx, tender_id, conn)
    settings = get_document_settings(tender_id, conn=conn)
    apply_docx_styles(document, settings=settings)
    _apply_header_footer(document, settings)
    if settings.get("include_cover"):
        _append_cover_page(document, tender, settings)
    if settings.get("include_toc"):
        _append_toc_page(document, entries, rich_blocks)
    title_paragraph = document.add_heading(settings.get("document_title") or tender["name"], level=0)
    if settings.get("include_toc"):
        title_paragraph.paragraph_format.page_break_before = True

    document.add_heading("技术标章节", level=1)
    for index, entry in enumerate(entries):
        if settings.get("section_page_break") and index > 0:
            document.add_page_break()
        if entry.get("missing"):
            document.add_heading(str(entry["section_title"]), level=1)
            document.add_paragraph("本章尚未生成。")
            continue
        safe_content = _client_safe_content(str(entry.get("content") or ""))
        content = _content_with_heading(safe_content, str(entry["section_title"]))
        _append_markdown_content(document, content)
        append_section_blocks(
            document,
            rich_blocks.get(str(entry["section_title"]), []),
            index + 1,
            formal=True,
            section_title=str(entry["section_title"]),
        )

    path = EXPORT_DIR / f"{_safe_name(tender['name'])}_{tender_id}_客户版.docx"
    path = _save_docx(document, path)
    audit = audit_docx(path, formal=True)
    final_hash = str(build_final_document(tender_id, include_content=False, conn=conn).get("summary", {}).get("hash") or "")
    audit = save_docx_audit(tender_id, audit, build_preflight_signature(tender_id, final_hash, conn))
    if not audit["ready"]:
        raise ValueError("客户DOCX结构审计未通过：" + "；".join(audit["blockers"]))
    artifact_audit = audit_docx_path(path)
    if not artifact_audit["ready"]:
        raise ValueError("客户DOCX最终产物审计未通过")
    if own_conn:
        conn.close()
    return {"path": str(path), "format": "client_docx", "layout_audit": audit, "artifact_audit": artifact_audit}


def export_review_docx(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")
    generated = export_client_docx(tender_id, conn=conn)
    source = Path(generated["path"])
    path = EXPORT_DIR / f"{_safe_name(tender['name'])}_{tender_id}_送审版.docx"
    shutil.copy2(source, path)
    artifact_audit = audit_docx_path(path)
    if not artifact_audit["ready"]:
        raise ValueError("送审DOCX最终产物审计未通过")
    result = {"path": str(path), "format": "review_docx", "layout_audit": generated.get("layout_audit") or {}, "artifact_audit": artifact_audit}
    if own_conn:
        conn.close()
    return result


def _convert_review_pdf(docx_path: Path) -> Path:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    windows_soffice = Path("C:/Program Files/LibreOffice/program/soffice.exe")
    if not executable and windows_soffice.exists():
        executable = str(windows_soffice)
    if not executable:
        raise ValueError("未安装LibreOffice，无法生成送审PDF")
    result = subprocess.run(
        [executable, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    pdf_path = docx_path.with_suffix(".pdf")
    if result.returncode or not pdf_path.is_file() or not pdf_path.stat().st_size:
        raise ValueError("LibreOffice生成送审PDF失败")
    return pdf_path


def export_review_pdf(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    docx_file = export_review_docx(tender_id, conn=conn)
    path = _convert_review_pdf(Path(docx_file["path"]))
    result = {"path": str(path), "format": "review_pdf", "source_docx": docx_file["path"], "artifact_audit": docx_file["artifact_audit"]}
    if own_conn:
        conn.close()
    return result


def export_review_package(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")
    pdf_file = export_review_pdf(tender_id, conn=conn)
    docx_path = Path(pdf_file["source_docx"])
    pdf_path = Path(pdf_file["path"])
    package_path = EXPORT_DIR / f"{_safe_name(tender['name'])}_{tender_id}_送审包.zip"
    note = "# 送审说明\n\n本包用于专业送审。正式投标前须补齐真实投标单位、专业复核人、合规确认、签章及人工定稿。\n"
    names = [docx_path.name, pdf_path.name, "送审说明.md", "文件清单.txt"]
    with zipfile.ZipFile(package_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(docx_path, docx_path.name)
        archive.write(pdf_path, pdf_path.name)
        archive.writestr("送审说明.md", note)
        archive.writestr("文件清单.txt", "\n".join(names) + "\n")
    package_audit = audit_zip_path(package_path, allowed_names=set(names), require_review_package_types=True)
    if not package_audit["ready"]:
        raise ValueError("送审包最终产物审计未通过")
    result = {
        "path": str(package_path),
        "format": "review_package",
        "components": {"docx": str(docx_path), "pdf": str(pdf_path), "customer_note": "送审说明.md", "file_list": "文件清单.txt"},
        "artifact_audit": package_audit,
    }
    result["delivery_record"] = record_delivery_export(tender_id, result, conn=conn)
    if own_conn:
        conn.close()
    return result


def render_client_delivery_note(tender_id: int, package_name: str = "", conn: sqlite3.Connection | None = None) -> str:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")
    task = get_production_task(tender_id, conn=conn)
    customer = task.get("customer_name") or "您好"
    deliverable_format = task.get("deliverable_format") or "DOCX + Markdown"
    deadline = task.get("deadline") or "按约定时间"
    notes = task.get("delivery_notes") or ""
    package_name = package_name or "客户发货包.zip"
    lines = [
        "# 客户发货说明",
        "",
        f"{customer}，您好。",
        "",
        f"{tender.get('name') or '本项目'} 的技术标初稿已整理完成，本次交付格式为 {deliverable_format}。",
        f"交付文件：{package_name}。",
        f"交付期限：{deadline}。",
        "",
        "## 文件清单",
        "",
        "- 技术标初稿_客户版.docx：可编辑 Word 初稿。",
        "- 技术标初稿_客户版.md：Markdown 备份稿，便于后续继续修改。",
        "- 客户发货说明.md：本说明文件。",
        "",
        "## 复核提示",
        "",
        "- 正式投标前，请结合最终招标文件核对项目名称、工期、质量安全目标、专用条款、页码目录和签章格式。",
        "- 如需调整目录、补充专项章节或修改格式，可以继续反馈修改要求。",
    ]
    if notes:
        lines.extend(["", f"备注：{notes}"])
    if own_conn:
        conn.close()
    return "\n".join(lines).strip() + "\n"


def export_client_package(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")

    markdown = export_client_markdown(tender_id, conn=conn)
    docx_file = export_client_docx(tender_id, conn=conn)
    base_name = f"{_safe_name(tender['name'])}_{tender_id}"
    package_path = EXPORT_DIR / f"{base_name}_客户发货包.zip"
    note = render_client_delivery_note(tender_id, package_name=package_path.name, conn=conn)
    file_list = "\n".join(
        [
            "客户发货包文件清单",
            "",
            "技术标初稿_客户版.docx",
            "技术标初稿_客户版.md",
            "客户发货说明.md",
            "README.txt",
            "",
            "本 ZIP 为客户交付版本，请以最终招标文件和投标格式要求进行人工复核。",
        ]
    )
    readme = "\n".join(
        [
            f"项目名称：{tender['name']}",
            f"项目 ID：{tender_id}",
            "",
            "本包为客户发货包，仅包含客户可接收的技术标初稿和说明。",
            "正式投标前请核对项目名称、章节目录、页码、格式、签章和招标专用条款。",
        ]
    )
    result = {
        "path": str(package_path),
        "format": "client_zip",
        "components": {
            "client_docx": docx_file["path"],
            "client_markdown": markdown["path"],
            "customer_note": "客户发货说明.md",
            "file_list": "客户文件清单.txt",
            "readme": "README.txt",
        },
        "client_safe": True,
    }
    with zipfile.ZipFile(package_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(docx_file["path"], "技术标初稿_客户版.docx")
        archive.write(markdown["path"], "技术标初稿_客户版.md")
        archive.writestr("客户发货说明.md", note)
        archive.writestr("客户文件清单.txt", file_list)
        archive.writestr("README.txt", readme)
    package_names = {"技术标初稿_客户版.docx", "技术标初稿_客户版.md", "客户发货说明.md", "客户文件清单.txt", "README.txt"}
    artifact_audit = audit_zip_path(package_path, allowed_names=package_names)
    if not artifact_audit["ready"]:
        raise ValueError("客户发货包最终产物审计未通过")
    result["artifact_audit"] = artifact_audit
    result["delivery_record"] = record_delivery_export(tender_id, result, conn=conn)
    if own_conn:
        conn.close()
    return result


def export_package(tender_id: int, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    own_conn = conn is None
    conn = conn or connect()
    tender = row_to_dict(conn.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone())
    if not tender:
        raise ValueError(f"Tender not found: {tender_id}")

    markdown = export_markdown(tender_id, conn=conn)
    docx_file = export_docx(tender_id, conn=conn)
    overview = build_project_overview(tender_id, conn=conn)
    timeline = build_project_timeline(tender_id, conn=conn)
    document_settings = get_document_settings(tender_id, conn=conn)
    enterprise_profile = get_enterprise_profile(conn=conn)
    bid_strategy = get_bid_strategy(tender_id, conn=conn)
    task = get_production_task(tender_id, conn=conn)
    profile = get_project_profile(tender_id, conn=conn)
    plan_audit = audit_section_plan(tender_id, conn=conn)
    response_matrix = build_response_matrix(tender_id, conn=conn)
    coverage = build_coverage_report(tender_id, conn=conn)
    source_audit = build_source_audit(tender_id, conn=conn)
    workflow = build_workflow_status(tender_id, conn=conn)
    production_starter = build_production_starter(tender_id, conn=conn)
    order_wizard = build_order_wizard(tender_id, conn=conn)
    channel_ops = build_channel_ops(tender_id, conn=conn)
    delivery_review = build_delivery_review(tender_id, conn=conn)
    intake_report = build_intake_assistant(tender_id, conn=conn)
    order_confirmation = build_order_confirmation(tender_id, conn=conn)
    communications = list_communications(tender_id, conn=conn)
    price_quote = build_price_quote(tender_id, conn=conn)
    payments = payment_summary(tender_id, conn=conn)
    materials = list_material_items(tender_id, conn=conn)
    quality_gate = build_quality_gate(tender_id, conn=conn)
    revision_tasks = sync_revision_tasks(tender_id, conn=conn)
    feedback_rework = build_feedback_rework(tender_id, conn=conn)
    final_checklist = build_final_checklist(tender_id, conn=conn)
    final_document = build_final_document(tender_id, include_content=False, conn=conn)
    replacement_records = list_replacement_records(tender_id, conn=conn)
    retrospective = build_project_retrospective(tender_id, conn=conn)
    case_assets = list_case_assets(tender_id=tender_id, limit=300, conn=conn)

    base_name = f"{_safe_name(tender['name'])}_{tender_id}"
    package_path = EXPORT_DIR / f"{base_name}_交付包.zip"
    readme = "\n".join(
        [
            f"项目名称：{tender['name']}",
            f"项目 ID：{tender_id}",
            f"章节数量：{coverage['summary'].get('total_sections', 0)}",
            f"已生成章节：{coverage['summary'].get('generated_sections', 0)}",
            f"高优先级缺口：{coverage['summary'].get('high_priority_missing', 0)}",
            f"审查问题：{coverage['summary'].get('review_findings', 0)}",
            f"是否具备导出初稿条件：{'是' if workflow['summary'].get('export_ready') else '否'}",
            "",
            "文件说明：",
            "- 技术标初稿.docx：可直接打开的 Word 初稿。",
            "- 技术标初稿.md：Markdown 备份稿。",
            "- 项目总览.md：项目当前状态、风险、指标和下一步动作。",
            "- 项目生产履历.md：项目从录入、报价、收款、生成、交付到复盘的时间线。",
            "- 新单生产向导.md：从客户询盘到接单确认、资料补齐、章节生成和人工交付的步骤路线。",
            "- 渠道运营助手.md：闲鱼、微信等渠道的商品说明、接单、催资料、确认、交付、售后话术和人工确认边界。",
            "- 生产启动包.md：接单判断、报价建议、资料缺口、渠道边界、客户回复和开工清单。",
            "- 成稿格式设置.md：封面、页眉页脚、目录、分页和附录导出选项。",
            "- 投标单位资料.md：投标单位、资质能力、人员设备、类似业绩和服务承诺。",
            "- 投标响应策略.md：项目卖点、风险边界、响应优先级和章节写作口径。",
            "- 生产任务.json：客户、来源平台、交付期限、交付格式和生产状态。",
            "- 项目资料.json：项目资料收集表。",
            "- 目录完整性审计.md：标准章节、施工工艺章节和未挂接招标要求审计。",
            "- 响应矩阵挂接报告.md：每条招标要求与目录章节、草稿响应之间的挂接关系。",
            "- 响应覆盖报告.json：条款覆盖、章节生成和审查结果。",
            "- 引用来源审计.md：章节引用、来源可定位性、案例资产引用和旧项目名风险审计。",
            "- 生产流程状态.json：从资料补齐到导出的流程状态。",
            "",
            "正式投标前请人工复核项目参数、招标专用条款、格式要求和全部引用来源。",
        ]
    )

    readme += "\n- 接单评估.md：接单判断、报价建议、资料清单和客户回复话术。\n- 接单评估.json：接单评估结构化数据。\n- 报价测算.md：价格区间、加价折扣明细、工作量和交付周期。\n- 报价测算.json：报价测算结构化数据和规则快照。\n- 收款记录.md：订单金额、已收款、待确认款和未收金额。\n- 收款记录.json：收款状态和明细结构化数据。\n- 订单确认单.md：费用、交付范围、资料前提、修改轮次和不包含范围。\n- 订单确认单.json：订单确认单结构化数据。\n- 资料清单.md：客户资料、招标资料、图纸清单和格式资料的收集状态。\n- 资料清单.json：资料清单结构化数据。\n- 质量门禁报告.md：终稿质量门禁、分项检查和修订任务清单。\n- 质量门禁报告.json：质量门禁结构化数据。\n- 修订任务台账.md：质量门禁生成的修订任务、负责人、状态和处理备注。\n- 修订任务台账.json：修订任务台账结构化数据。\n- 交付审查报告.md：交付前总审查、缺项、风险和处理建议。\n- 交付审查报告.json：交付审查结构化数据，便于后续系统读取。\n- 交付说明.md：客户发货说明、文件清单和内部交付检查。\n- 结案确认单.md：客户验收、反馈关闭和结案条件记录。\n- 结案确认单.json：结案确认结构化数据。\n- 项目复盘.md：成交、耗时、返工、风险和可复用经验台账。\n- 项目复盘.json：项目复盘结构化数据。"
    readme += "\n- 客户沟通记录.md：询盘、报价、资料催补、确认、交付和售后话术留痕。\n- 客户沟通记录.json：客户沟通结构化记录，便于后续复盘和归档。"
    readme += "\n- 案例资产.md：本项目已沉淀的可复用章节资产摘要。\n- 案例资产.json：案例资产结构化数据，可回流检索和后续生成。"
    readme += "\n- 最终核对清单.md：交付前自动审查和人工确认项。\n- 最终核对清单.json：最终核对结构化数据，可用于归档追责。"
    readme += "\n- 客户反馈返工处理单.md：客户修改意见、涉及章节、处理计划和客户回复话术。\n- 客户反馈返工处理单.json：客户反馈返工结构化数据。"
    readme += "\n- 成稿确认报告.md：整本预览、章节完成情况、风险提示和人工定稿确认记录。\n- 成稿确认报告.json：成稿确认结构化数据。"
    readme += "\n- 项目化校正记录.md：全稿查找替换、旧项目名和风险词清理记录。\n- 项目化校正记录.json：项目化校正结构化留痕。"
    readme += "\n- 项目可交付性评估.md：接单、开工和交付三阶段的阻断项、提醒项和下一步动作。\n- 项目可交付性评估.json：项目可交付性评估结构化数据。"
    readme += "\n- 交付放行单.md：汇总成稿确认、质量门禁、来源审计、交付审查和 ZIP 核验后的客户交付放行结论。\n- 交付放行单.json：交付放行结构化数据，可用于交付前总控。"
    readme += "\n- 交付包清单核验.md：交付包文件齐全性、ZIP 可读性和 JSON 可解析性核验报告。\n- 交付包清单核验.json：交付包核验结构化数据。"

    result = {
        "path": str(package_path),
        "format": "zip",
        "components": {
            "markdown": markdown["path"],
            "docx": docx_file["path"],
            "overview": "项目总览.md",
            "timeline": "项目生产履历.md",
            "document_settings": "成稿格式设置.md",
            "enterprise_profile": "投标单位资料.md",
            "bid_strategy": "投标响应策略.md",
            "task": "生产任务.json",
            "profile": "项目资料.json",
            "plan_audit": "目录完整性审计.md",
            "response_matrix": "响应矩阵挂接报告.md",
            "coverage": "响应覆盖报告.json",
            "source_audit": "引用来源审计.md",
            "workflow": "生产流程状态.json",
            "order_wizard": "新单生产向导.md",
            "channel_ops": "渠道运营助手.md",
            "production_starter": "生产启动包.md",
            "intake_assistant": "接单评估.md",
            "communications": "客户沟通记录.md",
            "price_quote": "报价测算.md",
            "payments": "收款记录.md",
            "order_confirmation": "订单确认单.md",
            "materials": "资料清单.md",
            "quality_gate": "质量门禁报告.md",
            "revision_tasks": "修订任务台账.md",
            "feedback_rework": "客户反馈返工处理单.md",
            "final_checklist": "最终核对清单.md",
            "final_document": "成稿确认报告.md",
            "replacement_records": "项目化校正记录.md",
            "delivery_review": "交付审查报告.md",
            "delivery_instruction": "交付说明.md",
            "production_readiness": "项目可交付性评估.md",
            "delivery_release": "交付放行单.md",
            "package_validation": "交付包清单核验.md",
            "closure_confirmation": "结案确认单.md",
            "retrospective": "项目复盘.md",
            "case_assets": "案例资产.md",
        },
    }
    closure_confirmation = build_closure_confirmation(tender_id, package_result=result, conn=conn)
    delivery_instruction = render_delivery_instruction(tender_id, package_result=result, conn=conn)

    with zipfile.ZipFile(package_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(docx_file["path"], "技术标初稿.docx")
        archive.write(markdown["path"], "技术标初稿.md")
        archive.writestr("项目总览.md", overview["markdown"])
        archive.writestr("项目总览.json", json.dumps(overview, ensure_ascii=False, indent=2))
        archive.writestr("项目生产履历.md", render_project_timeline_markdown(timeline))
        archive.writestr("项目生产履历.json", json.dumps(timeline, ensure_ascii=False, indent=2))
        archive.writestr("成稿格式设置.md", render_document_settings_markdown(tender_id, conn=conn))
        archive.writestr("成稿格式设置.json", json.dumps(document_settings, ensure_ascii=False, indent=2))
        archive.writestr("投标单位资料.md", render_enterprise_profile_markdown(conn=conn))
        archive.writestr("投标单位资料.json", json.dumps(enterprise_profile, ensure_ascii=False, indent=2))
        archive.writestr("投标响应策略.md", render_bid_strategy_markdown(tender_id, conn=conn))
        archive.writestr("投标响应策略.json", json.dumps(bid_strategy, ensure_ascii=False, indent=2))
        archive.writestr("生产任务.json", json.dumps(task, ensure_ascii=False, indent=2))
        archive.writestr("项目资料.json", json.dumps(profile, ensure_ascii=False, indent=2))
        archive.writestr("目录完整性审计.md", render_plan_audit_markdown(plan_audit))
        archive.writestr("目录完整性审计.json", json.dumps(plan_audit, ensure_ascii=False, indent=2))
        archive.writestr("响应矩阵挂接报告.md", render_response_matrix_markdown(response_matrix))
        archive.writestr("响应矩阵挂接报告.json", json.dumps(response_matrix, ensure_ascii=False, indent=2))
        archive.writestr("响应覆盖报告.json", json.dumps(coverage, ensure_ascii=False, indent=2))
        archive.writestr("引用来源审计.md", render_source_audit_markdown(source_audit))
        archive.writestr("引用来源审计.json", json.dumps(source_audit, ensure_ascii=False, indent=2))
        archive.writestr("生产流程状态.json", json.dumps(workflow, ensure_ascii=False, indent=2))
        archive.writestr("新单生产向导.md", render_order_wizard_markdown(order_wizard))
        archive.writestr("新单生产向导.json", json.dumps(order_wizard, ensure_ascii=False, indent=2))
        archive.writestr("渠道运营助手.md", render_channel_ops_markdown(channel_ops))
        archive.writestr("渠道运营助手.json", json.dumps(channel_ops, ensure_ascii=False, indent=2))
        archive.writestr("生产启动包.md", render_production_starter_markdown(production_starter))
        archive.writestr("生产启动包.json", json.dumps(production_starter, ensure_ascii=False, indent=2))
        archive.writestr("接单评估.md", render_intake_markdown(intake_report))
        archive.writestr("接单评估.json", json.dumps(intake_report, ensure_ascii=False, indent=2))
        archive.writestr("客户沟通记录.md", render_communications_markdown(tender_id, conn=conn))
        archive.writestr("客户沟通记录.json", json.dumps(communications, ensure_ascii=False, indent=2))
        archive.writestr("报价测算.md", price_quote["markdown"])
        archive.writestr("报价测算.json", json.dumps(price_quote, ensure_ascii=False, indent=2))
        archive.writestr("收款记录.md", render_payments_markdown(tender_id, conn=conn))
        archive.writestr("收款记录.json", json.dumps(payments, ensure_ascii=False, indent=2))
        archive.writestr("订单确认单.md", order_confirmation["markdown"])
        archive.writestr("订单确认单.json", json.dumps(order_confirmation, ensure_ascii=False, indent=2))
        archive.writestr("资料清单.md", render_materials_markdown(tender_id, conn=conn))
        archive.writestr("资料清单.json", json.dumps(materials, ensure_ascii=False, indent=2))
        archive.writestr("质量门禁报告.md", quality_gate["markdown"])
        archive.writestr("质量门禁报告.json", json.dumps(quality_gate, ensure_ascii=False, indent=2))
        archive.writestr("修订任务台账.md", render_revision_tasks_markdown(revision_tasks))
        archive.writestr("修订任务台账.json", json.dumps(revision_tasks, ensure_ascii=False, indent=2))
        archive.writestr("客户反馈返工处理单.md", render_feedback_rework_markdown(feedback_rework))
        archive.writestr("客户反馈返工处理单.json", json.dumps(feedback_rework, ensure_ascii=False, indent=2))
        archive.writestr("最终核对清单.md", final_checklist["markdown"])
        archive.writestr("最终核对清单.json", json.dumps(final_checklist, ensure_ascii=False, indent=2))
        archive.writestr("成稿确认报告.md", render_final_document_markdown(tender_id, conn=conn))
        archive.writestr("成稿确认报告.json", json.dumps(final_document, ensure_ascii=False, indent=2))
        archive.writestr("项目化校正记录.md", render_replacement_records_markdown(tender_id, conn=conn))
        archive.writestr("项目化校正记录.json", json.dumps(replacement_records, ensure_ascii=False, indent=2))
        archive.writestr("交付审查报告.md", delivery_review["markdown"])
        archive.writestr("交付审查报告.json", json.dumps(delivery_review, ensure_ascii=False, indent=2))
        archive.writestr("交付说明.md", delivery_instruction)
        archive.writestr("结案确认单.md", closure_confirmation["markdown"])
        archive.writestr("结案确认单.json", json.dumps(closure_confirmation, ensure_ascii=False, indent=2))
        archive.writestr("项目复盘.md", retrospective["markdown"])
        archive.writestr("项目复盘.json", json.dumps(retrospective, ensure_ascii=False, indent=2))
        archive.writestr("案例资产.md", render_case_assets_markdown(tender_id, conn=conn))
        archive.writestr("案例资产.json", json.dumps(case_assets, ensure_ascii=False, indent=2))
        archive.writestr("README.txt", readme)

    production_readiness = build_production_readiness(tender_id, package_path=str(package_path), conn=conn)
    result["production_readiness"] = production_readiness
    with zipfile.ZipFile(package_path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("项目可交付性评估.md", render_production_readiness_markdown(production_readiness))
        archive.writestr("项目可交付性评估.json", json.dumps(production_readiness, ensure_ascii=False, indent=2))

    package_validation = validate_package_path(str(package_path), tender_id=tender_id)
    result["package_validation"] = package_validation
    with zipfile.ZipFile(package_path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("交付包清单核验.md", render_package_validation_markdown(package_validation))
        archive.writestr("交付包清单核验.json", json.dumps(package_validation, ensure_ascii=False, indent=2))

    delivery_release = build_delivery_release(
        tender_id,
        package_path=str(package_path),
        production_readiness=production_readiness,
        package_validation=package_validation,
        conn=conn,
    )
    result["delivery_release"] = delivery_release
    with zipfile.ZipFile(package_path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("交付放行单.md", render_delivery_release_markdown(delivery_release))
        archive.writestr("交付放行单.json", json.dumps(delivery_release, ensure_ascii=False, indent=2))

    result["delivery_record"] = record_delivery_export(tender_id, result, conn=conn)
    if own_conn:
        conn.close()
    return result


def _set_paragraph_font(paragraph: Any, name: str, size: float | None = None, bold: bool = False) -> None:
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Pt  # type: ignore

    for run in paragraph.runs:
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        if size:
            run.font.size = Pt(size)
        run.font.bold = bold


_INLINE_MARKDOWN_RE = re.compile(r"(\*\*[^*]+?\*\*|__[^_]+?__|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")


def _append_inline_markdown(paragraph: Any, text: str) -> None:
    """Render the small inline Markdown subset emitted by chapter generation."""
    cursor = 0
    for match in _INLINE_MARKDOWN_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()].replace("**", "").replace("__", ""))
        token = match.group(0)
        if token.startswith(("**", "__")):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:].replace("**", "").replace("__", ""))


def _append_markdown_line(document: Any, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        return
    heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
    if heading_match:
        paragraph = document.add_heading("", level=len(heading_match.group(1)))
        _append_inline_markdown(paragraph, heading_match.group(2))
        return
    if stripped.startswith(("- ", "* ")):
        paragraph = document.add_paragraph(style="List Bullet")
        _append_inline_markdown(paragraph, stripped[2:].strip())
        return
    if stripped.startswith(">"):
        paragraph = document.add_paragraph()
        _append_inline_markdown(paragraph, stripped.lstrip("> "))
        for run in paragraph.runs:
            run.italic = True
        return
    paragraph = document.add_paragraph()
    _append_inline_markdown(paragraph, stripped)


def _split_markdown_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if not stripped or "|" not in stripped:
        return []
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|") and not stripped.endswith(r"\|"):
        stripped = stripped[:-1]
    cells = re.split(r"(?<!\\)\|", stripped)
    return [cell.strip().replace(r"\|", "|") for cell in cells]


def _is_markdown_table_separator(line: str, expected_columns: int) -> bool:
    cells = _split_markdown_table_row(line)
    return len(cells) == expected_columns and expected_columns >= 2 and all(
        re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) is not None for cell in cells
    )


def _is_any_markdown_table_separator(line: str) -> bool:
    cells = _split_markdown_table_row(line)
    return bool(cells) and all(
        re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) is not None for cell in cells
    )


def _append_markdown_content(document: Any, content: str) -> None:
    """Render generated Markdown by blocks so pipe tables become editable Word tables."""
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        columns = _split_markdown_table_row(lines[index])
        if columns and index + 1 < len(lines) and _is_markdown_table_separator(lines[index + 1], len(columns)):
            rows: list[list[str]] = []
            index += 2
            while index < len(lines):
                values = _split_markdown_table_row(lines[index])
                if len(values) != len(columns):
                    break
                rows.append(values)
                index += 1
            append_markdown_table(document, columns, rows)
            continue
        if _is_any_markdown_table_separator(lines[index]):
            index += 1
            continue
        if lines[index].strip().startswith("|") and len(columns) >= 2:
            block_rows: list[list[str]] = []
            saw_separator = False
            while index < len(lines) and lines[index].strip().startswith("|"):
                if _is_any_markdown_table_separator(lines[index]):
                    saw_separator = True
                else:
                    values = _split_markdown_table_row(lines[index])
                    if len(values) >= 2:
                        block_rows.append(values)
                index += 1
            if block_rows and (saw_separator or len(block_rows) >= 2):
                width = max(len(row) for row in block_rows)
                headers = [f"项目{column + 1}" for column in range(width)]
                normalized = [row + [""] * (width - len(row)) for row in block_rows]
                append_markdown_table(document, headers, normalized)
                continue
            for row in block_rows:
                _append_markdown_line(document, "；".join(value for value in row if value))
            continue
        _append_markdown_line(document, lines[index])
        index += 1


def _append_cover_page(document: Any, tender: dict[str, Any], settings: dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore

    title = str(settings.get("document_title") or tender.get("name") or "")
    subtitle = str(settings.get("document_subtitle") or "投标文件技术部分")
    document_type = str(settings.get("document_type") or "技术标")
    bidder = str(settings.get("bidder_name") or "")
    version = str(settings.get("version_label") or "")
    prepared_by = str(settings.get("prepared_by") or "")
    reviewed_by = str(settings.get("reviewed_by") or "")
    document_date = str(settings.get("document_date") or "")
    confidentiality = str(settings.get("confidentiality") or "")
    if any(term in confidentiality for term in ("内部", "系统生成", "人工复核", "审查")):
        confidentiality = ""
    font = str(settings.get("heading_font") or settings.get("body_font") or "Microsoft YaHei")

    for _ in range(3):
        document.add_paragraph("")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    _set_paragraph_font(p, font, size=22, bold=True)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(subtitle)
    _set_paragraph_font(p, font, size=15, bold=True)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(document_type)
    _set_paragraph_font(p, font, size=18, bold=True)

    for _ in range(5):
        document.add_paragraph("")
    info_lines = [
        ("投标单位", bidder),
        ("版本", version),
        ("编制人", prepared_by),
        ("复核人", reviewed_by),
        ("日期", document_date),
    ]
    for label, value in info_lines:
        if value:
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(f"{label}：{value}")
            _set_paragraph_font(p, str(settings.get("body_font") or "Microsoft YaHei"), size=11)
    if confidentiality:
        document.add_paragraph("")
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(confidentiality)
        _set_paragraph_font(p, str(settings.get("body_font") or "Microsoft YaHei"), size=10)
    document.add_page_break()


def _append_toc_page(
    document: Any,
    entries: list[dict[str, Any]],
    rich_blocks: dict[str, list[dict[str, Any]]] | None = None,
) -> None:
    document.add_heading("目录", level=1)
    for index, entry in enumerate(entries, 1):
        section_title = str(entry.get("section_title") or "")
        document.add_paragraph(f"{index}. {section_title}")
    caption_index = collect_caption_index(
        rich_blocks or {},
        [str(entry.get("section_title") or "") for entry in entries],
        formal=True,
    )
    if caption_index["tables"]:
        document.add_heading("表格目录", level=2)
        for item in caption_index["tables"]:
            document.add_paragraph(item)
    if caption_index["figures"]:
        document.add_heading("插图目录", level=2)
        for item in caption_index["figures"]:
            document.add_paragraph(item)


def _apply_header_footer(document: Any, settings: dict[str, Any]) -> None:
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    header_text = str(settings.get("header_text") or "")
    footer_text = str(settings.get("footer_text") or "")
    if any(term in footer_text for term in ("系统生成", "人工复核", "内部", "审查")):
        footer_text = ""
    body_font = str(settings.get("body_font") or "Microsoft YaHei")
    for section in document.sections:
        if header_text:
            paragraph = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
            paragraph.text = header_text
            _set_paragraph_font(paragraph, body_font, size=9)
        paragraph = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        paragraph.clear()
        if footer_text:
            paragraph.add_run(footer_text + "  ")
        paragraph.add_run("第 ")
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        value = OxmlElement("w:t")
        value.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run = paragraph.add_run()
        run._r.extend([begin, instruction, separate, value, end])
        paragraph.add_run(" 页")
        _set_paragraph_font(paragraph, body_font, size=9)


def _append_delivery_review_docx(document: Any, report: dict[str, Any]) -> None:
    summary = report.get("summary", {})
    document.add_heading("附录：交付审查摘要", level=1)
    document.add_paragraph(f"审查结论：{summary.get('readiness_label', '')}")
    document.add_paragraph(
        "；".join(
            [
                f"条款草稿覆盖率 {round(float(summary.get('draft_coverage_rate') or 0) * 100)}%",
                f"目录章节 {summary.get('generated_sections', 0)}/{summary.get('total_sections', 0)}",
                f"高优先级缺口 {summary.get('high_priority_missing', 0)}",
                f"审查问题 {summary.get('review_findings', 0)}",
            ]
        )
    )

    document.add_heading("交付检查清单", level=2)
    checklist = report.get("checklist", [])
    if checklist:
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for index, header in enumerate(["检查项", "状态", "情况", "建议动作"]):
            table.rows[0].cells[index].text = header
        for item in checklist:
            cells = table.add_row().cells
            cells[0].text = str(item.get("title") or "")
            cells[1].text = {"complete": "已完成", "warning": "需关注", "pending": "待处理"}.get(
                str(item.get("status") or ""),
                str(item.get("status") or ""),
            )
            cells[2].text = str(item.get("detail") or "")
            cells[3].text = str(item.get("action") or "")

    document.add_heading("主要阻碍", level=2)
    blockers = report.get("blockers", [])
    if blockers:
        for item in blockers:
            document.add_paragraph(f"[{item.get('severity')}] {item.get('title')}：{item.get('detail')}")
    else:
        document.add_paragraph("暂无阻碍项，可进入人工终审。")

    document.add_heading("处理建议", level=2)
    recommendations = report.get("recommendations", [])
    if recommendations:
        for item in recommendations:
            document.add_paragraph(str(item))
    else:
        document.add_paragraph("暂无额外建议。")


def apply_docx_styles(document: Any, settings: dict[str, Any] | None = None) -> None:
    from docx.shared import Cm, Pt  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    settings = settings or {}
    body_font = str(settings.get("body_font") or "Microsoft YaHei")
    heading_font = str(settings.get("heading_font") or body_font)
    try:
        body_size = float(settings.get("body_font_size") or 10.5)
    except (TypeError, ValueError):
        body_size = 10.5

    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.name = body_font
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
    normal.font.size = Pt(body_size)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(4)

    for style_name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = document.styles[style_name]
        style.font.name = heading_font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), heading_font)
        style.font.size = Pt(size)
        style.font.bold = True
