from __future__ import annotations

from pathlib import Path
from typing import Any
import unicodedata
import re

from .visual_assets import resolve_asset_path


def _set_cell_text(cell: Any, value: str, *, bold: bool = False, size: float = 9.0) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Pt  # type: ignore

    cell.text = str(value)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold or _display_width(str(value)) <= 16 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(size)
            run.font.bold = bold


def _display_width(value: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F", "A"} else 1 for char in str(value))


def _column_widths(columns: list[str], rows: list[Any], total_cm: float = 15.8) -> list[float]:
    counts = len(columns)
    if counts == 0:
        return []
    weights: list[float] = []
    for index, column in enumerate(columns):
        values = [str(column)]
        for raw_row in rows:
            row = list(raw_row) if isinstance(raw_row, (list, tuple)) else [raw_row]
            if index < len(row):
                values.append(str(row[index]))
        longest = max((_display_width(value) for value in values), default=4)
        weights.append(float(max(5, min(longest, 38))))
    minimum = 1.15 if counts >= 6 else 1.55 if counts >= 4 else 2.1
    available = max(0.1, total_cm - minimum * counts)
    total_weight = sum(weights) or float(counts)
    return [minimum + available * weight / total_weight for weight in weights]


def _set_cell_margins(cell: Any, *, top: int = 90, left: int = 110, bottom: int = 90, right: int = 110) -> None:
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table: Any, widths_cm: list[float]) -> None:
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.shared import Cm  # type: ignore

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    widths_twips = [int(Cm(width_cm).twips) for width_cm in widths_cm]
    total_twips = sum(widths_twips)
    table_width = tbl_pr.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        tbl_pr.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    table_indent = tbl_pr.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        tbl_pr.append(table_indent)
    table_indent.set(qn("w:w"), "0")
    table_indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_twips in widths_twips:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width_twips))
        grid.append(grid_col)
    for index, width_cm in enumerate(widths_cm):
        width = Cm(width_cm)
        table.columns[index].width = width
        for row in table.rows:
            cell = row.cells[index]
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width.twips)))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _repeat_table_header(row: Any) -> None:
    from docx.oxml import OxmlElement  # type: ignore

    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
    tr_pr.append(header)


def _prevent_row_split(row: Any) -> None:
    from docx.oxml import OxmlElement  # type: ignore

    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _shade_cell(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _caption(document: Any, text: str) -> Any:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Pt  # type: ignore

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(9.5)
    return paragraph


def append_native_table(document: Any, block: dict[str, Any], number: str) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # type: ignore

    data = block.get("data") or {}
    columns = [str(item) for item in (data.get("columns") or [])]
    rows = list(data.get("rows") or [])
    if not columns:
        return
    _caption(document, f"表 {number} {block.get('caption') or block.get('title') or ''}")
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    _repeat_table_header(header)
    for index, value in enumerate(columns):
        cell = header.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cell, "D9EAF3")
        _set_cell_text(cell, value, bold=True, size=9.0)
    for raw_row in rows:
        row = table.add_row()
        _prevent_row_split(row)
        values = list(raw_row) if isinstance(raw_row, (list, tuple)) else [raw_row]
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = values[index] if index < len(values) else ""
            _set_cell_text(cell, str(value), size=8.5)
    _set_table_geometry(table, _column_widths(columns, rows))


def append_markdown_table(document: Any, columns: list[str], rows: list[list[str]]) -> None:
    """Render a table found inside generated Markdown without adding a duplicate caption."""
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # type: ignore

    if not columns:
        return
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    _repeat_table_header(header)
    for index, value in enumerate(columns):
        cell = header.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cell, "D9EAF3")
        _set_cell_text(cell, value, bold=True, size=9.0)
    for values in rows:
        row = table.add_row()
        _prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_text(cell, values[index] if index < len(values) else "", size=8.5)
    _set_table_geometry(table, _column_widths(columns, rows))


def _formal_caption(value: str, section_title: str) -> str:
    text = re.sub(
        r"（(?:使用前人工核对适用性|非现场实景|AI生成[^）]*|岗位人员待确认|投标阶段建议计划)）",
        "",
        value,
    ).strip()
    if not text or "历史资料示例图" in text or "示例图片" in text:
        return f"{section_title or '施工技术'}示意图"
    return text.replace("历史资料", "").strip() or f"{section_title or '施工技术'}示意图"


def append_visual(
    document: Any,
    block: dict[str, Any],
    number: str,
    *,
    formal: bool = False,
    section_title: str = "",
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Cm, Pt  # type: ignore

    asset = block.get("asset") or {}
    value = str(asset.get("file_path") or block.get("source_path") or "")
    path = resolve_asset_path(value) if value else Path()
    if not value or not path.exists():
        paragraph = document.add_paragraph()
        paragraph.add_run(f"【待补图：{block.get('caption') or block.get('title') or block.get('block_type')}】").italic = True
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(15.8))
    caption_text = str(block.get("caption") or block.get("title") or asset.get("name") or "")
    if formal:
        caption_text = _formal_caption(caption_text, section_title)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run(f"图 {number} {caption_text}")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(9.5)
    disclaimer = str(asset.get("disclaimer") or (block.get("data") or {}).get("disclaimer") or "").strip()
    if formal and disclaimer and any(term in disclaimer for term in ("AI", "历史", "现场实景", "人工复核")):
        disclaimer = "本图为施工技术示意，具体实施以经审批方案、设计文件和现场条件为准。"
    if disclaimer:
        note = document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.space_after = Pt(8)
        note_run = note.add_run(disclaimer)
        note_run.italic = True
        note_run.font.name = "Microsoft YaHei"
        note_run.font.size = Pt(8.5)


def append_section_blocks(
    document: Any,
    blocks: list[dict[str, Any]],
    section_number: int,
    *,
    formal: bool = False,
    section_title: str = "",
) -> dict[str, int]:
    table_index = 0
    figure_index = 0
    rendered = 0
    for block in blocks:
        block_type = str(block.get("block_type") or "")
        if block_type in {"text", "page_break"}:
            if block_type == "page_break":
                document.add_page_break()
            continue
        if block_type == "table":
            table_index += 1
            append_native_table(document, block, f"{section_number}-{table_index}")
            rendered += 1
        elif block_type in {"organization_chart", "flow_chart", "gantt_chart", "image"}:
            figure_index += 1
            append_visual(
                document,
                block,
                f"{section_number}-{figure_index}",
                formal=formal,
                section_title=section_title,
            )
            rendered += 1
        elif block_type == "callout":
            text = str((block.get("data") or {}).get("text") or block.get("title") or "")
            paragraph = document.add_paragraph()
            paragraph.style = document.styles["Intense Quote"] if "Intense Quote" in document.styles else document.styles["Normal"]
            paragraph.add_run(text)
            rendered += 1
    return {"tables": table_index, "figures": figure_index, "rendered": rendered}


def collect_caption_index(
    blocks_by_section: dict[str, list[dict[str, Any]]],
    section_titles: list[str],
    *,
    formal: bool = False,
) -> dict[str, list[str]]:
    tables: list[str] = []
    figures: list[str] = []
    for section_number, section_title in enumerate(section_titles, 1):
        table_index = 0
        figure_index = 0
        for block in blocks_by_section.get(section_title, []):
            block_type = str(block.get("block_type") or "")
            caption = str(block.get("caption") or block.get("title") or "")
            if formal and block_type in {"organization_chart", "flow_chart", "gantt_chart", "image"}:
                caption = _formal_caption(caption, section_title)
            if block_type == "table":
                table_index += 1
                tables.append(f"表 {section_number}-{table_index} {caption}")
            elif block_type in {"organization_chart", "flow_chart", "gantt_chart", "image"}:
                figure_index += 1
                figures.append(f"图 {section_number}-{figure_index} {caption}")
    return {"tables": tables, "figures": figures}
