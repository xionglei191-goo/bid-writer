from __future__ import annotations

import json
import re
import shutil
import threading
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Literal

from pydantic import BaseModel, ConfigDict, Field

from ..ai_runtime import AiRuntime, KNOWLEDGE_FORMAL_REVIEW_PROMPT, KNOWLEDGE_REVIEW_PROMPT, KnowledgeReviewIssue, PromptSpec
from ..database import Database
from ..evaluation import RetrievalEvaluationService
from ..retrieval import HybridRetrievalService, tokenize
from ..utils import content_hash, normalize_text, parse_json, sha256_file, write_json_atomic, write_text_atomic
from .parsers import ParsedDocument
from .pipeline import KnowledgePipelineService
from .service import ARCHIVES, ASSET_EXTENSIONS, METADATA_ONLY, SUPPORTED_TEXT, KnowledgeService


TERMINAL_REASONS = {
    "duplicate",
    "metadata_only_excluded",
    "unsupported_excluded",
    "archive_expanded",
    "archive_empty_excluded",
    "unreadable_excluded",
    "encrypted_manual",
    "asset_archived_unlicensed",
    "asset_duplicate",
    "asset_registered",
    "content_duplicate",
    "knowledge_published",
    "no_reusable_knowledge",
    "manual_legal_pending",
}
SERVICE_ERROR_MARKERS = ("timeout", "timed out", "connection", "429", "500", "502", "503", "504", "http", "未配置")
LEGAL_PATTERNS = (
    ("confidentiality", r"(?:保密|机密|秘密|不得泄露)"),
    ("copyright", r"(?:著作权|版权|未经授权|禁止复制)"),
    ("binding_legal", r"(?:承担法律责任|违约责任|最终解释权|索赔|诉讼|仲裁)"),
)
BLOCK_PATTERNS = (
    ("personal_phone", r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    ("personal_email", r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b"),
    ("personal_id", r"(?<!\d)\d{17}[0-9Xx](?!\d)"),
    ("placeholder", r"(?:TODO|TBD|待确认|待补充|待提供|(?:\[|【)[^\]】]{0,40}(?:项目名称|公司名称|联系电话|填写|补充)[^\]】]{0,40}(?:\]|】))"),
    ("project_identity", r"(?:本项目|本工程|第[一二三四五六七八九十0-9]+标段|[\u4e00-\u9fffA-Za-z0-9（）()·—-]{6,60}(?:改扩建|新建|迁建|建设)项目)"),
    ("test_content", r"(?:测试投标单位|测试数据|示例公司|Lorem\s+ipsum)"),
    ("internal_path", r"(?:[A-Z]:\\|/workspace/|data/cache/archives/)"),
    ("provenance_disclosure", r"(?:AI生成|模型生成|历史项目素材|来源路径)"),
    ("absolute_commitment", r"(?:确保|保证).{0,12}(?:100%|零事故|绝不|全部)|完全避免"),
)
SCANNED_VISUAL_TITLE_PATTERN = re.compile(
    r"(?:封面|封皮|平面图|立面图|剖面图|布置图|系统图|示意图|配筋图|机构图|效果图|"
    r"曲线图|流程图|网络图|进度图|横道图|节点图|大样图|详图|总图|施工图)",
    re.IGNORECASE,
)
SCANNED_TEXTUAL_TITLE_PATTERN = re.compile(
    r"(?:审查|意见|记录|报告|说明|方案|文本|合同|要求|清单|验收|推荐|申报|表格)",
    re.IGNORECASE,
)
_RUN_LOCKS: dict[tuple[str, int, str], Any] = {}
_RUN_LOCKS_GUARD = threading.Lock()


RECOVERY_CONTRACT = "structured-body-summary-applicability-v2"


class RetiredPublicationRevision(BaseModel):
    model_config = ConfigDict(extra="forbid")

    candidate_index: Literal[0]
    decision: Literal["revise", "reject", "escalate"]
    confidence: float = Field(ge=0, le=1)
    issues: list[KnowledgeReviewIssue] = Field(max_length=20)
    corrected_content: str = Field(max_length=12000)
    corrected_summary: str = Field(max_length=600)
    corrected_applicability: str = Field(max_length=1500)


class RetiredPublicationRevisionOutput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    reviews: list[RetiredPublicationRevision] = Field(min_length=1, max_length=1)


RETIRED_PUBLICATION_REVISION_PROMPT = PromptSpec(
    key="knowledge.retired-publication-structured-revision", version="2.0.0",
    instructions=KNOWLEDGE_FORMAL_REVIEW_PROMPT.instructions,
    template=(
        "依据来源修订这份停用知识，逐项解决recovery_context中的原复核问题。失败说明只是复核证据，不能作为来源。"
        "只输出JSON对象reviews，且仅一条，完整字段为candidate_index=0、decision、confidence、issues、"
        "corrected_content、corrected_summary、corrected_applicability。可修复时decision=revise，并同时给出三个完整的新字段；"
        "证据不足时reject或escalate，三个修订字段均为空。corrected_content仅为技术正文，不能包裹标题、内容、正文、摘要、"
        "适用范围等字段标签或章节；corrected_summary必须是独立概括，不得复制正文或截取正文前200字；"
        "corrected_applicability必须具体且有来源支持，不得继承原来范围过宽的适用性。三者的条件和适用范围必须一致。"
        "保持输入标题，不要另造标题；禁止增加来源之外的事实、参数、承诺。\n\n"
        "来源章节：\n{section_text}\n\n待修订知识：\n{candidate_json}"
    ),
    output_model=RetiredPublicationRevisionOutput,
)


def _recovery_review_spec(formal: bool) -> PromptSpec:
    base = KNOWLEDGE_FORMAL_REVIEW_PROMPT if formal else KNOWLEDGE_REVIEW_PROMPT
    return PromptSpec(
        key="knowledge.retired-publication-formal" if formal else "knowledge.retired-publication-review",
        version="2.0.0", instructions=base.instructions,
        template=base.template + (
            "\n\n本次采用完整结构化修订契约v2。content是将同一summary、applicability与技术正文组合后的完整发布稿。"
            "独立核验这三个字段与来源及彼此一致，特别检查summary是否扩大范围、applicability是否遗漏条件；"
            "字段缺失、元数据与正文矛盾、摘要复制正文或范围扩大时不得pass。没有遗留旧摘要或旧范围；"
            "不得根据历史修订或通过记录推定本稿通过。"
        ),
        output_model=base.output_model,
    )


class FinalizationInterrupted(RuntimeError):
    """The operator stopped the run at a durable finalization checkpoint."""

    def __init__(self, message: str, *, cancel_requested: bool = False) -> None:
        super().__init__(message)
        self.cancel_requested = cancel_requested


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def iso_now() -> str:
    return utc_now().isoformat(timespec="seconds")


class CorpusCompletionService:
    def __init__(
        self,
        db: Database,
        knowledge: KnowledgeService,
        pipeline: KnowledgePipelineService,
        retrieval: HybridRetrievalService,
        ai_runtime: AiRuntime,
        evaluation: RetrievalEvaluationService | None = None,
    ) -> None:
        self.db = db
        self.knowledge = knowledge
        self.pipeline = pipeline
        self.retrieval = retrieval
        self.ai_runtime = ai_runtime
        self.evaluation = evaluation
        self.dispatch: Callable[[int, int, int, str], None] | None = None

    @contextmanager
    def _run_mutex(self, run_id: int, purpose: str):
        identity = self.db.database_url or str(self.db.path.resolve())
        with _RUN_LOCKS_GUARD:
            lock = _RUN_LOCKS.setdefault((identity, run_id, purpose), threading.Lock())
        if not lock.acquire(blocking=False):
            yield False
            return
        try:
            if self.db.backend == "postgresql":
                # A session lock spans the short transactions used by AI/index
                # checkpoints and is released even if the worker process exits.
                key = run_id * 100_000 + (9_138 if purpose == "finalize" else 9_139)
                with self.db.connect() as conn:
                    row = conn.execute("SELECT pg_try_advisory_lock(CAST(? AS BIGINT)) AS locked", (key,)).fetchone()
                    if not row or not row["locked"]:
                        yield False
                        return
                    try:
                        yield True
                    finally:
                        # Pooled connections must never retain a session lock.
                        conn.execute("SELECT pg_advisory_unlock(CAST(? AS BIGINT))", (key,))
            else:
                yield True
        finally:
            lock.release()

    def create_run(self, created_by: int | None = None) -> dict[str, Any]:
        active = self.db.row("SELECT id FROM corpus_runs WHERE status IN ('pending','running','paused') ORDER BY id DESC LIMIT 1")
        if active:
            return self.get_run(int(active["id"]))
        rows = self.db.rows(
            "SELECT id,sha256,extension,status,source_kind,duplicate_of,error_message FROM source_files ORDER BY id"
        )
        snapshot_hash = content_hash("|".join(f"{row['id']}:{row['sha256']}" for row in rows))
        policy = {
            "normalize_concurrency": 2,
            "ocr_concurrency": 1,
            "ai_concurrency": 1,
            "index_concurrency": 1,
            "max_attempts": 3,
            "retry_minutes": [1, 5, 15],
            "disk_min_bytes": 100 * 1024**3,
            "disk_min_ratio": 0.10,
            "service_consecutive_error_limit": 5,
            "recent_window": 20,
            "recent_failure_rate": 0.20,
            "ocr_provider": "paddleocr_aistudio_allowed",
            "technical_review": "ai",
            "defer_large_sources": True,
            "large_source_bytes": 150 * 1024 * 1024,
        }
        self._supersede_legacy_rule_units()
        with self.db.connect() as conn:
            cursor = conn.execute(
                "INSERT INTO corpus_runs(snapshot_hash,policy_json,created_by) VALUES (?,?,?)",
                (snapshot_hash, json.dumps(policy, ensure_ascii=False), created_by),
            )
            run_id = int(cursor.lastrowid)
            for row in rows:
                kind, status, stage, reason = self._initial_state(row)
                conn.execute(
                    """
                    INSERT INTO corpus_run_items(run_id,source_id,source_hash,item_kind,stage,status,terminal_reason,completed_at)
                    VALUES (?,?,?,?,?,?,?,?)
                    """,
                    (run_id, row["id"], row["sha256"], kind, stage, status, reason, iso_now() if status == "terminal" else None),
                )
            conn.execute(
                "UPDATE corpus_runs SET status='running',stage='normalize',started_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                (run_id,),
            )
        self._refresh(run_id)
        self._dispatch_next(run_id)
        return self.get_run(run_id)

    @staticmethod
    def _initial_state(row: dict[str, Any]) -> tuple[str, str, str, str]:
        extension = str(row.get("extension") or "").lower()
        source_status = str(row.get("status") or "")
        if row.get("duplicate_of") or source_status == "duplicate":
            return "duplicate", "terminal", "complete", "duplicate"
        if extension in SUPPORTED_TEXT:
            return "text", "pending", "normalize", ""
        if extension in ASSET_EXTENSIONS:
            return "asset", "pending", "governance", ""
        if extension in ARCHIVES:
            if source_status == "failed":
                return "archive", "terminal", "complete", "unreadable_excluded"
            return "archive", "terminal", "complete", "archive_expanded"
        if extension in METADATA_ONLY or source_status == "metadata_only":
            return "metadata", "terminal", "complete", "metadata_only_excluded"
        return "unsupported", "terminal", "complete", "unsupported_excluded"

    def rebuild_run_from_persisted_results(self, previous_run_id: int, created_by: int | None = None) -> dict[str, Any]:
        """Recreate a damaged run ledger without discarding durable outcomes.

        Standard documents, final AI candidates/publications, asset governance,
        and legal task dispositions are authoritative durable records.  Only
        unfinished stages are re-queued; the previous ledger remains auditable.
        """
        previous = self._raw_run(previous_run_id)
        if str(previous.get("status") or "") in {"pending", "running", "paused"}:
            with self.db.connect() as conn:
                conn.execute(
                    """
                    UPDATE corpus_runs SET status='cancelled',stage='operator_recovery',pause_reason=?,
                        completed_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP WHERE id=?
                    """,
                    ("运行台账操作异常；已依据持久化成果重建，源文件与成果未受影响", previous_run_id),
                )
        rows = self.db.rows(
            "SELECT id,sha256,extension,status,source_kind,duplicate_of,error_message FROM source_files ORDER BY id"
        )
        snapshot_hash = content_hash("|".join(f"{row['id']}:{row['sha256']}" for row in rows))
        policy = parse_json(previous.get("policy_json"), {})
        with self.db.connect() as conn:
            cursor = conn.execute(
                """
                INSERT INTO corpus_runs(snapshot_hash,status,stage,policy_json,checkpoint_json,created_by,
                    started_at,updated_at)
                VALUES (?,'running','normalize',?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)
                """,
                (
                    snapshot_hash,
                    json.dumps(policy, ensure_ascii=False),
                    json.dumps(
                        {
                            "recovered_from_run_id": previous_run_id,
                            "recovery_method": "persisted_result_reconciliation",
                            "recovered_at": iso_now(),
                            "previous_checkpoint": parse_json(previous.get("checkpoint_json"), {}),
                        },
                        ensure_ascii=False,
                    ),
                    created_by,
                ),
            )
            run_id = int(cursor.lastrowid)

        for row in rows:
            source_id = int(row["id"])
            kind, status, stage, reason = self._initial_state(row)
            processing_job_id: int | None = None
            document_id: int | None = None
            pipeline_run_id: int | None = None
            checkpoint: dict[str, Any] = {"reconciled_from_persisted_results": True}
            if kind == "text":
                document = self.db.row("SELECT id FROM standard_documents WHERE source_id=?", (source_id,))
                job = self.db.row("SELECT id,status FROM processing_jobs WHERE source_id=? ORDER BY id DESC LIMIT 1", (source_id,))
                processing_job_id = int(job["id"]) if job else None
                if document:
                    document_id = int(document["id"])
                    disposition = self.db.row(
                        """
                        SELECT pipeline_run_id,
                            COUNT(DISTINCT decision_stage) AS stages,
                            SUM(CASE WHEN decision='not_reusable' AND confidence>=0.95 THEN 1 ELSE 0 END) AS verified
                        FROM knowledge_source_dispositions
                        WHERE document_id=? AND decision_stage IN ('extraction','independent_review')
                        GROUP BY pipeline_run_id ORDER BY pipeline_run_id DESC LIMIT 1
                        """,
                        (document_id,),
                    )
                    candidate = self.db.row(
                        """
                        SELECT c.pipeline_run_id,
                            SUM(CASE WHEN c.status='accepted' THEN 1 ELSE 0 END) AS accepted,
                            SUM(CASE WHEN c.status='rejected' THEN 1 ELSE 0 END) AS rejected,
                            SUM(CASE WHEN c.status IN ('ready','needs_review') THEN 1 ELSE 0 END) AS unfinished
                        FROM knowledge_ai_candidates c
                        WHERE c.document_id=? AND c.status<>'superseded'
                        GROUP BY c.pipeline_run_id ORDER BY c.pipeline_run_id DESC LIMIT 1
                        """,
                        (document_id,),
                    )
                    if disposition and int(disposition.get("stages") or 0) == 2 and int(disposition.get("verified") or 0) == 2:
                        pipeline_run_id = int(disposition["pipeline_run_id"])
                        status, stage, reason = "terminal", "complete", "no_reusable_knowledge"
                        checkpoint.update({"reason": "two_stage_ai_verified_not_reusable", "source_disposition_recovered": True})
                    elif candidate and not int(candidate.get("unfinished") or 0):
                        pipeline_run_id = int(candidate["pipeline_run_id"])
                        status, stage = "terminal", "complete"
                        if int(candidate.get("accepted") or 0):
                            reason = "knowledge_published"
                        else:
                            manual = self.db.row(
                                "SELECT COUNT(*) AS count FROM governance_tasks WHERE source_id=? AND status='open'",
                                (source_id,),
                            ) or {}
                            reason = "manual_legal_pending" if int(manual.get("count") or 0) else "no_reusable_knowledge"
                        checkpoint.update(
                            {
                                "published": int(candidate.get("accepted") or 0),
                                "rejected": int(candidate.get("rejected") or 0),
                            }
                        )
                    else:
                        status, stage, reason = "pending", "ai", ""
                elif job and str(job.get("status") or "") == "waiting_ocr":
                    status, stage, reason = "pending", "ocr", ""
            elif kind == "asset":
                asset = self.db.row(
                    """
                    SELECT governance_status,COUNT(*) AS count FROM knowledge_assets
                    WHERE source_id=? GROUP BY governance_status ORDER BY COUNT(*) DESC LIMIT 1
                    """,
                    (source_id,),
                )
                if asset:
                    governance = str(asset.get("governance_status") or "")
                    status, stage = "terminal", "complete"
                    if governance == "duplicate":
                        reason = "asset_duplicate"
                    elif governance == "approved":
                        reason = "asset_registered"
                    elif governance == "pending_legal":
                        reason = "manual_legal_pending"
                    else:
                        reason = "asset_archived_unlicensed"
            with self.db.connect() as conn:
                conn.execute(
                    """
                    INSERT INTO corpus_run_items(
                        run_id,source_id,source_hash,item_kind,stage,status,terminal_reason,
                        processing_job_id,document_id,pipeline_run_id,checkpoint_json,completed_at
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        run_id,
                        source_id,
                        row["sha256"],
                        kind,
                        stage,
                        status,
                        reason,
                        processing_job_id,
                        document_id,
                        pipeline_run_id,
                        json.dumps(checkpoint, ensure_ascii=False),
                        iso_now() if status == "terminal" else None,
                    ),
                )
        old_tasks = self.db.rows("SELECT * FROM governance_tasks WHERE run_id=?", (previous_run_id,))
        with self.db.connect() as conn:
            for task in old_tasks:
                cursor = conn.execute(
                    """
                    INSERT INTO governance_tasks(
                        run_id,source_id,task_key,task_type,severity,status,title,message,requires_human,
                        resolution,resolved_by,created_at,resolved_at
                    )
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        run_id,
                        task.get("source_id"),
                        task["task_key"],
                        task["task_type"],
                        task["severity"],
                        task["status"],
                        task["title"],
                        task["message"],
                        task["requires_human"],
                        task.get("resolution") or "",
                        task.get("resolved_by"),
                        task["created_at"],
                        task.get("resolved_at"),
                    ),
                )
                new_task_id = int(cursor.lastrowid)
                for link in self.db.rows("SELECT source_id FROM governance_task_sources WHERE task_id=?", (task["id"],)):
                    conn.execute(
                        "INSERT INTO governance_task_sources(task_id,source_id) VALUES (?,?) ON CONFLICT(task_id,source_id) DO NOTHING",
                        (new_task_id, link["source_id"]),
                    )
        self._refresh(run_id)
        self._dispatch_next(run_id)
        return self.get_run(run_id)

    def run_tick(self, run_id: int, progress=None, cancelled=None, preferred_stage: str = "") -> dict[str, Any]:
        progress = progress or (lambda *_args, **_kwargs: None)
        cancelled = cancelled or (lambda: False)
        run = self._raw_run(run_id)
        if run["status"] == "paused" and str(run.get("pause_reason") or "").startswith("外部服务错误") and preferred_stage == "__recovery__":
            pause_started_at = str(run.get("updated_at") or iso_now())
            checkpoint = parse_json(run.get("checkpoint_json"), {})
            recovery_stage = str(checkpoint.get("service_pause_stage") or "")
            if not recovery_stage:
                failed_stage = self.db.row(
                    """
                    SELECT stage,COUNT(*) AS count FROM corpus_run_items
                    WHERE run_id=? AND status='retrying' AND error_message<>''
                    GROUP BY stage ORDER BY count DESC LIMIT 1
                    """,
                    (run_id,),
                ) or {}
                recovery_stage = str(failed_stage.get("stage") or "")
            self._recover_stale_items(run_id)
            if recovery_stage:
                with self.db.connect() as conn:
                    conn.execute(
                        """
                        UPDATE corpus_run_items SET status='retrying',error_code='service_recovery',
                            error_message='熔断前占位已回到检查点，等待外部服务恢复',
                            next_retry_at=NULL,updated_at=CURRENT_TIMESTAMP
                        WHERE run_id=? AND stage=? AND status='running' AND updated_at<=?
                        """,
                        (run_id, recovery_stage, pause_started_at),
                    )
            if recovery_stage == "ai":
                # 推理型模型会消耗推理 token，16 的预算会导致正文为空、探针永远失败。
                probe = self.ai_runtime.llm.generate(
                    "你是服务可用性探针。只返回 OK。",
                    "返回 OK。",
                    2048,
                )
                if probe.get("error") or not str(probe.get("content") or "").strip():
                    checkpoint["service_recovery_probe"] = {
                        "stage": recovery_stage,
                        "status": "failed",
                        "checked_at": iso_now(),
                        "error": str(probe.get("error") or "empty_model_response")[:300],
                    }
                    with self.db.connect() as conn:
                        conn.execute(
                            "UPDATE corpus_runs SET checkpoint_json=?,pause_reason=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                            (
                                json.dumps(checkpoint, ensure_ascii=False),
                                "外部服务错误达到熔断阈值: ai（恢复探针未通过）",
                                run_id,
                            ),
                        )
                    if self.dispatch:
                        self.dispatch(run_id, 15 * 60 * 1000, 1, "__recovery__")
                    return self.get_run(run_id)
                checkpoint["service_recovery_probe"] = {
                    "stage": recovery_stage,
                    "status": "passed",
                    "checked_at": iso_now(),
                    "model": str(probe.get("model") or ""),
                    "latency_ms": int(probe.get("latency_ms") or 0),
                }
            with self.db.connect() as conn:
                conn.execute(
                    "UPDATE corpus_runs SET status='running',pause_reason='',consecutive_errors=0,recent_results_json='[]',checkpoint_json=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                    (json.dumps(checkpoint, ensure_ascii=False), run_id),
                )
            run = self._raw_run(run_id)
            preferred_stage = recovery_stage
        if run["status"] not in {"pending", "running"}:
            return self.get_run(run_id)
        if cancelled():
            self._set_run(run_id, "cancelled", "cancelled", "用户取消")
            return self.get_run(run_id)
        self._recover_stale_items(run_id)
        self._requeue_uncovered_batch_sources(run_id)
        disk = shutil.disk_usage(self.knowledge.settings.knowledge_root)
        policy = parse_json(run.get("policy_json"), {})
        if disk.free < max(int(policy.get("disk_min_bytes", 100 * 1024**3)), int(disk.total * float(policy.get("disk_min_ratio", 0.1)))):
            self._set_run(run_id, "paused", str(run["stage"]), "磁盘空间低于安全阈值")
            return self.get_run(run_id)
        item = self._claim_next_item(run_id, preferred_stage)
        if not item:
            waiting = self.db.row("SELECT COUNT(*) AS count FROM corpus_run_items WHERE run_id=? AND status IN ('pending','retrying','running')", (run_id,))
            if int((waiting or {}).get("count", 0)):
                self._dispatch_next(run_id, 60_000, current_running_stage=preferred_stage)
                return self.get_run(run_id)
            return self._finalize_run(run_id, progress, cancelled)
        progress(str(item["stage"]), 5, f"处理 {item['file_name']}", {"source_id": item["source_id"]})
        active_items = [item]
        try:
            if item["stage"] == "normalize":
                self._normalize_item(item)
            elif item["stage"] == "ocr":
                self._ocr_item(item)
            elif item["stage"] == "governance":
                self._govern_asset(item)
            elif item["stage"] == "ai":
                active_items = self._claim_ai_batch(item)
                self._process_ai_batch(active_items, progress=progress, cancelled=cancelled)
            else:
                self._terminal(item, "no_reusable_knowledge")
            self._record_result(
                run_id,
                True,
                "",
                service_stage=str(item["stage"]) if str(item["stage"]) in {"ai", "ocr"} else "",
            )
        except Exception as exc:  # noqa: BLE001
            service_error = str(item["stage"]) in {"ai", "ocr"} and any(
                marker in f"{type(exc).__name__}: {exc}".lower() for marker in SERVICE_ERROR_MARKERS
            )
            for active_item in active_items:
                state = self.db.row("SELECT status FROM corpus_run_items WHERE id=?", (active_item["id"],)) or {}
                if state.get("status") == "running":
                    self._handle_failure(active_item, exc, service_error=service_error)
            self._record_result(
                run_id,
                not service_error,
                f"{type(exc).__name__}: {exc}",
                service_stage=str(item["stage"]) if service_error else "",
            )
        summary = self._refresh(run_id)
        progress(str(summary["stage"]), int(summary["progress"]), "全库批次检查点已保存", summary.get("counters"))
        current = self._raw_run(run_id)
        if current["status"] == "running":
            remaining = self.db.row(
                "SELECT COUNT(*) AS count FROM corpus_run_items WHERE run_id=? AND status IN ('pending','retrying','running')",
                (run_id,),
            )
            if int((remaining or {}).get("count", 0)) == 0:
                return self._finalize_run(run_id, progress, cancelled)
            self._dispatch_next(run_id, current_running_stage=str(item["stage"]))
        return self.get_run(run_id)

    def _claim_next_item(self, run_id: int, preferred_stage: str = "") -> dict[str, Any] | None:
        if preferred_stage in {"__recovery__", "__finalize__"}:
            preferred_stage = ""
        policy = parse_json(self._raw_run(run_id).get("policy_json"), {})
        capacities = {
            "normalize": int(policy.get("normalize_concurrency", 2)),
            "ocr": int(policy.get("ocr_concurrency", 1)),
            "governance": int(policy.get("normalize_concurrency", 2)),
            "ai": int(policy.get("ai_concurrency", 1)),
        }
        stages = [preferred_stage] if preferred_stage else ["normalize", "ocr", "governance", "ai"]
        for _attempt in range(5):
            with self.db.connect() as conn:
                if self.db.backend == "postgresql":
                    conn.execute("SELECT pg_advisory_xact_lock(CAST(? AS BIGINT))", (run_id * 100_000 + 9_137,))
                run_state = conn.execute("SELECT status FROM corpus_runs WHERE id=?", (run_id,)).fetchone()
                if not run_state or run_state["status"] not in {"pending", "running"}:
                    return None
                running = {
                    str(row["stage"]): int(row["count"])
                    for row in conn.execute(
                        "SELECT stage,COUNT(*) AS count FROM corpus_run_items WHERE run_id=? AND status='running' GROUP BY stage",
                        (run_id,),
                    ).fetchall()
                }
                for stage in stages:
                    if running.get(stage, 0) >= max(1, capacities.get(stage, 1)):
                        continue
                    item_row = conn.execute(
                        """
                        SELECT i.*,s.absolute_path,s.relative_path,s.file_name,s.extension,s.source_kind,s.parent_source_id,
                            s.status AS source_status,s.error_message AS source_error,s.industry,
                            d.char_count AS document_char_count
                        FROM corpus_run_items i JOIN source_files s ON s.id=i.source_id
                        LEFT JOIN standard_documents d ON d.id=i.document_id
                        WHERE i.run_id=? AND i.stage=? AND i.status IN ('pending','retrying')
                          AND (i.next_retry_at IS NULL OR i.next_retry_at<=?)
                        ORDER BY CASE WHEN i.stage IN ('normalize','ocr') THEN s.size_bytes ELSE 0 END,
                            CASE WHEN i.stage='ai' THEN COALESCE(d.char_count,2147483647) ELSE 0 END,i.id
                        LIMIT 1
                        """,
                        (run_id, stage, iso_now()),
                    ).fetchone()
                    if not item_row:
                        continue
                    item = dict(item_row)
                    claimed = conn.execute(
                        "UPDATE corpus_run_items SET status='running',attempt_count=attempt_count+1,updated_at=CURRENT_TIMESTAMP WHERE id=? AND status IN ('pending','retrying')",
                        (item["id"],),
                    )
                    if claimed.rowcount:
                        conn.execute("UPDATE corpus_runs SET stage=?,updated_at=CURRENT_TIMESTAMP WHERE id=?", (stage, run_id))
                        return item
        return None

    def _claim_ai_batch(self, item: dict[str, Any]) -> list[dict[str, Any]]:
        """Claim up to twelve short AI items without increasing AI concurrency.

        The items run in one model request pair and remain independently
        checkpointed.  Long documents keep the existing chunked path.
        """
        document_id = int(item.get("document_id") or 0)
        if not document_id:
            return [item]
        if parse_json(item.get("checkpoint_json"), {}).get("requires_individual_ai"):
            return [item]
        current = self.db.row("SELECT char_count FROM standard_documents WHERE id=?", (document_id,)) or {}
        current_chars = int(current.get("char_count") or 0)
        if current_chars <= 0 or current_chars > 24_000:
            return [item]
        claimed = [item]
        used_chars = current_chars
        candidates = self.db.rows(
            """
            SELECT i.*,s.absolute_path,s.relative_path,s.file_name,s.extension,s.source_kind,s.parent_source_id,
                s.status AS source_status,s.error_message AS source_error,s.industry,d.char_count AS document_char_count
            FROM corpus_run_items i
            JOIN source_files s ON s.id=i.source_id
            JOIN standard_documents d ON d.id=i.document_id
            WHERE i.run_id=? AND i.stage='ai' AND i.status IN ('pending','retrying')
              AND (i.next_retry_at IS NULL OR i.next_retry_at<=?)
              AND i.checkpoint_json NOT LIKE '%"requires_individual_ai": true%'
              AND d.char_count<=24000 AND i.id<>?
            ORDER BY i.id LIMIT 12
            """,
            (item["run_id"], iso_now(), item["id"]),
        )
        for candidate in candidates:
            if parse_json(candidate.get("checkpoint_json"), {}).get("requires_individual_ai"):
                continue
            candidate_chars = int(candidate.get("document_char_count") or 0)
            if candidate_chars <= 0 or used_chars + candidate_chars > 82_000:
                continue
            with self.db.connect() as conn:
                updated = conn.execute(
                    """
                    UPDATE corpus_run_items SET status='running',attempt_count=attempt_count+1,
                        updated_at=CURRENT_TIMESTAMP
                    WHERE id=? AND status IN ('pending','retrying')
                    """,
                    (candidate["id"],),
                )
            if updated.rowcount:
                claimed.append(candidate)
                used_chars += candidate_chars
            if len(claimed) >= 12:
                break
        return claimed

    def _dispatch_next(self, run_id: int, delay_ms: int = 0, *, current_running_stage: str = "") -> None:
        if not self.dispatch:
            return
        with self._run_mutex(run_id, "dispatch") as acquired:
            if acquired:
                self._dispatch_available(run_id, delay_ms)

    def _dispatch_available(self, run_id: int, delay_ms: int) -> None:
        run = self._raw_run(run_id)
        if run["status"] not in {"pending", "running"}:
            return
        pending = self.db.rows(
            """
            SELECT stage,COUNT(*) AS count FROM corpus_run_items
            WHERE run_id=? AND status IN ('pending','retrying')
            GROUP BY stage ORDER BY CASE stage WHEN 'normalize' THEN 1 WHEN 'ocr' THEN 2 WHEN 'governance' THEN 3 WHEN 'ai' THEN 4 ELSE 5 END
            """,
            (run_id,),
        )
        if not pending:
            unfinished = self.db.row(
                "SELECT COUNT(*) AS count FROM corpus_run_items WHERE run_id=? AND status<>'terminal'", (run_id,)
            ) or {}
            if int(unfinished.get("count") or 0):
                return
            active = self.db.row(
                """
                SELECT id FROM app_jobs WHERE job_type='knowledge.corpus.tick' AND target_id=?
                    AND status IN ('pending','retrying','running') LIMIT 1
                """,
                (f"{run_id}:__finalize__",),
            )
            if not active:
                latest_job = self._latest_corpus_job(run_id)
                checkpoint = parse_json(run.get("checkpoint_json"), {}).get("finalization", {})
                if (
                    latest_job and latest_job["status"] == "failed"
                    and int(checkpoint.get("acknowledged_failed_job_id") or 0) != int(latest_job["id"])
                ):
                    # A scheduler restart must not silently turn a failed final
                    # audit into thousands of fresh external-model requests.
                    error = {"type": latest_job.get("error_code") or "background_job_failed", "message": str(latest_job.get("error_message") or "后台任务失败")[:1000], "job_id": latest_job["id"]}
                    with self.db.connect() as conn:
                        conn.execute(
                            "UPDATE corpus_runs SET status='paused',stage='finalization_recovery',pause_reason=?,updated_at=CURRENT_TIMESTAMP WHERE id=? AND status IN ('pending','running')",
                            (f"收尾前发现上一后台任务失败，等待显式恢复：{error['message']}", run_id),
                        )
                    self._finalization_checkpoint(run_id, "finalization_recovery", 95, status="failed", error=error)
                    self._write_finalization_failure_report(run_id, {"error": error})
                    return
                self.dispatch(run_id, delay_ms, 1, "__finalize__")
            return
        policy = parse_json(run.get("policy_json"), {})
        for row in pending:
            stage = str(row["stage"])
            if stage == "__recovery__":
                continue
            capacity = {
                "normalize": int(policy.get("normalize_concurrency", 2)),
                "ocr": int(policy.get("ocr_concurrency", 1)),
                "governance": int(policy.get("normalize_concurrency", 2)),
                "ai": int(policy.get("ai_concurrency", 1)),
            }.get(stage, 1)
            desired = max(1, min(capacity, int(row["count"])))
            self.dispatch(run_id, delay_ms, desired, stage)

    def _latest_corpus_job(self, run_id: int) -> dict[str, Any] | None:
        return self.db.row(
            """
            SELECT id,status,error_code,error_message FROM app_jobs
            WHERE job_type='knowledge.corpus.tick' AND (target_id=? OR target_id LIKE ?)
            ORDER BY id DESC LIMIT 1
            """,
            (str(run_id), f"{run_id}:%"),
        )

    def _normalize_item(self, item: dict[str, Any]) -> None:
        if str(item.get("file_name") or "").startswith("~$"):
            with self.db.connect() as conn:
                conn.execute("UPDATE source_files SET status='metadata_only',updated_at=CURRENT_TIMESTAMP WHERE id=?", (item["source_id"],))
            self._terminal(item, "metadata_only_excluded", checkpoint={"reason": "office_lock_file"})
            return
        existing = self.db.row("SELECT * FROM standard_documents WHERE source_id=?", (item["source_id"],))
        if (
            existing
            and str(existing.get("source_hash") or "") == str(item["source_hash"])
            and str(existing.get("parser_version") or "") == "2"
        ):
            self._advance(item, "ai", document_id=int(existing["id"]))
            return
        job = self.db.row(
            "SELECT * FROM processing_jobs WHERE source_id=? ORDER BY id DESC LIMIT 1",
            (item["source_id"],),
        )
        if job and str(job.get("status") or "") == "skipped":
            checkpoint = parse_json(job.get("checkpoint_json"), {})
            if checkpoint.get("exclusion_reason") == "historical_single_page_visual_archived":
                self._terminal(
                    item,
                    "no_reusable_knowledge",
                    processing_job_id=int(job["id"]),
                    checkpoint={"reason": checkpoint["exclusion_reason"], "gate": "scanned_visual_archive"},
                )
                return
        if not job or str(job["status"]) in {"failed", "completed", "skipped", "cancelled"}:
            created = self.knowledge.create_jobs([int(item["source_id"])], 1)
            if created["job_ids"]:
                job = self.db.row("SELECT * FROM processing_jobs WHERE id=?", (created["job_ids"][0],))
            elif str((job or {}).get("status") or "") == "completed":
                document = self.db.row("SELECT id FROM standard_documents WHERE source_id=?", (item["source_id"],))
                if document:
                    self._advance(item, "ai", processing_job_id=int(job["id"]), document_id=int(document["id"]))
                    return
        if not job:
            raise RuntimeError("无法创建标准化任务")
        if str(job["status"]) == "waiting_ocr":
            self._advance(item, "ocr", processing_job_id=int(job["id"]))
            return
        result = self.knowledge.run_job(int(job["id"]), create_rule_units=False)
        if result.get("status") == "waiting_ocr":
            self._advance(item, "ocr", processing_job_id=int(job["id"]))
            return
        if result.get("status") != "completed":
            raise RuntimeError(str(result.get("error") or f"标准化状态为{result.get('status')}"))
        if result.get("duplicate_of"):
            self._terminal(item, "content_duplicate", processing_job_id=int(job["id"]), document_id=int(result["document_id"]))
            return
        self._advance(item, "ai", processing_job_id=int(job["id"]), document_id=int(result["document_id"]))

    def _ocr_item(self, item: dict[str, Any]) -> None:
        job_id = int(item.get("processing_job_id") or 0)
        if not job_id:
            job = self.db.row("SELECT id FROM processing_jobs WHERE source_id=? ORDER BY id DESC LIMIT 1", (item["source_id"],))
            if not job:
                raise RuntimeError("OCR任务不存在")
            job_id = int(job["id"])
        if self._is_historical_single_page_visual(item, job_id):
            checkpoint = {
                "page_count": 1,
                "exclusion_reason": "historical_single_page_visual_archived",
                "governance": "archived_not_searchable",
            }
            source = self.db.row("SELECT * FROM source_files WHERE id=?", (item["source_id"],)) or item
            path = self.knowledge.resolve_source_path(source)
            self._upsert_asset(
                source,
                path,
                "scanned_pdf_visual",
                str(source.get("sha256") or item.get("source_hash") or sha256_file(path)),
                "",
                0,
                0,
                {"extension": ".pdf", "page_count": 1, "classification": "historical_visual"},
                "archived",
                "历史单页图纸，无可复用正文且不进入检索",
            )
            with self.db.connect() as conn:
                conn.execute(
                    "UPDATE processing_jobs SET status='skipped',progress=100,current_step='archived_visual',checkpoint_json=?,error_message=NULL,finished_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                    (json.dumps(checkpoint, ensure_ascii=False), job_id),
                )
            self._terminal(
                item,
                "no_reusable_knowledge",
                processing_job_id=job_id,
                checkpoint={"reason": checkpoint["exclusion_reason"], "gate": "scanned_visual_archive"},
            )
            return
        result = self.knowledge.run_ocr(job_id, create_rule_units=False)
        if result.get("status") == "waiting_ocr":
            self._advance(
                item,
                "ocr",
                processing_job_id=job_id,
                completed_pages=int(result.get("completed_pages") or 0),
                page_count=int(result.get("page_count") or 0),
                completed_chunks=int(result.get("completed_chunks") or 0),
                total_chunks=int(result.get("total_chunks") or 0),
            )
            return
        if result.get("status") != "completed":
            raise RuntimeError(str(result.get("error") or f"OCR状态为{result.get('status')}"))
        if result.get("duplicate_of"):
            self._terminal(item, "content_duplicate", processing_job_id=job_id, document_id=int(result["document_id"]))
            return
        self._advance(item, "ai", processing_job_id=job_id, document_id=int(result["document_id"]))

    def _is_historical_single_page_visual(self, item: dict[str, Any], job_id: int) -> bool:
        job = self.db.row("SELECT checkpoint_json FROM processing_jobs WHERE id=?", (job_id,)) or {}
        checkpoint = parse_json(job.get("checkpoint_json"), {})
        if int(checkpoint.get("page_count") or 0) != 1:
            return False
        if str(item.get("source_kind") or "") == "system_generated":
            return False
        title = normalize_text(Path(str(item.get("file_name") or "")).stem)
        return bool(SCANNED_VISUAL_TITLE_PATTERN.search(title)) and not bool(SCANNED_TEXTUAL_TITLE_PATTERN.search(title))

    def _govern_asset(self, item: dict[str, Any]) -> None:
        source = self.db.row("SELECT * FROM source_files WHERE id=?", (item["source_id"],)) or item
        path = self.knowledge.resolve_source_path(source)
        if not path.exists():
            raise FileNotFoundError(str(path))
        extension = path.suffix.lower()
        metadata: dict[str, Any] = {"extension": extension, "size_bytes": path.stat().st_size, "classification": "internal"}
        digest = sha256_file(path)
        width = height = 0
        perceptual = ""
        asset_type = "document_asset"
        if extension in {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff"}:
            asset_type = "image"
            from PIL import Image, ImageFile, ImageStat
            import imagehash

            ImageFile.LOAD_TRUNCATED_IMAGES = True
            with Image.open(path) as image:
                width, height = image.size
                perceptual = str(imagehash.phash(image.convert("RGB")))
                stat = ImageStat.Stat(image.convert("L"))
                metadata.update({"sharpness_proxy": round(float(stat.var[0]), 4), "format": image.format or ""})
            duplicate = self.db.row(
                "SELECT id,source_id,perceptual_hash FROM knowledge_assets WHERE content_hash=? OR (perceptual_hash<>'' AND perceptual_hash=?) ORDER BY id LIMIT 1",
                (digest, perceptual),
            )
            if duplicate and int(duplicate.get("source_id") or 0) != int(source["id"]):
                self._upsert_asset(source, path, asset_type, digest, perceptual, width, height, metadata, "duplicate", f"与素材{duplicate['id']}重复")
                self._terminal(item, "asset_duplicate")
                return
        elif extension in {".xlsx", ".xls", ".pptx"}:
            text = self._extract_office_asset(path)
            metadata["extracted_chars"] = len(text)
            if len(text) >= 80 and extension in {".xlsx", ".pptx"}:
                parsed = ParsedDocument(path.stem, f"# {path.stem}\n\n{text}", f"asset-{extension[1:]}", 0)
                job = {**source, "source_id": source["id"], "industry": source.get("industry") or "通用", "sha256": source["sha256"]}
                result = self.knowledge._store_document(job, parsed, create_rule_units=False)
                self._upsert_asset(source, path, asset_type, digest, perceptual, width, height, metadata, "pending_legal", "来源授权待确认")
                self._legal_task(int(item["run_id"]), int(item["source_id"]), "asset_license", "素材来源授权待确认", "Office素材提取内容在授权确认前不进入正式检索。")
                self._terminal(item, "manual_legal_pending", document_id=int(result["document_id"]))
                return
        if str(source.get("source_kind") or "") == "system_generated":
            self._upsert_asset(source, path, asset_type, digest, perceptual, width, height, metadata, "approved", "程序生成且来源安全")
            self._terminal(item, "asset_registered")
            return
        self._upsert_asset(source, path, asset_type, digest, perceptual, width, height, metadata, "archived", "历史或来源授权不明")
        self._legal_task(int(item["run_id"]), int(item["source_id"]), "asset_license", "素材来源授权待确认", "历史或来源不明素材已归档，授权确认前不进入正式库。")
        self._terminal(item, "asset_archived_unlicensed")

    @staticmethod
    def _extract_office_asset(path: Path) -> str:
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"## {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [normalize_text(str(value)) for value in row if value is not None and value != ""]
                    if values:
                        lines.append(" | ".join(values))
            return normalize_text("\n".join(lines))
        if path.suffix.lower() == ".pptx":
            from pptx import Presentation

            lines = []
            for index, slide in enumerate(Presentation(path).slides, 1):
                text = [normalize_text(shape.text) for shape in slide.shapes if hasattr(shape, "text") and normalize_text(shape.text)]
                if text:
                    lines.append(f"## 幻灯片 {index}\n" + "\n".join(text))
            return normalize_text("\n\n".join(lines))
        return ""

    def _upsert_asset(self, source: dict[str, Any], path: Path, asset_type: str, digest: str, perceptual: str, width: int, height: int, metadata: dict[str, Any], governance: str, reason: str) -> None:
        existing = self.db.row("SELECT id FROM knowledge_assets WHERE source_id=? AND file_path=?", (source["id"], str(path)))
        with self.db.connect() as conn:
            values = (asset_type, path.name, str(path), json.dumps(metadata, ensure_ascii=False), "pending", digest, perceptual, width, height, governance, reason)
            if existing:
                conn.execute(
                    "UPDATE knowledge_assets SET asset_type=?,title=?,file_path=?,metadata_json=?,review_status=?,content_hash=?,perceptual_hash=?,width=?,height=?,governance_status=?,exclusion_reason=? WHERE id=?",
                    (*values, existing["id"]),
                )
            else:
                conn.execute(
                    "INSERT INTO knowledge_assets(source_id,asset_type,title,file_path,metadata_json,review_status,content_hash,perceptual_hash,width,height,governance_status,exclusion_reason) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    (source["id"], *values),
                )

    def _process_ai(self, item: dict[str, Any], progress=None, cancelled=None) -> None:
        document_id = int(item.get("document_id") or 0)
        if not document_id:
            row = self.db.row("SELECT id FROM standard_documents WHERE source_id=?", (item["source_id"],))
            if not row:
                raise RuntimeError("标准文档不存在")
            document_id = int(row["id"])
        non_reusable_reason = self._deterministic_non_reusable_reason(document_id)
        if non_reusable_reason:
            self._terminal(
                item,
                "no_reusable_knowledge",
                document_id=document_id,
                checkpoint={"reason": non_reusable_reason, "gate": "deterministic_non_text"},
            )
            return
        duplicate = self.db.row(
            "SELECT id,source_id FROM standard_documents WHERE text_fingerprint=(SELECT text_fingerprint FROM standard_documents WHERE id=?) AND id<>? ORDER BY id LIMIT 1",
            (document_id, document_id),
        )
        if duplicate:
            self._terminal(item, "content_duplicate", document_id=document_id)
            return
        representative_sections = self._prepare_representative_sections(int(item["run_id"]), document_id)
        if not representative_sections:
            self._terminal(item, "content_duplicate", document_id=document_id)
            return
        result = self.pipeline.process_document(
            document_id,
            max_candidates=20,
            auto_publish=False,
            section_ids=representative_sections,
            progress=progress,
            cancelled=cancelled,
        )
        if str(result.get("status") or "") != "completed" or int(result.get("failed_chunks") or 0):
            raise RuntimeError(
                f"model pipeline incomplete: status={result.get('status') or 'unknown'}; "
                f"failed_chunks={int(result.get('failed_chunks') or 0)}; {result.get('error_message') or ''}"
            )
        run_id = int(result["id"])
        published, legal, rejected = self._final_review(run_id, int(item["run_id"]))
        open_manual = self.db.row(
            "SELECT COUNT(*) AS count FROM governance_tasks WHERE run_id=? AND source_id=? AND status='open'",
            (item["run_id"], item["source_id"]),
        )
        reason = "knowledge_published" if published else ("manual_legal_pending" if legal or int((open_manual or {}).get("count", 0)) else "no_reusable_knowledge")
        self._terminal(item, reason, document_id=document_id, pipeline_run_id=run_id, checkpoint={"published": published, "legal": legal, "rejected": rejected})

    def _process_ai_batch(self, items: list[dict[str, Any]], progress=None, cancelled=None) -> None:
        if len(items) == 1:
            self._process_ai(items[0], progress=progress, cancelled=cancelled)
            return

        prepared: list[tuple[dict[str, Any], int, list[int]]] = []
        for item in items:
            document_id = int(item.get("document_id") or 0)
            if not document_id:
                row = self.db.row("SELECT id FROM standard_documents WHERE source_id=?", (item["source_id"],))
                if not row:
                    raise RuntimeError("standard document does not exist")
                document_id = int(row["id"])
            non_reusable_reason = self._deterministic_non_reusable_reason(document_id)
            if non_reusable_reason:
                self._terminal(
                    item,
                    "no_reusable_knowledge",
                    document_id=document_id,
                    checkpoint={"reason": non_reusable_reason, "gate": "deterministic_non_text"},
                )
                continue
            duplicate = self.db.row(
                """
                SELECT id,source_id FROM standard_documents
                WHERE text_fingerprint=(SELECT text_fingerprint FROM standard_documents WHERE id=?)
                  AND id<>? ORDER BY id LIMIT 1
                """,
                (document_id, document_id),
            )
            if duplicate:
                self._terminal(item, "content_duplicate", document_id=document_id)
                continue
            representative_sections = self._prepare_representative_sections(int(item["run_id"]), document_id)
            if not representative_sections:
                self._terminal(item, "content_duplicate", document_id=document_id)
                continue
            prepared.append((item, document_id, representative_sections))
        if not prepared:
            return
        if len(prepared) == 1:
            # Representatives were already persisted and are idempotent.
            self._process_ai(prepared[0][0], progress=progress, cancelled=cancelled)
            return

        result = self.pipeline.process_document_batch(
            [(document_id, section_ids) for _item, document_id, section_ids in prepared],
            max_candidates=20,
            progress=progress,
            cancelled=cancelled,
        )
        if str(result.get("status") or "") != "completed":
            raise RuntimeError(f"model batch pipeline incomplete: {result.get('error_message') or result.get('status')}")
        pipeline_run_id = int(result["id"])
        published, legal, rejected = self._final_review(pipeline_run_id, int(prepared[0][0]["run_id"]))
        for item, document_id, _section_ids in prepared:
            source_counts = self.db.row(
                """
                SELECT SUM(CASE WHEN status='accepted' THEN 1 ELSE 0 END) AS accepted,
                    SUM(CASE WHEN status='rejected' THEN 1 ELSE 0 END) AS rejected
                FROM knowledge_ai_candidates WHERE pipeline_run_id=? AND source_id=?
                """,
                (pipeline_run_id, item["source_id"]),
            ) or {}
            open_manual = self.db.row(
                "SELECT COUNT(*) AS count FROM governance_tasks WHERE run_id=? AND source_id=? AND status='open'",
                (item["run_id"], item["source_id"]),
            ) or {}
            source_published = int(source_counts.get("accepted") or 0)
            source_rejected = int(source_counts.get("rejected") or 0)
            source_legal = int(open_manual.get("count") or 0)
            source_dispositions = self.db.rows(
                """
                SELECT decision_stage,decision,confidence,reason,ai_run_id,prompt_key,prompt_version
                FROM knowledge_source_dispositions
                WHERE pipeline_run_id=? AND source_id=? ORDER BY id
                """,
                (pipeline_run_id, item["source_id"]),
            )
            disposition_map = {str(row["decision_stage"]): row for row in source_dispositions}
            extraction_disposition = disposition_map.get("extraction") or {}
            review_disposition = disposition_map.get("independent_review") or {}
            verified_not_reusable = (
                str(extraction_disposition.get("decision") or "") == "not_reusable"
                and float(extraction_disposition.get("confidence") or 0) >= 0.95
                and str(review_disposition.get("decision") or "") == "not_reusable"
                and float(review_disposition.get("confidence") or 0) >= 0.95
            )
            if not source_published and not source_rejected and not source_legal:
                if verified_not_reusable:
                    self._terminal(
                        item,
                        "no_reusable_knowledge",
                        document_id=document_id,
                        pipeline_run_id=pipeline_run_id,
                        checkpoint={
                            "reason": "two_stage_ai_verified_not_reusable",
                            "batch_documents": len(prepared),
                            "source_dispositions": source_dispositions,
                        },
                    )
                    continue
                self._advance(
                    item,
                    "ai",
                    document_id=document_id,
                    requires_individual_ai=True,
                    clear_pipeline_run_id=True,
                    coverage_repair="batch_source_returned_no_candidate",
                    batch_pipeline_run_id=pipeline_run_id,
                )
                continue
            reason = "knowledge_published" if source_published else ("manual_legal_pending" if source_legal else "no_reusable_knowledge")
            self._terminal(
                item,
                reason,
                document_id=document_id,
                pipeline_run_id=pipeline_run_id,
                checkpoint={
                    "batch_documents": len(prepared),
                    "published": source_published,
                    "legal": source_legal,
                    "rejected": source_rejected,
                    "batch_totals": {"published": published, "legal": legal, "rejected": rejected},
                },
            )

    def _deterministic_non_reusable_reason(self, document_id: int) -> str:
        document = self.db.row(
            """
            SELECT d.title,d.char_count,s.source_kind
            FROM standard_documents d JOIN source_files s ON s.id=d.source_id
            WHERE d.id=?
            """,
            (document_id,),
        ) or {}
        char_count = int(document.get("char_count") or 0)
        sections = self.db.rows(
            "SELECT heading,content FROM document_sections WHERE document_id=? ORDER BY order_no",
            (document_id,),
        )
        content = "\n".join(str(section.get("content") or "") for section in sections)
        has_visual_markup = bool(
            re.search(r"!\[[^\]]*\]\([^)]+\)|<img\b|/workspace/knowledge/.+?/image_\d+", content, re.IGNORECASE)
        )
        cleaned = re.sub(r"!\[[^\]]*\]\([^)]+\)", " ", content)
        cleaned = re.sub(r"<[^>]+>", " ", cleaned)
        cleaned = re.sub(r"/workspace/\S+", " ", cleaned)
        cleaned = re.sub(r"\s+", "", cleaned)
        meaningful_han = len(re.findall(r"[\u4e00-\u9fff]", cleaned))
        title = normalize_text(str(document.get("title") or ""))
        visual_title = bool(
            re.search(
                r"(?:^附图|(?:平面图|立面图|剖面图|布置图|系统图|示意图|配筋图|机构图|效果图|"
                r"曲线图|流程图|网络图|进度图|横道图|节点图|大样图|详图|总图|施工图)$)",
                title,
                re.IGNORECASE,
            )
        ) and not bool(SCANNED_TEXTUAL_TITLE_PATTERN.search(title))
        cover_or_directory = bool(
            re.search(
                r"(?:^|[-_：:])(?:封面|封皮|目录|COVER)(?:$|[-_：:（）()一二三四五六七八九十0-9])|(?:封面|封皮|目录)$",
                title,
                re.IGNORECASE,
            )
        )
        if has_visual_markup and meaningful_han < 80:
            return "visual_only_without_reusable_text"
        if str(document.get("source_kind") or "") != "system_generated" and (visual_title or cover_or_directory):
            return "historical_visual_or_front_matter_archived"
        if char_count > 700:
            return ""
        if visual_title and meaningful_han < 120:
            return "cover_directory_or_drawing_without_reusable_text"
        if char_count < 80 and meaningful_han < 40:
            return "insufficient_reusable_text"
        return ""

    def _prepare_representative_sections(self, run_id: int, document_id: int) -> list[int]:
        """Cluster one document with one batched embedding pass.

        The old implementation called /embed once per section (N+1 requests).
        We still shortlist lexically per section, but deduplicate all section and
        candidate texts before sending bounded batches to the GPU service.
        """
        sections = self.db.rows(
            "SELECT ds.*,d.source_id FROM document_sections ds JOIN standard_documents d ON d.id=ds.document_id WHERE ds.document_id=? AND LENGTH(TRIM(ds.content))>=30 ORDER BY ds.order_no",
            (document_id,),
        )
        plans: list[tuple[dict[str, Any], dict[str, Any] | None, dict[str, Any] | None, list[tuple[dict[str, Any], float]]]] = []
        vector_texts: dict[int, str] = {}
        for section in sections:
            existing = self.db.row(
                "SELECT representative_section_id,relation_type FROM corpus_section_clusters WHERE run_id=? AND member_section_id=?",
                (run_id, section["id"]),
            )
            exact = None
            shortlisted: list[tuple[dict[str, Any], float]] = []
            if not existing:
                exact = self.db.row(
                    """
                    SELECT c.representative_section_id,ds.content
                    FROM corpus_section_clusters c
                    JOIN document_sections ds ON ds.id=c.representative_section_id
                    WHERE c.run_id=? AND c.representative_section_id=c.member_section_id
                      AND ds.content_fingerprint=? AND ds.id<>?
                    ORDER BY ds.id LIMIT 1
                    """,
                    (run_id, section["content_fingerprint"], section["id"]),
                )
                if not exact:
                    shortlisted = self._near_section_shortlist(section, run_id)
            plans.append((section, existing, exact, shortlisted))
            # All current-document section vectors are prepared together. This
            # also preserves near-duplicate detection among sections first seen
            # in this same document.
            vector_texts[int(section["id"])] = str(section["content"])
            for candidate, _lexical in shortlisted[:32]:
                vector_texts.setdefault(int(candidate["id"]), str(candidate["content"]))

        vectors_by_id: dict[int, list[float]] = {}
        if vector_texts and self.retrieval.embedding.available:
            vector_ids = list(vector_texts)
            vectors = self.retrieval.embedding.embed([vector_texts[item_id] for item_id in vector_ids])
            if len(vectors) == len(vector_ids):
                vectors_by_id = dict(zip(vector_ids, vectors))

        representatives: list[int] = []
        current_representatives: list[dict[str, Any]] = []
        fingerprint_representatives: dict[str, int] = {}
        for section, existing, exact, shortlisted in plans:
            section_id = int(section["id"])
            if existing:
                if int(existing["representative_section_id"]) == section_id:
                    representatives.append(section_id)
                    current_representatives.append(section)
                    fingerprint_representatives[str(section["content_fingerprint"])] = section_id
                continue

            in_document_exact = fingerprint_representatives.get(str(section["content_fingerprint"]))
            representative_id: int | None = (
                in_document_exact
                or (int(exact["representative_section_id"]) if exact else None)
            )
            relation_type = "exact" if representative_id is not None else "representative"
            lexical_score = 1.0 if representative_id is not None else 0.0
            semantic_score = 1.0 if representative_id is not None else 0.0
            conflict = False
            if representative_id is None:
                combined = list(shortlisted)
                source_tokens = set(tokenize(section["content"]))
                source_length = max(1, len(normalize_text(section["content"])))
                for candidate in current_representatives:
                    candidate_length = len(normalize_text(candidate["content"]))
                    if not int(source_length * 0.8) <= candidate_length <= int(source_length * 1.2):
                        continue
                    other_tokens = set(tokenize(candidate["content"]))
                    lexical = len(source_tokens & other_tokens) / max(1, len(source_tokens | other_tokens))
                    if lexical >= 0.85:
                        combined.append((candidate, lexical))
                near = self._best_near_section(section, combined, vectors_by_id)
                if near:
                    representative_id = int(near["id"])
                    relation_type = "near"
                    lexical_score = float(near["lexical_score"])
                    semantic_score = float(near["semantic_score"])
                    conflict = bool(near["conflict_detected"])
            if representative_id is None or conflict:
                representative_id = section_id
                relation_type = "representative"
                lexical_score = semantic_score = 1.0
                representatives.append(representative_id)
                current_representatives.append(section)
                fingerprint_representatives[str(section["content_fingerprint"])] = representative_id
            else:
                self._reattach_section_sources(representative_id, section_id, int(section["source_id"]))
            with self.db.connect() as conn:
                conn.execute(
                    """
                    INSERT OR IGNORE INTO corpus_section_clusters(
                        run_id,representative_section_id,member_section_id,relation_type,
                        lexical_score,semantic_score,conflict_detected,explanation
                    ) VALUES (?,?,?,?,?,?,?,?)
                    """,
                    (
                        run_id,
                        representative_id,
                        section_id,
                        relation_type,
                        lexical_score,
                        semantic_score,
                        int(conflict),
                        "数值、规范编号或适用条件冲突，不自动合并" if conflict else "先去重后加工",
                    ),
                )
        return representatives

    def _near_section_shortlist(self, section: dict[str, Any], run_id: int) -> list[tuple[dict[str, Any], float]]:
        length = max(1, len(normalize_text(section["content"])))
        lower_length = int(length * 0.8)
        upper_length = int(length * 1.2)
        if self.db.backend == "postgresql":
            # Force PostgreSQL to use the content-length expression index first.
            # Without the materialized candidate set the planner underestimates
            # self-representatives and performs ~160k primary-key lookups per
            # section. Fetch full text only for the final 250 rows.
            candidates = self.db.rows(
                """
                WITH length_candidates AS MATERIALIZED (
                    SELECT id,heading FROM document_sections
                    WHERE id<>? AND LENGTH(content) BETWEEN ? AND ?
                ), picked AS MATERIALIZED (
                    SELECT ds.id,ds.heading,
                        CASE WHEN ds.heading=? THEN 0 ELSE 1 END AS heading_rank
                    FROM length_candidates ds
                    WHERE EXISTS (
                        SELECT 1 FROM corpus_section_clusters c
                        WHERE c.run_id=? AND c.member_section_id=ds.id
                          AND c.representative_section_id=ds.id
                        LIMIT 1 OFFSET 0
                    )
                    ORDER BY heading_rank,ds.id DESC LIMIT 250
                )
                SELECT p.id,p.heading,ds.content
                FROM picked p JOIN document_sections ds ON ds.id=p.id
                ORDER BY p.heading_rank,p.id DESC
                """,
                (section["id"], lower_length, upper_length, section["heading"], run_id),
            )
        else:
            candidates = self.db.rows(
                """
                SELECT ds.id,ds.heading,ds.content
                FROM corpus_section_clusters c JOIN document_sections ds ON ds.id=c.representative_section_id
                WHERE c.run_id=? AND c.representative_section_id=c.member_section_id AND ds.id<>?
                  AND LENGTH(ds.content) BETWEEN ? AND ?
                ORDER BY CASE WHEN ds.heading=? THEN 0 ELSE 1 END,ds.id DESC LIMIT 250
                """,
                (run_id, section["id"], lower_length, upper_length, section["heading"]),
            )
        source_tokens = set(tokenize(section["content"]))
        shortlisted: list[tuple[dict[str, Any], float]] = []
        for candidate in candidates:
            other_tokens = set(tokenize(candidate["content"]))
            lexical = len(source_tokens & other_tokens) / max(1, len(source_tokens | other_tokens))
            if lexical >= 0.85:
                shortlisted.append((candidate, lexical))
        return shortlisted[:32]

    def _best_near_section(
        self,
        section: dict[str, Any],
        shortlisted: list[tuple[dict[str, Any], float]],
        vectors_by_id: dict[int, list[float]],
    ) -> dict[str, Any] | None:
        source_vector = vectors_by_id.get(int(section["id"]))
        if not shortlisted or source_vector is None:
            return None
        source_numbers = self._number_signature(section["content"])
        best: dict[str, Any] | None = None
        seen: set[int] = set()
        for candidate, lexical in shortlisted:
            candidate_id = int(candidate["id"])
            if candidate_id in seen:
                continue
            seen.add(candidate_id)
            vector = vectors_by_id.get(candidate_id)
            if vector is None:
                continue
            semantic = self._cosine(source_vector, vector)
            if semantic < 0.94:
                continue
            conflict = source_numbers != self._number_signature(candidate["content"])
            proposal = {**candidate, "lexical_score": lexical, "semantic_score": semantic, "conflict_detected": conflict}
            if best is None or semantic > float(best["semantic_score"]):
                best = proposal
        return best

    @staticmethod
    def _number_signature(text: str) -> tuple[str, ...]:
        return tuple(sorted(set(re.findall(r"(?:GB|JGJ|CJJ|DBJ|GB/T)?\s*\d+(?:[./-]\d+)*(?:\s*(?:mm|cm|m|MPa|kN|%|天|人|台))?", text, re.IGNORECASE))))

    @staticmethod
    def _cosine(left: list[float], right: list[float]) -> float:
        numerator = sum(float(a) * float(b) for a, b in zip(left, right))
        left_norm = sum(float(value) ** 2 for value in left) ** 0.5
        right_norm = sum(float(value) ** 2 for value in right) ** 0.5
        return numerator / (left_norm * right_norm) if left_norm and right_norm else 0.0

    def _reattach_section_sources(self, representative_id: int, member_id: int, source_id: int) -> None:
        units = self.db.rows("SELECT unit_id,excerpt FROM knowledge_unit_sources WHERE section_id=?", (representative_id,))
        with self.db.connect() as conn:
            for unit in units:
                conn.execute(
                    "INSERT OR IGNORE INTO knowledge_unit_sources(unit_id,section_id,source_id,excerpt) VALUES (?,?,?,?)",
                    (unit["unit_id"], member_id, source_id, unit["excerpt"]),
                )

    def _attach_cluster_sources(self, run_id: int, representative_id: int, unit_id: int, excerpt: str) -> None:
        members = self.db.rows(
            """
            SELECT c.member_section_id,d.source_id
            FROM corpus_section_clusters c
            JOIN document_sections s ON s.id=c.member_section_id
            JOIN standard_documents d ON d.id=s.document_id
            WHERE c.run_id=? AND c.representative_section_id=?
            """,
            (run_id, representative_id),
        )
        with self.db.connect() as conn:
            for member in members:
                conn.execute(
                    "INSERT OR IGNORE INTO knowledge_unit_sources(unit_id,section_id,source_id,excerpt) VALUES (?,?,?,?)",
                    (unit_id, member["member_section_id"], member["source_id"], excerpt[:500]),
                )

    def _attach_exact_document_sources(self, source_id: int, section_id: int, unit_id: int, excerpt: str) -> None:
        duplicates = self.db.rows("SELECT id FROM source_files WHERE duplicate_of=? ORDER BY id", (source_id,))
        with self.db.connect() as conn:
            for duplicate in duplicates:
                exists = conn.execute(
                    "SELECT id FROM knowledge_unit_sources WHERE unit_id=? AND source_id=? LIMIT 1",
                    (unit_id, duplicate["id"]),
                ).fetchone()
                if not exists:
                    conn.execute(
                        "INSERT INTO knowledge_unit_sources(unit_id,section_id,source_id,excerpt) VALUES (?,?,?,?)",
                        (unit_id, section_id, duplicate["id"], excerpt[:500]),
                    )

    def _supersede_legacy_rule_units(self) -> None:
        """Keep legacy extracts auditable but remove them from the publishable queue."""
        with self.db.connect() as conn:
            conn.execute(
                """
                UPDATE knowledge_units SET status='superseded',updated_at=CURRENT_TIMESTAMP
                WHERE status<>'published' AND id IN (
                    SELECT unit_id FROM knowledge_versions WHERE origin='extracted'
                )
                """
            )
            conn.execute(
                """
                UPDATE knowledge_versions SET status='superseded'
                WHERE origin='extracted' AND status IN ('review_required','draft')
                """
            )
            conn.execute(
                """
                UPDATE knowledge_exception_tasks SET status='resolved',resolved_by='AI复核系统（非人工）',
                    resolution='全库统一复核运行已接管',resolved_at=CURRENT_TIMESTAMP
                WHERE status='open'
                """
            )
            conn.execute(
                "UPDATE knowledge_ai_candidates SET status='superseded' WHERE status IN ('ready','needs_review')"
            )

    def _final_review(self, pipeline_run_id: int, corpus_run_id: int) -> tuple[int, int, int]:
        candidates = self.pipeline._candidate_rows(
            "WHERE c.pipeline_run_id=? AND c.status NOT IN ('superseded','accepted','rejected')", (pipeline_run_id,), 500
        )
        existing = self.db.row(
            "SELECT SUM(CASE WHEN status='accepted' THEN 1 ELSE 0 END) AS accepted,SUM(CASE WHEN status='rejected' THEN 1 ELSE 0 END) AS rejected FROM knowledge_ai_candidates WHERE pipeline_run_id=?",
            (pipeline_run_id,),
        ) or {}
        published = int(existing.get("accepted") or 0)
        legal = 0
        rejected = int(existing.get("rejected") or 0)
        publication_changed = self.knowledge.on_publication_changed
        self.knowledge.on_publication_changed = None
        try:
            survivors: list[tuple[dict[str, Any], float, bool]] = []
            revision_queue: list[tuple[dict[str, Any], float]] = []
            for candidate in candidates:
                legal_hits = [code for code, pattern in LEGAL_PATTERNS if re.search(pattern, candidate["content"])]
                deterministic = [code for code, pattern in BLOCK_PATTERNS if re.search(pattern, candidate["content"], re.IGNORECASE)]
                findings = list(candidate.get("review_issues") or []) + list(candidate.get("rule_findings") or [])
                source_block = any(item.get("code") in {"quote_mismatch", "invalid_section", "review_missing"} for item in findings)
                if legal_hits:
                    legal += 1
                    self._legal_task(corpus_run_id, int(candidate["source_id"]), "legal_content", candidate["title"], "、".join(legal_hits))
                    self._reject_candidate(candidate, "AI法律边界路由：等待人工确认")
                    self._decision(corpus_run_id, candidate, "legal_route", "manual_required", 1.0, legal_hits)
                    continue
                if source_block:
                    rejected += 1
                    reason = "确定性规则阻断：source_mismatch"
                    self._reject_candidate(candidate, reason)
                    self._decision(corpus_run_id, candidate, "rule_gate", "reject", 1.0, ["source_mismatch"])
                    continue
                risk = str(candidate.get("risk_level") or "medium")
                threshold = {"low": 0.92, "medium": 0.95, "high": 0.97}.get(risk, 0.95)
                blocking_findings = [
                    item for item in findings
                    if item.get("severity") in {"medium", "high"}
                    and not (risk == "high" and item.get("code") in {"candidate_risk", "numeric_parameter"})
                ]
                first_pass = str(candidate.get("review_decision")) in {"pass", "revise"} and float(candidate.get("review_confidence") or 0) >= threshold and not blocking_findings
                self._decision(corpus_run_id, candidate, "independent_review", str(candidate.get("review_decision")), float(candidate.get("review_confidence") or 0), findings, candidate.get("review_ai_run_id"))
                if deterministic or not first_pass:
                    revision_queue.append((candidate, threshold))
                    continue
                survivors.append((candidate, threshold, False))

            repair_reviews = self._ai_review_batch(
                [candidate for candidate, _threshold in revision_queue],
                "knowledge_technical_revision",
                formal=True,
            )
            revised_candidates: list[tuple[dict[str, Any], float]] = []
            for candidate, threshold in revision_queue:
                    repair = repair_reviews[int(candidate["id"])]
                    self._decision(corpus_run_id, candidate, "automatic_revision", repair["decision"], repair["confidence"], repair["issues"], repair.get("run_id"))
                    corrected = normalize_text(repair.get("corrected_content") or "")
                    if repair["decision"] != "revise" or not corrected:
                        rejected += 1
                        self._reject_candidate(candidate, "自动修订未能消除确定性风险或达到独立复核阈值")
                        continue
                    candidate = self._apply_revision(candidate, corrected)
                    revised_candidates.append((candidate, threshold))

            post_reviews = self._ai_review_batch(
                [candidate for candidate, _threshold in revised_candidates],
                "knowledge_post_revision_review",
            )
            for candidate, threshold in revised_candidates:
                    post = post_reviews[int(candidate["id"])]
                    self._decision(corpus_run_id, candidate, "post_revision_review", post["decision"], post["confidence"], post["issues"], post.get("run_id"))
                    if not self._review_passes(post, threshold) or self._blocking_codes(candidate):
                        rejected += 1
                        self._reject_candidate(candidate, "自动修订后的独立复核未达到风险阈值")
                        continue
                    survivors.append((candidate, threshold, True))

            high_risk = [candidate for candidate, _threshold, _revision_used in survivors if str(candidate.get("risk_level")) == "high"]
            second_reviews = self._ai_review_batch(high_risk, "high_risk_second_review")
            adjudication_ready: list[tuple[dict[str, Any], float, bool]] = []
            for candidate, threshold, revision_used in survivors:
                if str(candidate.get("risk_level")) == "high":
                    second = second_reviews[int(candidate["id"])]
                    self._decision(corpus_run_id, candidate, "second_review", second["decision"], second["confidence"], second["issues"], second.get("run_id"))
                    if not self._review_passes(second, threshold):
                        rejected += 1
                        self._reject_candidate(candidate, "高风险第二次独立复核未通过")
                        continue
                adjudication_ready.append((candidate, threshold, revision_used))

            final_reviews = self._ai_review_batch(
                [candidate for candidate, _threshold, _revision_used in adjudication_ready],
                "formal_adjudication",
                formal=True,
            )
            final_revision_queue: list[tuple[dict[str, Any], float]] = []
            for candidate, threshold, revision_used in adjudication_ready:
                final = final_reviews[int(candidate["id"])]
                self._decision(corpus_run_id, candidate, "final_adjudication", final["decision"], final["confidence"], final["issues"], final.get("run_id"))
                if final["decision"] == "revise" and not revision_used and normalize_text(final.get("corrected_content") or ""):
                    candidate = self._apply_revision(candidate, final["corrected_content"])
                    final_revision_queue.append((candidate, threshold))
                    continue
                if not self._review_passes(final, threshold) or self._blocking_codes(candidate):
                    rejected += 1
                    self._reject_candidate(candidate, "最终裁决未达到风险阈值")
                    continue
                promoted = self.knowledge.create_unit_from_ai_candidate(candidate)
                unit_id = int(promoted["unit_id"])
                self._attach_cluster_sources(corpus_run_id, int(candidate["source_section_id"]), unit_id, str(candidate["source_quote"]))
                self._attach_exact_document_sources(int(candidate["source_id"]), int(candidate["source_section_id"]), unit_id, str(candidate["source_quote"]))
                detail = self.knowledge.get_unit(unit_id)
                if detail["status"] != "published":
                    version = detail["versions"][0]
                    if version["status"] != "approved":
                        self.knowledge.review_unit(unit_id, int(version["id"]), "approve", "AI复核系统（非人工）", "全库收口独立复核通过")
                    self.knowledge.publish_unit(unit_id, int(version["id"]), "AI复核系统（非人工）")
                with self.db.connect() as conn:
                    conn.execute("UPDATE knowledge_ai_candidates SET status='accepted',unit_id=?,reviewed_by='AI复核系统（非人工）',review_notes='全库收口自动复核通过',reviewed_at=CURRENT_TIMESTAMP WHERE id=?", (unit_id, candidate["id"]))
                    conn.execute("UPDATE knowledge_exception_tasks SET status='resolved',resolved_by='AI复核系统（非人工）',resolution='自动技术复核闭环',resolved_at=CURRENT_TIMESTAMP WHERE candidate_id=? AND status='open'", (candidate["id"],))
                self._decision(corpus_run_id, {**candidate, "unit_id": unit_id}, "publication", "publish", final["confidence"], [])
                published += 1

            revised_finals = self._ai_review_batch(
                [candidate for candidate, _threshold in final_revision_queue],
                "formal_adjudication_after_revision",
                formal=True,
            )
            for candidate, threshold in final_revision_queue:
                final = revised_finals[int(candidate["id"])]
                self._decision(corpus_run_id, candidate, "final_adjudication_after_revision", final["decision"], final["confidence"], final["issues"], final.get("run_id"))
                if not self._review_passes(final, threshold) or self._blocking_codes(candidate):
                    rejected += 1
                    self._reject_candidate(candidate, "最终修订裁决未达到风险阈值")
                    continue
                promoted = self.knowledge.create_unit_from_ai_candidate(candidate)
                unit_id = int(promoted["unit_id"])
                self._attach_cluster_sources(corpus_run_id, int(candidate["source_section_id"]), unit_id, str(candidate["source_quote"]))
                self._attach_exact_document_sources(int(candidate["source_id"]), int(candidate["source_section_id"]), unit_id, str(candidate["source_quote"]))
                detail = self.knowledge.get_unit(unit_id)
                if detail["status"] != "published":
                    version = detail["versions"][0]
                    if version["status"] != "approved":
                        self.knowledge.review_unit(unit_id, int(version["id"]), "approve", "AI复核系统（非人工）", "全库收口独立复核通过")
                    self.knowledge.publish_unit(unit_id, int(version["id"]), "AI复核系统（非人工）")
                with self.db.connect() as conn:
                    conn.execute("UPDATE knowledge_ai_candidates SET status='accepted',unit_id=?,reviewed_by='AI复核系统（非人工）',review_notes='全库收口自动复核通过',reviewed_at=CURRENT_TIMESTAMP WHERE id=?", (unit_id, candidate["id"]))
                    conn.execute("UPDATE knowledge_exception_tasks SET status='resolved',resolved_by='AI复核系统（非人工）',resolution='自动技术复核闭环',resolved_at=CURRENT_TIMESTAMP WHERE candidate_id=? AND status='open'", (candidate["id"],))
                self._decision(corpus_run_id, {**candidate, "unit_id": unit_id}, "publication", "publish", final["confidence"], [])
                published += 1
        finally:
            self.knowledge.on_publication_changed = publication_changed
        return published, legal, rejected

    def _ai_review_batch(self, candidates: list[dict[str, Any]], task_type: str, formal: bool = False) -> dict[int, dict[str, Any]]:
        if not candidates:
            return {}
        if len(candidates) > 12:
            combined: dict[int, dict[str, Any]] = {}
            for start in range(0, len(candidates), 12):
                combined.update(self._ai_review_batch(candidates[start:start + 12], task_type, formal=formal))
            return combined
        spec = KNOWLEDGE_FORMAL_REVIEW_PROMPT if formal else KNOWLEDGE_REVIEW_PROMPT
        section_ids = sorted({int(candidate.get("source_section_id") or 0) for candidate in candidates if candidate.get("source_section_id")})
        sections: dict[int, dict[str, Any]] = {}
        if section_ids:
            placeholders = ",".join("?" for _ in section_ids)
            sections = {
                int(row["id"]): row
                for row in self.db.rows(
                    f"SELECT id,heading,content FROM document_sections WHERE id IN ({placeholders}) ORDER BY id",
                    section_ids,
                )
            }
        section_text = self._review_source_context(candidates, sections, max_chars=48_000)
        candidate_json = []
        for index, candidate in enumerate(candidates):
            candidate_json.append({
                "candidate_index": index,
                **{key: candidate.get(key) for key in ("title", "content", "summary", "applicability", "risk_level", "source_quote", "source_section_id")},
            })
        result = self.ai_runtime.execute(
            spec,
            spec.render(section_text=section_text, candidate_json=json.dumps(candidate_json, ensure_ascii=False)),
            {
                "candidate_ids": [int(candidate["id"]) for candidate in candidates],
                "content_hashes": [content_hash(str(candidate["content"])) for candidate in candidates],
            },
            task_type=task_type,
            target_type="knowledge_candidate_batch",
            target_id=int(candidates[0].get("pipeline_run_id") or candidates[0]["id"]),
            use_cache=False,
            max_output_tokens=12000,
        )
        reviews = (result.get("payload") or {}).get("reviews") or []
        by_index = {int(review["candidate_index"]): review for review in reviews}
        if len(by_index) != len(candidates) or any(index not in by_index for index in range(len(candidates))):
            raise RuntimeError(f"model batch review incomplete: {result.get('error') or f'{len(by_index)}/{len(candidates)}'}")
        return {
            int(candidate["id"]): {**by_index[index], "run_id": result.get("run_id")}
            for index, candidate in enumerate(candidates)
        }

    @staticmethod
    def _review_source_context(
        candidates: list[dict[str, Any]],
        sections: dict[int, dict[str, Any]],
        max_chars: int = 48_000,
    ) -> str:
        """Build bounded, quote-centred source windows for independent review.

        Candidate provenance has already passed the deterministic full-section
        quote gate.  The review model therefore needs the exact quote and its
        surrounding source text, not hundreds of thousands of unrelated chars.
        """
        by_section: dict[int, list[dict[str, Any]]] = {}
        for candidate in candidates:
            section_id = int(candidate.get("source_section_id") or 0)
            by_section.setdefault(section_id, []).append(candidate)
        blocks: list[str] = []
        used = 0
        for section_id in sorted(by_section):
            section = sections.get(section_id, {})
            content = normalize_text(str(section.get("content") or ""))
            windows: list[tuple[int, int]] = []
            for candidate in by_section[section_id]:
                quote = normalize_text(str(candidate.get("source_quote") or ""))
                quote_lines = [line.strip() for line in quote.splitlines() if len(line.strip()) >= 8]
                anchors = quote_lines or ([quote] if quote else [])
                found = False
                for anchor in anchors:
                    position = content.find(anchor)
                    if position < 0:
                        continue
                    start = max(0, position - 900)
                    end = min(len(content), position + len(anchor) + 900)
                    windows.append((start, end))
                    found = True
                if not found and content:
                    windows.append((0, min(len(content), 3500)))
            merged: list[tuple[int, int]] = []
            for start, end in sorted(windows):
                if merged and start <= merged[-1][1] + 200:
                    merged[-1] = (merged[-1][0], max(merged[-1][1], end))
                else:
                    merged.append((start, end))
            source_window = "\n...\n".join(content[start:end] for start, end in merged)
            quotes = "\n".join(
                f"[CANDIDATE:{candidate['id']}] SOURCE_QUOTE: {normalize_text(str(candidate.get('source_quote') or ''))}"
                for candidate in by_section[section_id]
            )
            block = f"[SECTION:{section_id}] {section.get('heading', '')}\n{source_window}\n{quotes}"
            if used + len(block) > max_chars:
                remaining = max_chars - used
                if remaining >= 500:
                    blocks.append(block[:remaining])
                break
            blocks.append(block)
            used += len(block)
        return "\n\n".join(blocks)

    def _ai_review(self, candidate: dict[str, Any], task_type: str, formal: bool = False, *, recovery_context: dict[str, Any] | None = None) -> dict[str, Any]:
        section_id = int(candidate.get("source_section_id") or 0)
        section = self.db.row("SELECT heading,content FROM document_sections WHERE id=?", (section_id,)) or {}
        spec = KNOWLEDGE_FORMAL_REVIEW_PROMPT if formal else KNOWLEDGE_REVIEW_PROMPT
        structured_recovery = recovery_context is not None or candidate.get("recovery_contract") == RECOVERY_CONTRACT
        if structured_recovery:
            spec = RETIRED_PUBLICATION_REVISION_PROMPT if task_type == "retired_publication_revision" else _recovery_review_spec(formal)
        review_input = {"candidate_index": 0, **{key: candidate.get(key) for key in ("title", "content", "summary", "applicability", "risk_level", "source_quote", "source_section_id")}}
        payload = {"candidate_id": candidate["id"], "content_hash": content_hash(candidate["content"])}
        if structured_recovery:
            review_input["recovery_contract"] = RECOVERY_CONTRACT
            payload["recovery_contract"] = RECOVERY_CONTRACT
        if recovery_context is not None:
            # Previous findings are review evidence, never part of a source quote.
            review_input["recovery_context"] = recovery_context
            payload["recovery_context"] = recovery_context
        prompt = spec.render(
            section_text=f"[SECTION:{section_id}] {section.get('heading','')}\n{section.get('content','')}",
            candidate_json=json.dumps([review_input], ensure_ascii=False),
        )
        result = self.ai_runtime.execute(
            spec, prompt, payload,
            task_type=task_type, target_type="knowledge_candidate", target_id=int(candidate.get("id") or candidate.get("unit_id") or 0), use_cache=False,
        )
        reviews = ((result.get("payload") or {}).get("reviews") or [])
        if not reviews or (structured_recovery and (len(reviews) != 1 or reviews[0].get("candidate_index") != 0)):
            raise RuntimeError(f"model review failed: {result.get('error') or '复核无结果'}")
        review = reviews[0]
        return {**review, "run_id": result.get("run_id")}

    @staticmethod
    def _recovery_artifact(content: str, summary: str, applicability: str) -> dict[str, str]:
        fields = {"body": content, "summary": summary, "applicability": applicability}
        if any(not isinstance(value, str) for value in fields.values()):
            raise ValueError("recovery_metadata_fields_must_be_strings")
        fields = {key: normalize_text(value) for key, value in fields.items()}
        if any(not value for value in fields.values()):
            raise ValueError("recovery_requires_complete_body_summary_applicability")
        if len(fields["body"]) > 12000 or len(fields["summary"]) > 600 or len(fields["applicability"]) > 1500:
            raise ValueError("recovery_metadata_field_too_long")
        # A labelled whole document in corrected_content used to leave stale
        # summary/applicability in the subsequent model input. Require actual
        # separate fields; never silently extract or invent missing metadata.
        wrapper = r"(?:^|\n)\s*(?:#{1,6}\s*)?(?:标题|内容|正文|摘要|适用范围|title|content|body|summary|applicability)\s*(?:[:：]|\n|$)"
        if any(re.search(wrapper, value, re.IGNORECASE) for value in fields.values()):
            raise ValueError("recovery_metadata_contains_wrapped_or_duplicate_fields")
        if fields["summary"] in {fields["body"], fields["body"][:200]}:
            raise ValueError("recovery_summary_must_not_copy_or_truncate_body")
        artifact = f"## 摘要\n{fields['summary']}\n\n## 适用范围\n{fields['applicability']}\n\n## 正文\n{fields['body']}"
        return {**fields, "content": artifact}

    @classmethod
    def _read_recovery_artifact(cls, content: str, summary: str) -> dict[str, str]:
        match = re.fullmatch(r"## 摘要\n(.+?)\n\n## 适用范围\n(.+?)\n\n## 正文\n(.+)", str(content), re.DOTALL)
        if not match:
            raise ValueError("recovery_version_missing_structured_metadata")
        embedded_summary, applicability, body = match.groups()
        artifact = cls._recovery_artifact(body, embedded_summary, applicability)
        if artifact["content"] != content or artifact["summary"] != summary:
            raise ValueError("recovery_version_metadata_mismatch")
        return artifact

    def _apply_revision(self, candidate: dict[str, Any], corrected_content: str) -> dict[str, Any]:
        corrected = normalize_text(corrected_content)
        with self.db.connect() as conn:
            conn.execute(
                "UPDATE knowledge_ai_candidates SET content=?,review_decision='revise',review_notes='AI自动修订一次' WHERE id=?",
                (corrected, candidate["id"]),
            )
        return {**candidate, "content": corrected, "review_decision": "revise"}

    def _accepted_without_publication(self, run_id: int) -> list[dict[str, Any]]:
        # Do not inner-join provenance here: missing provenance is a blocker,
        # not a reason for an accepted candidate to disappear from this queue.
        return self.db.rows(
            """
            SELECT c.id AS candidate_id,c.unit_id,c.title FROM knowledge_ai_candidates c
            WHERE c.status='accepted'
                AND EXISTS (SELECT 1 FROM corpus_run_items i WHERE i.run_id=? AND i.pipeline_run_id=c.pipeline_run_id)
                AND NOT EXISTS (SELECT 1 FROM knowledge_publications p WHERE p.unit_id=c.unit_id AND p.status='published')
            ORDER BY c.id
            """, (run_id,),
        )

    def _retired_recovery_state(self, candidate_id: int) -> dict[str, Any]:
        candidate = self.db.row("SELECT * FROM knowledge_ai_candidates WHERE id=?", (candidate_id,))
        if not candidate or candidate["status"] != "accepted" or not candidate.get("unit_id"):
            raise ValueError("candidate_not_accepted_or_missing_unit")
        unit_id = int(candidate["unit_id"])
        if self.db.row("SELECT id FROM knowledge_publications WHERE unit_id=? AND status='published'", (unit_id,)):
            raise ValueError("publication_state_changed")
        unit = self.db.row("SELECT * FROM knowledge_units WHERE id=?", (unit_id,))
        publication = self.db.row(
            """SELECT p.*,v.content AS version_content,v.summary AS version_summary,v.origin AS version_origin FROM knowledge_publications p
            JOIN knowledge_versions v ON v.id=p.version_id WHERE p.unit_id=? AND p.status='retired' ORDER BY p.id DESC LIMIT 1""",
            (unit_id,),
        )
        candidate_source = self.db.row(
            """SELECT ds.id,ds.document_id,ds.heading,ds.content,d.source_id,s.sha256
            FROM document_sections ds JOIN standard_documents d ON d.id=ds.document_id
            JOIN source_files s ON s.id=d.source_id WHERE ds.id=?""", (candidate.get("source_section_id"),),
        )
        if not unit or not publication:
            raise ValueError("missing_retired_publication")
        if (not candidate_source or not normalize_text(str(candidate_source.get("content") or ""))
                or candidate_source["document_id"] != candidate["document_id"] or candidate_source["source_id"] != candidate["source_id"]):
            raise ValueError("missing_or_mismatched_source_provenance")
        # Publication audit chooses the first unit-source link. Use that same
        # source for recovery, even when this accepted candidate came from a
        # later source linked to a shared unit; do not rewrite source links.
        source = self.db.row(
            """SELECT ds.id,ds.document_id,ds.heading,ds.content,us.source_id,s.sha256,us.id AS source_link_id
            FROM knowledge_unit_sources us LEFT JOIN document_sections ds ON ds.id=us.section_id
            LEFT JOIN source_files s ON s.id=us.source_id WHERE us.unit_id=? ORDER BY us.id LIMIT 1""", (unit_id,),
        )
        if not source or not source.get("id") or not source.get("sha256") or not normalize_text(str(source.get("content") or "")):
            raise ValueError("missing_publication_audit_source")
        retirement = self.db.row(
            """SELECT * FROM knowledge_review_decisions WHERE unit_id=?
            AND decision_stage='published_artifact_audit' AND decision='retire' ORDER BY id DESC LIMIT 1""", (unit_id,),
        )
        if not retirement:
            raise ValueError("missing_retirement_review_evidence")
        try:
            retirement_time = datetime.fromisoformat(str(retirement["created_at"]).replace("Z", "+00:00"))
            publication_time = datetime.fromisoformat(str(publication["retired_at"]).replace("Z", "+00:00"))
            retirement_time = retirement_time.replace(tzinfo=retirement_time.tzinfo or timezone.utc)
            publication_time = publication_time.replace(tzinfo=publication_time.tzinfo or timezone.utc)
            if retirement_time < publication_time:
                raise ValueError("retirement_review_does_not_match_latest_publication")
        except (TypeError, ValueError) as exc:
            raise ValueError("retirement_review_does_not_match_latest_publication") from exc
        previous_review = self.db.row(
            """SELECT d.*,a.input_json AS ai_input_json,a.status AS ai_status,a.prompt_hash AS ai_prompt_hash
            FROM knowledge_review_decisions d LEFT JOIN ai_runs a ON a.id=d.ai_run_id WHERE d.unit_id=? AND d.id<?
            AND d.decision_stage='published_revalidation' ORDER BY d.id DESC LIMIT 1""", (unit_id, retirement["id"]),
        )
        identity = {"publication_id": int(publication["id"]), "version_id": int(publication["version_id"]),
                    "publication_version": int(publication["publication_version"]), "content_hash": publication["content_hash"],
                    "body_hash": content_hash(publication["version_content"])}
        bound = False
        review_bound = False
        for decision in (retirement, previous_review):
            if not decision:
                continue
            findings = parse_json(decision.get("findings_json"), [])
            saved = next((finding for finding in findings if isinstance(finding, dict) and finding.get("code") == "publication_binding"), None)
            if saved and all(saved.get(key) == value for key, value in identity.items()):
                bound = True
                if decision is previous_review:
                    review_bound = True
        if not review_bound and previous_review:
            # Earlier audits predate explicit bindings, but their immutable AI
            # input records the exact publication ID and body hash.
            ai_input = parse_json(previous_review.get("ai_input_json"), {})
            payload = ai_input.get("payload") or {}
            review_bound = (previous_review.get("ai_status") == "succeeded"
                            and previous_review.get("ai_prompt_hash") == KNOWLEDGE_FORMAL_REVIEW_PROMPT.prompt_hash
                            and payload.get("candidate_id") == identity["publication_id"]
                            and payload.get("content_hash") == identity["body_hash"])
            bound = bound or review_bound
        if not bound:
            raise ValueError("retirement_evidence_not_bound_to_publication")
        if not review_bound:
            previous_review = None
        context = {
            "candidate_id": candidate_id, "unit_id": unit_id,
            "old_publication_id": publication["id"], "old_version_id": publication["version_id"],
            "old_content_hash": publication["content_hash"], "old_body_hash": content_hash(publication["version_content"]),
            "retirement_decision_id": retirement["id"], "retirement_findings": parse_json(retirement["findings_json"], []),
            "prior_review_id": previous_review["id"] if previous_review else None,
            "prior_findings": parse_json(previous_review["findings_json"], []) if previous_review else [],
            "prior_decision": previous_review["decision"] if previous_review else None,
            "prior_confidence": previous_review["confidence"] if previous_review else None,
            "source_section_id": source["id"],
            "source_hash": content_hash(json.dumps(source, ensure_ascii=False, sort_keys=True)),
            "candidate_source_hash": content_hash(json.dumps(candidate_source, ensure_ascii=False, sort_keys=True)),
            "candidate_hash": content_hash(json.dumps(candidate, ensure_ascii=False, sort_keys=True)),
        }
        return {"candidate": candidate, "unit": unit, "publication": publication, "source": source, "context": context}

    def _recovery_blocked(self, run_id: int, candidate_id: int, unit_id: int | None, reason: str, context: dict[str, Any]) -> dict[str, Any]:
        details = {"code": "retired_publication_recovery", "severity": "high", "message": reason[:1000], **context}
        self._decision(run_id, {"id": candidate_id, "unit_id": unit_id}, "retired_recovery", "blocked", 0, [details])
        # The accepted candidate stays accepted, so this open issue cannot make
        # the formal-eligible denominator fall or masquerade as a rejection.
        with self.db.connect() as conn:
            task = conn.execute(
                "SELECT id FROM knowledge_exception_tasks WHERE candidate_id=? AND issue_code='retired_publication_recovery' AND status='open'",
                (candidate_id,),
            ).fetchone()
            if task:
                conn.execute("UPDATE knowledge_exception_tasks SET message=? WHERE id=?", (reason[:1000], task["id"]))
            else:
                conn.execute(
                    """INSERT INTO knowledge_exception_tasks(candidate_id,issue_code,severity,message,source)
                    VALUES (?,'retired_publication_recovery','high',?,'corpus_recovery')""", (candidate_id, reason[:1000]),
                )
        return {"candidate_id": candidate_id, "unit_id": unit_id, "status": "blocked", "reason": reason[:1000]}

    def _audit_recovery_publications(self, run_id: int, progress=None, cancelled=None) -> dict[str, Any]:
        # Version origin survives a crash between publish_unit and recording the
        # recovery decision. Audit these even when another recovery is blocked.
        targets = self.db.rows(
            """SELECT DISTINCT p.id AS publication_id,p.unit_id FROM knowledge_publications p
            JOIN knowledge_versions v ON v.id=p.version_id WHERE p.status='published'
            AND v.origin='corpus_retired_recovery' AND EXISTS (
                SELECT 1 FROM knowledge_ai_candidates c JOIN corpus_run_items i ON i.pipeline_run_id=c.pipeline_run_id
                WHERE c.unit_id=p.unit_id AND c.status='accepted' AND i.run_id=?) ORDER BY p.id""", (run_id,),
        )
        if not targets:
            return {"scanned": 0, "revalidated_publication_ids": []}
        audit = self._audit_active_publications(run_id, progress=progress, cancelled=cancelled,
                                                publication_ids=[int(item["publication_id"]) for item in targets])
        revalidated = []
        for target in targets:
            row = next((row for row in self._published_audit_rows(int(target["unit_id"])) if row["id"] == target["publication_id"]), None)
            decision = self.db.row(
                "SELECT * FROM knowledge_review_decisions WHERE unit_id=? AND decision_stage='published_revalidation' ORDER BY id DESC LIMIT 1",
                (target["unit_id"],),
            )
            if not row or not decision or not self._publication_review_reusable(row, [decision]):
                continue
            self._check_finalization(run_id, cancelled or (lambda: False))
            with self.db.connect() as conn:
                conn.execute(
                    """UPDATE knowledge_exception_tasks SET status='resolved',resolved_by='AI复核系统（非人工）',
                    resolution=?,resolved_at=CURRENT_TIMESTAMP
                    WHERE issue_code='retired_publication_recovery' AND status='open'
                    AND candidate_id IN (SELECT id FROM knowledge_ai_candidates WHERE unit_id=?)""",
                    (f"定向修订的新发布{target['publication_id']}已通过独立发布复核；decision={decision['id']}", target["unit_id"]),
                )
            revalidated.append(target["publication_id"])
        return {**audit, "revalidated_publication_ids": revalidated}

    def _recover_retired_candidates(self, run_id: int, progress=None, cancelled=None, *, candidate_ids: list[int] | None = None, limit: int = 50) -> dict[str, Any]:
        """One bounded recovery pass; only an explicit resume starts another.

        Keep accepted candidates and retired publications immutable. A repaired
        body lives in a new version of the SAME unit, and its new publication
        still has to pass the normal publication audit before final activation.
        """
        progress = progress or (lambda *_args, **_kwargs: None)
        cancelled = cancelled or (lambda: False)
        if not 1 <= limit <= 50:
            raise ValueError("recovery limit must be between 1 and 50")
        selected = None if candidate_ids is None else {int(value) for value in candidate_ids}
        if selected is not None:
            if not selected or len(selected) > 50:
                raise ValueError("select between 1 and 50 candidate IDs")
            for candidate_id in selected:
                if not self.db.row(
                    """SELECT c.id FROM knowledge_ai_candidates c WHERE c.id=? AND c.status='accepted'
                    AND EXISTS (SELECT 1 FROM corpus_run_items i WHERE i.pipeline_run_id=c.pipeline_run_id AND i.run_id=?)""",
                    (candidate_id, run_id),
                ):
                    raise ValueError(f"candidate {candidate_id} is not accepted in this corpus run")
        targets = {}
        for item in self._accepted_without_publication(run_id):
            if selected is None or item["candidate_id"] in selected:
                targets.setdefault(item["unit_id"] or -item["candidate_id"], item)
        outcome: dict[str, Any] = {"scanned": 0, "published": [], "blocked": [], "remaining": []}
        callback = self.knowledge.on_publication_changed
        self.knowledge.on_publication_changed = None
        try:
            for item in list(targets.values())[:limit]:
                self._check_finalization(run_id, cancelled)
                candidate_id, unit_id = int(item["candidate_id"]), item["unit_id"]
                context: dict[str, Any] = {"candidate_id": candidate_id, "unit_id": unit_id}
                outcome["scanned"] += 1
                progress("retired_recovery", 95, f"恢复停用知识 {outcome['scanned']}/{min(len(targets), limit)}", context)
                try:
                    state = self._retired_recovery_state(candidate_id)
                    original, unit, retired, source, context = (state[key] for key in ("candidate", "unit", "publication", "source", "context"))
                    risk = "high" if "high" in {original.get("risk_level"), unit.get("risk_level")} else "medium"
                    threshold = 0.97 if risk == "high" else 0.95
                    # The failed publication is the revision target. Replace an
                    # unfaithful historical quote only with actual source text;
                    # preserve the historical candidate/quote in the database.
                    verified_source = normalize_text(source["content"])
                    context["historical_source_quote"] = str(original.get("source_quote") or "")
                    old_body = retired["version_content"]
                    old_summary = str(retired.get("version_summary") or original.get("summary") or "")
                    old_applicability = str(original.get("applicability") or "")
                    if retired.get("version_origin") == "corpus_retired_recovery":
                        old_artifact = self._read_recovery_artifact(retired["version_content"], retired.get("version_summary"))
                        old_body, old_summary, old_applicability = (old_artifact[key] for key in ("body", "summary", "applicability"))
                    candidate = {**original, "title": unit["title"], "content": retired["version_content"], "source_section_id": source["id"],
                                 "summary": old_summary, "applicability": old_applicability,
                                 "risk_level": risk, "source_quote": verified_source, "recovery_contract": RECOVERY_CONTRACT}
                    context["contract"] = RECOVERY_CONTRACT
                    context["instruction"] = "依据原始来源逐项处理上次复核问题；完整修订正文、摘要及适用范围，不能补造事实、仅修改标签或恢复旧发布。失败说明是复核证据，不是来源原文。"

                    def review(task: str, target: dict[str, Any], *, formal: bool = False) -> dict[str, Any]:
                        self._check_finalization(run_id, cancelled)
                        result = self._ai_review(target, task, formal=formal, recovery_context={**context, "stage": task})
                        self._check_finalization(run_id, cancelled)
                        self._decision(run_id, target, task, result["decision"], result["confidence"],
                                       [*(result.get("issues") or []), {"code": "recovery_binding", "severity": "info", **context,
                                        "draft_hash": content_hash(target["content"])}], result.get("run_id"))
                        return result

                    repair = review("retired_publication_revision", candidate, formal=True)
                    if repair.get("decision") != "revise":
                        raise ValueError("revision_requires_a_changed_nonempty_body")
                    artifact = self._recovery_artifact(repair.get("corrected_content"), repair.get("corrected_summary"), repair.get("corrected_applicability"))
                    if content_hash(artifact["body"]) == content_hash(normalize_text(old_body)):
                        raise ValueError("revision_requires_a_changed_nonempty_body")
                    corrected = artifact["content"]
                    candidate = {**candidate, **{key: artifact[key] for key in ("content", "summary", "applicability")}}
                    context["instruction"] = "独立复核这份完整修订稿的正文、摘要及适用范围，核查三者彼此一致且忠实来源；缺失或不一致必须拒绝。只有上次问题已解决且无中高风险事项才可pass，不得降低门槛。"
                    if self._blocking_codes(candidate) or any(re.search(pattern, corrected) for _, pattern in LEGAL_PATTERNS):
                        raise ValueError("revised_body_has_deterministic_or_legal_blockers")
                    post = review("retired_publication_post_revision_review", candidate)
                    if not self._review_passes(post, threshold):
                        raise ValueError("post_revision_review_not_passed")
                    if risk == "high":
                        second = review("retired_publication_second_review", candidate)
                        if not self._review_passes(second, threshold):
                            raise ValueError("high_risk_second_review_not_passed")
                    final = review("retired_publication_formal_adjudication", candidate, formal=True)
                    if not self._review_passes(final, threshold):
                        raise ValueError("formal_adjudication_not_passed")
                    fresh = self._retired_recovery_state(candidate_id)
                    for key in ("old_publication_id", "old_version_id", "old_content_hash", "old_body_hash", "retirement_decision_id", "prior_review_id", "source_hash", "candidate_source_hash", "candidate_hash"):
                        if fresh["context"].get(key) != context.get(key):
                            raise ValueError(f"recovery_input_changed:{key}")
                    self._check_finalization(run_id, cancelled)
                    digest = content_hash(corrected)
                    # Reuse an unpublished recovery version of this exact body
                    # after a publication/write interruption; re-review above
                    # still occurs, and no historical published version is reused.
                    with self.db.connect() as conn:
                        version = conn.execute(
                            """SELECT v.id,v.version_no FROM knowledge_versions v WHERE v.unit_id=?
                            AND v.origin='corpus_retired_recovery' AND v.content_hash=? AND v.content=? AND v.summary=?
                            AND NOT EXISTS (SELECT 1 FROM knowledge_publications p WHERE p.version_id=v.id)
                            ORDER BY v.id DESC LIMIT 1""", (unit_id, digest, corrected, candidate["summary"]),
                        ).fetchone()
                        if version:
                            version_id, version_no = int(version["id"]), int(version["version_no"])
                        else:
                            version_no = int(conn.execute("SELECT COALESCE(MAX(version_no),0)+1 FROM knowledge_versions WHERE unit_id=?", (unit_id,)).fetchone()[0])
                            version_id = int(conn.execute(
                                """INSERT INTO knowledge_versions(unit_id,version_no,content,summary,origin,content_hash,status)
                                VALUES (?,?,?,?,'corpus_retired_recovery',?,'review_required')""", (unit_id, version_no, corrected, candidate["summary"], digest),
                            ).lastrowid)
                        conn.execute("UPDATE knowledge_units SET cleaned_content=?,content_fingerprint=?,status='review_required',updated_at=CURRENT_TIMESTAMP WHERE id=?", (corrected, digest, unit_id))
                    draft_path = self.knowledge.settings.knowledge_directories["drafts"] / f"{unit_id:07d}" / f"v{version_no}.md"
                    write_text_atomic(draft_path, corrected)
                    self._decision(run_id, original, "retired_recovery_version", "prepared", final["confidence"],
                                   [{"code": "recovery_binding", "severity": "info", **context, "version_id": version_id, "draft_hash": digest}], final.get("run_id"))
                    self.knowledge.review_unit(int(unit_id), version_id, "approve", "AI复核系统（非人工）",
                                               f"停用发布{retired['id']}定向修订；新独立复核与正式裁决通过，candidate={candidate_id}，原停用记录保留。")
                    self._check_finalization(run_id, cancelled)
                    fresh = self._retired_recovery_state(candidate_id)
                    if any(fresh["context"].get(key) != context.get(key) for key in ("old_publication_id", "source_hash", "candidate_source_hash", "candidate_hash", "retirement_decision_id")):
                        raise ValueError("recovery_input_changed_before_publication")
                    self._check_finalization(run_id, cancelled)
                    publication = self.knowledge.publish_unit(int(unit_id), version_id, "AI复核系统（非人工）")
                    published = {"candidate_id": candidate_id, "unit_id": unit_id, "old_publication_id": retired["id"],
                                 "version_id": version_id, "publication_id": publication["publication_id"], "status": "pending_revalidation"}
                    # Not the 'publication' stage: that would create a reusable
                    # approval binding and skip the new publication's own audit.
                    self._decision(run_id, original, "retired_recovery_publication", "publish", final["confidence"],
                                   [{"code": "recovery_binding", "severity": "info", **context, **published}], final.get("run_id"))
                    outcome["published"].append(published)
                except FinalizationInterrupted:
                    raise
                except Exception as exc:  # noqa: BLE001
                    outcome["blocked"].append(self._recovery_blocked(run_id, candidate_id, unit_id, str(exc), context))
                self._finalization_checkpoint(run_id, "retired_recovery", 95, retired_recovery=outcome)
        finally:
            self.knowledge.on_publication_changed = callback
        outcome["remaining"] = self._accepted_without_publication(run_id)
        return outcome

    @staticmethod
    def _blocking_codes(candidate: dict[str, Any]) -> list[str]:
        content = str(candidate.get("content") or "")
        return [code for code, pattern in BLOCK_PATTERNS if re.search(pattern, content, re.IGNORECASE)]

    @staticmethod
    def _review_passes(review: dict[str, Any], threshold: float) -> bool:
        return (
            str(review.get("decision") or "") == "pass"
            and float(review.get("confidence") or 0) >= threshold
            and not [item for item in (review.get("issues") or []) if item.get("severity") in {"medium", "high"}]
        )

    def _decision(self, run_id: int, candidate: dict[str, Any], stage: str, decision: str, confidence: float, findings: Any, ai_run_id: int | None = None, *, publication_binding: dict[str, Any] | None = None) -> None:
        if stage == "publication" and candidate.get("unit_id") and publication_binding is None:
            published = self._published_audit_rows(int(candidate["unit_id"]))
            if len(published) == 1:
                publication_binding = self._publication_review_binding(published[0])
        if publication_binding:
            findings = [*(findings or []), {"code": "publication_binding", "severity": "info", **publication_binding}]
        with self.db.connect() as conn:
            conn.execute(
                """
                INSERT INTO knowledge_review_decisions(run_id,candidate_id,unit_id,decision_stage,actor_type,actor_id,decision,confidence,findings_json,ai_run_id,prompt_key,prompt_version)
                VALUES (?,?,?,?,'ai','AI复核系统（非人工）',?,?,?,?,?,?)
                """,
                (run_id, candidate.get("id"), candidate.get("unit_id"), stage, decision, float(confidence or 0), json.dumps(findings or [], ensure_ascii=False), ai_run_id, "knowledge.formal-review", "1.0.0"),
            )

    def _reject_candidate(self, candidate: dict[str, Any], reason: str) -> None:
        with self.db.connect() as conn:
            conn.execute("UPDATE knowledge_ai_candidates SET status='rejected',reviewed_by='AI复核系统（非人工）',review_notes=?,reviewed_at=CURRENT_TIMESTAMP WHERE id=?", (reason, candidate["id"]))
            conn.execute("UPDATE knowledge_exception_tasks SET status='resolved',resolved_by='AI复核系统（非人工）',resolution=?,resolved_at=CURRENT_TIMESTAMP WHERE candidate_id=? AND status='open'", (reason, candidate["id"]))

    def _legal_task(self, run_id: int, source_id: int, task_type: str, title: str, message: str) -> None:
        if source_id <= 0:
            return
        source = self.db.row("SELECT family_key,relative_path,parent_source_id FROM source_files WHERE id=?", (source_id,)) or {}
        key = self._governance_task_key(source, source_id, task_type)
        with self.db.connect() as conn:
            conn.execute(
                "INSERT OR IGNORE INTO governance_tasks(run_id,source_id,task_key,task_type,severity,title,message) VALUES (?,?,?,?,?,?,?)",
                (run_id, source_id, key, task_type, "high", title[:300], message[:2000]),
            )
            task = conn.execute("SELECT id FROM governance_tasks WHERE run_id=? AND task_key=?", (run_id, key)).fetchone()
            if task:
                conn.execute("INSERT OR IGNORE INTO governance_task_sources(task_id,source_id) VALUES (?,?)", (task["id"], source_id))
        # Keep the human queue grouped while governance continues to discover
        # additional members of an already-seen source batch.
        self._consolidate_manual_tasks(run_id)

    @staticmethod
    def _governance_task_key(source: dict[str, Any], source_id: int, task_type: str) -> str:
        if source.get("parent_source_id"):
            group = f"archive:{source['parent_source_id']}"
        else:
            # Source inventories were originally produced on Windows, while the
            # production worker runs on Linux.  Normalize separators explicitly.
            relative = PurePosixPath(str(source.get("relative_path") or source_id).replace("\\", "/"))
            parts = [part for part in relative.parent.parts if part not in {".", "..", "/"}]
            group = "/".join(parts[:3]) or str(source.get("family_key") or source_id)
        return f"{task_type}:{content_hash(group)[:16]}"

    def _consolidate_manual_tasks(self, run_id: int) -> None:
        tasks = self.db.rows(
            "SELECT * FROM governance_tasks WHERE run_id=? AND status='open' ORDER BY id",
            (run_id,),
        )
        groups: dict[str, list[dict[str, Any]]] = {}
        task_sources: dict[int, list[int]] = {}
        for task in tasks:
            sources = self.db.rows(
                "SELECT source_id FROM governance_task_sources WHERE task_id=? ORDER BY source_id",
                (task["id"],),
            )
            source_ids = [int(row["source_id"]) for row in sources]
            if not source_ids and task.get("source_id"):
                source_ids = [int(task["source_id"])]
            task_sources[int(task["id"])] = source_ids
            source = self.db.row(
                "SELECT family_key,relative_path,parent_source_id FROM source_files WHERE id=?",
                (source_ids[0],),
            ) if source_ids else {}
            desired = self._governance_task_key(source or {}, source_ids[0] if source_ids else int(task["id"]), str(task["task_type"]))
            groups.setdefault(desired, []).append(task)
        with self.db.connect() as conn:
            for desired, members in groups.items():
                primary = next((task for task in members if task["task_key"] == desired), members[0])
                if primary["task_key"] != desired:
                    conn.execute("UPDATE governance_tasks SET task_key=? WHERE id=?", (desired, primary["id"]))
                for task in members:
                    for source_id in task_sources[int(task["id"])]:
                        conn.execute(
                            "INSERT OR IGNORE INTO governance_task_sources(task_id,source_id) VALUES (?,?)",
                            (primary["id"], source_id),
                        )
                    if int(task["id"]) != int(primary["id"]):
                        conn.execute(
                            "UPDATE governance_tasks SET status='resolved',resolution=?,resolved_at=CURRENT_TIMESTAMP WHERE id=?",
                            (f"superseded_by_group_task:{primary['id']}", task["id"]),
                        )

    def _advance(self, item: dict[str, Any], stage: str, **values: Any) -> None:
        checkpoint = {
            key: value
            for key, value in values.items()
            if value is not None and key not in {"clear_pipeline_run_id"}
        }
        with self.db.connect() as conn:
            conn.execute(
                "UPDATE corpus_run_items SET stage=?,status='pending',attempt_count=0,processing_job_id=COALESCE(?,processing_job_id),document_id=COALESCE(?,document_id),pipeline_run_id=CASE WHEN ? THEN NULL ELSE pipeline_run_id END,checkpoint_json=?,error_code='',error_message='',next_retry_at=NULL,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                (
                    stage,
                    values.get("processing_job_id"),
                    values.get("document_id"),
                    bool(values.get("clear_pipeline_run_id")),
                    json.dumps(checkpoint, ensure_ascii=False),
                    item["id"],
                ),
            )

    def _terminal(self, item: dict[str, Any], reason: str, **values: Any) -> None:
        if reason not in TERMINAL_REASONS:
            raise ValueError(f"未知终态原因: {reason}")
        checkpoint = values.pop("checkpoint", {})
        with self.db.connect() as conn:
            conn.execute(
                """
                UPDATE corpus_run_items SET stage='complete',status='terminal',terminal_reason=?,
                    processing_job_id=COALESCE(?,processing_job_id),document_id=COALESCE(?,document_id),
                    pipeline_run_id=COALESCE(?,pipeline_run_id),checkpoint_json=?,error_code='',error_message='',
                    next_retry_at=NULL,completed_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP WHERE id=?
                """,
                (reason, values.get("processing_job_id"), values.get("document_id"), values.get("pipeline_run_id"), json.dumps(checkpoint, ensure_ascii=False), item["id"]),
            )

    def _handle_failure(self, item: dict[str, Any], exc: Exception, *, service_error: bool = False) -> None:
        attempts = int(item.get("attempt_count") or 0) + 1
        message = f"{type(exc).__name__}: {exc}"[:1000]
        # Deployment restarts are operational checkpoints, not source failures,
        # and must never consume the three real processing attempts.
        checkpoint_restarts = int(item.get("attempt_count") or 0) if str(item.get("error_code") or "").startswith(("image_rollout", "deployment_restart", "worker_restart")) else 0
        effective_attempts = max(1, attempts - checkpoint_restarts)
        encrypted = any(word in message.lower() for word in ("password", "encrypted", "密码", "加密"))
        if encrypted:
            self._legal_task(int(item["run_id"]), int(item["source_id"]), "credentials", item["file_name"], message)
            self._terminal(item, "encrypted_manual")
            return
        if effective_attempts >= 3 and not service_error:
            if str(item.get("stage") or "") == "ai":
                self._terminal(
                    item,
                    "no_reusable_knowledge",
                    checkpoint={"reason": "insufficient_evidence_after_model_failures", "error": message, "attempts": effective_attempts},
                )
            else:
                self._terminal(item, "unreadable_excluded", checkpoint={"error": message, "attempts": effective_attempts})
            return
        retry_index = min(effective_attempts - 1, 2) if not service_error else min(effective_attempts - 1, 2)
        delay = [1, 5, 15][retry_index]
        retry_at = (utc_now() + timedelta(minutes=delay)).isoformat(timespec="seconds")
        with self.db.connect() as conn:
            conn.execute(
                "UPDATE corpus_run_items SET status='retrying',attempt_count=?,error_code=?,error_message=?,next_retry_at=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                (effective_attempts, type(exc).__name__.lower(), message, retry_at, item["id"]),
            )

    def _recover_stale_items(self, run_id: int) -> None:
        # A model request may legitimately spend up to an hour across transport
        # and one schema-repair retry.  Use a conservative two-hour orphan
        # window so concurrent API work is never closed while still active.
        # If the host clock moved backwards, only close a future-dated record
        # after a newer request for the same target has demonstrably taken over.
        ai_cutoff = (utc_now() - timedelta(hours=2)).strftime("%Y-%m-%d %H:%M:%S")
        future_cutoff = (utc_now() + timedelta(minutes=30)).strftime("%Y-%m-%d %H:%M:%S")
        cutoff = (utc_now() - timedelta(hours=7)).strftime("%Y-%m-%d %H:%M:%S")
        with self.db.connect() as conn:
            conn.execute(
                """
                UPDATE ai_runs SET status='failed',error_code='worker_restart',
                    error_message='候选流水线已闭环且无可执行候选，孤立模型调用已自动闭环',
                    completed_at=CURRENT_TIMESTAMP
                WHERE status='running' AND target_type='knowledge_candidate_batch'
                  AND EXISTS (
                      SELECT 1 FROM knowledge_ai_pipeline_runs p
                      WHERE p.id=ai_runs.target_id AND p.status<>'running'
                  )
                  AND NOT EXISTS (
                      SELECT 1 FROM knowledge_ai_candidates c
                      WHERE c.pipeline_run_id=ai_runs.target_id
                        AND c.status IN ('ready','needs_review')
                  )
                """
            )
            conn.execute(
                """
                UPDATE ai_runs SET status='failed',error_code='worker_restart',
                    error_message='后台进程中断，模型调用记录已自动闭环',completed_at=CURRENT_TIMESTAMP
                WHERE status='running' AND (
                    created_at<? OR (
                        created_at>? AND EXISTS (
                            SELECT 1 FROM ai_runs newer
                            WHERE newer.id>ai_runs.id
                              AND newer.status IN ('running','succeeded','failed')
                              AND newer.task_type=ai_runs.task_type
                              AND COALESCE(newer.target_type,'')=COALESCE(ai_runs.target_type,'')
                              AND COALESCE(newer.target_id,-1)=COALESCE(ai_runs.target_id,-1)
                              AND newer.created_at<=?
                        )
                    )
                )
                """,
                (ai_cutoff, future_cutoff, future_cutoff),
            )
            conn.execute(
                """
                UPDATE knowledge_ai_pipeline_runs SET status='completed_with_exceptions',
                    error_message='后台进程中断，流水线记录已自动闭环',completed_at=CURRENT_TIMESTAMP
                WHERE status='running' AND (
                    created_at<? OR (
                        created_at>? AND EXISTS (
                            SELECT 1 FROM knowledge_ai_pipeline_runs newer
                            WHERE newer.id>knowledge_ai_pipeline_runs.id
                              AND newer.document_id=knowledge_ai_pipeline_runs.document_id
                              AND newer.status IN ('running','completed')
                              AND newer.created_at<=?
                        )
                    )
                )
                  AND NOT EXISTS (
                      SELECT 1 FROM ai_runs a
                      WHERE a.status='running' AND a.id IN (
                          knowledge_ai_pipeline_runs.extraction_ai_run_id,
                          knowledge_ai_pipeline_runs.review_ai_run_id,
                          knowledge_ai_pipeline_runs.adjudication_ai_run_id
                      )
                  )
                """,
                (ai_cutoff, future_cutoff, future_cutoff),
            )
            conn.execute(
                """
                UPDATE corpus_run_items SET status='retrying',error_code='stale_worker',
                    error_message='后台任务中断，已从最近检查点恢复',next_retry_at=NULL,updated_at=CURRENT_TIMESTAMP
                WHERE run_id=? AND status='running' AND updated_at<?
                """,
                (run_id, cutoff),
            )

    def _requeue_uncovered_batch_sources(self, run_id: int) -> int:
        """Repair old batch outcomes that had no source-level model decision.

        A shared extraction may legitimately find no reusable content for one
        source, but silence is not an auditable decision.  Such sources get one
        individual extraction/review route before they can be closed.
        """
        with self.db.connect() as conn:
            updated = conn.execute(
                """
                UPDATE corpus_run_items SET stage='ai',status='pending',attempt_count=0,terminal_reason='',pipeline_run_id=NULL,
                    checkpoint_json=?,completed_at=NULL,next_retry_at=NULL,updated_at=CURRENT_TIMESTAMP
                WHERE run_id=? AND status='terminal' AND terminal_reason='no_reusable_knowledge'
                  AND pipeline_run_id IS NOT NULL
                  AND checkpoint_json LIKE '%"batch_documents"%'
                  AND EXISTS (
                      SELECT 1 FROM knowledge_ai_pipeline_runs p
                      WHERE p.id=corpus_run_items.pipeline_run_id AND p.chunk_count>1
                  )
                  AND NOT EXISTS (
                      SELECT 1 FROM knowledge_ai_candidates c
                      WHERE c.pipeline_run_id=corpus_run_items.pipeline_run_id
                        AND c.source_id=corpus_run_items.source_id
                  )
                  AND NOT EXISTS (
                      SELECT 1 FROM knowledge_source_dispositions d
                      WHERE d.pipeline_run_id=corpus_run_items.pipeline_run_id
                        AND d.source_id=corpus_run_items.source_id
                  )
                """,
                (
                    json.dumps(
                        {
                            "requires_individual_ai": True,
                            "coverage_repair": "historical_batch_source_returned_no_candidate",
                        },
                        ensure_ascii=False,
                    ),
                    run_id,
                ),
            )
            return int(updated.rowcount or 0)

    def _record_result(self, run_id: int, succeeded: bool, error: str, *, service_stage: str = "") -> None:
        run = self._raw_run(run_id)
        policy = parse_json(run.get("policy_json"), {})
        checkpoint = parse_json(run.get("checkpoint_json"), {})
        if not service_stage:
            return
        health = checkpoint.setdefault("service_health", {})
        stage_health = health.setdefault(service_stage, {})
        recent = list(stage_health.get("recent_results") or [])
        recent.append({"ok": bool(succeeded), "at": iso_now(), "error": error[:300]})
        recent = recent[-int(policy.get("recent_window", 20)):]
        consecutive = 0 if succeeded else int(stage_health.get("consecutive_errors") or 0) + 1
        failures = sum(not item["ok"] for item in recent)
        rate = failures / len(recent) if recent else 0
        should_pause = consecutive >= int(policy.get("service_consecutive_error_limit", 5)) or (len(recent) >= 20 and rate > float(policy.get("recent_failure_rate", 0.2)))
        stage_health.update(
            {
                "consecutive_errors": consecutive,
                "recent_results": recent,
                "failure_rate": round(rate, 4),
                "updated_at": iso_now(),
            }
        )
        if should_pause and service_stage:
            checkpoint["service_pause_stage"] = service_stage
            checkpoint["service_recovery_probe"] = {
                "stage": service_stage,
                "status": "scheduled",
                "checked_at": iso_now(),
            }
        with self.db.connect() as conn:
            conn.execute(
                "UPDATE corpus_runs SET consecutive_errors=?,recent_results_json=?,checkpoint_json=?,status=?,pause_reason=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                (
                    consecutive,
                    json.dumps(recent, ensure_ascii=False),
                    json.dumps(checkpoint, ensure_ascii=False),
                    "paused" if should_pause else run["status"],
                    f"外部服务错误达到熔断阈值: {service_stage or 'unknown'}" if should_pause else run.get("pause_reason", ""),
                    run_id,
                ),
            )
        if should_pause and self.dispatch:
            self.dispatch(run_id, 15 * 60 * 1000, 1, "__recovery__")

    def _check_finalization(self, run_id: int, cancelled: Callable[[], bool]) -> None:
        if cancelled():
            # Unwind evaluation's own transaction before writing cancellation;
            # otherwise a SQLite writer can deadlock against its callback.
            raise FinalizationInterrupted("用户取消", cancel_requested=True)
        run = self._raw_run(run_id)
        if run["status"] not in {"pending", "running"}:
            raise FinalizationInterrupted(str(run.get("pause_reason") or run["status"]))

    def _finalization_checkpoint(self, run_id: int, phase: str, percent: int, **details: Any) -> None:
        checkpoint = parse_json(self._raw_run(run_id).get("checkpoint_json"), {})
        finalization = checkpoint.setdefault("finalization", {})
        finalization.update(details)
        finalization.update({"phase": phase, "updated_at": iso_now()})
        with self.db.connect() as conn:
            conn.execute(
                """
                UPDATE corpus_runs SET checkpoint_json=?,
                    stage=CASE WHEN status IN ('pending','running') THEN ? ELSE stage END,
                    progress=CASE WHEN status IN ('pending','running') THEN ? ELSE progress END,
                    updated_at=CURRENT_TIMESTAMP WHERE id=?
                """,
                (json.dumps(checkpoint, ensure_ascii=False), phase, percent, run_id),
            )

    def _publication_snapshot(self) -> str:
        rows = self.db.rows(
            "SELECT id,unit_id,version_id,publication_version,content_hash FROM knowledge_publications WHERE status='published' ORDER BY id"
        )
        return content_hash(json.dumps(rows, sort_keys=True))

    def _finalize_run(self, run_id: int, progress=None, cancelled=None) -> dict[str, Any]:
        progress = progress or (lambda *_args, **_kwargs: None)
        cancelled = cancelled or (lambda: False)
        with self._run_mutex(run_id, "finalize") as acquired:
            if not acquired or self._raw_run(run_id)["status"] not in {"pending", "running"}:
                return self.get_run(run_id)
            remaining = self.db.row(
                "SELECT COUNT(*) AS count FROM corpus_run_items WHERE run_id=? AND status<>'terminal'", (run_id,)
            ) or {}
            if int(remaining.get("count") or 0):
                return self.get_run(run_id)
            previous = parse_json(self._raw_run(run_id).get("checkpoint_json"), {}).get("finalization", {})
            phase = "retired_recovery"
            percentages = {"retired_recovery": 95, "publication_audit": 95, "index": 96, "evaluation": 97, "activation": 99, "reports": 99}
            extra: dict[str, Any] = {}
            completed_here = False

            def report(step: str, value: int, message: str = "", details=None) -> None:
                self._check_finalization(run_id, cancelled)
                self._finalization_checkpoint(
                    run_id, phase, percentages[phase], status="running", message=message,
                    step_stage=step, step_progress=value, details=details or {},
                )
                progress(phase, percentages[phase], message, details or {})

            def check_cancelled() -> bool:
                self._check_finalization(run_id, cancelled)
                return False

            try:
                self._check_finalization(run_id, cancelled)
                self._finalization_checkpoint(
                    run_id, phase, 95, status="running", attempt=int(previous.get("attempt") or 0) + 1,
                    started_at=iso_now(), error=None, previous_error=previous.get("error"),
                )
                report(phase, 0, "核对已接纳但发布停用的知识")
                recovery = self._recover_retired_candidates(run_id, progress=report, cancelled=cancelled)
                extra["retired_recovery"] = recovery
                recovery_audit = self._audit_recovery_publications(run_id, progress=report, cancelled=cancelled)
                extra["retired_recovery_audit"] = recovery_audit
                for item in recovery["published"]:
                    item["status"] = "revalidated" if item["publication_id"] in recovery_audit["revalidated_publication_ids"] else "blocked"
                recovery["remaining"] = self._accepted_without_publication(run_id)
                self._finalization_checkpoint(run_id, phase, 95, retired_recovery=recovery)
                if recovery["blocked"] or recovery["remaining"]:
                    ids = sorted({int(item["candidate_id"]) for item in [*recovery["blocked"], *recovery["remaining"]]})
                    raise RuntimeError(f"停用知识恢复仍待复核，已保留候选和停用记录；candidate_ids={ids}")
                phase = "publication_audit"
                report(phase, 0, "复核当前已发布知识")
                publication_audit = self._audit_active_publications(run_id, progress=report, cancelled=cancelled)
                extra["publication_audit"] = publication_audit
                missing = self._accepted_without_publication(run_id)
                if missing:
                    extra["publication_recovery_pending"] = missing
                    raise RuntimeError(f"发布复核后仍有已接纳知识未发布，需显式恢复后定向修订；candidate_ids={[item['candidate_id'] for item in missing]}")
                snapshot = self._publication_snapshot()
                phase = "index"
                report(phase, 0, "构建最终检索索引")
                index = None
                if previous.get("publication_snapshot") == snapshot and previous.get("candidate_index_id"):
                    index = self.db.row(
                        "SELECT * FROM retrieval_indexes WHERE id=? AND status IN ('building','active')",
                        (int(previous["candidate_index_id"]),),
                    )
                if not index:
                    index = self.retrieval.build_index(progress=report, cancelled=check_cancelled, activate=False)
                self._check_finalization(run_id, cancelled)
                if not index.get("id") or index.get("cancelled"):
                    raise RuntimeError("候选索引未完成构建")
                self._finalization_checkpoint(
                    run_id, phase, 96, candidate_index_id=int(index["id"]),
                    publication_snapshot=snapshot, publication_audit=publication_audit,
                )
                extra["candidate_index"] = index
                phase = "evaluation"
                report(phase, 0, "执行最终检索评测")
                evaluation = self._build_and_run_evaluation(run_id, report, int(index["id"]), cancelled=cancelled)
                extra["evaluation"] = evaluation
                self._finalization_checkpoint(
                    run_id, phase, 97, evaluation={key: value for key, value in evaluation.items() if key != "results"},
                )
                self._check_finalization(run_id, cancelled)
                if not evaluation.get("passed"):
                    raise RuntimeError("检索评测未达到发布阈值，已保留上一活动索引")
                if self._publication_snapshot() != snapshot:
                    raise RuntimeError("评测期间已发布知识发生变化，已保留上一活动索引")
                completion = self.completion(run_id)
                if completion["corpus_terminal_rate"] < 1 or completion["formal_eligible_publication_rate"] < 1:
                    raise RuntimeError("来源闭环或正式可用知识发布尚未完成，已保留上一活动索引")
                phase = "activation"
                report(phase, 0, "验证并激活已通过评测的候选索引")
                extra["index"] = self.retrieval.activate_index(int(index["id"]))
                self._check_finalization(run_id, cancelled)
                completion = self.completion(run_id)
                if not completion["index_consistent"]:
                    raise RuntimeError("活动索引与发布表不一致")
                if completion["corpus_terminal_rate"] < 1 or completion["formal_eligible_publication_rate"] < 1:
                    raise RuntimeError("来源闭环或正式可用知识发布尚未完成")
                phase = "reports"
                report(phase, 0, "写入最终验收报告与清单")
                # Materialize deliverables before declaring the run completed.
                self.write_reports(run_id, extra)
                self._check_finalization(run_id, cancelled)
                self._finalization_checkpoint(run_id, phase, 99, status="completed", completed_at=iso_now())
                with self.db.connect() as conn:
                    changed = conn.execute(
                        "UPDATE corpus_runs SET status='completed',stage='completed',progress=100,completed_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP WHERE id=? AND status IN ('pending','running')",
                        (run_id,),
                    )
                if not changed.rowcount:
                    raise FinalizationInterrupted("运行状态已改变")
                completed_here = True
                self.write_reports(run_id, extra)
            except FinalizationInterrupted as exc:
                if exc.cancel_requested:
                    with self.db.connect() as conn:
                        conn.execute(
                            "UPDATE corpus_runs SET status='cancelled',stage='cancelled',pause_reason='用户取消',updated_at=CURRENT_TIMESTAMP WHERE id=? AND status IN ('pending','running')",
                            (run_id,),
                        )
                state = self._raw_run(run_id)
                self._finalization_checkpoint(run_id, phase, percentages[phase], status=state["status"], interrupted_at=iso_now())
                self._write_finalization_failure_report(run_id, extra)
            except Exception as exc:  # noqa: BLE001
                error = {"type": type(exc).__name__, "message": str(exc)[:1000], "phase": phase, "at": iso_now()}
                extra["error"] = error
                with self.db.connect() as conn:
                    conn.execute(
                        """
                        UPDATE corpus_runs SET status='paused',stage=?,pause_reason=?,completed_at=NULL,updated_at=CURRENT_TIMESTAMP
                        WHERE id=? AND (status IN ('pending','running') OR (status='completed' AND ?=1))
                        """,
                        (phase, f"最终收口失败（{phase}）：{error['type']}: {error['message']}", run_id, int(completed_here)),
                    )
                self._finalization_checkpoint(run_id, phase, percentages[phase], status="failed", error=error)
                self._write_finalization_failure_report(run_id, extra)
            return self.get_run(run_id)

    def _write_finalization_failure_report(self, run_id: int, extra: dict[str, Any]) -> None:
        try:
            self.write_reports(run_id, extra)
        except Exception as exc:  # noqa: BLE001
            state = self._raw_run(run_id)
            self._finalization_checkpoint(
                run_id, str(state["stage"]), int(state["progress"]),
                report_error={"type": type(exc).__name__, "message": str(exc)[:1000]},
            )

    def _published_audit_rows(self, unit_id: int | None = None) -> list[dict[str, Any]]:
        rows = self.db.rows(
            f"""
            SELECT p.id,p.unit_id,p.version_id,p.publication_version,p.content_hash,p.published_at,
                v.content,v.summary,v.origin,u.title,u.risk_level,us.section_id,s.id AS source_id,
                ds.heading AS source_heading,ds.content AS source_content
            FROM knowledge_publications p
            JOIN knowledge_versions v ON v.id=p.version_id JOIN knowledge_units u ON u.id=p.unit_id
            LEFT JOIN knowledge_unit_sources us ON us.unit_id=u.id
            LEFT JOIN source_files s ON s.id=us.source_id
            LEFT JOIN document_sections ds ON ds.id=us.section_id
            WHERE p.status='published' {"AND p.unit_id=?" if unit_id is not None else ""} ORDER BY p.id,us.id
            """,
            (unit_id,) if unit_id is not None else (),
        )
        unique: dict[int, dict[str, Any]] = {}
        for row in rows:
            unique.setdefault(int(row["id"]), row)
        return list(unique.values())

    @classmethod
    def _published_review_candidate(cls, row: dict[str, Any]) -> dict[str, Any]:
        candidate = {
            "id": int(row["id"]), "source_section_id": int(row.get("section_id") or 0),
            "title": row["title"], "content": row["content"], "summary": str(row["content"])[:200],
            "applicability": "按来源章节适用条件执行",
            "risk_level": "high" if str(row.get("risk_level")) == "high" else "medium",
            "source_quote": str(row["content"])[:1000],
        }
        if row.get("origin") == "corpus_retired_recovery":
            candidate.update(summary=str(row.get("summary") or ""), applicability="",
                             source_quote=normalize_text(str(row.get("source_content") or "")), recovery_contract=RECOVERY_CONTRACT)
            try:
                artifact = cls._read_recovery_artifact(row["content"], row.get("summary"))
                candidate.update({key: artifact[key] for key in ("content", "summary", "applicability")})
            except ValueError as exc:
                candidate["recovery_metadata_error"] = str(exc)
        return candidate

    def _published_review_prompt(self, row: dict[str, Any]) -> str:
        candidate = self._published_review_candidate(row)
        spec = _recovery_review_spec(True) if candidate.get("recovery_contract") else KNOWLEDGE_FORMAL_REVIEW_PROMPT
        review_input = {
            "candidate_index": 0,
            **{key: candidate.get(key) for key in ("title", "content", "summary", "applicability", "risk_level", "source_quote", "source_section_id")},
        }
        if candidate.get("recovery_contract"):
            review_input["recovery_contract"] = RECOVERY_CONTRACT
        return spec.render(
            section_text=f"[SECTION:{candidate['source_section_id']}] {row.get('source_heading') or ''}\n{row.get('source_content') or ''}",
            candidate_json=json.dumps([review_input], ensure_ascii=False),
        )

    def _publication_review_binding(self, row: dict[str, Any]) -> dict[str, Any]:
        return {
            "publication_id": int(row["id"]), "version_id": int(row["version_id"]),
            "publication_version": int(row["publication_version"]),
            "content_hash": row["content_hash"], "body_hash": content_hash(str(row["content"])),
            "review_input_hash": content_hash(self._published_review_prompt(row)),
        }

    def _publication_review_reusable(self, row: dict[str, Any], decisions: list[dict[str, Any]]) -> bool:
        if self._published_review_candidate(row).get("recovery_metadata_error"):
            return False
        binding = self._publication_review_binding(row)
        if binding["content_hash"] != binding["body_hash"]:
            return False
        for decision in decisions:
            publication_stage = decision["decision_stage"] == "publication"
            expected_decision = "publish" if publication_stage else "pass"
            threshold = {"low": 0.92, "high": 0.97}.get(str(row.get("risk_level")), 0.95) if publication_stage else (0.97 if row.get("risk_level") == "high" else 0.95)
            issues = parse_json(decision.get("findings_json"), [])
            if decision["decision"] != expected_decision or float(decision.get("confidence") or 0) < threshold:
                continue
            if not isinstance(issues, list) or any(isinstance(item, dict) and item.get("severity") in {"medium", "high"} for item in issues):
                continue
            saved = next((item for item in issues if isinstance(item, dict) and item.get("code") == "publication_binding"), None)
            if saved:
                if all(saved.get(key) == value for key, value in binding.items()):
                    return True
                continue
            if not publication_stage:
                # Legacy successful revalidations already contain exact source
                # context and the publication ID in their immutable AI input.
                ai_input = parse_json(decision.get("ai_input_json"), {})
                payload = ai_input.get("payload") or {}
                if (
                    decision.get("ai_status") == "succeeded"
                    and decision.get("ai_prompt_hash") == KNOWLEDGE_FORMAL_REVIEW_PROMPT.prompt_hash
                    and payload.get("candidate_id") == int(row["id"])
                    and payload.get("content_hash") == binding["body_hash"]
                    and ai_input.get("rendered_prompt") == self._published_review_prompt(row)
                ):
                    return True
            elif (
                decision.get("candidate_status") == "accepted"
                and decision.get("candidate_unit_id") == row["unit_id"]
                and content_hash(str(decision.get("candidate_content") or "")) == binding["body_hash"]
            ):
                # A legacy publication decision is usable only for a version
                # already published when that exact accepted candidate was signed.
                try:
                    reviewed_at = datetime.fromisoformat(str(decision["created_at"]).replace("Z", "+00:00"))
                    published_at = datetime.fromisoformat(str(row["published_at"]).replace("Z", "+00:00"))
                    reviewed_at = reviewed_at.replace(tzinfo=reviewed_at.tzinfo or timezone.utc)
                    published_at = published_at.replace(tzinfo=published_at.tzinfo or timezone.utc)
                    if reviewed_at >= published_at:
                        return True
                except (TypeError, ValueError):
                    pass
        return False

    def _audit_active_publications(self, run_id: int, progress=None, cancelled=None, *, publication_ids: list[int] | None = None) -> dict[str, Any]:
        progress = progress or (lambda *_args, **_kwargs: None)
        cancelled = cancelled or (lambda: False)
        rows = self._published_audit_rows()
        if publication_ids is not None:
            selected = set(publication_ids)
            rows = [row for row in rows if row["id"] in selected]
        decisions_by_unit: dict[int, list[dict[str, Any]]] = {}
        for decision in self.db.rows(
            """
            SELECT d.*,c.content AS candidate_content,c.status AS candidate_status,c.unit_id AS candidate_unit_id,
                a.input_json AS ai_input_json,a.status AS ai_status,a.prompt_hash AS ai_prompt_hash
            FROM knowledge_review_decisions d
            LEFT JOIN knowledge_ai_candidates c ON c.id=d.candidate_id LEFT JOIN ai_runs a ON a.id=d.ai_run_id
            WHERE d.run_id=? AND d.decision_stage IN ('publication','published_revalidation')
                AND d.unit_id IN (SELECT unit_id FROM knowledge_publications WHERE status='published')
            ORDER BY d.id DESC
            """,
            (run_id,),
        ):
            decisions_by_unit.setdefault(int(decision["unit_id"]), []).append(decision)
        retired: list[dict[str, Any]] = []
        reused = 0
        publication_callback = self.knowledge.on_publication_changed
        self.knowledge.on_publication_changed = None
        try:
            for ordinal, row in enumerate(rows, 1):
                self._check_finalization(run_id, cancelled)
                blockers = [code for code, pattern in BLOCK_PATTERNS if re.search(pattern, row["content"], re.IGNORECASE)]
                if row["content_hash"] != content_hash(str(row["content"])):
                    blockers.append("published_content_hash_mismatch")
                metadata_error = self._published_review_candidate(row).get("recovery_metadata_error")
                if metadata_error:
                    blockers.append(metadata_error)
                already_reviewed = self._publication_review_reusable(row, decisions_by_unit.get(int(row["unit_id"]), []))
                if already_reviewed:
                    reused += 1
                if not already_reviewed or ordinal % 25 == 0:
                    progress("publication_audit", 95, f"复核已发布知识 {ordinal}/{len(rows)}", {"publication_id": row["id"], "reused": reused, "retired": len(retired)})
                if not blockers and not already_reviewed and row.get("section_id"):
                    legacy = self._published_review_candidate(row)
                    review = self._ai_review(legacy, "published_knowledge_revalidation", formal=True)
                    self._check_finalization(run_id, cancelled)
                    threshold = 0.97 if legacy["risk_level"] == "high" else 0.95
                    self._decision(run_id, {"unit_id": row["unit_id"]}, "published_revalidation", review["decision"], review["confidence"], review["issues"], review.get("run_id"), publication_binding=self._publication_review_binding(row))
                    if not self._review_passes(review, threshold):
                        blockers.append("published_revalidation_failed")
                if blockers:
                    self.knowledge.retire_publication(int(row["id"]), "AI内容安全审计（非人工）")
                    retired.append({"publication_id": row["id"], "unit_id": row["unit_id"], "blockers": blockers})
                    self._decision(run_id, {"unit_id": row["unit_id"]}, "published_artifact_audit", "retire", 1.0, blockers,
                                   publication_binding=self._publication_review_binding(row))
        finally:
            self.knowledge.on_publication_changed = publication_callback
        return {"scanned": len(rows), "retired": retired, "passed": len(rows) - len(retired), "reused": reused}

    def _build_and_run_evaluation(self, run_id: int, progress, candidate_index_id: int, cancelled=None) -> dict[str, Any]:
        cancelled = cancelled or (lambda: False)
        self._check_finalization(run_id, cancelled)
        if not self.evaluation:
            return {"status": "not_configured"}
        dataset = f"corpus-{run_id}"
        publications = self.db.rows(
            "SELECT DISTINCT unit_id FROM knowledge_publications WHERE status='published' ORDER BY unit_id"
        )
        failures = self._ensure_evaluation_case_coverage(dataset, publications, progress)
        if failures:
            return {"passed": False, "technical_failures": failures, "dataset_name": dataset}
        def search_fn(query, industry, unit_type, top_k):
            self._check_finalization(run_id, cancelled)
            return self.retrieval.search_index(candidate_index_id, query, industry, unit_type, top_k)

        self._check_finalization(run_id, cancelled)
        self.evaluation.repair_confusing_negatives(dataset, search_fn=search_fn)
        self._check_finalization(run_id, cancelled)
        result = self.evaluation.run(
            dataset,
            "silver",
            10,
            search_fn=search_fn,
        )
        metrics = result["metrics"]
        thresholds = {
            "recall_at_k": 0.90,
            "mrr": 0.80,
            "ndcg_at_k": 0.80,
            "negative_rejection_rate": 0.85,
        }
        failed = {key: {"actual": metrics.get(key), "required": value} for key, value in thresholds.items() if metrics.get(key) is None or float(metrics[key]) < value}
        result["thresholds"] = thresholds
        coverage = self._evaluation_coverage(dataset, result)
        result["coverage"] = coverage
        if not coverage["passed"]:
            failed["coverage"] = coverage
        result["passed"] = not failed
        result["failed_thresholds"] = failed
        return result

    def _ensure_evaluation_case_coverage(
        self,
        dataset: str,
        publications: list[dict[str, Any]],
        progress,
    ) -> list[dict[str, Any]]:
        """Generate and independently review evaluation cases with bounded repair.

        Direct and synonym coverage is checked after every review round.  A
        rejected or incomplete model response therefore triggers a genuinely
        new (non-cached) generation round instead of immediately pausing the
        corpus run or leaving a technical task for a person.
        """
        if not self.evaluation:
            return [{"error": "evaluation service is not configured"}]
        publication_ids = [int(item["unit_id"]) for item in publications]
        errors: dict[int, str] = {}
        for generation_round in range(1, 4):
            coverage = self._evaluation_unit_case_coverage(dataset, publication_ids)
            missing = [unit_id for unit_id in publication_ids if not self._unit_eval_covered(coverage.get(unit_id, {}))]
            if not missing:
                return []
            for index, unit_id in enumerate(missing, 1):
                state = coverage.get(unit_id, {})
                kinds = set(state.get("kinds") or [])
                required = sorted({"direct", "synonym"} - kinds)
                progress(
                    "evaluation",
                    97,
                    f"生成评测问题（补齐轮次 {generation_round}/3）{index}/{len(missing)}",
                    {"unit_id": unit_id, "required_query_kinds": required},
                )
                try:
                    self.evaluation.generate_silver_cases(
                        unit_id,
                        8,
                        dataset,
                        generation_round=generation_round,
                        required_query_kinds=required,
                    )
                    errors.pop(unit_id, None)
                except Exception as exc:  # noqa: BLE001
                    errors[unit_id] = f"{type(exc).__name__}: {exc}"

            # Review proposed cases in bounded groups.  The service itself
            # caps each request at 100, so repeat until no candidates remain.
            while True:
                proposed = self.db.rows(
                    """
                    SELECT id FROM retrieval_eval_cases
                    WHERE dataset_name=? AND status='proposed' AND source_type='silver'
                      AND ai_review_status='pending'
                    ORDER BY id LIMIT 100
                    """,
                    (dataset,),
                )
                if not proposed:
                    break
                progress("evaluation", 97, "独立复核评测问题", {"pending_cases": len(proposed)})
                try:
                    result = self.evaluation.review_silver_cases_ai(
                        dataset,
                        [int(item["id"]) for item in proposed],
                    )
                except Exception as exc:  # noqa: BLE001
                    for unit_id in missing:
                        errors[unit_id] = f"{type(exc).__name__}: {exc}"
                    break
                if int(result.get("case_count") or 0) == 0:
                    break

        coverage = self._evaluation_unit_case_coverage(dataset, publication_ids)
        failures: list[dict[str, Any]] = []
        for unit_id in publication_ids:
            state = coverage.get(unit_id, {})
            if self._unit_eval_covered(state):
                continue
            failures.append(
                {
                    "unit_id": unit_id,
                    "error": errors.get(unit_id, "evaluation_case_coverage_incomplete_after_three_rounds"),
                    "approved_count": int(state.get("count") or 0),
                    "approved_query_kinds": sorted(state.get("kinds") or []),
                }
            )
        return failures

    def _evaluation_unit_case_coverage(self, dataset: str, unit_ids: list[int]) -> dict[int, dict[str, Any]]:
        if not unit_ids:
            return {}
        placeholders = ",".join("?" for _ in unit_ids)
        rows = self.db.rows(
            f"""
            SELECT source_unit_id,query_kind,COUNT(*) AS count
            FROM retrieval_eval_cases
            WHERE dataset_name=? AND status='approved' AND source_type='silver'
              AND source_unit_id IN ({placeholders})
            GROUP BY source_unit_id,query_kind
            """,
            (dataset, *unit_ids),
        )
        result: dict[int, dict[str, Any]] = {}
        for row in rows:
            unit_id = int(row["source_unit_id"])
            state = result.setdefault(unit_id, {"count": 0, "kinds": set()})
            state["count"] += int(row["count"])
            state["kinds"].add(str(row["query_kind"]))
        return result

    @staticmethod
    def _unit_eval_covered(state: dict[str, Any]) -> bool:
        return int(state.get("count") or 0) >= 4 and {"direct", "synonym"}.issubset(set(state.get("kinds") or []))

    def _evaluation_coverage(self, dataset: str, result: dict[str, Any]) -> dict[str, Any]:
        cases = self.db.rows(
            "SELECT id,source_unit_id,query_kind,industry,unit_type,expected_unit_ids_json FROM retrieval_eval_cases WHERE dataset_name=? AND status='approved'",
            (dataset,),
        )
        result_map = {int(item["case_id"]): item for item in result.get("results") or []}
        publications = {int(row["unit_id"]) for row in self.db.rows("SELECT unit_id FROM knowledge_publications WHERE status='published'")}
        by_unit: dict[int, set[str]] = {}
        segments: dict[str, list[bool]] = {}
        for case in cases:
            unit_id = int(case.get("source_unit_id") or 0)
            by_unit.setdefault(unit_id, set()).add(str(case["query_kind"]))
            expected = parse_json(case.get("expected_unit_ids_json"), [])
            if expected:
                passed = bool((result_map.get(int(case["id"])) or {}).get("passed"))
                for label in (f"industry:{case['industry'] or '通用'}", f"type:{case['unit_type'] or '通用'}"):
                    segments.setdefault(label, []).append(passed)
        missing = sorted(unit_id for unit_id in publications if not {"direct", "synonym"}.issubset(by_unit.get(unit_id, set())))
        segment_metrics = {
            label: {"positive_cases": len(values), "recall_at_10": round(sum(values) / len(values), 4)}
            for label, values in segments.items()
        }
        failed_segments = {label: value for label, value in segment_metrics.items() if value["positive_cases"] >= 5 and value["recall_at_10"] < 0.85}
        return {
            "published_units": len(publications),
            "units_with_direct_and_synonym": len(publications) - len(missing),
            "missing_unit_ids": missing,
            "segments": segment_metrics,
            "failed_major_segments": failed_segments,
            "passed": not missing and not failed_segments,
        }

    def pause(self, run_id: int, reason: str = "管理员暂停") -> dict[str, Any]:
        self._set_run(run_id, "paused", str(self._raw_run(run_id)["stage"]), reason)
        return self.get_run(run_id)

    def resume(self, run_id: int) -> dict[str, Any]:
        run = self._raw_run(run_id)
        if run["status"] in {"completed", "cancelled"}:
            return self.get_run(run_id)
        self._consolidate_manual_tasks(run_id)
        latest_job = self._latest_corpus_job(run_id)
        if latest_job and latest_job["status"] == "failed":
            self._finalization_checkpoint(
                run_id, str(run["stage"]), int(run["progress"]),
                acknowledged_failed_job_id=int(latest_job["id"]), resumed_at=iso_now(),
            )
        with self.db.connect() as conn:
            conn.execute("UPDATE corpus_runs SET status='running',pause_reason='',consecutive_errors=0,recent_results_json='[]',updated_at=CURRENT_TIMESTAMP WHERE id=? AND status IN ('pending','running','paused')", (run_id,))
        self._dispatch_next(run_id)
        return self.get_run(run_id)

    def cancel(self, run_id: int) -> dict[str, Any]:
        self._set_run(run_id, "cancelled", "cancelled", "管理员取消")
        return self.get_run(run_id)

    def _set_run(self, run_id: int, status: str, stage: str, reason: str) -> None:
        with self.db.connect() as conn:
            conn.execute("UPDATE corpus_runs SET status=?,stage=?,pause_reason=?,updated_at=CURRENT_TIMESTAMP WHERE id=?", (status, stage, reason, run_id))

    def _raw_run(self, run_id: int) -> dict[str, Any]:
        row = self.db.row("SELECT * FROM corpus_runs WHERE id=?", (run_id,))
        if not row:
            raise KeyError("全库收口任务不存在")
        return row

    def _refresh(self, run_id: int) -> dict[str, Any]:
        counts = self.db.rows("SELECT status,stage,item_kind,terminal_reason,COUNT(*) AS count FROM corpus_run_items WHERE run_id=? GROUP BY status,stage,item_kind,terminal_reason", (run_id,))
        total = sum(int(row["count"]) for row in counts)
        terminal = sum(int(row["count"]) for row in counts if row["status"] == "terminal")
        reasons: Counter[str] = Counter()
        stages: Counter[str] = Counter()
        kinds: Counter[str] = Counter()
        for row in counts:
            count = int(row["count"])
            stages[f"{row['stage']}:{row['status']}"] += count
            kinds[f"{row['item_kind']}:{row['status']}"] += count
            if row["terminal_reason"]:
                reasons[str(row["terminal_reason"])] += count
        counters = {
            "total": total,
            "terminal": terminal,
            "remaining": total - terminal,
            "by_reason": dict(reasons),
            "by_stage": dict(stages),
            "by_kind": dict(kinds),
        }
        progress = int(terminal * 95 / total) if total else 95
        with self.db.connect() as conn:
            conn.execute(
                """
                UPDATE corpus_runs SET progress=CASE WHEN status='completed' THEN 100
                    WHEN stage IN ('publication_audit','index','evaluation','activation','reports') AND progress>? THEN progress
                    ELSE ? END,counters_json=?,updated_at=CURRENT_TIMESTAMP WHERE id=?
                """,
                (progress, progress, json.dumps(counters, ensure_ascii=False), run_id),
            )
        return self.get_run(run_id)

    def get_run(self, run_id: int) -> dict[str, Any]:
        row = self._raw_run(run_id)
        row["policy"] = parse_json(row.pop("policy_json", "{}"), {})
        row["counters"] = parse_json(row.pop("counters_json", "{}"), {})
        row["checkpoint"] = parse_json(row.pop("checkpoint_json", "{}"), {})
        row["recent_results"] = parse_json(row.pop("recent_results_json", "[]"), [])
        row["completion"] = self.completion(run_id)
        return row

    def completion(self, run_id: int | None = None) -> dict[str, Any]:
        if run_id is None:
            latest = self.db.row("SELECT id FROM corpus_runs ORDER BY id DESC LIMIT 1")
            if not latest:
                return {"corpus_terminal_rate": 0.0, "formal_eligible_publication_rate": 0.0, "manual_legal_open": 0, "legal_clearance_ready": False, "index_consistent": False}
            run_id = int(latest["id"])
        totals = self.db.row("SELECT COUNT(*) AS total,SUM(CASE WHEN status='terminal' THEN 1 ELSE 0 END) AS terminal FROM corpus_run_items WHERE run_id=?", (run_id,)) or {}
        candidate = self.db.row(
            """
            SELECT COUNT(DISTINCT CASE WHEN c.status IN ('accepted','rejected') THEN c.id END) AS terminal,
                COUNT(DISTINCT CASE WHEN c.status='accepted' AND p.id IS NOT NULL THEN c.id END) AS published,
                COUNT(DISTINCT CASE WHEN c.status='accepted' THEN c.id END) AS eligible
            FROM knowledge_ai_candidates c
            JOIN knowledge_ai_pipeline_runs r ON r.id=c.pipeline_run_id
            JOIN corpus_run_items i ON i.pipeline_run_id=r.id AND i.run_id=?
            LEFT JOIN knowledge_publications p ON p.unit_id=c.unit_id AND p.status='published'
            WHERE c.status<>'superseded'
            """,
            (run_id,),
        ) or {"published": 0, "eligible": 0}
        manual = self.db.row("SELECT COUNT(*) AS count FROM governance_tasks WHERE run_id=? AND status='open' AND requires_human=1", (run_id,)) or {}
        published_count = int((self.db.row("SELECT COUNT(*) AS count FROM knowledge_publications WHERE status='published'") or {}).get("count", 0))
        index_status = self.retrieval.status()
        index = index_status.get("active_index") or {}
        total = int(totals.get("total") or 0)
        terminal = int(totals.get("terminal") or 0)
        eligible = int(candidate.get("eligible") or 0)
        published = int(candidate.get("published") or 0)
        return {
            "run_id": run_id,
            "corpus_total": total,
            "corpus_terminal": terminal,
            "corpus_terminal_rate": round(terminal / total, 6) if total else 0.0,
            "formal_eligible": eligible,
            "formal_published": published,
            "formal_eligible_publication_rate": round(published / eligible, 6) if eligible else (1.0 if terminal == total and total else 0.0),
            "manual_legal_open": int(manual.get("count") or 0),
            "legal_clearance_ready": int(manual.get("count") or 0) == 0,
            "published_total": published_count,
            "active_index_publications": int(index.get("publication_count") or 0),
            "index_consistent": bool(index_status.get("index_consistent")),
        }

    def list_manual_tasks(self, status: str = "open", limit: int = 500, run_id: int | None = None) -> list[dict[str, Any]]:
        if run_id is None:
            latest = self.db.row("SELECT id FROM corpus_runs ORDER BY id DESC LIMIT 1")
            run_id = int(latest["id"]) if latest else 0
        clauses = ["t.run_id=?"]
        params: list[Any] = [run_id]
        if status:
            clauses.append("t.status=?")
            params.append(status)
        where = "WHERE " + " AND ".join(clauses)
        return self.db.rows(f"SELECT t.*,s.file_name,s.relative_path FROM governance_tasks t LEFT JOIN source_files s ON s.id=t.source_id {where} ORDER BY t.id DESC LIMIT ?", [*params, max(1, min(limit, 500))])

    def resolve_manual_task(self, task_id: int, action: str, resolution: str, user_id: int | None) -> dict[str, Any]:
        if action not in {"approve", "exclude"}:
            raise ValueError("action必须为approve或exclude")
        if not resolution.strip():
            raise ValueError("处理说明不能为空")
        task = self.db.row("SELECT * FROM governance_tasks WHERE id=?", (task_id,))
        if not task:
            raise KeyError("人工事项不存在")
        source_rows = self.db.rows("SELECT source_id FROM governance_task_sources WHERE task_id=?", (task_id,))
        source_ids = [int(row["source_id"]) for row in source_rows] or ([int(task["source_id"])] if task.get("source_id") else [])
        with self.db.connect() as conn:
            conn.execute("UPDATE governance_tasks SET status='resolved',resolution=?,resolved_by=?,resolved_at=CURRENT_TIMESTAMP WHERE id=?", (f"{action}: {resolution.strip()}", user_id, task_id))
            if task["task_type"] == "asset_license" and source_ids:
                placeholders = ",".join("?" for _ in source_ids)
                conn.execute(
                    f"UPDATE knowledge_assets SET governance_status=?,review_status=?,exclusion_reason=? WHERE source_id IN ({placeholders})",
                    ("approved" if action == "approve" else "excluded", "approved" if action == "approve" else "rejected", resolution.strip(), *source_ids),
                )
        row = self.db.row("SELECT * FROM governance_tasks WHERE id=?", (task_id,))
        if action == "approve" and task["task_type"] == "asset_license":
            self._resume_authorized_office_sources(int(task["run_id"]), source_ids)
        return row

    def _resume_authorized_office_sources(self, run_id: int, source_ids: list[int]) -> None:
        if not source_ids:
            return
        placeholders = ",".join("?" for _ in source_ids)
        with self.db.connect() as conn:
            conn.execute(
                f"""
                UPDATE corpus_run_items SET stage='ai',status='pending',terminal_reason='',completed_at=NULL,updated_at=CURRENT_TIMESTAMP
                WHERE run_id=? AND source_id IN ({placeholders}) AND document_id IS NOT NULL
                  AND status='terminal' AND terminal_reason='manual_legal_pending'
                """,
                (run_id, *source_ids),
            )
        run = self._raw_run(run_id)
        if run["status"] == "completed":
            with self.db.connect() as conn:
                conn.execute("UPDATE corpus_runs SET status='running',stage='ai',completed_at=NULL,updated_at=CURRENT_TIMESTAMP WHERE id=?", (run_id,))
        self._dispatch_next(run_id)

    def write_reports(self, run_id: int, extra: dict[str, Any] | None = None) -> dict[str, str]:
        root = self.knowledge.settings.qa_root / "knowledge"
        root.mkdir(parents=True, exist_ok=True)
        completion = self.completion(run_id)
        run = self.get_run(run_id)
        sources = self.db.rows("SELECT i.source_id,s.relative_path,i.item_kind,i.status,i.terminal_reason,i.error_code,i.error_message FROM corpus_run_items i JOIN source_files s ON s.id=i.source_id WHERE i.run_id=? ORDER BY i.source_id", (run_id,))
        publications = self.db.rows("SELECT p.id,p.unit_id,u.title,u.unit_type,u.industry,p.publication_version,p.content_hash FROM knowledge_publications p JOIN knowledge_units u ON u.id=p.unit_id WHERE p.status='published' ORDER BY p.id", ())
        exclusions = [item for item in sources if item["terminal_reason"] not in {"knowledge_published", "asset_registered"}]
        manual = self.db.rows(
            "SELECT t.*,s.file_name,s.relative_path FROM governance_tasks t LEFT JOIN source_files s ON s.id=t.source_id WHERE t.run_id=? ORDER BY t.id",
            (run_id,),
        )
        assets = self.db.rows(
            "SELECT a.id,a.source_id,s.relative_path,a.asset_type,a.content_hash,a.perceptual_hash,a.width,a.height,a.governance_status,a.exclusion_reason FROM knowledge_assets a LEFT JOIN source_files s ON s.id=a.source_id WHERE a.source_id IN (SELECT source_id FROM corpus_run_items WHERE run_id=?) ORDER BY a.id",
            (run_id,),
        )
        payload = {"schema_version": "1.0", "generated_at": iso_now(), "run": run, "completion": completion, "extra": extra or {}, "counts": {"sources": len(sources), "publications": len(publications), "exclusions": len(exclusions), "manual_tasks": len(manual)}}
        write_json_atomic(root / "knowledge_completion.json", payload)
        finalization = run.get("checkpoint", {}).get("finalization", {})
        evaluation = (extra or {}).get("evaluation") or finalization.get("evaluation") or {}
        evaluation_status = "通过" if evaluation.get("passed") else ("未通过" if evaluation else "尚未完成")
        lines = ["# 知识库全量整理验收", "", f"- 运行状态：{run['status']}", f"- 当前阶段：{run['stage']}", f"- 检索评测：{evaluation_status}", f"- 暂停或失败原因：{run.get('pause_reason') or '无'}", f"- 全资料闭环率：{completion['corpus_terminal_rate']:.2%}", f"- 正式可用发布率：{completion['formal_eligible_publication_rate']:.2%}", f"- 正式发布知识：{completion['published_total']} 条", f"- 人工法律/授权事项：{completion['manual_legal_open']} 项", f"- 活动索引一致：{'是' if completion['index_consistent'] else '否'}", "", "人工事项不进入正式检索，只有全部解决后 `legal_clearance_ready` 才为 true。"]
        write_text_atomic(root / "knowledge_completion.md", "\n".join(lines) + "\n")
        for name, rows in (("source_disposition.json", sources), ("publication_manifest.json", publications), ("exclusion_manifest.json", exclusions), ("manual_tasks.json", manual), ("asset_governance.json", assets)):
            write_json_atomic(root / name, rows)
        hashes = []
        for path in sorted(root.glob("*")):
            if path.is_file() and path.name != "checksums.json":
                hashes.append({"file": path.name, "size_bytes": path.stat().st_size, "sha256": sha256_file(path)})
        write_json_atomic(root / "checksums.json", hashes)
        return {"json": str(root / "knowledge_completion.json"), "markdown": str(root / "knowledge_completion.md")}
