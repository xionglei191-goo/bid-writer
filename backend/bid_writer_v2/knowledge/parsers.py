from __future__ import annotations

import re
import shutil
import subprocess
import zipfile
from xml.etree import ElementTree
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from ..settings import Settings
from ..utils import normalize_text


class OcrRequired(RuntimeError):
    def __init__(self, page_count: int) -> None:
        super().__init__("PDF文本量过低，需要OCR")
        self.page_count = page_count


@dataclass
class ParsedDocument:
    title: str
    markdown: str
    parser: str
    page_count: int


def parse_document(path: Path, source_id: int, settings: Settings) -> ParsedDocument:
    extension = path.suffix.lower()
    if extension == ".docx":
        return _parse_docx(path, source_id, settings)
    if extension == ".pdf":
        return _parse_pdf(path)
    if extension == ".doc":
        converted = _convert_doc(path, settings.cache_root / "doc_conversion" / str(source_id))
        return _parse_docx(converted, source_id, settings)
    if extension in {".md", ".txt"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return ParsedDocument(path.stem, normalize_text(text), "plain_text", 0)
    raise ValueError(f"暂不支持正文解析：{extension or '无扩展名'}")


def _parse_docx(path: Path, source_id: int, settings: Settings) -> ParsedDocument:
    blocks: list[str] = [f"# {path.stem}"]
    try:
        document = Document(path)
        for paragraph in document.paragraphs:
            text = normalize_text(paragraph.text)
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style else "").lower()
            match = re.search(r"heading\s*(\d+)", style)
            if match:
                level = max(1, min(int(match.group(1)), 6))
                blocks.append(f"{'#' * level} {text}")
            else:
                blocks.append(text)
        for table_index, table in enumerate(document.tables, 1):
            rows = [[normalize_text(cell.text).replace("|", "\\|") for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            blocks.append(f"## 表格 {table_index}")
            blocks.append("| " + " | ".join(rows[0]) + " |")
            blocks.append("| " + " | ".join(["---"] * width) + " |")
            blocks.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    except KeyError as exc:
        # Some legacy Word files contain a broken relationship whose target is
        # literally "NULL".  The main document XML is still usable and is read
        # directly so that one bad embedded object does not discard the text.
        if "NULL" not in str(exc):
            raise
        blocks.extend(_parse_docx_xml_fallback(path))

    asset_dir = settings.knowledge_directories["assets"] / "导入图片" / str(source_id)
    with zipfile.ZipFile(path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        for index, name in enumerate(media, 1):
            suffix = Path(name).suffix.lower() or ".bin"
            destination = asset_dir / f"image_{index:03d}{suffix}"
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_bytes(archive.read(name))
            blocks.append(f"![原文图片{index}]({destination.as_posix()})")
    return ParsedDocument(path.stem, normalize_text("\n\n".join(blocks)), "python-docx", 0)


def _parse_docx_xml_fallback(path: Path) -> list[str]:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    blocks: list[str] = []
    body = root.find(f"{namespace}body")
    if body is None:
        return blocks
    table_index = 0
    for child in body:
        if child.tag == f"{namespace}p":
            text = normalize_text("".join(node.text or "" for node in child.iter(f"{namespace}t")))
            if text:
                blocks.append(text)
        elif child.tag == f"{namespace}tbl":
            rows: list[list[str]] = []
            for row in child.findall(f"{namespace}tr"):
                values = []
                for cell in row.findall(f"{namespace}tc"):
                    values.append(normalize_text("".join(node.text or "" for node in cell.iter(f"{namespace}t"))).replace("|", "\\|"))
                if values:
                    rows.append(values)
            if rows:
                table_index += 1
                width = max(len(row) for row in rows)
                rows = [row + [""] * (width - len(row)) for row in rows]
                blocks.extend([
                    f"## 表格 {table_index}",
                    "| " + " | ".join(rows[0]) + " |",
                    "| " + " | ".join(["---"] * width) + " |",
                    *("| " + " | ".join(row) + " |" for row in rows[1:]),
                ])
    return blocks


def _parse_pdf(path: Path) -> ParsedDocument:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, 1):
        text = normalize_text(page.extract_text() or "")
        pages.append(f"<!-- page:{index} -->\n{text}")
    combined = "\n\n".join(pages)
    if len(re.sub(r"\s+", "", combined)) < max(200, len(reader.pages) * 40):
        raise OcrRequired(len(reader.pages))
    return ParsedDocument(path.stem, f"# {path.stem}\n\n{combined}", "pypdf", len(reader.pages))


def _convert_doc(path: Path, target_dir: Path) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    libreoffice = shutil.which("soffice") or shutil.which("libreoffice")
    if libreoffice:
        subprocess.run(
            [libreoffice, "--headless", "--convert-to", "docx", "--outdir", str(target_dir), str(path)],
            check=True,
            capture_output=True,
            timeout=180,
        )
        output = target_dir / f"{path.stem}.docx"
        if output.exists():
            return output
    if shutil.which("powershell"):
        script = (
            "$word=New-Object -ComObject Word.Application;"
            "$word.Visible=$false;"
            f"$doc=$word.Documents.Open('{str(path).replace("'", "''")}');"
            f"$doc.SaveAs2('{str(target_dir / (path.stem + '.docx')).replace("'", "''")}',16);"
            "$doc.Close();$word.Quit()"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, timeout=180)
        output = target_dir / f"{path.stem}.docx"
        if output.exists():
            return output
    raise RuntimeError("无法转换DOC：未找到LibreOffice且Microsoft Word自动化不可用")
