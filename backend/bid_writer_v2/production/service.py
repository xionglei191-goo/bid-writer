from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
import hashlib
import mimetypes
import os
import zipfile
from contextlib import nullcontext
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from bid_writer.artifact_audit import audit_docx_path, audit_text, audit_zip_path

from ..ai_runtime import AiRuntime, SECTION_DRAFT_PROMPT
from ..database import Database
from ..audit import AuditService
from ..evidence import EvidenceService
from ..project_evidence import factual_profile
from ..knowledge.service import KnowledgeService
from ..llm import LlmClient
from ..settings import Settings
from ..storage import ObjectStorage
from ..utils import content_hash, normalize_text, parse_json, write_text_atomic
from .generation_state import context_fingerprint, part_fingerprint, public_generation, read_generation
from .requirements_workflow import RequirementsWorkflow
from .workbench import build_workbench, response_metrics, human_confirmation_indices, generation_incomplete
from .response_text import copies_requirement


DEFAULT_OUTLINE = [
    "工程概况与编制依据",
    "施工总体部署",
    "施工准备与资源配置",
    "主要施工方案与技术措施",
    "工程重点难点分析及对策",
    "质量保证体系与措施",
    "安全文明施工与环境保护",
    "施工进度计划及保证措施",
    "施工总平面布置",
    "季节性施工与应急预案",
    "BIM及新技术应用",
]


class ProductionService:
    def __init__(
        self,
        db: Database,
        settings: Settings,
        knowledge: KnowledgeService,
        llm: LlmClient | None = None,
        ai_runtime: AiRuntime | None = None,
        evidence: EvidenceService | None = None,
        storage: ObjectStorage | None = None,
        audit: AuditService | None = None,
    ) -> None:
        self.db = db
        self.settings = settings
        self.knowledge = knowledge
        self.llm = llm or LlmClient()
        self.ai_runtime = ai_runtime or AiRuntime(db, self.llm)
        self.evidence = evidence or EvidenceService(db)
        self.storage = storage or ObjectStorage(db, settings)
        self.audit = audit or AuditService(db)
        self.requirements_workflow = RequirementsWorkflow(self)

    def create_project(self, payload: dict[str, Any]) -> dict[str, Any]:
        name = str(payload.get("name") or "").strip()
        if not name:
            raise ValueError("项目名称不能为空")
        profile = payload.get("profile") if isinstance(payload.get("profile"), dict) else {}
        profile = {key: value for key, value in profile.items() if key != "delivery_confirmation"}
        with self.db.connect() as conn:
            cursor = conn.execute(
                "INSERT INTO projects(name,industry,project_type,region,source_text,profile_json) VALUES (?,?,?,?,?,?)",
                (
                    name,
                    payload.get("industry") or "通用",
                    payload.get("project_type") or "",
                    payload.get("region") or "",
                    payload.get("source_text") or "",
                    json.dumps(profile, ensure_ascii=False),
                ),
            )
            project_id = int(cursor.lastrowid)
        return self.get_project(project_id)

    def list_projects(self) -> list[dict[str, Any]]:
        with self.db.connect() as conn:
            rows = conn.execute(
                """
                SELECT p.*,
                    (SELECT COUNT(*) FROM project_requirements r WHERE r.project_id=p.id) AS requirement_count,
                    (SELECT COUNT(*) FROM project_sections s WHERE s.project_id=p.id) AS section_count,
                    (SELECT COUNT(*) FROM project_drafts d WHERE d.project_id=p.id) AS draft_count
                FROM projects p ORDER BY p.id DESC
                """
            ).fetchall()
        items = [dict(row) for row in rows]
        for item in items:
            item["profile"] = parse_json(item.pop("profile_json", "{}"), {})
        return items

    def get_project(self, project_id: int) -> dict[str, Any]:
        with self.db.connect() as conn:
            return self._get_project_from_connection(conn, project_id)

    def _get_project_from_connection(self, conn, project_id: int) -> dict[str, Any]:
        row = conn.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
        if not row:
            raise KeyError("项目不存在")
        item = dict(row)
        item["profile"] = parse_json(item.pop("profile_json", "{}"), {})
        item["requirements"] = [dict(value) for value in conn.execute("SELECT * FROM project_requirements WHERE project_id=? ORDER BY id", (project_id,))]
        sections = [dict(value) for value in conn.execute("SELECT * FROM project_sections WHERE project_id=? ORDER BY order_no", (project_id,))]
        for section in sections:
            section["requirement_ids"] = parse_json(section.pop("requirement_ids_json", "[]"), [])
            draft = conn.execute("SELECT * FROM project_drafts WHERE section_id=? ORDER BY version_no DESC LIMIT 1", (section["id"],)).fetchone()
            section["draft"] = dict(draft) if draft else None
            if section["draft"]:
                section["draft"]["citations"] = parse_json(section["draft"].pop("citations_json", "[]"), [])
                section["draft"]["confirmations"] = parse_json(section["draft"].pop("confirmations_json", "[]"), [])
                section["draft"]["confirmation_required_indices"] = human_confirmation_indices(section["draft"])
                # Approval fingerprints describe the saved body, including when
                # a legacy or interrupted write left its cached hash stale.
                section["draft"]["content_hash"] = content_hash(section["draft"]["content"])
                section["draft"]["generation"] = public_generation(section["draft"])
                section["draft"].pop("generation_json", None)
                section["draft"]["confirmation_resolutions"] = [
                    dict(value)
                    for value in conn.execute(
                        "SELECT confirmation_index,confirmation_text,resolution,resolver,created_at "
                        "FROM confirmation_resolutions WHERE draft_id=? AND target_hash=? ORDER BY confirmation_index",
                        (section["draft"]["id"], section["draft"]["content_hash"]),
                    )
                ]
                section["draft"]["confirmation_resolutions"] = [resolution for resolution in section["draft"]["confirmation_resolutions"]
                    if 0 <= resolution["confirmation_index"] < len(section["draft"]["confirmations"])
                    and resolution["confirmation_text"] == section["draft"]["confirmations"][resolution["confirmation_index"]]]
                section["draft"]["claims"] = self.evidence.list_claims(int(section["draft"]["id"]), conn=conn)
        item["sections"] = sections
        item["quality_issues"] = [
            dict(value)
            for value in conn.execute(
                "SELECT * FROM quality_issues WHERE project_id=? ORDER BY CASE severity WHEN 'high' THEN 1 WHEN 'medium' THEN 2 ELSE 3 END,id DESC",
                (project_id,),
            )
        ]
        item["evidence_source"] = self.evidence.project_source_status(project_id, conn=conn)
        self.requirements_workflow.enrich(item, conn)
        response_pairs = set(response_metrics(self, item, conn=conn)["response_pairs"])
        technical_ids = set(item["requirements_workflow"]["formal_technical_ids"])
        for section in sections:
            if section.get("draft"):
                section["draft"]["missing_response_requirement_ids"] = [rid for rid in section["requirement_ids"]
                    if rid in technical_ids and (section["draft"]["id"], rid) not in response_pairs]
        return item

    def _lock_project_for_update(self, conn, project_id: int) -> None:
        if self.db.backend == "sqlite":
            if not conn.in_transaction:
                conn.execute("BEGIN IMMEDIATE")
            return
        if not conn.execute("SELECT id FROM projects WHERE id=? FOR UPDATE", (project_id,)).fetchone():
            raise KeyError("项目不存在")
        # Match the existing draft-edit order (draft before section), while the
        # parent row lock also prevents insertion of new child versions.
        for table in ("project_drafts", "project_sections", "project_requirements"):
            conn.execute(f"SELECT id FROM {table} WHERE project_id=? ORDER BY id FOR UPDATE", (project_id,)).fetchall()

    def update_project(self, project_id: int, payload: dict[str, Any], expected_project_hash: str | None = None) -> dict[str, Any]:
        with self.db.connect() as conn:
            self._lock_project_for_update(conn, project_id)
            current = self._get_project_from_connection(conn, project_id)
            if expected_project_hash is not None and expected_project_hash != self._project_hash(current):
                raise ValueError("项目内容已变化，请重新预览整本后确认")
            profile = dict(current["profile"])
            if isinstance(payload.get("profile"), dict):
                profile.update({key: value for key, value in payload["profile"].items() if key != "delivery_confirmation"})
            updated = {**current, **{key: payload[key] for key in ("name", "industry", "project_type", "region", "source_text") if key in payload}, "profile": profile}
            profile_update = payload.get("profile") or {}
            # Keep the explicit profile confirmation API compatible, binding
            # every approval to the same locked snapshot as its write.
            if any(profile_update.get(key) is True for key in ("compliance_confirmed", "manual_finalized", "final_approved")):
                previous_confirmation = self._delivery_confirmation(current)
                reviewer = normalize_text(str(profile.get("professional_reviewer") or profile.get("reviewed_by") or ""))
                same_content = self._project_hash(current) == self._project_hash(updated) and reviewer == previous_confirmation["professional_reviewer"]
                profile["delivery_confirmation"] = {
                    "target_hash": self._project_hash(updated),
                    "professional_reviewer": reviewer,
                    "compliance_confirmed": profile_update.get("compliance_confirmed") is True or (same_content and previous_confirmation["compliance_confirmed"]),
                    "manual_finalized": profile_update.get("manual_finalized") is True or profile_update.get("final_approved") is True or (same_content and previous_confirmation["manual_finalized"]),
                    "confirmed_at": datetime.now(timezone.utc).isoformat(),
                }
            if all(updated.get(key) == current.get(key) for key in ("name", "industry", "project_type", "region", "source_text", "profile")):
                return current
            conn.execute(
                """
                UPDATE projects SET name=?,industry=?,project_type=?,region=?,source_text=?,profile_json=?,updated_at=CURRENT_TIMESTAMP
                WHERE id=?
                """,
                (
                    payload.get("name", current["name"]),
                    payload.get("industry", current["industry"]),
                    payload.get("project_type", current["project_type"]),
                    payload.get("region", current["region"]),
                    payload.get("source_text", current["source_text"]),
                    json.dumps(profile, ensure_ascii=False),
                    project_id,
                ),
            )
            conn.execute("UPDATE delivery_manifests SET invalidated_at=CURRENT_TIMESTAMP,status='invalidated' WHERE project_id=? AND invalidated_at IS NULL", (project_id,))
            if self._project_hash(current) != self._project_hash(updated):
                conn.execute("UPDATE generation_runs SET status='invalidated' WHERE project_id=? AND status='completed'", (project_id,))
        return self.get_project(project_id)

    def preview_project(self, project_id: int) -> dict[str, Any]:
        project = self.get_project(project_id)
        return {
            "project_id": project_id,
            "project_hash": self._project_hash(project),
            "markdown": self._project_markdown(project),
            "confirmation": self._delivery_confirmation(project),
        }

    def workbench(self, project_id: int) -> dict[str, Any]:
        return build_workbench(self, self.get_project(project_id))

    def refresh_project_evidence(self, project_id: int, progress=None, cancelled=None) -> dict[str, Any]:
        progress = progress or (lambda *_args, **_kwargs: None)
        cancelled = cancelled or (lambda: False)
        project = self.get_project(project_id)
        drafts = [section["draft"] for section in project["sections"] if section.get("draft")]
        results = []
        for index, draft in enumerate(drafts):
            if cancelled():
                return {"project_id": project_id, "cancelled": True, "analyzed_count": len(results), "drafts": results}
            progress("checking", int(index / max(1, len(drafts)) * 95), f"重新核验第{index + 1}/{len(drafts)}章证据", {"draft_id": draft["id"]})
            with self.db.connect() as conn:
                self._lock_project_for_update(conn, project_id)
                refreshed = self.evidence.refresh_project_evidence(project_id, draft_ids=[draft["id"]], conn=conn)
                analysis = refreshed["drafts"][0]
                row = conn.execute("SELECT content,confirmations_json FROM project_drafts WHERE id=?", (draft["id"],)).fetchone()
                if content_hash(row["content"]) != analysis["current_content_hash"] or analysis.get("stale"):
                    raise ValueError("核验期间正文或来源发生变化，请重新检查")
                confirmations = [str(text) for text in parse_json(row["confirmations_json"], [])
                                 if not str(text).startswith("高风险表述缺少充分证据：")]
                confirmations.extend(f"高风险表述缺少充分证据：{text[:160]}" for text in analysis["unsupported_high"])
                conn.execute("UPDATE project_drafts SET confirmations_json=?,status='draft' WHERE id=?",
                             (json.dumps(list(dict.fromkeys(confirmations)), ensure_ascii=False), draft["id"]))
                conn.execute("UPDATE draft_review_decisions SET decision='invalidated' WHERE draft_id=? AND decision='approved'", (draft["id"],))
                conn.execute("UPDATE requirement_responses SET review_status='invalidated' WHERE draft_id=?", (draft["id"],))
                conn.execute("UPDATE project_sections SET status='drafted' WHERE id=(SELECT section_id FROM project_drafts WHERE id=?)", (draft["id"],))
                conn.execute("UPDATE delivery_manifests SET status='invalidated',invalidated_at=CURRENT_TIMESTAMP WHERE project_id=? AND invalidated_at IS NULL", (project_id,))
                results.append(analysis)
        progress("completed", 100, "已重新核验证据，正文保持原文，章节需重新签审", {"analyzed_count": len(results)})
        return {"project_id": project_id, "analyzed_count": len(results), "drafts": results,
                "source_status": self.evidence.project_source_status(project_id)}

    def confirm_final_review(self, project_id: int, project_hash: str, professional_reviewer: str,
                             compliance_confirmed: bool, manual_finalized: bool) -> dict[str, Any]:
        if not normalize_text(professional_reviewer):
            raise ValueError("专业复核人不能为空")
        if not compliance_confirmed or not manual_finalized:
            raise ValueError("请由责任人完成合规确认和整本人工定稿")
        self.update_project(project_id, {"profile": {
            "professional_reviewer": normalize_text(professional_reviewer),
            "compliance_confirmed": True,
            "manual_finalized": True,
        }}, expected_project_hash=project_hash)
        self.audit.record("delivery.final_review", "project", project_id,
                          details={"project_hash": project_hash, "professional_reviewer": normalize_text(professional_reviewer)})
        return self.preview_project(project_id)

    def _delivery_confirmation(self, project: dict[str, Any]) -> dict[str, Any]:
        profile = project.get("profile") or {}
        confirmation = profile.get("delivery_confirmation") or {}
        reviewer = normalize_text(str(profile.get("professional_reviewer") or profile.get("reviewed_by") or ""))
        valid = bool(confirmation.get("target_hash") == self._project_hash(project)
                     and reviewer and confirmation.get("professional_reviewer") == reviewer)
        compliance_confirmed = valid and profile.get("compliance_confirmed") is True and confirmation.get("compliance_confirmed") is True
        manual_finalized = valid and (profile.get("manual_finalized") is True or profile.get("final_approved") is True) and confirmation.get("manual_finalized") is True
        return {
            "valid": compliance_confirmed and manual_finalized,
            "professional_reviewer": reviewer,
            "target_hash": confirmation.get("target_hash"),
            "confirmed_at": confirmation.get("confirmed_at"),
            "compliance_confirmed": compliance_confirmed,
            "manual_finalized": manual_finalized,
        }

    def parse_requirements(self, project_id: int) -> dict[str, Any]:
        from .requirements_workflow import RequirementsWorkflow

        project = self.get_project(project_id)
        if any(section.get("draft") for section in project.get("sections", [])):
            raise ValueError("项目已有章节草稿或签审，不能重新解析要求；请新建项目后重新解析，避免覆盖已有成果")
        if project["requirements"]:
            # Existing source clauses are stable records. Reclassification must
            # not recreate their IDs or delete existing review provenance.
            snapshot = RequirementsWorkflow(self).refresh_suggestions(project_id)
            return {"project_id": project_id, "requirements": snapshot["items"], "retained_originals": True}
        source_text = normalize_text(project.get("source_text") or "")
        candidates = self._requirement_paragraphs(source_text, str(project.get("name") or ""))
        requirements: list[dict[str, Any]] = []
        seen: set[tuple[str, int | None]] = set()
        keywords = ["项目", "工程", "应", "须", "不得", "禁止", "严禁", "评分", "要求", "工期", "质量", "安全", "施工", "技术", "方案", "响应"]
        for candidate in candidates:
            line = candidate["content"]
            key = (line, candidate["source_page"])
            if key in seen:
                continue
            # "不得分" is a scoring result and "不得力" describes performance;
            # neither is a prohibition. Keep actual prohibitions such as 不得分包.
            veto = bool(candidate.get("inherited_veto") or re.search(
                r"废标|否决|禁止|严禁|不得(?!力|分(?:$|[\s。；，：！？、,.!?;:）)]|的|者|或|并|且|则|时))", line
            ))
            if not candidate.get("scoring") and not candidate.get("inherited_veto") and ((len(line) < 8 and not re.search(r"(?:不得|禁止|严禁|必须).+", line)) or not any(keyword in line for keyword in keywords)):
                continue
            seen.add(key)
            priority = "high" if candidate.get("scoring") or veto or any(keyword in line for keyword in ["评分", "不得分", "必须"]) else "normal"
            if candidate.get("scoring"):
                kind = "scoring"
            elif veto:
                kind = "veto"
            elif any(keyword in line for keyword in ["评分", "得分", "分值"]):
                kind = "scoring"
            elif any(keyword in line for keyword in ["资质", "资格", "业绩", "证书"]):
                kind = "qualification"
            elif any(keyword in line for keyword in ["承诺", "保证"]):
                kind = "commitment"
            elif any(keyword in line for keyword in ["参数", "不低于", "不少于", "不大于", "工期"]):
                kind = "parameter"
            elif any(keyword in line for keyword in ["提交", "交付", "附件", "表格"]):
                kind = "deliverable"
            elif any(keyword in line for keyword in ["待定", "另行", "未明确"]):
                kind = "confirmation"
            else:
                kind = "technical"
            requirements.append({"content": line, "priority": priority, "kind": kind, "source_page": candidate["source_page"]})
        if not requirements and source_text and not re.search(r"\[第\s*\d+\s*页\]", source_text):
            requirements.append({"content": source_text, "priority": "normal", "kind": "technical", "source_page": None})
        with self.db.connect() as conn:
            self._lock_project_for_update(conn, project_id)
            current = self._get_project_from_connection(conn, project_id)
            if current["requirements"] or any(section.get("draft") for section in current["sections"]):
                raise ValueError("项目条款或草稿已变化，请刷新后重新操作，原始条款将保留")
            if normalize_text(current.get("source_text") or "") != source_text:
                raise ValueError("招标原文已变化，请刷新后重新解析")
            for index, requirement in enumerate(requirements, 1):
                conn.execute(
                    "INSERT INTO project_requirements(project_id,requirement_key,kind,content,priority,source_page) VALUES (?,?,?,?,?,?)",
                    (project_id, f"REQ-{index:04d}", requirement["kind"], requirement["content"], requirement["priority"], requirement["source_page"]),
                )
        snapshot = RequirementsWorkflow(self).refresh_suggestions(project_id)
        return {"project_id": project_id, "requirements": snapshot["items"]}

    @staticmethod
    def _requirement_paragraphs(source_text: str, project_name: str = "") -> list[dict[str, Any]]:
        """Retain PDF page provenance and join wrapped prose/table score rows."""
        entries: list[dict[str, Any]] = []
        page = None
        for raw_line in source_text.splitlines():
            line = raw_line.strip()
            marker = re.fullmatch(r"\[第\s*(\d+)\s*页\]", line)
            if marker:
                page = int(marker.group(1))
                continue
            if line:
                entries.append({"text": line, "page": page})
        # Printed page numbers may appear first or last in extracted text.
        page_positions: dict[int | None, list[int]] = {}
        for index, entry in enumerate(entries):
            page_positions.setdefault(entry["page"], []).append(index)
        ignored: set[int] = set()
        for page_number, positions in page_positions.items():
            if page_number is not None:
                ignored.update(index for index in {positions[0], positions[-1]} if entries[index]["text"].isdigit())
        compact_name = re.sub(r"\s+", "", project_name)
        header_pattern = r"^(?:第[一二三四五六七八九十百\d]+章.{0,30}|目\s*录|招\s*标\s*文\s*件|招标公告|投标人须知(?:前附表)?|评标办法(?:前附表)?|条款号\s*条款名称\s*编列内容|评分因素\s*参考评分标准|评审项目\s*标准分\s*评分因素)$"
        for index, entry in enumerate(entries):
            line = entry["text"]
            compact = re.sub(r"\s+", "", line)
            if re.match(header_pattern, line) or (compact_name and len(compact) >= 8 and compact in compact_name):
                ignored.add(index)

        def title_fragment(line: str) -> bool:
            return len(line) <= 20 and bool(re.fullmatch(r"[\u4e00-\u9fffA-Za-z][\u4e00-\u9fffA-Za-z、及与的（）() /-]*", line))

        def closes_score_table(line: str) -> bool:
            return bool(re.match(r"^(?:第[一二三四五六七八九十百\d]+章|\d+[、．.]\s*[^\d\s])", line) or re.fullmatch(r"\d+(?:\.\d+)+(?:\s*[（(]\d+[)）])?", line))

        # A table header enables score recognition. Numeric prose outside a
        # scoring table must not become a score or consume unrelated headings.
        anchors = []
        score_table = False
        table_start = 0
        for index, entry in enumerate(entries):
            line = entry["text"]
            if re.search(r"评分因素|参考评分标准|评审项目.*标准分", line):
                score_table = True
                table_start = index + 1
                continue
            if closes_score_table(line):
                score_table = False
            if not score_table or index in ignored:
                continue
            inline = re.fullmatch(r"([\u4e00-\u9fffA-Za-z][\u4e00-\u9fffA-Za-z、及与的（）() /-]{1,40}?)\s+(\d{1,3}(?:\.\d+)?)\s*(?:分)?(?:\s+(.*))?", line)
            weight = re.fullmatch(r"(\d{1,3}(?:\.\d+)?)\s*(?:分)?(?:\s+(.*))?", line)
            if inline:
                start, title, points, description = index, inline.group(1), inline.group(2), inline.group(3) or ""
            elif weight:
                start = index
                fragments = []
                while start > table_start and len(fragments) < 6:
                    previous = start - 1
                    if previous in ignored or not title_fragment(entries[previous]["text"]):
                        break
                    fragments.insert(0, entries[previous]["text"])
                    start = previous
                if not fragments:
                    continue
                title, points, description = "".join(fragments), weight.group(1), weight.group(2) or ""
            else:
                continue
            if float(points) > 100:
                continue
            anchors.append({"start": start, "weight_index": index, "title": title, "points": points, "description": description})

        paragraphs: list[tuple[int, dict[str, Any]]] = []
        consumed: set[int] = set()
        for anchor_index, anchor in enumerate(anchors):
            end = anchors[anchor_index + 1]["start"] if anchor_index + 1 < len(anchors) else len(entries)
            for position in range(anchor["weight_index"] + 1, end):
                if closes_score_table(entries[position]["text"]) or re.search(r"评分因素|参考评分标准|评审项目.*标准分", entries[position]["text"]):
                    end = position
                    break
            description = anchor["description"] + "".join(entries[position]["text"] for position in range(anchor["weight_index"] + 1, end) if position not in ignored)
            text = f"{anchor['title']}（{anchor['points']}分）：{description}"
            paragraphs.append((anchor["start"], {"content": text, "source_page": entries[anchor["start"]]["page"], "scoring": True}))
            consumed.update(range(anchor["start"], end))

        parts: list[str] = []
        start = 0
        current_page = None

        def flush() -> None:
            if parts:
                paragraphs.append((start, {"content": "".join(parts), "source_page": current_page, "scoring": False}))
                parts.clear()

        for index, entry in enumerate(entries):
            line = entry["text"]
            if index in ignored or index in consumed:
                flush()
                continue
            numbered = bool(re.match(r"^(?:\d+(?:\.\d+)*(?:[、．.]|\s)|[（(][一二三四五六七八九十\d]+[)）]|[一二三四五六七八九十]+、)", line))
            if re.match(r"^\d+(?:\.\d+)?\s*(?:日历天|天|日|人|台|套|万元|元|平方米|米|%|m[²2]?)", line, re.IGNORECASE):
                numbered = False
            if parts and (entry["page"] != current_page or numbered or re.search(r"[。；！？]$", parts[-1]) or sum(map(len, parts)) + len(line) > 1200):
                flush()
            if not parts:
                start, current_page = index, entry["page"]
            parts.append(line)
        flush()
        # Exclusion lists inherit their parent's meaning even if the parent is
        # on the preceding PDF page and the child contains no keyword itself.
        # A new clause/list style or a broken sequence ends that inheritance.
        result: list[dict[str, Any]] = []
        exclusion_list = False
        list_style = None
        next_number = 1
        previous_position = -1

        def list_number(text: str) -> tuple[str, int] | None:
            match = re.match(r"^(?:[（(](\d+|[一二三四五六七八九十]+)[)）]|(\d+|[一二三四五六七八九十]+)[、．.](?!\d))\s*", text)
            if not match:
                return None
            value = match.group(1) or match.group(2)
            if value.isdigit():
                number = int(value)
            else:
                digits = {character: index for index, character in enumerate("一二三四五六七八九", 1)}
                if "十" in value:
                    tens, ones = value.split("十", 1)
                    number = digits.get(tens, 1) * 10 + digits.get(ones, 0)
                else:
                    number = digits.get(value, 0)
            return ("parenthesized" if match.group(1) else "plain", number)

        for position, paragraph in sorted(paragraphs, key=lambda item: item[0]):
            text = paragraph["content"]
            chapter_boundary = any(re.match(r"^第[一二三四五六七八九十百\d]+章", entries[index]["text"])
                                   for index in range(previous_position + 1, position))
            previous_position = position
            if chapter_boundary:
                exclusion_list = False
            if not paragraph.get("scoring") and re.search(r"(?:不得|禁止|严禁)[^。；！？]{0,12}(?:下列|以下)(?:情形|情况|行为)", text):
                exclusion_list, list_style, next_number = True, None, 1
            elif exclusion_list:
                number = list_number(text)
                if number and number[1] == next_number and (list_style is None or number[0] == list_style):
                    paragraph["inherited_veto"] = True
                    list_style, next_number = number[0], number[1] + 1
                elif (not number and not paragraph.get("scoring") and result and result[-1].get("inherited_veto")
                      and not re.match(r"^\d+(?:\.\d+)+\b", text) and not re.search(r"[。；！？]$", result[-1]["content"])):
                    # A child itself can wrap over a page boundary. Its first
                    # source page remains the provenance of the joined clause.
                    result[-1]["content"] += text
                    continue
                else:
                    exclusion_list = False
            result.append(paragraph)
        return result

    def build_outline(self, project_id: int) -> dict[str, Any]:
        from .requirements_workflow import RequirementsWorkflow

        return RequirementsWorkflow(self).build_outline(project_id, DEFAULT_OUTLINE, self._requirement_matches)

    def coverage_matrix(self, project_id: int) -> dict[str, Any]:
        from .requirements_workflow import RequirementsWorkflow

        project = RequirementsWorkflow(self).enrich(self.get_project(project_id))
        rows = []
        for requirement in project["requirements"]:
            matched = [
                {"section_id": section["id"], "title": section["title"]}
                for section in project["sections"]
                if requirement["id"] in section["requirement_ids"]
            ]
            rows.append({"requirement": requirement, "sections": matched, "covered": bool(matched),
                         "planning_technical": requirement["planning_category"] in {"technical", "unclassified"},
                         "formal_technical": requirement["formal_technical"]})
        workflow = project["requirements_workflow"]
        return {
            "project_id": project_id,
            "items": rows,
            "covered": sum(1 for row in rows if row["covered"]),
            "total": len(rows),
            "planning_total": workflow["planning_technical_total"],
            "planning_covered": workflow["mapped_technical"],
            "formal_technical_total": workflow["formal_technical_total"],
            "classification_pending": workflow["classification_pending"],
            "checklist_pending": workflow["checklist_pending"],
            "unmapped_high": workflow["unmapped_high"],
        }

    @staticmethod
    def _requirement_matches(title: str, content: str) -> bool:
        mappings = {
            "质量": ["质量", "验收", "竣工资料", "竣工文件", "移交资料", "技术档案"],
            "安全": ["安全", "文明", "环保"],
            "进度": ["进度", "工期"],
            "资源": ["资源", "人员", "机械", "材料"],
            "施工方案": ["施工", "工艺", "技术", "水性漆", "水性涂料", "涂装", "涂料"],
            "重点难点": ["重点", "难点", "风险"],
            "平面": ["平面", "临建", "场地"],
        }
        for key, keywords in mappings.items():
            if key in title and any(keyword in content for keyword in keywords):
                return True
        return False

    def generate_section(self, project_id: int, section_id: int, progress=None, cancelled=None,
                         *, repair_only: bool = False, target_hash: str | None = None) -> dict[str, Any]:
        progress = progress or (lambda *_args, **_kwargs: None)
        cancelled = cancelled or (lambda: False)
        generation_run_id = None
        batch_results: list[dict[str, Any]] = []

        class CancelledBeforeSave(Exception):
            pass

        def cancellation_result() -> dict[str, Any]:
            if generation_run_id is not None:
                with self.db.connect() as conn:
                    conn.execute("UPDATE generation_runs SET status='cancelled',completed_at=CURRENT_TIMESTAMP WHERE id=?", (generation_run_id,))
            progress("cancelled", 0, "章节生成已取消，未保存章节草稿", {"section_id": section_id, "completed_batches": len(batch_results)})
            return {"section_id": section_id, "cancelled": True, "generation_status": "cancelled",
                    "generation_run_id": generation_run_id, "generation_batches": batch_results}

        if cancelled():
            return cancellation_result()
        project = self.get_project(project_id)
        section = next((item for item in project["sections"] if item["id"] == section_id), None)
        if not section:
            raise KeyError("章节不存在")
        requirement_map = {item["id"]: item for item in project["requirements"]}
        requirements = [requirement_map[value] for value in section["requirement_ids"] if value in requirement_map]
        starting_review_fingerprints = [(item["id"], item.get("requirement_fingerprint")) for item in requirements]
        starting_draft = section.get("draft")
        starting_version = (starting_draft["id"], starting_draft["content_hash"]) if starting_draft else None
        frozen = {}
        starting_metadata_hash = None
        if repair_only:
            if not starting_draft or not target_hash or target_hash != starting_draft["content_hash"]:
                raise ValueError("当前章节版本已变化，请重新载入后修复")
            with self.db.connect() as conn:
                raw_draft = dict(conn.execute("SELECT * FROM project_drafts WHERE id=?", (starting_draft["id"],)).fetchone())
            if not public_generation(raw_draft)["can_repair"]:
                raise ValueError("当前章节没有可恢复的分段记录或正文已经修改，请重新生成此章")
            starting_metadata_hash = content_hash(raw_draft.get("generation_json") or "{}")
            frozen = read_generation(raw_draft)
        query = " ".join([section["title"], *[item["content"] for item in requirements[:5]]])
        sources = frozen.get("sources") if repair_only else self.knowledge.search(query, project.get("industry") or "", limit=8)
        if not repair_only and not sources and query != section["title"]:
            sources = self.knowledge.search(section["title"], project.get("industry") or "", limit=8)
        if not sources:
            raise ValueError("已发布知识库没有匹配内容，章节生成已阻断")
        scope_fingerprint = context_fingerprint(project, section, requirements, sources)
        if repair_only and frozen.get("context_fingerprint") != scope_fingerprint:
            raise ValueError("项目资料、条款或章节映射已变化，请重新生成此章")
        with self.db.connect() as conn:
            self._check_generation_sources(conn, sources)
        if cancelled():
            return cancellation_result()
        source_text = "\n\n".join(
            f"[知识{index}] {item['title']} / 发布v{item['publication_version']}\n{item['content'][:4000]}"
            for index, item in enumerate(sources, 1)
        )
        profile = factual_profile(project["profile"])
        retrieval_run_id = int(sources[0]["retrieval_run_id"]) if sources and sources[0].get("retrieval_run_id") else None
        generation_input_hash = content_hash(
            json.dumps(
                {
                    "project": {"id": project_id, "updated_at": project.get("updated_at"), "profile": profile},
                    "section": {"id": section_id, "title": section["title"]},
                    "requirements": requirements,
                    "knowledge": [item["content_hash"] for item in sources],
                },
                ensure_ascii=False,
                sort_keys=True,
            )
        )
        generation_run_id = self.evidence.create_generation_run(
            project_id,
            section_id,
            generation_input_hash,
            str(self.llm.settings().get("model") or ""),
            SECTION_DRAFT_PROMPT.version,
            "evidence-2.0.0",
            retrieval_run_id,
            [
                {
                    "unit_id": item["unit_id"],
                    "publication_id": item["publication_id"],
                    "publication_version": item["publication_version"],
                    "content_hash": item["content_hash"],
                    "index_version": item.get("index_version", ""),
                }
                for item in sources
            ],
        )
        source_text = source_text[:28000]
        batches = self._section_requirement_batches(requirements)
        if repair_only and [part["requirement_ids"] for part in frozen["parts"]] != [[item["id"] for item in batch] for batch in batches]:
            raise ValueError("条款分段规则已变化，请重新生成此章")
        saved_parts: list[dict[str, Any]] = []
        content_parts = [f"# {section['title']}"]
        confirmations: list[str] = []
        evidence: list[dict[str, Any]] = []
        for batch_index, batch_requirements in enumerate(batches, 1):
            if cancelled():
                return cancellation_result()
            content_budget = min(3000, max(1000, len(batch_requirements) * 240)) if batch_requirements else 2400
            batch = {"index": batch_index, "count": len(batches), "content_budget": content_budget}
            prior_part = frozen["parts"][batch_index - 1] if repair_only else None
            if prior_part and prior_part["status"] == "ai" and not prior_part.get("missing_requirement_ids"):
                preserved = {**prior_part, "reused": True}
                saved_parts.append(preserved)
                content_parts.append(preserved["content"])
                evidence.extend(preserved["evidence"])
                confirmations.extend(preserved["confirmations"])
                batch_results.append({key: value for key, value in preserved.items() if key not in {"content", "evidence", "confirmations", "part_hash"}})
                progress("generating", 10 + int(batch_index / len(batches) * 75), f"保留第{batch_index}/{len(batches)}部分原文", {"section_id": section_id, **batch, "reused": True})
                continue
            progress("generating", 10 + int((batch_index - 1) / len(batches) * 75),
                     f"生成章节第{batch_index}/{len(batches)}部分", {"section_id": section_id, **batch})
            if cancelled():
                return cancellation_result()
            prompt_requirements = [{key: item.get(key) for key in ("id", "requirement_key", "content", "kind", "priority", "source_page")} for item in batch_requirements]
            prompt = SECTION_DRAFT_PROMPT.render(
                project_name=project["name"], industry=project["industry"],
                profile=json.dumps(profile, ensure_ascii=False), section_title=section["title"],
                requirements=json.dumps(prompt_requirements, ensure_ascii=False), source_text=source_text,
                batch_index=str(batch_index), batch_count=str(len(batches)), content_budget=str(content_budget),
            )
            if prior_part:
                prompt += "\n本次只重写此部分。上次缺少可核验响应的要求ID：" + ",".join(map(str, prior_part.get("missing_requirement_ids") or prior_part["requirement_ids"])) + "。逐项提供正文原句作为evidence，不得把招标要求原文直接当作已完成的施工响应。"
            try:
                result = self.ai_runtime.execute(
                    SECTION_DRAFT_PROMPT, prompt,
                    {
                        "project": {"id": project_id, "name": project["name"], "industry": project["industry"], "profile": profile},
                        "section": {"id": section_id, "title": section["title"]},
                        "batch": batch, "requirements": prompt_requirements, "context_fingerprint": scope_fingerprint,
                        "sources": [
                            {"publication_id": item["publication_id"], "content_hash": item["content_hash"], "content": item["content"][:4000]}
                            for item in sources
                        ],
                    },
                    task_type="section_draft", target_type="project_section", target_id=section_id,
                    max_output_tokens=8000,
                    use_cache=not repair_only,
                    cancelled=cancelled,
                )
            except Exception as exc:  # noqa: BLE001
                result = {"error": f"{type(exc).__name__}: {exc}"}
            if cancelled():
                return cancellation_result()
            payload = result.get("payload") or {}
            body = normalize_text(str(payload.get("content") or ""))
            succeeded = bool(body)
            part_confirmations: list[str] = []
            if succeeded:
                part_confirmations.extend(
                    normalize_text(str(item)) for item in (payload.get("confirmations") or []) if normalize_text(str(item))
                )
                batch_evidence = payload.get("evidence") or []
            else:
                # Multi-part fallback lists only this part's requirements, so it
                # cannot duplicate an entire chapter for every failed request.
                body = self._fallback_section(project, section, batch_requirements, sources if len(batches) == 1 else [])
                if len(batches) == 1:
                    part_confirmations.append("大模型不可用，当前章节为基于已审核知识的结构化草稿，需人工复核")
                else:
                    part_confirmations.append(f"第{batch_index}/{len(batches)}部分AI编写失败，当前部分仅保留招标要求，须补充技术响应并人工复核")
                batch_evidence = [
                    {"requirement_id": item["id"], "text": f"{item['requirement_key']}：{item['content']}"}
                    for item in batch_requirements
                ]
            body = body.replace("[项目名称]", project["name"])
            body_lines = []
            for line in body.splitlines():
                if line.strip() == f"# {section['title']}":
                    continue
                if len(batches) > 1 and line == f"本章结合{project['name']}的招标要求和已审核知识进行编制。":
                    continue
                body_lines.append(re.sub(r"^#\s+", "## ", line))
            batch_body = normalize_text("\n".join(body_lines))
            content_parts.append(batch_body)
            batch_ids = {int(item["id"]) for item in batch_requirements}
            covered_ids: set[int] = set()
            accepted_evidence: list[dict[str, Any]] = []
            for item in batch_evidence:
                requirement_id = int(item.get("requirement_id") or 0)
                # Legacy ordinal output can only map within its own part; a real
                # ID belonging to another part must never be silently reassigned.
                if requirement_id not in requirement_map and 1 <= requirement_id <= len(batch_requirements):
                    requirement_id = int(batch_requirements[requirement_id - 1]["id"])
                text = normalize_text(str(item.get("text") or "")).replace("[项目名称]", project["name"])
                # Evidence must quote this part's final body. A quotation that
                # only occurs in another part must not establish coverage here.
                if requirement_id in batch_ids and text and text in batch_body and (not succeeded or not copies_requirement(text, requirement_map[requirement_id])):
                    accepted_evidence.append({"requirement_id": requirement_id, "text": text, "coverage_score": 1.0 if succeeded else 0.0})
                    covered_ids.add(requirement_id)
            missing_ids = sorted(batch_ids - covered_ids) if succeeded else []
            if missing_ids:
                part_confirmations.append(f"第{batch_index}/{len(batches)}部分有{len(missing_ids)}项要求缺少可核验的正文响应（要求ID：{','.join(map(str, missing_ids))}），须补充后复核")
            batch_results.append({
                **batch, "requirement_ids": [item["id"] for item in batch_requirements],
                "status": ("ai_incomplete" if missing_ids else "ai") if succeeded else "fallback", "ai_run_id": result.get("run_id"),
                "missing_requirement_ids": missing_ids,
                "model": result.get("model"), "cached": bool(result.get("cached")),
                "error": ("部分要求缺少可核验的正文响应" if missing_ids else "") if succeeded else str(result.get("error") or "模型未返回有效正文"),
            })
            part = {**batch_results[-1], "content": batch_body, "evidence": accepted_evidence,
                    "confirmations": part_confirmations, "reused": False}
            part["part_hash"] = part_fingerprint(part)
            saved_parts.append(part)
            evidence.extend(accepted_evidence)
            confirmations.extend(part_confirmations)
            progress("generating", 10 + int(batch_index / len(batches) * 75),
                     f"第{batch_index}/{len(batches)}部分已处理" + ("，需要补充技术响应并复核" if not succeeded or missing_ids else ""),
                     {"section_id": section_id, **batch, "status": batch_results[-1]["status"]})
        content = "\n\n".join(part for part in content_parts if part)
        confirmations = list(dict.fromkeys(confirmations))
        fallback_batches = sum(item["status"] == "fallback" for item in batch_results)
        incomplete_batches = sum(item["status"] == "ai_incomplete" for item in batch_results)
        generation_status = ("incomplete" if incomplete_batches else "ai") if not fallback_batches else "fallback" if fallback_batches == len(batches) else "partial_fallback"
        models = list(dict.fromkeys(str(item["model"]) for item in batch_results if item["status"] in {"ai", "ai_incomplete"} and item.get("model")))
        citations = [
            {
                "unit_id": item["unit_id"],
                "publication_id": item["publication_id"],
                "publication_version": item["publication_version"],
                "title": item["title"],
                "content_hash": item["content_hash"],
                "sources": item["sources"],
            }
            for item in sources
        ]
        draft_hash = content_hash(content)
        generation_metadata = {"schema_version": 1, "context_fingerprint": scope_fingerprint,
                               "draft_hash": draft_hash, "status": generation_status, "sources": sources,
                               "parts": saved_parts, "repaired_from_draft_id": starting_draft["id"] if repair_only else None}
        progress("saving", 95, "合并各部分并核验证据", {"section_id": section_id, "batch_count": len(batches), "fallback_batches": fallback_batches, "incomplete_batches": incomplete_batches})
        if cancelled():
            return cancellation_result()
        try:
            with self.db.connect() as conn:
                self._lock_project_for_update(conn, project_id)
                if cancelled():
                    raise CancelledBeforeSave()
                current = self._get_project_from_connection(conn, project_id)
                current_section = next((item for item in current["sections"] if item["id"] == section_id), None)
                current_draft = current_section.get("draft") if current_section else None
                current_version = (current_draft["id"], current_draft["content_hash"]) if current_draft else None
                current_requirement_map = {item["id"]: item for item in current["requirements"]}
                current_requirements = [current_requirement_map[rid] for rid in current_section["requirement_ids"] if rid in current_requirement_map] if current_section else []
                current_review_fingerprints = [(item["id"], item.get("requirement_fingerprint")) for item in current_requirements]
                current_metadata_hash = None
                if repair_only and current_draft:
                    current_metadata_hash = content_hash(conn.execute("SELECT generation_json FROM project_drafts WHERE id=?", (current_draft["id"],)).fetchone()[0] or "{}")
                if (not current_section or current_version != starting_version or current_review_fingerprints != starting_review_fingerprints
                    or (repair_only and current_metadata_hash != starting_metadata_hash)
                    or context_fingerprint(current, current_section, current_requirements, sources) != scope_fingerprint):
                    conn.execute("UPDATE generation_runs SET status='invalidated',completed_at=CURRENT_TIMESTAMP WHERE id=?", (generation_run_id,))
                    raise ValueError("生成期间项目资料、条款映射或章节正文发生变化，本次结果未覆盖现有章节，请重新载入")
                self._check_generation_sources(conn, sources, lock=True)
                version_no = int(conn.execute("SELECT COALESCE(MAX(version_no),0)+1 FROM project_drafts WHERE section_id=?", (section_id,)).fetchone()[0])
                cursor = conn.execute(
                    "INSERT INTO project_drafts(project_id,section_id,content,citations_json,confirmations_json,version_no,status,content_hash,generation_json) VALUES (?,?,?,?,?,?,'draft',?,?)",
                    (project_id, section_id, content, json.dumps(citations, ensure_ascii=False), json.dumps(confirmations, ensure_ascii=False), version_no, draft_hash, json.dumps(generation_metadata, ensure_ascii=False)),
                )
                draft_id = int(cursor.lastrowid)
                conn.execute("UPDATE project_sections SET status='drafted' WHERE id=?", (section_id,))
                conn.execute("UPDATE delivery_manifests SET invalidated_at=CURRENT_TIMESTAMP,status='invalidated' WHERE project_id=? AND invalidated_at IS NULL", (project_id,))
                for item in evidence:
                    requirement_id = int(item.get("requirement_id") or 0)
                    text = normalize_text(str(item.get("text") or ""))
                    coverage_score = float(item["coverage_score"])
                    conn.execute(
                        "INSERT INTO requirement_responses(requirement_id,draft_id,evidence_text,coverage_score,content_fingerprint) VALUES (?,?,?,?,?)",
                        (requirement_id, draft_id, text, coverage_score, draft_hash),
                    )
                if cancelled():
                    raise CancelledBeforeSave()
        except CancelledBeforeSave:
            return cancellation_result()
        except Exception:
            with self.db.connect() as conn:
                conn.execute("UPDATE generation_runs SET status='failed',completed_at=CURRENT_TIMESTAMP WHERE id=?", (generation_run_id,))
            raise
        evidence_sources = [
            *sources,
            *[
                {
                    "title": item["requirement_key"],
                    "content": item["content"],
                    "content_hash": content_hash(item["content"]),
                    "sources": [{"page_start": item.get("source_page")}],
                }
                for item in requirements
            ],
        ]
        evidence_report = self.evidence.analyze_draft(generation_run_id, draft_id, content, evidence_sources)
        with self.db.connect() as conn:
            self._lock_project_for_update(conn, project_id)
            saved = dict(conn.execute("SELECT * FROM project_drafts WHERE id=?", (draft_id,)).fetchone())
            confirmations = list(parse_json(saved["confirmations_json"], []))
            # A human can bind a missing response or resolve a claim after the
            # analysis transaction. Merge current claims into current notices;
            # never restore the generation task's older confirmation list.
            if content_hash(saved["content"]) == draft_hash and not evidence_report.get("stale"):
                current_claims = self.evidence.list_claims(draft_id, conn=conn)
                confirmations = [text for text in confirmations if not str(text).startswith("高风险表述缺少充分证据：")]
                confirmations.extend(f"高风险表述缺少充分证据：{claim['text'][:160]}" for claim in current_claims
                                     if claim["risk_level"] == "high" and claim["support_status"] != "supported")
                confirmations = list(dict.fromkeys(confirmations))
                conn.execute(
                    "UPDATE project_drafts SET confirmations_json=? WHERE id=?",
                    (json.dumps(confirmations, ensure_ascii=False), draft_id),
                )
        needs_review = fallback_batches + incomplete_batches
        progress("completed", 100, "章节生成完成" + (f"，含{needs_review}个待补充复核部分" if needs_review else ""),
                 {"section_id": section_id, "generation_status": generation_status, "fallback_batches": fallback_batches, "incomplete_batches": incomplete_batches})
        return {
            "draft_id": draft_id,
            "section_id": section_id,
            "version_no": version_no,
            "content": content,
            "citations": citations,
            "confirmations": confirmations,
            "model": ", ".join(models) if models else "structured-fallback",
            "ai_run_id": batch_results[0]["ai_run_id"] if len(batch_results) == 1 else None,
            "ai_run_ids": [item["ai_run_id"] for item in batch_results if item["ai_run_id"] is not None],
            "generation_status": generation_status, "generation_batches": batch_results,
            "ai_batches": len(batches) - fallback_batches, "fallback_batches": fallback_batches,
            "incomplete_batches": incomplete_batches,
            "cached": not fallback_batches and not incomplete_batches and all(item["cached"] for item in batch_results),
            "error": "；".join(f"第{item['index']}部分：{item['error']}" for item in batch_results if item["error"]),
            "generation_run_id": generation_run_id,
            "evidence": evidence_report,
            "generation": public_generation({"content": content, "generation": generation_metadata}),
            "reused_parts": sum(bool(item.get("reused")) for item in saved_parts),
            "claims": self.evidence.list_claims(draft_id),
        }

    def _check_generation_sources(self, conn, sources: list[dict], *, lock: bool = False) -> None:
        suffix = " FOR SHARE OF p,v" if lock and self.db.backend == "postgresql" else ""
        for source in sorted(sources, key=lambda item: int(item["publication_id"])):
            publication = conn.execute("SELECT p.*,v.content FROM knowledge_publications p JOIN knowledge_versions v ON v.id=p.version_id WHERE p.id=?" + suffix, (source["publication_id"],)).fetchone()
            if not publication or publication["status"] != "published" or publication["content_hash"] != source["content_hash"] or int(publication["publication_version"]) != int(source["publication_version"]) or publication["content"] != source["content"]:
                raise ValueError("章节所引用的知识版本已失效，请重新生成此章以选取有效来源")

    @staticmethod
    def _section_requirement_batches(requirements: list[dict[str, Any]], max_items: int = 12, max_chars: int = 6000) -> list[list[dict[str, Any]]]:
        batches: list[list[dict[str, Any]]] = []
        current: list[dict[str, Any]] = []
        current_chars = 0
        for requirement in requirements:
            size = len(str(requirement.get("content") or ""))
            if current and (len(current) >= max_items or current_chars + size > max_chars):
                batches.append(current)
                current, current_chars = [], 0
            # A single oversized source clause stays intact in its own part;
            # truncating it here would silently discard tender requirements.
            current.append(requirement)
            current_chars += size
        if current:
            batches.append(current)
        return batches or [[]]

    def generate_all(self, project_id: int, progress=None, cancelled=None) -> dict[str, Any]:
        progress = progress or (lambda *_args, **_kwargs: None)
        cancelled = cancelled or (lambda: False)
        project = self.get_project(project_id)
        if not project["sections"]:
            raise ValueError("请先生成目录")
        matrix = self.coverage_matrix(project_id)
        if matrix["unmapped_high"]:
            raise ValueError(f"高优先级要求尚未映射章节: {matrix['unmapped_high']}")
        results: list[dict[str, Any]] = []
        for index, section in enumerate(project["sections"]):
            if cancelled():
                break
            progress("generating", int(index / len(project["sections"]) * 95), f"生成章节 {index + 1}/{len(project['sections'])}", {"section_id": section["id"]})
            try:
                def section_progress(stage, value, message="", details=None):
                    progress("cancelled" if stage == "cancelled" else "generating",
                             int((index + value / 100) / len(project["sections"]) * 95),
                             f"第{index + 1}章：{message}", details or {})

                results.append(self.generate_section(project_id, int(section["id"]), progress=section_progress, cancelled=cancelled))
                if results[-1].get("cancelled"):
                    break
            except Exception as exc:  # noqa: BLE001
                results.append({"section_id": section["id"], "section_title": section["title"], "error": f"{type(exc).__name__}: {exc}"})
        output = {
            "project_id": project_id,
            "generated": sum(1 for item in results if item.get("draft_id")),
            "failed": sum(1 for item in results if not item.get("draft_id") and not item.get("cancelled")),
            "fallback_chapters": sum(1 for item in results if item.get("fallback_batches")),
            "incomplete_chapters": sum(1 for item in results if item.get("incomplete_batches")),
            "cancelled": bool(cancelled()) or any(item.get("cancelled") for item in results),
            "results": results,
        }
        progress("cancelled" if output["cancelled"] else "completed", 100,
                 "批量章节生成已取消" if output["cancelled"] else f"批量章节生成完成，{sum(bool(item.get('fallback_batches') or item.get('incomplete_batches')) for item in results)}章含待补充复核内容",
                 {"generated": output["generated"], "failed": output["failed"], "fallback_chapters": output["fallback_chapters"], "incomplete_chapters": output["incomplete_chapters"]})
        return output

    @staticmethod
    def _fallback_section(project: dict[str, Any], section: dict[str, Any], requirements: list[dict[str, Any]], sources: list[dict[str, Any]]) -> str:
        parts = [f"# {section['title']}", f"本章结合{project['name']}的招标要求和已审核知识进行编制。"]
        if requirements:
            parts.append("## 条款响应")
            parts.extend(f"- {item['requirement_key']}：{item['content']}" for item in requirements)
        for source in sources[:5]:
            parts.append(f"## {source['title']}")
            parts.append(source["content"])
        return "\n\n".join(parts)

    def update_draft(self, draft_id: int, content: str) -> dict[str, Any]:
        normalized = normalize_text(content)
        updated_hash = content_hash(normalized)
        with self.db.connect() as conn:
            if self.db.backend == "sqlite":
                conn.execute("BEGIN IMMEDIATE")
            owner = conn.execute("SELECT project_id FROM project_drafts WHERE id=?", (draft_id,)).fetchone()
            if not owner:
                raise KeyError("章节草稿不存在")
            self._lock_project_for_update(conn, int(owner["project_id"]))
            row = conn.execute("SELECT * FROM project_drafts WHERE id=?", (draft_id,)).fetchone()
            if not row:
                raise KeyError("章节草稿不存在")
            if conn.execute("SELECT 1 FROM project_drafts WHERE section_id=? AND version_no>? LIMIT 1", (row["section_id"], row["version_no"])).fetchone():
                raise ValueError("只能修改当前最新章节版本，请重新载入")
            current_hash = content_hash(row["content"])
            if current_hash == updated_hash:
                return {"draft_id": draft_id, "status": row["status"], "changed": False, "quality_confirmations_invalidated": False}
            generation_metadata = read_generation(dict(row))
            generation_metadata["manual_content_hash"] = updated_hash
            conn.execute(
                "UPDATE project_drafts SET content=?,content_hash=?,generation_json=?,status='draft',evidence_status='pending',updated_at=CURRENT_TIMESTAMP WHERE id=?",
                (normalized, updated_hash, json.dumps(generation_metadata, ensure_ascii=False), draft_id),
            )
            conn.execute("UPDATE project_sections SET status='drafted' WHERE id=?", (row["section_id"],))
            conn.execute("UPDATE requirement_responses SET review_status='invalidated' WHERE draft_id=?", (draft_id,))
            conn.execute("UPDATE claims SET support_status='invalidated',updated_at=CURRENT_TIMESTAMP WHERE draft_id=?", (draft_id,))
            conn.execute("UPDATE generation_runs SET status='invalidated' WHERE draft_id=?", (draft_id,))
            conn.execute("UPDATE delivery_manifests SET status='invalidated',invalidated_at=CURRENT_TIMESTAMP WHERE project_id=? AND invalidated_at IS NULL", (row["project_id"],))
            responses = conn.execute("SELECT id,evidence_text,coverage_score FROM requirement_responses WHERE draft_id=?", (draft_id,)).fetchall()
            for response in responses:
                evidence_text = normalize_text(response["evidence_text"] or "")
                conn.execute("UPDATE requirement_responses SET content_fingerprint=?,coverage_score=? WHERE id=?",
                             (updated_hash, float(response["coverage_score"]) if response["coverage_score"] >= .8 and evidence_text and evidence_text in normalized else 0.0, response["id"]))
        # A saved edit must be evaluated against the frozen sources again;
        # invalidated claims must never disappear from the high-risk gate.
        citations = parse_json(row["citations_json"], [])
        evidence_sources = []
        for citation in citations:
            source = self.db.row(
                "SELECT p.unit_id,p.id AS publication_id,p.content_hash,v.content "
                "FROM knowledge_publications p JOIN knowledge_versions v ON v.id=p.version_id "
                "WHERE p.id=? AND p.status='published' AND p.content_hash=?",
                (citation.get("publication_id"), citation.get("content_hash")),
            )
            if source:
                evidence_sources.append({**source, "sources": citation.get("sources") or []})
        requirements = self.db.rows(
            "SELECT DISTINCT r.content,r.source_page FROM project_requirements r "
            "JOIN requirement_responses rr ON rr.requirement_id=r.id WHERE rr.draft_id=?", (draft_id,),
        )
        evidence_sources.extend({"content": item["content"], "content_hash": content_hash(item["content"]),
                                 "sources": [{"page_start": item.get("source_page")}]} for item in requirements)
        run_id = self.evidence.create_generation_run(int(row["project_id"]), int(row["section_id"]), updated_hash,
                                                      "manual-edit", "manual-edit-1", "evidence-2.0.0", None, citations)
        analysis = self.evidence.analyze_draft(run_id, draft_id, normalized, evidence_sources)
        return {"draft_id": draft_id, "status": "draft", "changed": True, "quality_confirmations_invalidated": True,
                "evidence_status": analysis["evidence_status"], "analysis_stale": bool(analysis.get("stale"))}

    def confirm_draft(
        self,
        draft_id: int,
        reviewer: str,
        resolutions: list[dict[str, Any]] | None = None,
        notes: str = "",
        target_hash: str | None = None,
    ) -> dict[str, Any]:
        if not reviewer.strip():
            raise ValueError("确认人不能为空")
        with self.db.connect() as conn:
            if self.db.backend == "sqlite":
                conn.execute("BEGIN IMMEDIATE")
            owner = conn.execute("SELECT project_id FROM project_drafts WHERE id=?", (draft_id,)).fetchone()
            if not owner:
                raise KeyError("章节草稿不存在")
            self._lock_project_for_update(conn, int(owner["project_id"]))
            lock_clause = " FOR UPDATE" if self.db.backend == "postgresql" else ""
            row = conn.execute(f"SELECT * FROM project_drafts WHERE id=?{lock_clause}", (draft_id,)).fetchone()
            if not row:
                raise KeyError("章节草稿不存在")
            if conn.execute("SELECT 1 FROM project_drafts WHERE section_id=? AND version_no>? LIMIT 1", (row["section_id"], row["version_no"])).fetchone():
                raise ValueError("只能签审当前最新章节版本，请重新载入")
            current_hash = content_hash(row["content"])
            if target_hash is not None and target_hash != current_hash:
                raise ValueError("章节内容已变化，请重新查看后签审")
            confirmations = parse_json(row["confirmations_json"], [])
            source_status = self.evidence.project_source_status(int(owner["project_id"]), conn=conn)
            if draft_id in source_status["stale_draft_ids"] or row["evidence_status"] in {"pending", "invalidated", "failed"}:
                raise ValueError("当前章节证据尚未绑定最新项目资料，请先重新核验证据")
            if source_status["profile_conflicts"]:
                raise ValueError("项目资料与招标原文存在冲突，请先核对项目资料")
            if generation_incomplete({"generation": public_generation(dict(row)), "confirmations": confirmations}):
                raise ValueError("当前正文含未完成的生成部分，请先局部修复或编辑正文并补录响应，不能用处理结论代替正文")
            project = self._get_project_from_connection(conn, int(owner["project_id"]))
            section = next(section for section in project["sections"] if section["id"] == row["section_id"])
            missing_responses = section["draft"]["missing_response_requirement_ids"]
            if missing_responses:
                raise ValueError(f"本章还有{len(missing_responses)}条未绑定实际正文响应，请补写或定位现有正文后再签审")
            unsupported_high = int(
                conn.execute(
                    "SELECT COUNT(*) FROM claims WHERE draft_id=? AND risk_level='high' AND support_status NOT IN ('supported','confirmed')",
                    (draft_id,),
                ).fetchone()[0]
            )
            if unsupported_high:
                raise ValueError(f"仍有{unsupported_high}条高风险Claim缺少证据，不能签审")
            target_hash = current_hash
            supplied: dict[int, str] = {}
            required_indices = human_confirmation_indices({"confirmations": confirmations})
            for item in resolutions or []:
                index = int(item.get("index", -1))
                resolution = normalize_text(str(item.get("resolution") or ""))
                if index < 0 or index >= len(confirmations):
                    raise ValueError("待确认项序号无效")
                if not resolution:
                    raise ValueError(f"第{index + 1}项待确认事项缺少处理结论")
                if index in supplied:
                    raise ValueError(f"第{index + 1}项待确认事项重复提交")
                supplied[index] = resolution
            missing = [index + 1 for index in required_indices if index not in supplied]
            if missing:
                raise ValueError(f"待确认事项必须逐项处理：缺少第{','.join(map(str, missing))}项")
            conn.execute("UPDATE project_drafts SET content_hash=? WHERE id=?", (target_hash, draft_id))
            for index, confirmation in enumerate(confirmations):
                if index not in required_indices:
                    continue
                conn.execute(
                    """
                    INSERT INTO confirmation_resolutions(
                        draft_id,confirmation_index,confirmation_text,resolution,resolver,target_hash
                    ) VALUES (?,?,?,?,?,?)
                    ON CONFLICT(draft_id,confirmation_index,target_hash) DO UPDATE SET
                        confirmation_text=excluded.confirmation_text,
                        resolution=excluded.resolution,
                        resolver=excluded.resolver,
                        created_at=CURRENT_TIMESTAMP
                    """,
                    (draft_id, index, str(confirmation), supplied[index], reviewer.strip(), target_hash),
                )
            conn.execute(
                "INSERT INTO draft_review_decisions(draft_id,reviewer,target_hash,decision,notes) VALUES (?,?,?,'approved',?)",
                (draft_id, reviewer.strip(), target_hash, normalize_text(notes)),
            )
            conn.execute(
                "UPDATE project_drafts SET status='reviewed',updated_at=CURRENT_TIMESTAMP WHERE id=?",
                (draft_id,),
            )
            conn.execute("UPDATE project_sections SET status='reviewed' WHERE id=?", (row["section_id"],))
            conn.execute("UPDATE requirement_responses SET review_status='invalidated' WHERE draft_id=?", (draft_id,))
            conn.execute(
                """
                UPDATE requirement_responses SET review_status='confirmed'
                WHERE draft_id=? AND content_fingerprint=? AND coverage_score>=0.8
                    AND TRIM(COALESCE(evidence_text,''))<>''
                """,
                (draft_id, target_hash),
            )
        return {
            "draft_id": draft_id,
            "status": "reviewed",
            "reviewer": reviewer.strip(),
            "target_hash": target_hash,
            "resolved_confirmations": len(required_indices),
        }

    def quality_gate(self, project_id: int, *, conn=None, sync_issues: bool = True) -> dict[str, Any]:
        project = self._get_project_from_connection(conn, project_id) if conn is not None else self.get_project(project_id)
        sections = project["sections"]
        drafts = [item["draft"] for item in sections if item.get("draft")]
        text = "\n".join(item["content"] for item in drafts)
        review_blockers: list[dict[str, str]] = []
        formal_only_blockers: list[dict[str, str]] = []
        if not sections or len(drafts) < len(sections) or any(not draft["content"].strip() for draft in drafts):
            review_blockers.append({"key": "missing_sections", "title": "章节未全部生成", "detail": f"{len(drafts)}/{len(sections)}章已有草稿"})
        confirmation_count = 0
        unreviewed_drafts = 0
        with (nullcontext(conn) if conn is not None else self.db.connect()) as connection:
            for draft in drafts:
                target_hash = draft["content_hash"]
                resolved = {int(row["confirmation_index"]) for row in draft["confirmation_resolutions"]}
                confirmation_count += sum(1 for index in human_confirmation_indices(draft) if index not in resolved)
                reviewed = connection.execute(
                    "SELECT 1 FROM draft_review_decisions WHERE draft_id=? AND target_hash=? AND decision='approved' LIMIT 1",
                    (draft["id"], target_hash),
                ).fetchone()
                if not reviewed:
                    unreviewed_drafts += 1
        if confirmation_count:
            review_blockers.append({"key": "confirmations", "title": "存在待确认事项", "detail": f"共{confirmation_count}项"})
        if unreviewed_drafts:
            formal_only_blockers.append({"key": "unreviewed_drafts", "title": "章节尚未绑定当前版本签审", "detail": f"共{unreviewed_drafts}章"})
        invalid_citations = 0
        with (nullcontext(conn) if conn is not None else self.db.connect()) as connection:
            for draft in drafts:
                for citation in draft.get("citations") or []:
                    valid = connection.execute(
                        "SELECT 1 FROM knowledge_publications WHERE id=? AND status='published' AND content_hash=?",
                        (citation.get("publication_id"), citation.get("content_hash")),
                    ).fetchone()
                    if not valid:
                        invalid_citations += 1
        if invalid_citations:
            review_blockers.append({"key": "invalid_citations", "title": "引用已失效", "detail": f"共{invalid_citations}项"})
        project_evidence = project["evidence_source"]
        if project_evidence["stale_draft_ids"]:
            review_blockers.append({"key": "project_evidence_stale", "title": "项目来源证据需要重新核验", "detail": f"共{len(project_evidence['stale_draft_ids'])}章尚未绑定当前原文或项目资料"})
        if project_evidence["profile_conflicts"]:
            review_blockers.append({"key": "project_profile_conflicts", "title": "项目资料与原文存在冲突", "detail": f"共{len(project_evidence['profile_conflicts'])}处，请核对面积、工期等资料"})
        incomplete_generation = sum(generation_incomplete(draft) for draft in drafts)
        if incomplete_generation:
            review_blockers.append({"key": "incomplete_generation", "title": "正文仍有未完成的生成部分", "detail": f"共{incomplete_generation}章，请先局部修复或补充实际正文响应"})
        stale_evidence = sum(
            1 for draft in drafts
            if draft.get("evidence_status") in {"pending", "invalidated", "failed"}
            or any(claim.get("support_status") == "invalidated" for claim in draft.get("claims", []))
        )
        if stale_evidence:
            review_blockers.append({"key": "stale_evidence", "title": "章节证据检查未完成", "detail": f"共{stale_evidence}章证据检查未完成或已失效"})
        evidence_metrics = self.evidence.metrics(project_id, conn=conn)
        if evidence_metrics["high_unsupported"]:
            review_blockers.append({"key": "unsupported_high_claims", "title": "高风险表述缺少证据", "detail": f"共{evidence_metrics['high_unsupported']}项"})
        artifact_audit = audit_text(self._project_markdown(project), source="assembled_markdown")
        if not artifact_audit["ready"]:
            detail = "、".join(f"{key}={value}" for key, value in artifact_audit["counts"].items() if value)
            review_blockers.append({"key": "artifact_audit", "title": "最终产物审计未通过", "detail": detail or "发现禁止内容"})
        workflow_metrics = response_metrics(self, project, conn=conn)
        coverage_total = workflow_metrics["technical_total"]
        covered = workflow_metrics["signed"]
        review_blockers.extend(project["requirements_workflow"]["blockers"])
        coverage = round(covered / coverage_total * 100, 2) if coverage_total else 100.0
        if coverage < 95:
            review_blockers.append({"key": "coverage", "title": "技术条款签审覆盖不足", "detail": f"当前{coverage}%，{covered}/{coverage_total}条已签审；未确认的非技术分类仍计入分母"})
        consistency_issues = self._consistency_issues(project)
        warnings = [item for item in consistency_issues if item["key"] == "duplicate_sections"]
        review_blockers.extend(item for item in consistency_issues if item["key"] != "duplicate_sections")
        profile = project.get("profile") or {}
        bidder_name = normalize_text(str(profile.get("bidder_name") or profile.get("tenderer_name") or profile.get("bidder") or ""))
        if not bidder_name or audit_text(bidder_name, source="bidder_name")["counts"]["test_data"]:
            formal_only_blockers.append({"key": "bidder_identity", "title": "缺少真实投标单位", "detail": "正式版必须登记真实投标单位"})
        reviewer = normalize_text(str(profile.get("professional_reviewer") or profile.get("reviewed_by") or ""))
        if not reviewer:
            formal_only_blockers.append({"key": "professional_reviewer", "title": "缺少专业复核人", "detail": "正式版必须由专业复核人署名"})
        confirmation = self._delivery_confirmation(project)
        if not confirmation["compliance_confirmed"]:
            formal_only_blockers.append({"key": "compliance_confirmation", "title": "合规确认未完成", "detail": "正式版必须针对当前内容完成合规清单确认；修改后须重新确认"})
        if not confirmation["manual_finalized"]:
            formal_only_blockers.append({"key": "manual_finalization", "title": "人工定稿未完成", "detail": "正式版必须由责任人对当前内容完成整本定稿；修改后须重新确认"})
        formal_blockers = [*review_blockers, *formal_only_blockers]
        review_score = max(0, 100 - len(review_blockers) * 12)
        score = max(0, 100 - len(formal_blockers) * 12)
        review_ready = not review_blockers and review_score >= 90
        formal_ready = review_ready and not formal_only_blockers and score >= 90
        report = {
            "project_id": project_id,
            "ready": formal_ready,
            "review_ready": review_ready,
            "formal_ready": formal_ready,
            "readiness_level": "formal" if formal_ready else "review" if review_ready else "blocked",
            "score": score,
            "review_score": review_score,
            "blockers": formal_blockers,
            "review_blockers": review_blockers,
            "formal_blockers": formal_only_blockers,
            "warnings": warnings,
            "artifact_audit": artifact_audit,
            "metrics": {
                "sections": len(sections),
                "drafts": len(drafts),
                "requirements": coverage_total,
                "covered_requirements": covered,
                "coverage": coverage,
                "mapped_requirements": workflow_metrics["mapped"],
                "responded_requirements": workflow_metrics["responded"],
                "total_requirements": workflow_metrics["total"],
                "technical_requirements": workflow_metrics["technical_total"],
                "classification_pending": project["requirements_workflow"]["classification_pending"],
                "confirmations": confirmation_count,
                "unreviewed_drafts": unreviewed_drafts,
                "invalid_citations": invalid_citations,
                "claims": evidence_metrics,
                "artifact_findings": len(artifact_audit["findings"]),
            },
        }
        if sync_issues:
            self._sync_quality_issues(project_id, [*formal_blockers, *warnings])
        return report

    @staticmethod
    def _consistency_issues(project: dict[str, Any]) -> list[dict[str, str]]:
        issues: list[dict[str, str]] = []
        drafts = [section["draft"] for section in project["sections"] if section.get("draft")]
        combined = re.sub(r"[*_`]", "", "\n".join(draft["content"] for draft in drafts))
        label_patterns = {
            "工期": r"工期.{0,12}?(\d+\s*(?:天|日历天))",
            "质量目标": r"质量目标[ \t]*(?:[:：]|为|是)[ \t]*([^，。；\n]{2,40})",
            "安全目标": r"安全目标[ \t]*(?:[:：]|为|是)[ \t]*([^，。；\n]{2,40})",
        }
        for label, pattern in label_patterns.items():
            values = {normalize_text(value) for value in re.findall(pattern, combined, flags=re.IGNORECASE)}
            if label == "工期":
                values = {re.sub(r"\s+", "", value) for value in values}
            if len(values) > 1:
                issues.append({"key": f"conflict_{label}", "title": f"{label}存在冲突", "detail": "、".join(sorted(values))[:300]})
        normalized = [set(re.findall(r"[\u4e00-\u9fff]{2,4}", draft["content"])) for draft in drafts]
        for left in range(len(normalized)):
            for right in range(left + 1, len(normalized)):
                union = normalized[left] | normalized[right]
                similarity = len(normalized[left] & normalized[right]) / len(union) if union else 0
                if similarity >= 0.92:
                    issues.append({"key": "duplicate_sections", "title": "章节内容高度重复", "detail": f"第{left + 1}章与第{right + 1}章相似度{similarity:.0%}"})
        return issues

    def _sync_quality_issues(self, project_id: int, blockers: list[dict[str, str]]) -> None:
        active_fingerprints: set[str] = set()
        with self.db.connect() as conn:
            for blocker in blockers:
                fingerprint = content_hash(f"{project_id}|{blocker['key']}|{blocker['detail']}")
                active_fingerprints.add(fingerprint)
                exists = conn.execute(
                    "SELECT id FROM quality_issues WHERE project_id=? AND fingerprint=? AND status='open'",
                    (project_id, fingerprint),
                ).fetchone()
                if not exists:
                    severity = "high" if blocker["key"] in {"unsupported_high_claims", "invalid_citations", "missing_sections", "coverage"} or blocker["key"].startswith("conflict_") else "medium"
                    conn.execute(
                        """
                        INSERT INTO quality_issues(project_id,issue_code,severity,source,message,fingerprint)
                        VALUES (?,?,?,?,?,?)
                        """,
                        (project_id, blocker["key"], severity, "quality_gate", f"{blocker['title']}：{blocker['detail']}", fingerprint),
                    )
            open_rows = conn.execute("SELECT id,fingerprint FROM quality_issues WHERE project_id=? AND source='quality_gate' AND status='open'", (project_id,)).fetchall()
            for row in open_rows:
                if row["fingerprint"] not in active_fingerprints:
                    conn.execute(
                        "UPDATE quality_issues SET status='resolved',resolution='后续质量检查已不再复现',resolved_at=CURRENT_TIMESTAMP WHERE id=?",
                        (row["id"],),
                    )

    def resolve_quality_issue(self, issue_id: int, resolution: str, user_id: int | None = None) -> dict[str, Any]:
        issue = self.db.row("SELECT * FROM quality_issues WHERE id=?", (issue_id,))
        if not issue:
            raise KeyError("质量问题不存在")
        if not normalize_text(resolution):
            raise ValueError("处理结论不能为空")
        with self.db.connect() as conn:
            conn.execute(
                """
                UPDATE quality_issues SET status='resolved',resolution=?,resolved_by=?,resolved_at=CURRENT_TIMESTAMP
                WHERE id=?
                """,
                (normalize_text(resolution), user_id, issue_id),
            )
        self.audit.record("quality.resolve", "quality_issue", issue_id, actor_user_id=user_id, details={"project_id": issue["project_id"]})
        return self.db.row("SELECT * FROM quality_issues WHERE id=?", (issue_id,)) or {}

    def list_manifests(self, project_id: int) -> list[dict[str, Any]]:
        rows = self.db.rows("SELECT * FROM delivery_manifests WHERE project_id=? ORDER BY id DESC", (project_id,))
        for row in rows:
            row["manifest"] = parse_json(row.pop("manifest_json", "{}"), {})
        return rows

    def list_deliveries(self, project_id: int) -> list[dict[str, Any]]:
        project = self.get_project(project_id)
        project_hash = self._project_hash(project)
        confirmation = self._delivery_confirmation(project)
        readiness = self.quality_gate(project_id, sync_issues=False)
        rows = self.db.rows(
            "SELECT d.*,m.id AS manifest_id,m.manifest_hash,m.project_hash,m.status AS manifest_status,m.manifest_json "
            "FROM deliveries d LEFT JOIN delivery_manifests m ON m.delivery_id=d.id "
            "WHERE d.project_id=? ORDER BY d.id DESC", (project_id,),
        )
        items = []
        for row in rows:
            metadata = self._delivery_metadata(row)
            mode, separator, file_format = row["format"].partition("_")
            frozen_confirmation = parse_json(row.get("manifest_json"), {}).get("delivery_confirmation")
            items.append({
                "delivery_id": row["id"], "project_id": project_id,
                "format": file_format if separator else row["format"], "mode": mode if separator else "unknown",
                "created_at": row["created_at"], "file_name": metadata.get("file_name") or Path(row["file_path"]).name,
                "manifest_id": row.get("manifest_id"), "manifest_hash": row.get("manifest_hash"),
                "status": row.get("manifest_status") or "unverified",
                "current": (row.get("manifest_status") == "frozen" and row.get("project_hash") == project_hash
                            and bool(readiness.get(f"{mode}_ready")) and frozen_confirmation == confirmation),
                "available": bool(metadata.get("sha256") and (metadata.get("object_key") or self._safe_delivery_path(row["file_path"]))),
                "download_url": f"/api/projects/{project_id}/deliveries/{row['id']}/download",
            })
        return items

    def _safe_delivery_path(self, value: str) -> Path | None:
        path = Path(value).resolve()
        if not path.is_relative_to(self.settings.export_root.resolve()) or not path.is_file():
            return None
        return path

    def _delivery_metadata(self, row: dict[str, Any]) -> dict[str, Any]:
        quality = parse_json(row.get("quality_json"), {})
        metadata = quality.get("_delivery_file") or {}
        if metadata:
            return metadata
        # Earlier releases stored an editable export path. Only offer those
        # artifacts when the manifest supplies an independently frozen hash.
        manifest = parse_json(row.get("manifest_json"), {})
        file_name = Path(row["file_path"]).name
        if row["format"].endswith("_package") and not file_name.endswith(".zip"):
            return {}
        item = next((item for item in manifest.get("files", []) if item.get("name") == file_name), None)
        if not item:
            return {}
        object_key = f"deliveries/{row['project_id']}/{str(row.get('manifest_hash') or '')[:12]}/{file_name}"
        record = self.db.row("SELECT sha256 FROM object_records WHERE object_key=?", (object_key,))
        return {"file_name": file_name, "sha256": item.get("sha256"),
                "object_key": object_key if record and record["sha256"] == item.get("sha256") else ""}

    def download_delivery(self, project_id: int, delivery_id: int) -> dict[str, Any]:
        row = self.db.row(
            "SELECT d.*,m.manifest_hash,m.manifest_json FROM deliveries d "
            "LEFT JOIN delivery_manifests m ON m.delivery_id=d.id WHERE d.project_id=? AND d.id=?",
            (project_id, delivery_id),
        )
        if not row:
            raise KeyError("交付记录不存在")
        metadata = self._delivery_metadata(row)
        expected_hash = metadata.get("sha256")
        if not expected_hash:
            raise ValueError("历史交付文件缺少可核验的冻结记录，请重新导出")
        data = None
        object_key = str(metadata.get("object_key") or "")
        if object_key:
            if not object_key.startswith(f"deliveries/{project_id}/") or any(part in {"", ".", ".."} for part in object_key.replace("\\", "/").split("/")):
                raise ValueError("交付文件存储位置无效")
            record = self.db.row("SELECT sha256 FROM object_records WHERE object_key=?", (object_key,))
            if record and record["sha256"] == expected_hash:
                try:
                    with self.storage.open(object_key) as handle:
                        data = handle.read()
                except FileNotFoundError:
                    pass
        if data is None:
            path = self._safe_delivery_path(row["file_path"])
            if not path:
                raise KeyError("交付文件不存在")
            data = path.read_bytes()
        if hashlib.sha256(data).hexdigest() != expected_hash:
            raise ValueError("交付文件与冻结记录不一致，请重新导出")
        file_name = Path(str(metadata.get("file_name") or row["file_path"])).name
        return {"data": data, "file_name": file_name,
                "media_type": mimetypes.guess_type(file_name)[0] or "application/octet-stream"}

    def export_project(self, project_id: int, file_format: str, mode: str = "formal", progress=None) -> dict[str, Any]:
        progress = progress or (lambda *_args, **_kwargs: None)
        if mode not in {"review", "formal"}:
            raise ValueError("导出mode仅支持review或formal")
        project = self.get_project(project_id)
        quality = self.quality_gate(project_id)
        if not quality[f"{mode}_ready"]:
            label = "送审" if mode == "review" else "正式"
            raise ValueError(f"质量门禁未通过，不能导出{label}版")
        safe_name = re.sub(r"[\\/:*?\"<>|]", "_", project["name"])
        export_id = uuid4().hex
        export_directory = self.settings.export_root / str(project_id) / export_id
        export_directory.mkdir(parents=True, exist_ok=False)
        progress("rendering", 20, "生成交付文件")
        files: list[Path] = []
        if file_format == "markdown":
            path = export_directory / f"{safe_name}_{project_id}_{mode}.md"
            content = self._project_markdown(project)
            write_text_atomic(path, content)
            files.append(path)
        elif file_format == "docx":
            path = export_directory / f"{safe_name}_{project_id}_{mode}.docx"
            self._write_docx(project, path, mode=mode)
            self._preflight_docx(path)
            files.append(path)
        elif file_format in {"pdf", "package"}:
            docx_path = export_directory / f"{safe_name}_{project_id}_{mode}.docx"
            self._write_docx(project, docx_path, mode=mode)
            self._preflight_docx(docx_path)
            pdf_path = self._convert_pdf(docx_path)
            self._preflight_docx(docx_path)
            self._preflight_pdf(pdf_path)
            files.extend([docx_path, pdf_path])
            path = pdf_path
        else:
            raise ValueError("仅支持markdown、docx、pdf或package")
        project_hash = self._project_hash(project)
        file_items = [
            {"name": item.name, "size_bytes": item.stat().st_size, "sha256": hashlib.sha256(item.read_bytes()).hexdigest()}
            for item in files
        ]
        manifest = {
            "schema_version": "1.0",
            "export_id": export_id,
            "mode": mode,
            "format": file_format,
            "project_id": project_id,
            "project_name": project["name"],
            "project_hash": project_hash,
            "delivery_confirmation": self._delivery_confirmation(project),
            "bidder_name": str((project.get("profile") or {}).get("bidder_name") or (project.get("profile") or {}).get("tenderer_name") or (project.get("profile") or {}).get("bidder") or ""),
            "quality": quality,
            "files": file_items,
            "drafts": [
                {"draft_id": section["draft"]["id"], "content_hash": section["draft"]["content_hash"], "evidence_status": section["draft"].get("evidence_status")}
                for section in project["sections"] if section.get("draft")
            ],
            "knowledge": sorted(
                {
                    (citation["publication_id"], citation["publication_version"], citation["content_hash"])
                    for section in project["sections"] if section.get("draft")
                    for citation in section["draft"].get("citations") or []
                }
            ),
        }
        manifest_json = json.dumps(manifest, ensure_ascii=False, indent=2)
        manifest_hash = content_hash(manifest_json)
        manifest_path = export_directory / f"{safe_name}_{project_id}_{mode}_manifest.json"
        write_text_atomic(manifest_path, manifest_json)
        files.append(manifest_path)
        if file_format == "package":
            package_path = export_directory / f"{safe_name}_{project_id}_{'review' if mode == 'review' else 'delivery'}.zip"
            with zipfile.ZipFile(package_path, "w", zipfile.ZIP_DEFLATED) as archive:
                if mode == "review":
                    deliverables = [item for item in files if item.suffix.lower() in {".docx", ".pdf"}]
                    for item in deliverables:
                        archive.write(item, item.name)
                    note = "# 送审说明\n\n本包用于专业送审。正式投标前须补齐真实投标单位、专业复核人、合规确认、签章及人工定稿。\n"
                    names = [item.name for item in deliverables]
                    archive.writestr("送审说明.md", note)
                    archive.writestr("文件清单.txt", "\n".join([*names, "送审说明.md", "文件清单.txt"]) + "\n")
                else:
                    for item in files:
                        archive.write(item, item.name)
                    archive.writestr("质量门禁报告.json", json.dumps(quality, ensure_ascii=False, indent=2))
                    archive.writestr("要求覆盖矩阵.json", json.dumps(self.coverage_matrix(project_id), ensure_ascii=False, indent=2))
                    archive.writestr("证据审计.json", json.dumps({section["title"]: section["draft"].get("claims", []) for section in project["sections"] if section.get("draft")}, ensure_ascii=False, indent=2))
            if mode == "review":
                package_audit = audit_zip_path(
                    package_path,
                    allowed_names={*[item.name for item in deliverables], "送审说明.md", "文件清单.txt"},
                    require_review_package_types=True,
                )
                if not package_audit["ready"]:
                    raise ValueError("送审包产物审计失败")
            path = package_path
            file_items.append({"name": package_path.name, "size_bytes": package_path.stat().st_size, "sha256": hashlib.sha256(package_path.read_bytes()).hexdigest()})
        if self._project_hash(self.get_project(project_id)) != project_hash:
            raise ValueError("导出期间项目内容已变化，请检查后重新导出")
        if not self.quality_gate(project_id)[f"{mode}_ready"]:
            raise ValueError("导出期间质量或签审状态已变化，请检查后重新导出")
        stored = {}
        for item in files + ([path] if path not in files else []):
            stored[item.name] = self.storage.put_file(f"deliveries/{project_id}/{export_id}/{item.name}", item)
        artifact = {"file_name": path.name, "object_key": stored[path.name]["object_key"], "sha256": stored[path.name]["sha256"]}
        with self.db.connect() as conn:
            # File upload can take minutes. Recheck under the same project lock
            # used by edits/approvals, immediately before the immutable record.
            self._lock_project_for_update(conn, project_id)
            if self.db.backend == "postgresql":
                for table in ("claims", "requirement_responses", "confirmation_resolutions", "draft_review_decisions"):
                    conn.execute(f"SELECT id FROM {table} WHERE draft_id IN (SELECT id FROM project_drafts WHERE project_id=?) ORDER BY id FOR UPDATE", (project_id,)).fetchall()
                publication_ids = sorted({item[0] for item in manifest["knowledge"]})
                for publication_id in publication_ids:
                    conn.execute("SELECT id FROM knowledge_publications WHERE id=? FOR SHARE", (publication_id,)).fetchone()
            frozen_project = self._get_project_from_connection(conn, project_id)
            if self._project_hash(frozen_project) != project_hash:
                raise ValueError("导出期间项目内容已变化，请检查后重新导出")
            if self._delivery_confirmation(frozen_project) != manifest["delivery_confirmation"]:
                raise ValueError("导出期间质量或签审状态已变化，请检查后重新导出")
            frozen_quality = self.quality_gate(project_id, conn=conn, sync_issues=False)
            if not frozen_quality[f"{mode}_ready"]:
                raise ValueError("导出期间质量或签审状态已变化，请检查后重新导出")
            cursor = conn.execute(
                "INSERT INTO deliveries(project_id,format,file_path,quality_json) VALUES (?,?,?,?)",
                (project_id, f"{mode}_{file_format}", str(path), json.dumps({**quality, "_delivery_file": artifact}, ensure_ascii=False)),
            )
            delivery_id = int(cursor.lastrowid)
            cursor = conn.execute(
                """
                INSERT INTO delivery_manifests(project_id,delivery_id,project_hash,manifest_json,manifest_hash)
                VALUES (?,?,?,?,?)
                """,
                (project_id, delivery_id, project_hash, manifest_json, manifest_hash),
            )
            manifest_id = int(cursor.lastrowid)
        self.audit.record("delivery.freeze", "project", project_id, details={"delivery_id": delivery_id, "manifest_id": manifest_id, "manifest_hash": manifest_hash})
        progress("completed", 100, "交付包已冻结", {"manifest_hash": manifest_hash})
        return {
            "project_id": project_id,
            "delivery_id": delivery_id,
            "manifest_id": manifest_id,
            "manifest_hash": manifest_hash,
            "format": file_format,
            "mode": mode,
            "file_path": str(path),
            "object_key": artifact["object_key"],
            "file_name": path.name,
            "download_url": f"/api/projects/{project_id}/deliveries/{delivery_id}/download",
            "files": file_items,
            "quality": quality,
        }

    @staticmethod
    def _project_hash(project: dict[str, Any]) -> str:
        confirmation_fields = {"delivery_confirmation", "professional_reviewer", "reviewed_by", "compliance_confirmed", "manual_finalized", "final_approved"}
        return content_hash(
            json.dumps(
                {
                    "id": project["id"],
                    "name": project.get("name"),
                    "industry": project.get("industry"),
                    "project_type": project.get("project_type"),
                    "region": project.get("region"),
                    "source_text": project.get("source_text"),
                    "profile": {key: value for key, value in (project.get("profile") or {}).items() if key not in confirmation_fields},
                    "requirements": [(item["id"], item["content"], item["kind"], item.get("requirement_fingerprint"), item.get("classification")) for item in project["requirements"]],
                    "sections": [(section["id"], section["title"], section.get("order_no"), section.get("requirement_ids")) for section in project["sections"]],
                    "drafts": [(section["draft"]["id"], section["draft"]["content_hash"], section["draft"].get("citations")) for section in project["sections"] if section.get("draft")],
                },
                ensure_ascii=False,
                sort_keys=True,
            )
        )

    @staticmethod
    def _preflight_docx(path: Path) -> dict[str, Any]:
        document = Document(path)
        issues = []
        if not any(paragraph.text.strip() for paragraph in document.paragraphs):
            issues.append("文档正文为空")
        audit = audit_docx_path(path)
        if not audit["ready"]:
            issues.append("文档最终产物审计未通过")
        if not document.sections:
            issues.append("文档缺少页面设置")
        if issues:
            raise ValueError("DOCX预检失败：" + "；".join(issues))
        return {"paragraphs": len(document.paragraphs), "tables": len(document.tables), "size_bytes": path.stat().st_size, "artifact_audit": audit}

    @staticmethod
    def _convert_pdf(docx_path: Path) -> Path:
        executable = shutil.which("libreoffice") or shutil.which("soffice")
        if not executable:
            windows_soffice = Path("C:/Program Files/LibreOffice/program/soffice.exe")
            if windows_soffice.exists():
                executable = str(windows_soffice)
        if not executable:
            raise ValueError("未安装LibreOffice，无法生成和预检PDF")
        uno_python = os.environ.get("BID_WRITER_LIBREOFFICE_PYTHON") or "/usr/bin/python3"
        if not Path(uno_python).is_file():
            raise ValueError("缺少LibreOffice目录更新组件，请安装python3-uno或设置BID_WRITER_LIBREOFFICE_PYTHON")
        pdf_path = docx_path.with_suffix(".pdf")
        result = subprocess.run(
            [uno_python, str(Path(__file__).with_name("render_pdf.py")), "--soffice", executable,
             "--input", str(docx_path.resolve()), "--output", str(pdf_path.resolve())],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        if result.returncode or not pdf_path.exists():
            detail = (result.stderr or result.stdout or "未生成文件").strip().splitlines()[-1][:240]
            raise ValueError(f"LibreOffice生成PDF或更新目录失败：{detail}")
        return pdf_path

    @staticmethod
    def _preflight_pdf(path: Path) -> dict[str, Any]:
        from pypdf import PdfReader

        reader = PdfReader(path)
        if not reader.pages:
            raise ValueError("PDF预检失败：没有页面")
        blank_pages = [index + 1 for index, page in enumerate(reader.pages) if not (page.extract_text() or "").strip()]
        if blank_pages:
            raise ValueError(f"PDF预检失败：发现空白页{blank_pages[:10]}")
        return {"pages": len(reader.pages), "size_bytes": path.stat().st_size}

    @staticmethod
    def _project_markdown(project: dict[str, Any]) -> str:
        parts = [f"# {project['name']}"]
        for section in project["sections"]:
            if section.get("draft"):
                parts.append(section["draft"]["content"])
        return "\n\n".join(parts)

    def _write_docx(self, project: dict[str, Any], path: Path, mode: str = "formal") -> None:
        if mode not in {"review", "formal"}:
            raise ValueError("封面模式仅支持review或formal")
        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.5)
        normal = document.styles["Normal"]
        normal.font.name = "宋体"
        normal.font.size = Pt(12)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        for style_name in ["Title", "Subtitle", "TOC Heading", *[f"Heading {level}" for level in range(1, 10)]]:
            style = document.styles[style_name]
            style.font.color.rgb = RGBColor(0, 0, 0)
            for borders in style._element.xpath("./w:pPr/w:pBdr"):
                borders.getparent().remove(borders)
        title = document.add_heading(project["name"], level=0)
        title.alignment = 1
        title.paragraph_format.space_before = Pt(72)
        title.paragraph_format.space_after = Pt(20)
        for run in title.runs:
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
        subtitle = document.add_paragraph("技术标")
        subtitle.alignment = 1
        subtitle.paragraph_format.space_after = Pt(12)
        subtitle.runs[0].font.size = Pt(18)
        edition = document.add_paragraph("专业送审版" if mode == "review" else "正式版")
        edition.alignment = 1
        edition.paragraph_format.space_after = Pt(48)
        edition.runs[0].font.size = Pt(14)
        profile = project.get("profile") or {}
        cover_fields = [
            ("投标单位", profile.get("bidder_name") or profile.get("tenderer_name") or profile.get("bidder")),
            ("授权签字人", profile.get("authorized_signatory") or profile.get("signatory")),
            ("专业复核人", profile.get("professional_reviewer") or profile.get("reviewed_by")),
        ]
        confirmed_at = (profile.get("delivery_confirmation") or {}).get("confirmed_at")
        if mode == "formal" and confirmed_at:
            try:
                confirmed = datetime.fromisoformat(str(confirmed_at).replace("Z", "+00:00"))
                if confirmed.tzinfo is None:
                    confirmed = confirmed.replace(tzinfo=timezone.utc)
                # Confirmation timestamps are stored in UTC; bid documents use China local dates.
                cover_fields.append(("定稿日期", confirmed.astimezone(timezone(timedelta(hours=8))).strftime("%Y年%m月%d日")))
            except ValueError:
                pass
        elif profile.get("document_date"):
            cover_fields.append(("编制日期", profile["document_date"]))
        for label, value in cover_fields:
            text = normalize_text(str(value or ""))
            if not text:
                continue
            paragraph = document.add_paragraph(f"{label}：{text}")
            paragraph.alignment = 1
            paragraph.paragraph_format.space_after = Pt(12)
            paragraph.runs[0].font.size = Pt(13)
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        toc_title = document.add_paragraph("目录", style="TOC Heading")
        toc_title.alignment = 0
        toc = document.add_paragraph()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin.set(qn("w:dirty"), "true")
        toc.add_run()._r.append(begin)
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
        toc.add_run()._r.append(instruction)
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        toc.add_run()._r.append(separate)
        # Native TOC fields need cached result paragraphs. These are actual titles,
        # never guessed page numbers; LibreOffice refreshes pagination for PDF export.
        for index, item in enumerate(section for section in project["sections"] if section.get("draft")):
            if index:
                toc = document.add_paragraph()
            toc.add_run(item["title"])
            toc.paragraph_format.space_after = Pt(8)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        toc.add_run()._r.append(end)
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        document.settings._element.append(update_fields)
        document.add_page_break()
        for section_index, item in enumerate(project["sections"], 1):
            draft = item.get("draft")
            if not draft:
                continue
            if section_index > 1:
                document.add_page_break()
            document.add_heading(item["title"], level=1)
            self._append_markdown(document, draft["content"], item["title"])
        for table in document.tables:
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement(f"w:{edge}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    cell_properties = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for edge in ("top", "left", "bottom", "right"):
                        margin = OxmlElement(f"w:{edge}")
                        margin.set(qn("w:w"), "80" if edge in {"top", "bottom"} else "100")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    cell_properties.append(margins)
                    if row_index == 0:
                        for shading in cell_properties.findall(qn("w:shd")):
                            shading.set(qn("w:fill"), "DDEBF7")
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.line_spacing = 1.15
                        paragraph.paragraph_format.space_before = Pt(2)
                        paragraph.paragraph_format.space_after = Pt(2)
                        for run in paragraph.runs:
                            run.font.size = Pt(10.5)
        footer = document.sections[0].footer.paragraphs[0]
        footer.alignment = 1
        footer.add_run(f"{project['name']}  |  ")
        page_field = OxmlElement("w:fldSimple")
        page_field.set(qn("w:instr"), "PAGE")
        footer._p.append(page_field)
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)

    def _append_markdown(self, document: Document, markdown: str, section_title: str) -> None:
        lines = markdown.splitlines()
        index = 0
        first_heading_skipped = False
        while index < len(lines):
            line = lines[index].strip()
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading:
                title = heading.group(2).strip()
                if not first_heading_skipped and title == section_title:
                    first_heading_skipped = True
                else:
                    document.add_heading(title, level=min(len(heading.group(1)) + 1, 9))
                index += 1
                continue
            image = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", line)
            if image:
                image_path = Path(image.group(2).strip())
                if image_path.exists() and image_path.is_file():
                    paragraph = document.add_paragraph()
                    paragraph.alignment = 1
                    paragraph.add_run().add_picture(str(image_path), width=Cm(15.5))
                    caption = document.add_paragraph(f"图 {image.group(1).strip() or '施工技术示意图'}")
                    caption.alignment = 1
                index += 1
                continue
            if self._is_table_start(lines, index):
                index = self._append_markdown_table(document, lines, index)
                continue
            if line.startswith("- ") or line.startswith("* "):
                document.add_paragraph(line[2:].strip(), style="List Bullet")
            elif re.match(r"^\d+[.、]\s*", line):
                text = re.sub(r"^\d+[.、]\s*", "", line)
                document.add_paragraph(text, style="List Number")
            elif line in {"---pagebreak---", "<!-- pagebreak -->"}:
                document.add_page_break()
            elif line and not line.startswith("<!--"):
                paragraph = document.add_paragraph(line)
                paragraph.paragraph_format.first_line_indent = Cm(0.74)
                paragraph.paragraph_format.line_spacing = 1.5
                paragraph.paragraph_format.space_after = Pt(3)
            index += 1

    @staticmethod
    def _is_table_start(lines: list[str], index: int) -> bool:
        separator_index = index + 1
        while separator_index < len(lines) and not lines[separator_index].strip():
            separator_index += 1
        if separator_index >= len(lines):
            return False
        current = lines[index].strip()
        separator = lines[separator_index].strip()
        return current.startswith("|") and current.endswith("|") and bool(re.match(r"^\|[\s:|-]+\|$", separator))

    @staticmethod
    def _table_cells(line: str) -> list[str]:
        return [value.strip().replace("\\|", "|") for value in line.strip().strip("|").split("|")]

    def _append_markdown_table(self, document: Document, lines: list[str], index: int) -> int:
        headers = self._table_cells(lines[index])
        rows: list[list[str]] = []
        separator_index = index + 1
        while separator_index < len(lines) and not lines[separator_index].strip():
            separator_index += 1
        cursor = separator_index + 1
        while cursor < len(lines):
            line = lines[cursor].strip()
            if not line:
                next_index = cursor + 1
                while next_index < len(lines) and not lines[next_index].strip():
                    next_index += 1
                if next_index < len(lines) and lines[next_index].strip().startswith("|"):
                    cursor = next_index
                    line = lines[cursor].strip()
                else:
                    break
            if not (line.startswith("|") and line.endswith("|")):
                break
            rows.append(self._table_cells(line))
            cursor += 1
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        header_row = table.rows[0]
        header_properties = header_row._tr.get_or_add_trPr()
        header_properties.append(OxmlElement("w:tblHeader"))
        for cell_index, value in enumerate(headers):
            cell = header_row.cells[cell_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "DDEBE8")
            cell._tc.get_or_add_tcPr().append(shading)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)
        for values in rows:
            row = table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = values[cell_index] if cell_index < len(values) else ""
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(8.5)
        document.add_paragraph()
        return cursor
